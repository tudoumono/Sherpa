"""読み取り部品（KB 検索・グラフ探索・原本取得）。
`run_tool(name, args, world, scope_paths, ...)` が道具ディスパッチの正本。MCP（`mcp_server.py`）・チャットの回答ループ・簡易チャット（`/ext/v1/answer`）が `sherpa.tool_dispatch.run_tool` 経由で共有する。
権限（資料フォルダ・scope_paths・layer・personal 除外）は呼び出し側が確定済みの値を渡す契約で、本モジュール自身は権限を持たず、書き込みも行わない。
設計: docs/design/interfaces.md「読み取り部品（境界としての切り出し済み・手順②）」、docs/design/codex.md「MCP の道具」
"""
from __future__ import annotations

import json
import logging
import os
import re
import stat
from pathlib import Path

from ... import citations, es_index, graph_coverage, grep_tool, redact_keys, scope_infer, worlds
from ...env_int import env_int
from ... import layer as layer_mod
from ... import scope as scope_mod
from ... import text_encoding
from ...ingest import importance, text_kind
from ...ingest.analyzers import registry as _analyzer_registry
from ...safe_open import open_file_nofollow_walk as _open_file_nofollow_walk  # TOCTOU 耐性のファイル open（`safe_open.py`・`ext_api.py` と共用）


# `simple_chat.py`／`ext_api.py` 等と同じ共有ロガー
_log = logging.getLogger("sherpa")

# grep／es_search 1 回あたりのヒット数の既定（管理画面の基準値が未設定のときに使う）
MAX_HITS = 45

# ヒット数の絶対上限。調べる深さ（`depth_profile.scaled_ratio`）が倍率適用後に 1 度だけ適用する絶対上限として grep／ES が共有する
MAX_HITS_ABS_MAX = 1000

# read_around の精読窓（行数）の既定（管理画面の基準値が未設定のときに使う）。窓のハード上限（`max(200, READ_WINDOW)`）はこの値が 200 を超えたときだけ追随する
READ_WINDOW = 60

_OFFICE_MD = {".docx", ".xlsx", ".pptx", ".pdf", ".doc", ".xls", ".ppt",
              # ラスタ画像（OCR アーム）も本文は派生 MD 側（`image.png.md`）にある。`office_md.IMAGE_EXT` が真実源で、grep と read_around が一致する
              ".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}  # 本文は派生 MD 側にある

# 旧形式（.doc/.xls/.ppt）は前段変換した OOXML が MD 化され、解決規約は新形式と同じ「原本 rel + .md」。`_OFFICE_MD` に加えるだけで grep_search と read_around が一致する。
# `_READABLE_EXT` は read_around が読める本文種別の代表集合（`_FILE_HEAD_KINDS` の材料）。`_safe_doc_path` は拡張子の事前フィルタには使わず `classify_document`／`reachable_as_text` の確定判定に委ねる
_READABLE_EXT = ({".md", ".markdown", ".txt"} | _analyzer_registry.registered_extensions() | _OFFICE_MD
                | text_kind.CODE_EXT | text_kind.DOCUMENT_EXT)

# tool result（外部 LLM へ渡る）から明らかな秘密を伏せる（多層防御）
_SECRET_RE = re.compile(
    r"(sk-[A-Za-z0-9_-]{16,}|AIza[0-9A-Za-z_-]{20,}|gh[pousr]_[A-Za-z0-9]{20,}|xox[baprs]-[A-Za-z0-9-]{10,}"
    r"|AKIA[0-9A-Z]{16}|eyJ[A-Za-z0-9_-]{8,}\.[A-Za-z0-9_-]{8,}\.[A-Za-z0-9_-]{6,}"
    r"|-----BEGIN[^-]+PRIVATE KEY(?: BLOCK)?-----[\s\S]*?-----END[^-]+PRIVATE KEY(?: BLOCK)?-----)")

_KV_SECRET_RE = re.compile(r"(?i)\b(pass(?:word|wd)?|secret|api[_-]?key|token|authorization)\b(\s*[=:]\s*)(\S+)")

def _redact_plain(text: str) -> str:
    t = _SECRET_RE.sub("[REDACTED]", text or "")
    return _KV_SECRET_RE.sub(r"\1\2[REDACTED]", t)

def _redact(text: str) -> str:
    """秘密値の伏せ字。BEGIN だけ・END だけ・本文の base64 連続行だけの断片（窓や切り詰めで片側が無い鍵）も伏せる。"""
    return redact_keys.mask_orphan_key_body(redact_keys.KeyBlockRedactor(_redact_plain)(text or ""))

def _walk_redact(obj, redactor):
    if isinstance(obj, str):
        return redactor(obj)
    if isinstance(obj, list):
        return [_walk_redact(v, redactor) for v in obj]
    if isinstance(obj, dict):
        return {k: _walk_redact(v, redactor) for k, v in obj.items()}
    return obj

def _redact_deep(obj):
    """`_redact` を dict／list を再帰的に辿って全ての文字列値へ適用する（原本読取ツール専用）。
    `redact_keys.KeyBlockRedactor` に委譲し、出現順を 1 本の状態付きスキャンとして扱う（PEM 秘密鍵の BEGIN／END が別要素にまたがっても取りこぼさない）。`doc_readers` が切り詰め前に適用済みの後の多層防御。数値・真偽値・None は素通し。
    """
    return _walk_redact(obj, redact_keys.KeyBlockRedactor(_redact))

# `read_around` の出力は行数でしか絞られないため、(a) 返却テキストの UTF-8 バイト上限で切り詰め、(b) 生バイトを `_READ_AROUND_FILE_CAP_BYTES` までに制限して読む。
# 既定は精度優先で、管理画面（`agentic_budget.per_result`）が唯一の真実源のため env フォールバックは持たない（settings 未設定時のコード既定として `effective_tool_result_max_bytes()` が使う）
TOOL_RESULT_MAX_BYTES = 262144

# read_around／read_doc／doc_outline／verify_citation がディスクから読む生バイト数の上限。`grep_tool._GREP_FILE_CAP_BYTES` と揃える（揃えないと grep が cap より後ろでヒットを見つけても read 側が読めない）
_READ_AROUND_FILE_CAP_BYTES = env_int(
    "SHERPA_READ_AROUND_FILE_CAP_BYTES", 64 * 1024 * 1024, 65536, 64 * 1024 * 1024)

# read 側の単一巨大行への安全弁（`grep_tool._GREP_LINE_MAX_BYTES` と同じ役割）
_READ_LINE_MAX_BYTES = 2 * 1024 * 1024

# `ripgrep_search` の tool result に載せる「打切りで探せていない文書」の件数上限（ツール結果のバイト予算を圧迫しないため）
_TRUNCATED_DOCS_MAX = 20

# ripgrep_search／es_search: ヒット単位のバイト上限の下限（日本語で約 170 文字＝見出し＋冒頭数行）。件数ではなく本文量をヒット単位で頭打ちにする
_HIT_TEXT_MIN_BYTES = 512

# es_search の方式。hybrid は語の一致＋意味の近さ・keyword は語の一致だけ・vector は意味の近さだけ
_ES_MODES = ("hybrid", "keyword", "vector")
# 意味検索（埋め込み）が使えず BM25 だけで返したことを示す縮退理由（`mode_used` を keyword にする）
_ES_VECTOR_FALLBACK_REASONS = frozenset({
    "embedding_not_configured", "embedding_cloud_unavailable", "vector_feature_mismatch",
    "query_embed_failed", "hybrid_query_failed"})
# 親返し: es_search のヒットを doc_id で束ね、予算内なら rag.md の領域（P2）を返す（常時 ON）。
# P2 の対象チャンク集合を ES から引く 1 クエリあたりの取得上限（`es_index.chunk_ids_for_parent` の `limit`）。env 化しない
_PARENT_RETURN_REGION_CHUNKS_MAX = 5000

def _clip_utf8_bytes(s: str, max_bytes: int) -> str:
    """UTF-8 エンコード後のバイト数が `max_bytes` を超えないよう `s` を切り詰める（マルチバイト文字の境界で壊れた文字が残らないよう `errors="ignore"` で再デコードする）。"""
    b = s.encode("utf-8")
    if len(b) <= max_bytes:
        return s
    return b[:max_bytes].decode("utf-8", errors="ignore")

# 直列化不能時のフォールバック値。どの上限設定よりも大きくし、「測定不能＝上限超過扱い」にする
_UNMEASURABLE_SIZE = 1 << 40

def _result_byte_size(result) -> int:
    """JSON 化した際の概算 UTF-8 バイト数（1 run 累計上限の判定に使う）。`run_tool` の戻り値の 1 つ目（tool result dict）と 4 つ目（`cards` サイドカー）に使う。
    測定不能（bytes・非 JSON 型・不正 Unicode 等）は特大（`_UNMEASURABLE_SIZE`）として扱う（fail-closed。0 扱いだと個別上限・累計上限をすり抜ける）。
    """
    try:
        return len(json.dumps(result, ensure_ascii=False).encode("utf-8"))
    except Exception:
        return _UNMEASURABLE_SIZE

# graph_neighbors のカード件数上限（troubleshoot UI 用サイドカー・env 化しない）
_GRAPH_CARDS_MAX = 30

def _clip_cards(cards: list, max_count: int = _GRAPH_CARDS_MAX, max_bytes: int = TOOL_RESULT_MAX_BYTES) -> list:
    """`cards`（`graph_neighbors` のカード・troubleshoot UI 用サイドカー）を件数＋直列化バイト上限で切り詰める（超過分は捨てる・fail-closed）。
    各カードを仮に追加した候補リスト全体の実直列化バイト数が `max_bytes` を超えるなら、そのカードを追加せず打ち切る（巨大な単一カードも弾く）。
    """
    out: list = []
    for c in cards[:max_count]:
        candidate = out + [c]
        if _result_byte_size(candidate) > max_bytes:
            break
        out = candidate
    return out

# glob_search の返却上限（打ち切りは明示する）。env 化しない
_GLOB_MAX_RESULTS = 200

# `_重要度.txt` の glob と同じ長さ上限を流用する（`importance._match_segment_glob` を共有するため）
_GLOB_PATTERN_MAX_LEN = importance._MAX_PATTERN_LEN

# doc_outline の見出し検出（ATX 形式・レベル 1〜3 のみ）。構造の当たり付けに要る大枠だけを返し、細部は read_doc／read_around に委ねる。`\s+` を要求して `#!/bin/sh` 等を誤検出しない
_HEADING_RE = re.compile(r"^(#{1,3})\s+(.+)$")

# 1 回の返却件数上限（見出し数の増幅を防ぐ）
_OUTLINE_MAX_HEADINGS = 200

_OUTLINE_TITLE_MAX_CHARS = 300

def _validate_glob_pattern(raw) -> str | None:
    """`glob_search` の `pattern` 引数を検証する（無効なら None）。`doc_id`（`_safe_doc_path`）と同じトラバーサル拒否（絶対パス・バックスラッシュ・NUL・`..`／空セグメント）を適用する。"""
    if not isinstance(raw, str):
        return None
    pattern = raw.strip()
    if not pattern or len(pattern) > _GLOB_PATTERN_MAX_LEN:
        return None
    if pattern.startswith("/") or "\\" in pattern or "\x00" in pattern:
        return None
    parts = pattern.split("/")
    if ".." in parts or "" in parts:
        return None
    return pattern

def _glob_match_pattern(pattern: str) -> str:
    """スラッシュを含まないパターンは「どの階層のファイル名にも一致」とみなし `**/` を前置する（ripgrep の `--glob` と同じ慣習）。スラッシュを含むパターンはそのまま（資料フォルダ root からの絞り込み）。"""
    return pattern if "/" in pattern else f"**/{pattern}"

def _safe_doc_path(world: str, doc_id: str, *, layer=None):
    """doc_id（rel_path）→ `(root, lexical_rel, 読み取り可能な実パス)`（無効／範囲外／秘匿種別は None）。
    - `layer`（省略可・既定 None＝層チェックしない）: 指定時は `classify_document` の確定結果（`layer_mod.in_layer_code`）で層一致も確認する。Office／画像は常に `"docs"` 側。
    - 検査: トラバーサル（`..`／絶対／空セグメント／バックスラッシュ／NUL）拒否、本文種別のみ、解決後に許可ルート（Office＝派生 MD root／その他＝資料フォルダ root）配下に閉じることを realpath で確認（symlink 脱出も拒否）。
    - 拡張子の事前フィルタは持たない。Office／画像（`_OFFICE_MD`）だけ派生 MD root へ分岐して実在確認のみ行い、それ以外は `corpus_docs.classify_document`（`_classify_verdict_reachable`）で最終確定する（grep／ES／list_docs と同じ契約）。
    - rag／legacy の優先順位（`grep_tool.preferred_derived_name`）はここで 1 回だけ解決する。返す `lexical_rel` は doc_id から機械的に導いた値で、呼び出し元（`run_tool` の read_around）は `root`／`lexical_rel` を後段の nofollow walk へそのまま渡し、再解決しない。
    - 順序: 封じ込め・symlink 拒否・regular file 確認を先に行い、通過した実パスに対してだけ `classify_document`（実ファイルを読む）を呼ぶ。
    - symlink 拒否は字面パス（解決済み root＋`lexical_rel` の連結）と `resolve()` 結果を突き合わせ、不一致なら `cand` 自身か祖先に symlink があったとして拒否する（root 内を指す symlink も一律拒否）。
    """
    if not doc_id or doc_id.startswith("/") or "\\" in doc_id or "\x00" in doc_id:
        return None
    parts = doc_id.split("/")
    if ".." in parts or "" in parts:
        return None
    if any(scope_infer.is_vcs_dir_name(p) for p in parts[:-1]):  # 版管理の記録のフォルダの中は取り込みの対象外
        return None
    ext = Path(doc_id).suffix.lower()
    if importance.is_importance_control_path(doc_id):  # 重要度設定ファイル自体は精読対象外
        return None
    if text_kind.is_sensitive(Path(doc_id).name, ext):
        # 秘匿名は Office／画像（`classify_document` を通らない）も含めてここで先に弾く
        _log.warning("read_around: 秘匿名のため対象外にしました（ext=%s）", ext)
        return None
    is_office = ext in _OFFICE_MD
    if is_office:
        # rag（RAG 正本）／md（人間用・legacy 縮退）は別ディレクトリ。`preferred_derived_name` が返す名前が `.rag.md` で終わるかで物理ルートを判別する
        der_rag = worlds.derived_rag_dir(world)
        lexical_rel = grep_tool.preferred_derived_name(der_rag, doc_id)
        root = der_rag if lexical_rel.endswith(grep_tool._RAG_SUFFIX) else worlds.derived_md_dir(world)
    else:
        root = worlds.world_dir(world)
        lexical_rel = doc_id
    if not root:
        return None
    root = Path(root)
    cand = root / lexical_rel
    try:
        rr = root.resolve()
        rp = cand.resolve()
        if not (rp == rr or rp.is_relative_to(rr)):
            return None
        if rp != rr / lexical_rel:  # 字面パスと不一致＝経路上のどこかに symlink があった
            return None
        if not rp.is_file():  # FIFO／ソケット等の非 regular も拒否
            return None
    except OSError:
        return None
    is_code = False
    if not is_office:
        from ... import corpus_docs
        # `text_quality` を渡し、文字コードを判別できない原本を対象外にする（grep／scan_report と同じ `_classify_verdict_reachable` 判定を共有）
        verdict = corpus_docs.classify_document(
            doc_id, ext, lambda p=rp, size=4096: corpus_docs._read_head(p, size),
            text_quality=lambda p=rp: corpus_docs._text_quality_for(p))
        if not corpus_docs._classify_verdict_reachable(verdict):
            return None
        is_code = verdict["kind"] == "code"
    if layer is not None and not layer_mod.in_layer_code(is_code, layer):
        return None
    return root, lexical_rel, rp

# 原本読取ツール（`doc_readers.py`）専用の doc_id 解決に許す拡張子（小文字・ドット付き）。`.xlsm` は台帳が文書種別として扱わないため対象外（`.xlsx` のみ）
_XLSX_KINDS = frozenset({".xlsx"})

_DOCX_KINDS = frozenset({".docx"})

_PPTX_KINDS = frozenset({".pptx"})

_PDF_KINDS = frozenset({".pdf"})

# file_head はテキスト・コードのみ（Office／PDF／画像は専用ツール）。`_safe_original_path` は `kinds` がこの定数と同じオブジェクトかで file_head を見分け、実際に読めるかは `reachable_as_text` に委ねる
_FILE_HEAD_KINDS = frozenset(_READABLE_EXT - _OFFICE_MD)

def _safe_original_path(world: str, doc_id: str, scope_paths, *, kinds: frozenset, layer=None):
    """原本読取ツール（xlsx／docx／pptx／pdf／file_head）専用の doc_id → `(root, doc_id, 実パス, stat)` 解決。
    `_safe_doc_path` と同じ検査（トラバーサル拒否・重要度制御ファイル除外・秘匿名除外・realpath 封じ込め・symlink 拒否・regular file）を共有するが、Office／PDF も派生 MD でなく資料フォルダ root の原本へ解決する。
    `kinds`（必須）: このツールが扱える拡張子集合。Office／PDF の 4 ツールは拡張子の一致を見たうえで、実在・文書種別・範囲を `verify_doc_exists` に委ねる。file_head（`kinds is _FILE_HEAD_KINDS`）は Office／PDF／画像を拒み、それ以外は `corpus_docs.reachable_as_text`（`classify_document` 経由）で可否を決める（`verify_doc_exists` は内容 sniff をしないため第 2 段の文書を弾いてしまう）。
    `layer`（省略可）は file_head 専用で、指定時は `layer_mod.in_layer_code` で層一致も見る。
    無効・範囲外・拡張子不一致・秘匿名・重要度制御・traversal・symlink・非 regular・未実在・読み取り不可はすべて `None`。
    戻り値の 4 つ目 `stat` は検査直後の `rp.stat()`。呼び出し元は `open()` までに差し替えられていないかを `os.fstat` の (st_dev, st_ino) と突き合わせる（`_open_verified_original` 参照）。
    """
    if not doc_id or doc_id.startswith("/") or "\\" in doc_id or "\x00" in doc_id:
        return None
    parts = doc_id.split("/")
    if ".." in parts or "" in parts:
        return None
    if any(scope_infer.is_vcs_dir_name(p) for p in parts[:-1]):  # 版管理の記録のフォルダの中は取り込みの対象外
        return None
    ext = Path(doc_id).suffix.lower()
    is_file_head = kinds is _FILE_HEAD_KINDS
    if is_file_head:
        if ext in _OFFICE_MD:  # 専用ツール（xlsx／docx／pptx／pdf）が扱う——file_head は対象外
            return None
    elif ext not in kinds:
        return None
    if importance.is_importance_control_path(doc_id):
        return None
    if text_kind.is_sensitive(Path(doc_id).name, ext):
        _log.warning("read_original: 秘匿名のため対象外にしました（ext=%s）", ext)
        return None
    if not scope_mod.in_scope(doc_id, scope_paths):
        return None
    root = worlds.world_dir(world)
    if not root:
        return None
    root = Path(root)
    cand = root / doc_id
    try:
        rr = root.resolve()
        rp = cand.resolve()
        if not (rp == rr or rp.is_relative_to(rr)):
            return None
        if rp != rr / doc_id:  # 字面パスと不一致＝経路上のどこかに symlink があった
            return None
        if not rp.is_file():  # FIFO／ソケット等の非 regular も拒否
            return None
    except OSError:
        return None
    from ... import corpus_docs
    if is_file_head:
        # `_safe_doc_path` と同じ理由（文字コードを判別できない原本は対象外）
        verdict = corpus_docs.classify_document(
            doc_id, ext, lambda p=rp, size=4096: corpus_docs._read_head(p, size),
            text_quality=lambda p=rp: corpus_docs._text_quality_for(p))
        if not corpus_docs._classify_verdict_reachable(verdict):
            return None
        is_code = verdict["kind"] == "code"
        if layer is not None and not layer_mod.in_layer_code(is_code, layer):
            return None
    else:
        if not verify_doc_exists(doc_id, world, scope_paths):  # 実在・文書種別・範囲の確定判定を共有
            return None
    try:
        st = rp.stat()
    except OSError:
        return None
    return root, doc_id, rp, st

def _open_verified_original(root: Path, doc_id: str, expected_st) -> tuple:
    """`_safe_original_path` が検査した `doc_id` を、`root`（信頼済みアンカー）から `_open_file_nofollow_walk`（各階層を `O_DIRECTORY|O_NOFOLLOW` で 1 段ずつ辿る）で開き直し、検査直後の `expected_st` と dev／inode が一致することを確認する。
    最終要素だけの `O_NOFOLLOW` では祖先ディレクトリの symlink 差し替えを防げないため、各段を辿る。fstat 突合は仕上げの二重の安全弁。
    戻り値 `(f, error)`。成功時 `f` は open 済みバイナリファイルで、所有権は `doc_readers` へ引き継がれる。失敗時 `(None, {"error": ...})`。
    """
    rel_parts = Path(doc_id).parts
    if not rel_parts:
        return None, {"error": "読み取りに失敗しました", "error_code": "read_io_failed"}
    try:
        fd = _open_file_nofollow_walk(root, rel_parts)
    except OSError:
        return None, {"error": "読み取りに失敗しました", "error_code": "read_io_failed"}
    try:
        post = os.fstat(fd)
        if not stat.S_ISREG(post.st_mode):
            os.close(fd)
            return None, {"error": "読み取りに失敗しました", "error_code": "read_io_failed"}
        if (post.st_dev, post.st_ino) != (expected_st.st_dev, expected_st.st_ino):
            os.close(fd)
            return None, {"error": "読み取りに失敗しました", "error_code": "read_io_failed"}
        f = os.fdopen(fd, "rb")
    except OSError:
        try:
            os.close(fd)
        except OSError:
            pass
        return None, {"error": "読み取りに失敗しました", "error_code": "read_io_failed"}
    return f, None

def _close_quiet_local(f) -> None:
    """`_open_verified_original` が開いたが `doc_readers` へ渡さずに終わる分岐で fd を漏らさず閉じる。"""
    try:
        f.close()
    except OSError:
        pass

def _open_doc_stream(world: str, doc_id: str, sp, layer) -> tuple:
    """`doc_id` を安全に解決し、読み取り用に open 済みのバイナリファイルオブジェクトを返す（`read_around`／`read_doc`／`doc_outline` が共有する土台）。
    範囲・層フィルタと symlink TOCTOU 対策（`_safe_doc_path` の結果を信頼アンカーに、`lexical_rel` を `_open_file_nofollow_walk` で 1 段ずつ open）を 3 ツールで共通にする。
    戻り値 `(f, error)`。成功時 `f` は呼び出し元が close する（ストリーミング走査の間じゅう開いておく）。失敗時 `(None, {"error": ...})`。全文の一括ロードはせず、`grep_tool._CappedStreamReader`／`_logical_lines` で bounded に走査する（`_stream_doc_lines` 参照）。
    """
    if not scope_mod.in_scope(doc_id, sp):  # 範囲外は読まない
        return None, {"error": "指定 doc_id は対象範囲外です"}
    resolved = _safe_doc_path(world, doc_id, layer=layer)
    if resolved is None:
        return None, {"error": "doc_id が無効、または読み取り対象外です"}
    root, lexical_rel, _validated_path = resolved
    rel_parts = Path(lexical_rel).parts
    if not rel_parts:
        return None, {"error": "読み取りに失敗しました"}
    try:
        fd = _open_file_nofollow_walk(root, rel_parts)
    except OSError:
        # 読取 I/O 失敗は固定理由コード（`error_code`）を付ける（`run_tool` 境界が `backend_failures["read_io"]` へ反映する）
        return None, {"error": "読み取りに失敗しました", "error_code": "read_io_failed"}
    fd_owned = True
    try:
        if not stat.S_ISREG(os.fstat(fd).st_mode):
            return None, {"error": "読み取りに失敗しました"}
        f = os.fdopen(fd, "rb")
        fd_owned = False  # 以後の close は呼び出し元（`f.close()`）が引き受ける
    except OSError:
        return None, {"error": "読み取りに失敗しました", "error_code": "read_io_failed"}
    finally:
        if fd_owned:
            try:
                os.close(fd)
            except OSError:
                pass
    return f, None

def _stream_doc_lines(f):
    """open 済み `f`（`_open_doc_stream` が返すバイナリファイル）を `_READ_AROUND_FILE_CAP_BYTES`／`_READ_LINE_MAX_BYTES` で bounded にストリーミング走査し、`(reader, 行イテレータ, encoding_caution)` を返す。
    符号化の判定と行分割（`_logical_lines`＝`str.splitlines()` と同じ論理行）は grep と共有し、検索・精読・引用の行番号を揃える。呼び出し元はイテレータを消費し終えた後に `reader.truncated`／`reader.line_overflowed` を見て `file_truncated` を判定する。
    `encoding_caution`: 符号化が不確実（`quality_of` が "partial"）なら短い注意文、それ以外は `None`（fd のキャッシュを使い、二重 open を避ける）。
    """
    from ... import corpus_docs
    reader = grep_tool._CappedStreamReader(f, line_max_bytes=_READ_LINE_MAX_BYTES)
    enc, ratio, majority_garbled = text_encoding.detect_fd_quality(f.fileno())
    caution = corpus_docs.encoding_caution_for_ratio(ratio, majority_garbled)
    return reader, grep_tool._logical_lines(reader, _READ_AROUND_FILE_CAP_BYTES, encoding=enc), caution

def _rag_md_region_text(world: str, doc_id: str, sp, layer, target_chunk_ids, byte_cap: int,
                        info: dict | None = None) -> str | None:
    """親返し P2: rag.md をアンカー（`<!-- chunk:{chunk_id} -->`）単位でストリーミング走査し、`target_chunk_ids` に属するチャンクの本文だけを集める。
    `info`（省略可）には、集められなかった対象チャンクの数を `info["missing_chunks"]` として返す（rag.md に無い・走査が cap で終わった等）。
    対象外のチャンク本文は保持せず、全件そろうか `byte_cap` 超過で打ち切る。`byte_cap` を超えたら None を返し、集めた部分的な本文は使わない（途中で打ち切ったものを完全なものとして返さない）。
    """
    if not target_chunk_ids:
        return None
    f, err = _open_doc_stream(world, doc_id, sp, layer)
    if err is not None:
        return None
    remaining = set(target_chunk_ids)
    collected: dict = {}
    order: list = []
    cur_id = None
    cur_buf: list = []
    total_bytes = 0
    over = False

    def _close(cid: str, buf: list) -> None:
        nonlocal total_bytes, over
        body = "\n".join(buf).strip()
        collected[cid] = body
        order.append(cid)
        remaining.discard(cid)
        total_bytes += len(body.encode("utf-8"))
        if total_bytes > byte_cap:
            over = True

    try:
        _reader, it, _enc_caution = _stream_doc_lines(f)  # 親返しの抜粋には注意文を付けない
        for line in it:
            anchor_id = es_index.rag_md_anchor_chunk_id(line)
            if anchor_id is not None:
                if cur_id is not None and cur_id in remaining:
                    _close(cur_id, cur_buf)
                    if over:
                        break
                if not remaining:
                    break
                cur_id, cur_buf = anchor_id, []
                continue
            if cur_id is not None and cur_id in remaining:
                cur_buf.append(line)
        else:
            # EOF（break していない）＝最後のアンカーの本文が未確定なら確定させる
            if cur_id is not None and cur_id in remaining:
                _close(cur_id, cur_buf)
    finally:
        f.close()
    if over or not collected:
        return None
    if info is not None:
        info["missing_chunks"] = len(remaining)
    return "\n\n".join(collected[cid] for cid in order)

_hit_scores: list = []


def _set_hit_scores(scores: list) -> None:
    """直前の `es_search` が返したヒットの点数（結果の `hits` と同じ並び）を控える。"""
    global _hit_scores
    _hit_scores = list(scores)


_hit_ranks: list = []


def _set_hit_ranks(ranks: list) -> None:
    """直前の検索が返したヒットの元の順位（秘匿の除外・資料ごとの束ねの前の並びでの順位・結果の `hits` と同じ並び）を控える。"""
    global _hit_ranks
    _hit_ranks = list(ranks)


def pop_hit_ranks() -> list:
    """`_set_hit_ranks` で控えた順位を取り出して空にする（道具の呼び出しの記録が、結果を変えずに順位を読むため）。"""
    global _hit_ranks
    out, _hit_ranks = _hit_ranks, []
    return out


def pop_hit_scores() -> list:
    """`_set_hit_scores` で控えた点数を取り出して空にする（道具の呼び出しの記録が、結果を変えずに点数を読むため）。"""
    global _hit_scores
    out, _hit_scores = _hit_scores, []
    return out


def _resolve_parent_return(world: str, rag_groups: dict, sp, layer, budget_for_rag: int) -> list:
    """親返し本体: doc_id ごとに束ねた rag チャンクのヒットを P2（領域）／chunk（子のみ）へ振り分ける。決定的な貪欲法:
    ① 全 doc の最低保証（子チャンク本文の合計＝`baseline`）を `budget_for_rag` から先に確保する（先頭の巨大文書が予算を食い尽くして子チャンクが消えるのを防ぐ）。
    ② 残り予算をベストスコア順（同点は doc_id 昇順）に使い、領域が「共有の残り予算」と「1 文書あたりの上限（`per_doc_cap`）」の両方に収まれば P2、収まらなければ chunk のままにする。`per_doc_cap` は総予算を doc 数で均等割り（下限 `_HIT_TEXT_MIN_BYTES`）した値。
    ③ 各 doc は必ず 1 エントリを返し、`tier` を必ず申告する。`per_doc_cap` に阻まれて領域を取れなかった doc には `text_truncated: true` を立てる（read_doc で補える余地を示す）。
    `rag_groups`: `{doc_id: [{"chunk_id", "parent_id", "locator", "score", "text"}, ...]}`（`text` は redaction 済みの子チャンク本文）。`budget_for_rag`: この tool result のうち rag doc 群に残っている予算（呼び出し元が計算する）。
    """
    groups = []
    for doc_id, items in rag_groups.items():
        baseline = sum(len(it["text"].encode("utf-8")) for it in items)
        best_score = max(float(it.get("score") or 0) for it in items)
        groups.append((doc_id, items, baseline, best_score))
    remaining = max(0, budget_for_rag - sum(g[2] for g in groups))
    groups.sort(key=lambda g: (-g[3], g[0]))
    per_doc_cap = max(_HIT_TEXT_MIN_BYTES, budget_for_rag // max(1, len(groups)))

    out = []
    for doc_id, items, baseline, _best_score in groups:
        chunk_ids = [it["chunk_id"] for it in items]
        tier = "chunk"
        text = "\n\n".join(it["text"] for it in items)  # 最低保証（redaction／クリップ済み）
        text_truncated = False
        region_capped = False
        parent_unfetched = False
        region_info = {}
        parent_ids = sorted({it["parent_id"] for it in items if it.get("parent_id")})
        if parent_ids:
            es_ids = es_index.chunk_ids_for_parent(
                world, doc_id, parent_ids, limit=_PARENT_RETURN_REGION_CHUNKS_MAX)
            # ES の取得が上限ちょうどまで返ったら、領域のチャンクが上限で打ち切られている可能性がある
            region_capped = len(es_ids) >= _PARENT_RETURN_REGION_CHUNKS_MAX
            parent_unfetched = not es_ids  # 親の領域のチャンク一覧を取れなかった（ES 不達・失敗）＝ヒット自身だけでは領域とは言えない
            target_ids = set(es_ids)
            target_ids |= set(chunk_ids)  # ヒット自身のチャンクは必ず含める（ES 反映漏れの安全弁）
            shared_cap = baseline + remaining
            per_doc_ceiling = max(per_doc_cap, baseline)  # baseline（最低保証）自体は削らない
            region_cap = min(shared_cap, per_doc_ceiling)
            region_info = {}
            region_text = _rag_md_region_text(world, doc_id, sp, layer, target_ids, region_cap, region_info)
            if region_text is not None:
                delta_region = len(region_text.encode("utf-8")) - baseline
                if delta_region <= remaining:
                    text = _redact(region_text)
                    tier = "region"
                    remaining -= delta_region
                    if region_info.get("missing_chunks"):
                        text_truncated = True
            if tier != "region" and per_doc_ceiling < shared_cap:
                # 1 文書あたりの上限が共有予算より先に効いた＝この文書は上限で頭打ちにされた
                text_truncated = True
        entry = {"doc_id": doc_id, "tier": tier, "text": text,
                 "chunks": [{"chunk_id": it["chunk_id"],
                             **({"locator": it["locator"]} if it.get("locator") is not None else {})}
                            for it in items]}
        if text_truncated:
            entry["text_truncated"] = True
        if tier == "chunk":
            entry["fragment"] = True  # 子チャンクだけ＝文書の一部（全文は read_doc／read_around）
        if parent_unfetched:
            entry["fragment"] = True
            entry["text_truncated"] = True
            entry["parent_chunks_unfetched"] = True  # 親の領域のチャンクを取得できなかった（本文はヒットのチャンクだけ）
        if region_capped:
            entry["region_chunks_capped"] = True  # 領域のチャンク取得が ES の取得上限（5000）に当たった
        if tier == "region" and region_info.get("missing_chunks"):
            entry["region_missing_chunks"] = region_info["missing_chunks"]  # 領域の一部のチャンクが集まらなかった
        if any(it.get("keyword_match") is not None for it in items):
            entry["keyword_match"] = any(it.get("keyword_match") for it in items)  # 束ねた子チャンクのどれかに語が一致していれば True
        out.append(entry)
    return out

# ツール名 → 切り詰め対象フィールド
_READER_CLIP_FIELD = {
    "xlsx_sheets": "sheets", "xlsx_range": "rows", "docx_paragraphs": "paragraphs",
    "pptx_slides": "slides", "pdf_pages": "pages", "file_head": "text",
}

def _row_start_from_a1_range(range_a1: str) -> int | None:
    """`"B3:D10"` 等の A1 range 文字列から開始行番号（3）を取り出す（解析できなければ None）。"""
    m = re.match(r"^[A-Za-z]+(\d+)", (range_a1 or "").split(":")[0])
    return int(m.group(1)) if m else None

def _doc_reader_text_locator(name: str, result: dict) -> tuple[str | None, str | None]:
    """原本読取ツール（6 本）の結果から `read_evidence`／根拠ゲートに載せる `text`（rows／paragraphs／slides／pages を 1 本の本文に連結・位置情報を行頭に付ける）と `locator` を組む。エラー・中身なしは `(None, None)`（空の read evidence を作らない）。"""
    if not isinstance(result, dict) or result.get("error"):
        return None, None
    if name == "xlsx_sheets":
        sheets = result.get("sheets") or []
        if not sheets:
            return None, None
        def _dims(s):
            if s.get("dims_estimated"):
                if s.get("max_row") is None or s.get("max_col") is None:
                    return "大きさ不明（時間内に数えられず）"
                return f"約{s.get('max_row')}行×{s.get('max_col')}列（記録値・推定）"
            return f"{s.get('max_row', 0)}行×{s.get('max_col', 0)}列"
        text = "\n".join(f"{s.get('name')}: {_dims(s)}"
                         for s in sheets if isinstance(s, dict))
        return (text or None), "sheets"
    if name == "xlsx_range":
        rows = result.get("rows") or []
        if not rows:
            return None, None
        sheet = result.get("sheet") or ""
        rng = result.get("range") or ""
        locator = f"{sheet}!{rng}" if sheet else (rng or "range")
        start_row = _row_start_from_a1_range(rng)
        lines = []
        for i, row in enumerate(rows):
            no = start_row + i if start_row is not None else i
            cells = row if isinstance(row, list) else []
            lines.append(f"{no}: " + "\t".join(str(c) for c in cells))
        return "\n".join(lines), locator
    if name == "docx_paragraphs":
        # 表（`tables`）も本文合成の対象にする（表しか無い docx で read_evidence が空にならないように）
        paras = [p for p in (result.get("paragraphs") or []) if isinstance(p, dict)]
        tables = [t for t in (result.get("tables") or []) if isinstance(t, dict)]
        if not paras and not tables:
            return None, None
        lines = [f"段落{p.get('i')}: {p.get('text', '')}" for p in paras]
        for t in tables:
            ti = t.get("i")
            row_start = t.get("row_start") or 0
            for ri, row in enumerate(t.get("rows") or []):
                cells = row if isinstance(row, list) else []
                lines.append(f"表{ti}行{row_start + ri}: " + "\t".join(str(c) for c in cells))
        for h in result.get("headers_footers") or []:
            lines.append(f"{'ヘッダー' if 'header' in str(h.get('kind')) else 'フッター'}: {h.get('text', '')}")
        lines += [f"テキストボックス: {t}" for t in result.get("textboxes") or []]
        lines += [f"脚注{n.get('id')}: {n.get('text', '')}" for n in result.get("footnotes") or []]
        lines += [f"文末脚注{n.get('id')}: {n.get('text', '')}" for n in result.get("endnotes") or []]
        text = "\n".join(lines)
        # locator は表の行範囲も含めて「実際に返した範囲」を表す（`paragraphs[s-e];tables[ts-te]rows[rs-re]`）。`InvestigationState._find` が locator 文字列で同一性を判定するため
        ids = [p.get("i") for p in paras]
        parts = []
        if ids:
            parts.append(f"paragraphs[{ids[0]}-{ids[-1]}]")
        if tables:
            t_ids = [t.get("i") for t in tables if isinstance(t.get("i"), int)]
            row_ranges = [(t.get("row_start"), len(t.get("rows") or [])) for t in tables]
            row_lo = [rs for rs, n in row_ranges if isinstance(rs, int) and n > 0]
            row_hi = [rs + n - 1 for rs, n in row_ranges if isinstance(rs, int) and n > 0]
            if t_ids:
                tables_part = f"tables[{min(t_ids)}-{max(t_ids)}]"
                if row_lo and row_hi:
                    tables_part += f"rows[{min(row_lo)}-{max(row_hi)}]"
                parts.append(tables_part)
        locator = ";".join(parts) if parts else "paragraphs"
        return (text or None), locator
    if name == "pptx_slides":
        # 表・ノートも本文合成の対象にする
        slides = [s for s in (result.get("slides") or []) if isinstance(s, dict)]
        if not slides:
            return None, None
        lines = []
        for s in slides:
            no = s.get("no")
            parts = [f"スライド{no}: " + " / ".join(s.get('texts') or [])]
            for ti, table in enumerate(s.get("tables") or []):
                for ri, row in enumerate(table if isinstance(table, list) else []):
                    cells = row if isinstance(row, list) else []
                    parts.append(f"スライド{no}表{ti}行{ri}: " + "\t".join(str(c) for c in cells))
            notes = s.get("notes")
            if notes:
                parts.append(f"スライド{no}ノート: {notes}")
            lines.append("\n".join(parts))
        text = "\n".join(lines)
        nos = [s.get("no") for s in slides]
        locator = f"slides[{','.join(str(n) for n in nos)}]" if nos else "slides"
        return text, locator
    if name == "pdf_pages":
        pages = [p for p in (result.get("pages") or []) if isinstance(p, dict)]
        if not pages:
            return None, None
        text = "\n".join(f"ページ{p.get('no')}: {p.get('text') or p.get('note') or ''}" for p in pages)
        nos = [p.get("no") for p in pages]
        locator = f"pages[{','.join(str(n) for n in nos)}]" if nos else "pages"
        return text, locator
    if name == "file_head":
        text = result.get("text")
        if not text:
            return None, None
        return text, "head"
    return None, None

def _shrink_xlsx_range_field(orig_range: str | None, n_rows: int) -> str | None:
    """`orig_range`（`xlsx_range` が返した A1 レンジ）を、行が `n_rows` 行へ削減された場合の範囲に更新する（列は不変・終了行だけ詰める）。解析できなければ元の値のまま返す。"""
    if not orig_range or n_rows <= 0:
        return orig_range
    try:
        from openpyxl.utils.cell import range_boundaries
        from openpyxl.utils import get_column_letter
        min_col, min_row, max_col, _max_row = range_boundaries(orig_range)
        return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{min_row + n_rows - 1}"
    except Exception:
        return orig_range

# 二分探索で 1 件も残せない場合に「先頭 1 件の text を切り詰めて残す」対象のツール（dict 要素が str の `text` を持つ形）。それ以外は 1 件も入らなければ空のまま返す
_SINGLE_ITEM_TEXT_FIELDS = frozenset({"pdf_pages"})

def _shrink_single_item_result(name: str, result: dict, doc_id: str, field: str,
                               item, tr_max_bytes: int) -> dict | None:
    """1 件も残せないとき、先頭 1 件だけを予算内へ切り詰めて `text_truncated: true` を立てて残す（番号と locator は保つ）。"""
    if not isinstance(item, dict):
        return None
    text0 = item.get("text")
    if not isinstance(text0, str) or not text0:
        return None

    def _build_one(txt: str) -> dict:
        it = {**item, "text": txt, "text_truncated": True}
        r = dict(result)
        r[field] = [it]
        # 鍵ブロックの伏せ字は呼び出し元（`run_tool` の `_redact_deep`）が適用済みなので、ここでは掛け直さない
        text, locator = _doc_reader_text_locator(name, r)
        if text:
            r = {**r, "doc_id": doc_id, "text": text, "locator": locator}
        return r

    lo, hi, best_text = 0, len(text0), ""
    while lo <= hi:
        mid = (lo + hi) // 2
        if _result_byte_size(_build_one(text0[:mid])) <= tr_max_bytes:
            best_text = text0[:mid]
            lo = mid + 1
        else:
            hi = mid - 1
    return _build_one(best_text)

# バイト予算で切ったことを申告する印 `"byte_clipped": true` の JSON 上の増分。切り詰め側は先にこの分を差し引く
_BYTE_CLIP_MARK_BYTES = len(', "byte_clipped": true')

_DOCX_EXTRA_LIST_KEYS = ("headers_footers", "textboxes", "footnotes", "endnotes")

def _fit_docx_extras(result: dict, budget: int) -> dict:
    """`docx_paragraphs` の追加欄（ヘッダー・フッター・テキストボックス・脚注）を `budget` バイトに収める。末尾の項目から落とし、落とした件数を `extras_clipped` に足して `truncated` を立てる。"""
    keys = [k for k in _DOCX_EXTRA_LIST_KEYS if result.get(k)]
    if not keys or _result_byte_size({k: result[k] for k in keys}) <= budget:
        return result
    r = dict(result)
    cut = 0
    while keys and _result_byte_size({k: r[k] for k in keys}) > budget:
        k = max(keys, key=lambda x: len(r[x]))
        r[k] = r[k][:-1]
        cut += 1
        if not r[k]:
            del r[k]
            keys.remove(k)
    r["extras_clipped"] = int(r.get("extras_clipped") or 0) + cut
    r["truncated"] = True
    return r

def _finish_docx_paragraphs_result(result: dict, doc_id: str, tr_max_bytes: int) -> dict:
    """`docx_paragraphs` 専用の仕上げ。段落だけでなく表（表→行）もバイト予算の削減対象にする。
    予算超過時は段落を先に確保（表 0 行で段落数を二分探索）し、余った予算で表の行を先頭の表から順に埋める（途中で打ち切ると `row_truncated`）。
    段落が 1 件も入らない場合は、まず表を `best_n_rows` に固定したまま先頭 1 段落を切り詰めて試す。それでも段落が空文字になる・予算を超える場合だけ、表を 0 行にして段落を切り詰め直し、残り予算で表の行数を改めて決める。
    """
    result = _fit_docx_extras(result, tr_max_bytes // 2)
    paras = [p for p in (result.get("paragraphs") or []) if isinstance(p, dict)]
    tables_orig = [t for t in (result.get("tables") or []) if isinstance(t, dict)]
    flat_rows: list[tuple[int, object]] = []
    for ti, t in enumerate(tables_orig):
        for row in (t.get("rows") or []):
            flat_rows.append((ti, row))

    def _tables_for(n_rows_keep: int) -> list[dict]:
        if n_rows_keep <= 0:
            return []
        counts: dict[int, int] = {}
        for ti, _row in flat_rows[:n_rows_keep]:
            counts[ti] = counts.get(ti, 0) + 1
        out = []
        for ti, t in enumerate(tables_orig):
            n = counts.get(ti, 0)
            if n <= 0:
                continue
            orig_rows = t.get("rows") or []
            nt = {**t, "rows": orig_rows[:n]}
            if n < len(orig_rows):
                nt["row_truncated"] = True
            out.append(nt)
        return out

    def _build(n_paras: int, n_rows_keep: int) -> dict:
        r = dict(result)
        r["paragraphs"] = paras[:n_paras]
        r["tables"] = _tables_for(n_rows_keep)
        # 合成元は呼び出し元（`run_tool`）の `_redact_deep` を通過済みなので、合成後に伏せ字を掛け直さない
        text, locator = _doc_reader_text_locator("docx_paragraphs", r)
        if text:
            r = {**r, "doc_id": doc_id, "text": text, "locator": locator}
        return r

    full = _build(len(paras), len(flat_rows))
    if _result_byte_size(full) <= tr_max_bytes:
        return full
    tr_max_bytes = max(1, tr_max_bytes - _BYTE_CLIP_MARK_BYTES)  # 計測用の印の分を予算から先に引く

    # 段落を先に確保する（表 0 行で段落数を二分探索）→ 余った予算で表の行を埋める（表を満量のまま段落を削ると、大きな表を持つ文書で段落が 1 件も返らなくなる）
    lo, hi, best_n_paras = 0, len(paras), 0
    while lo <= hi:
        mid = (lo + hi) // 2
        if _result_byte_size(_build(mid, 0)) <= tr_max_bytes:
            best_n_paras = mid
            lo = mid + 1
        else:
            hi = mid - 1
    lo, hi, best_n_rows = 0, len(flat_rows), 0
    while lo <= hi:
        mid = (lo + hi) // 2
        if _result_byte_size(_build(best_n_paras, mid)) <= tr_max_bytes:
            best_n_rows = mid
            lo = mid + 1
        else:
            hi = mid - 1
    if best_n_paras > 0:
        r = _build(best_n_paras, best_n_rows)
        r["truncated"] = True
        r["byte_clipped"] = True  # バイト予算由来（計測用の印）
        return r

    # 段落が 1 件も入らない場合は、小さな表が入っていても早期に返さず、下の段落の救済へ進む（表しか無い docx は対象外）
    if paras:
        # まず表を `best_n_rows` に固定したまま段落を救済する。それが「段落が空文字」または「予算超過」に終わる場合だけ、表を 0 行にして救済し直す
        fixed_tables_base = dict(result)
        fixed_tables_base["tables"] = _tables_for(best_n_rows)
        shrunk = _shrink_single_item_result("docx_paragraphs", fixed_tables_base, doc_id,
                                            "paragraphs", paras[0], tr_max_bytes)
        shrunk_ok = (shrunk is not None and (shrunk.get("paragraphs") or [{}])[0].get("text")
                    and _result_byte_size(shrunk) <= tr_max_bytes)
        if not shrunk_ok:
            zero_tables_base = dict(result)
            zero_tables_base["tables"] = []
            shrunk = _shrink_single_item_result("docx_paragraphs", zero_tables_base, doc_id,
                                                "paragraphs", paras[0], tr_max_bytes)
        if shrunk is not None:
            # 救済した段落を固定した上で、表の行数を残り予算に合わせて二分探索し直す
            def _with_tables(n_rows_keep: int) -> dict:
                r = dict(shrunk)
                r["tables"] = _tables_for(n_rows_keep)
                text, locator = _doc_reader_text_locator("docx_paragraphs", r)
                if text:
                    r = {**r, "doc_id": doc_id, "text": text, "locator": locator}
                return r

            lo, hi, rescued_n_rows = 0, len(flat_rows), 0
            while lo <= hi:
                mid = (lo + hi) // 2
                if _result_byte_size(_with_tables(mid)) <= tr_max_bytes:
                    rescued_n_rows = mid
                    lo = mid + 1
                else:
                    hi = mid - 1
            final = _with_tables(rescued_n_rows)
            final["truncated"] = True
            final["byte_clipped"] = True  # バイト予算由来（件数上限と区別する計測用の印）
            return final
    r = _build(0, best_n_rows)
    r["truncated"] = True
    r["byte_clipped"] = True
    return r

def _finish_reader_result(name: str, result: dict, doc_id: str, tr_max_bytes: int) -> dict:
    """原本読取ツール 6 本共通の仕上げ（`_redact_deep` の後に呼ぶ）。`doc_id`／`text`／`locator` の合成とバイト上限クリップを同時に行う（`text` は `result[field]` から毎回作り直すので、クリップ後も構造と食い違わない）。
    二分探索は `doc_id`／`text`／`locator` を含めた最終形の JSON バイト数で判定する。エラー結果はそのまま。`docx_paragraphs` は表も削減対象にするため `_finish_docx_paragraphs_result` に委譲する。
    """
    if not (isinstance(result, dict) and not result.get("error")):
        return result
    if name == "docx_paragraphs":
        return _finish_docx_paragraphs_result(result, doc_id, tr_max_bytes)
    field = _READER_CLIP_FIELD.get(name)

    def _build(seq):
        r = dict(result)
        if field is not None and seq is not None:
            r[field] = seq
            if name == "xlsx_range":
                # 行を削った分だけ `range`（延いては locator）も実際の範囲へ合わせる
                r["range"] = _shrink_xlsx_range_field(result.get("range"), len(seq))
        # 合成元は `_redact_deep` を通過済みなので、合成後に伏せ字を掛け直さない
        text, locator = _doc_reader_text_locator(name, r)
        if text:
            r = {**r, "doc_id": doc_id, "text": text, "locator": locator}
        return r

    full_seq = result.get(field) if field is not None else None
    full = _build(full_seq)
    if field is None or not isinstance(full_seq, (list, str)) or _result_byte_size(full) <= tr_max_bytes:
        return full
    tr_max_bytes = max(1, tr_max_bytes - _BYTE_CLIP_MARK_BYTES)  # 計測用の印の分を予算から先に引く（印を足しても予算内）
    lo, hi, best = 0, len(full_seq), _build(full_seq[:0])
    while lo <= hi:
        mid = (lo + hi) // 2
        cand = _build(full_seq[:mid])
        if _result_byte_size(cand) <= tr_max_bytes:
            best = cand
            lo = mid + 1
        else:
            hi = mid - 1
    # 1 件も残せない場合は、対応済みのツール（`_SINGLE_ITEM_TEXT_FIELDS`）に限り先頭 1 件の text を予算内へ切り詰めて残す
    if (isinstance(full_seq, list) and full_seq and name in _SINGLE_ITEM_TEXT_FIELDS
            and not (best.get(field) if field is not None else None)):
        shrunk = _shrink_single_item_result(name, result, doc_id, field, full_seq[0], tr_max_bytes)
        if shrunk is not None:
            shrunk["truncated"] = True
            shrunk["byte_clipped"] = True  # バイト予算由来（件数上限と区別する計測用の印）
            return shrunk
    best["truncated"] = True
    best["byte_clipped"] = True
    return best

# グラフのスキーマ世代不一致（`GraphSchemaEraError`）のツール結果コード。API 経路（`run_tool` の `graph_neighbors`）と MCP 経路が共有する閉じたコード
GRAPH_REINGEST_ERROR_CODE = "graph_reingest_required"

# read_doc／read_around 自身の引数検証エラー（整数変換失敗・行範囲が総行数超過等）の固定コード。`_open_doc_stream` の「読めない」とは別枠で、coverage（項目ごとの未確認）には記録しない
_READ_INVALID_ARGS_ERROR_CODE = "read_invalid_args"


# ---- ツール実行（読み取り専用・資料フォルダ＋範囲に限定）----

def run_tool(name: str, args: dict, world: str, scope_paths,
            deadline: float | None = None, layer=None,
            max_hits: int | None = None, window_cap: int | None = None,
            tool_result_max_bytes: int | None = None,
            graph_only: bool = False) -> tuple[dict, set, list, list]:
    """ツールを実行し `(結果, 触れた doc_id 集合, 引用候補, 候補カード)` を返す。範囲外・未解決・秘匿は安全に error にする。tool result の本文は秘密を伏せて返す。
    引用候補＝`{doc_id, span, quote, ext}`（grep／ES ヒット由来）。候補カード＝`graph_neighbors` 由来の原因候補（troubleshoot の UI／エクスポート用）。
    - `layer`（省略可・`"docs"|"code"|"both"`・既定 both）: `scope_paths` と同じ会話ターン全体の硬いフィルタ。`ripgrep_search`／`glob_search`／`es_search`／`list_docs` の対象を絞り、`read_around`／`read_doc`／`doc_outline` は層外の doc_id を範囲外と同じく拒否する。`graph_neighbors` は層で結果を絞れない（言及エッジが木を跨ぐ）ため、層が限定されている間はツール自体を拒否する（層外の名前・経路が漏れる迂回路になるため）。
    - `max_hits`（省略可）: 調べる深さが計算した grep／ES のヒット上限の実効値（`grep_search` の `max_hits`／`es_index.search` の `k` へ渡す）。
    - `window_cap`（省略可）: `read_around` の読み取り窓の既定値と `max(200, READ_WINDOW)` クランプの `READ_WINDOW` 部分を置き換える（200 の下限は維持）。`read_doc` の 1 ページ幅にも `max(200, window_cap or READ_WINDOW)` を使う。read_doc は 1 行ずつバイト予算を累積し、超える直前の行で止めてそこを `end_line` にする（超過時は `text_truncated: true`）。`doc_outline` は見出し件数（`_OUTLINE_MAX_HEADINGS`）とタイトルの累積バイト数の両方で打ち切る（`truncated`）。読み込みが `_READ_AROUND_FILE_CAP_BYTES` に達したら `file_truncated: true` を返す。`ripgrep_search` も `_GREP_FILE_CAP_BYTES` で打ち切ったファイル由来のヒットに `hits[i].file_truncated: true` を付ける。
    - `ripgrep_search` のページング: `args["offset"]`（省略時 0・負値は 0）を `grep_tool.grep_search` へそのまま渡す（巻き戻さない）。1 ページのヒット数 `used_max_hits` は `max_hits` と `MAX_HITS_ABS_MAX - offset` の小さい方（`offset` はクランプせず、件数を縮める側で絶対上限を守る）。`offset >= MAX_HITS_ABS_MAX` なら検索せず空を返す。`next_offset` は `offset+used_max_hits < MAX_HITS_ABS_MAX` のときだけ付く。`es_search` は offset ページングを持たず（続きが要るなら `max_hits` を増やす）、`truncated:true` のみ。
    - ヒットの本文: 両ツールとも `hits[i].text` は `tool_result_max_bytes` をヒット数で均等割りしたバイト数（下限 `_HIT_TEXT_MIN_BYTES`）で末尾クリップし、切ったヒットに `text_truncated: true` を付ける（親返しの rag doc 群は別予算配分で、`_resolve_parent_return` が独自に付ける）。`view` を組んだ後に直列化後の実バイト数を測り、`tr_max_bytes` 超過なら per_hit を詰めて再構築する（最大 4 回）。それでも収まらなければ呼び出し元の外側クリップに委ねる。
    - `deadline`（省略可・`time.monotonic()` 系の絶対期限）: ツリー走査を伴うツール（`ripgrep_search`／`list_docs`／`glob_search`／`es_search`）へそのまま渡し、実行中の呼び出しを打ち切る。超過時は `grep_tool.GrepDeadlineExceeded`／`scope_infer.ScopeWalkDeadlineExceeded` を送出する（呼び出し元が 504 相当へ再分類する）。
    - `tool_result_max_bytes`（省略可）: run 開始時に 1 回解決したツール結果 1 件あたりのバイト予算。本関数内のバイトクリップは全てこの実効値を使う。
    - `graph_only`（省略可）: 真のとき `graph_neighbors` は `lens_service.neighbor_cards_graph_only`（grep をせず、起点は名前の一致で直接引く）を使う（MCP が `SHERPA_MCP_TOOLSET=plain` のときだけ）。
    設計: docs/design/codex.md「MCP の道具」
    """
    sp = scope_mod.normalize_scope_paths(scope_paths) or None
    args = args or {}
    docs: set = set()
    cites: list = []
    cards: list = []
    # run 単位で snapshot 済みの値（無ければコード既定）
    tr_max_bytes = tool_result_max_bytes if tool_result_max_bytes is not None else TOOL_RESULT_MAX_BYTES
    if name == "list_docs":
        from ... import doc_ledger  # 台帳＝資料フォルダのフォルダ木を走査（常に live）
        prefix = str(args.get("path_prefix") or "").strip().strip("/")
        pattern = str(args.get("name_pattern") or "").strip().lower()
        doctype_filter = str(args.get("doctype") or "").strip().lower()
        state_filter = str(args.get("state") or "").strip().lower()
        try:
            limit = int(args.get("limit") or 200)
        except (TypeError, ValueError):
            limit = 200
        limit = max(1, min(limit, 500))  # `limit` は 1〜500 にクランプ
        try:
            offset = int(args.get("offset") or 0)
        except (TypeError, ValueError):
            offset = 0
        offset = max(0, offset)
        # 層判定は `doc_ledger`（`classify_document` 確定済み）の `branch=="source"` を使う（grep／ES と同じ確定判定）
        rows = [r for r in doc_ledger.documents_for(world, deadline=deadline)
               if scope_mod.in_scope(r["name"], sp)
               and layer_mod.in_layer_code(r.get("branch") == "source", layer)]
        if prefix:
            rows = [r for r in rows if scope_mod.in_scope(r["name"], [prefix])]  # 同じ prefix 一致ロジックを再利用
        if pattern:
            rows = [r for r in rows if pattern in r["name"].lower()]
        if doctype_filter:
            rows = [r for r in rows if str(r.get("doctype") or "").lower() == doctype_filter]
        # state（"ready"／"unreadable"／"unknown"）も通す（読み取り不可な文書を「使える」文書と同列に見せない）
        if state_filter:
            rows = [r for r in rows if str(r.get("state", "ready")).lower() == state_filter]
        rows.sort(key=lambda r: r["name"])  # rel_path 昇順固定（同条件なら offset が安定する）
        page = rows[offset:offset + limit]
        out = [{"rel_path": r["name"], "doctype": r.get("doctype"), "state": r.get("state", "ready")}
              for r in page]
        for d in out:
            docs.add(d["rel_path"])  # 一覧に出した分だけ出典（sources）に載せる
        shown_end = offset + len(out)
        truncated = shown_end < len(rows)
        return ({"count": len(rows), "offset": offset, "docs": out, "truncated": truncated,
                "next_offset": shown_end if truncated else None}, docs, cites, cards)
    if name == "folder_tree":
        # フォルダは文書ではない（`docs`＝doc_id 集合には何も足さない・出典／引用の対象外）
        from ... import folder_tree as folder_tree_mod
        result = folder_tree_mod.build(world, args, scope_paths=sp, deadline=deadline, layer=layer,
                                       tool_result_max_bytes=tr_max_bytes)
        return (result, docs, cites, cards)
    if name == "glob_search":
        from ... import doc_ledger  # list_docs と同じ台帳走査（常に live）
        pattern = _validate_glob_pattern(args.get("pattern"))
        if pattern is None:
            return ({"error": "pattern が不正です（絶対パス・`..`・NUL・空文字・長すぎるパターンは使えません）"},
                    docs, cites, cards)
        match_pattern = _glob_match_pattern(pattern).lower()
        # 範囲・層は list_docs と同じ確定判定（`branch=="source"`）を使う
        matched = [r["name"] for r in doc_ledger.documents_for(world, deadline=deadline)
                  if scope_mod.in_scope(r["name"], sp)
                  and layer_mod.in_layer_code(r.get("branch") == "source", layer)
                  and importance._match_segment_glob(match_pattern, r["name"].lower())]
        shown = matched[:_GLOB_MAX_RESULTS]
        for p in shown:
            docs.add(p)  # 出した分だけ出典に載せる
        return ({"count": len(matched), "paths": shown, "truncated": len(matched) > len(shown)},
                docs, cites, cards)
    if name in ("ripgrep_search", "es_search"):
        q = str(args.get("query") or "")
        degrade_reason = None
        truncated_docs: list = []  # ripgrep_search のみ（es_search は空のまま）
        grep_stats: dict = {}  # ripgrep_search のみ（読めずに飛ばしたファイルの件数）
        excluded = {"not_current": 0, "withheld": 0}  # es_search のみ（結果から除いたヒットの件数）
        cap_reached = False  # ヒット上限に達した（母集団の一部しか見ていない）
        offset = 0  # ripgrep_search のみ意味を持つ（es_search は常に 0＝ページングしない）
        if name == "es_search":
            # 全文検索（ES hybrid）は offset ページングを持たない（kNN の候補集合がページを跨いで固定できず、順位の入れ替わりで重複と欠落が起きる）。続きが要るなら max_hits を増やす
            used_max_hits = max_hits or MAX_HITS
            from ... import documents  # ES ヒットは現資料フォルダに実在する doc だけ採用
            # 古い ES 索引由来の 404／別内容リンクを引用／出典に出さない。実在集合は 1 回だけ作る（per-hit のツリー走査を避ける）
            valid = documents.world_rel_set(world, deadline=deadline)
            # `es_index.search()` は (hits, degrade_reason) を返す。reason は tool result に生値のまま載せ、各ループが `_degrade_result_node()` で思考ノードへ変換する。
            # `k_ceiling=MAX_HITS_ABS_MAX`（grep と共通の絶対上限）で es_index 側の env 由来の再クランプを迂回する
            # `mode`: hybrid（既定）／keyword（語の一致だけ）／vector（意味の近さだけ）。不正値は hybrid。実際に使った方式を `mode_used` で返す
            mode = str(args.get("mode") or "hybrid").strip().lower()
            if mode not in _ES_MODES:
                mode = "hybrid"
            mode_used = mode
            es_hits, degrade_reason = [], None
            if mode == "vector":
                es_hits, degrade_reason = es_index.search_knn_only(world, q, scope_paths=sp, k=used_max_hits,
                                                                   layer=layer, k_ceiling=MAX_HITS_ABS_MAX)
                if degrade_reason in _ES_VECTOR_FALLBACK_REASONS:
                    # ベクトルが使えない理由のときだけ語の一致（BM25）へ倒し、理由を返す（黙って変えない）。ES クエリ自体の失敗・拒否は倒さず、理由付きの空でそのまま返す
                    reason, mode_used = degrade_reason, "keyword"
                    es_hits, bm25_reason = es_index.search(world, q, scope_paths=sp, k=used_max_hits, layer=layer,
                                                           vector=False, k_ceiling=MAX_HITS_ABS_MAX)
                    degrade_reason = bm25_reason or reason
            elif mode == "keyword":
                es_hits, degrade_reason = es_index.search(world, q, scope_paths=sp, k=used_max_hits, layer=layer,
                                                          vector=False, k_ceiling=MAX_HITS_ABS_MAX)
            else:
                es_hits, degrade_reason = es_index.search(world, q, scope_paths=sp,
                                                          k=used_max_hits, layer=layer,
                                                          k_ceiling=MAX_HITS_ABS_MAX)
                if degrade_reason in _ES_VECTOR_FALLBACK_REASONS:
                    mode_used = "keyword"
            # 上限到達は実在チェック・秘匿名除外の前（生ヒット数）で判定する
            cap_reached = len(es_hits) >= used_max_hits
            # 秘匿名文書（秘匿判定の導入前に索引化されたヒット）はここで一律に弾き、件数（`count`／`docs`）にも含めない（`_safe_doc_path` の秘匿ガードは es_search を通らないため）
            valid_es_hits = []
            hit_ranks = []  # valid_es_hits と同じ並びの、除外前の並びでの順位
            for rank0, h in enumerate(es_hits, 1):
                did = h.get("doc_id")
                if not did or did not in valid:
                    excluded["not_current"] += 1
                    continue
                if text_kind.is_sensitive_doc_id(did):
                    _log.warning("es_search: 秘匿名のため対象外にしました（ext=%s）", Path(did).suffix.lower())
                    excluded["withheld"] += 1
                    continue
                valid_es_hits.append(h)
                hit_ranks.append(rank0)
            hits = [{"doc_id": h["doc_id"], "line": h.get("line"), "text": h.get("text", ""),
                     "span": [h.get("line"), h.get("line")], "ext": h.get("ext"),
                     "score": h.get("score"),  # 親返しの並び順にのみ使う・LLM 出力へは出さない
                     **({"keyword_match": h["keyword_match"]} if h.get("keyword_match") is not None else {}),
                     **({"locator": h["locator"]} if h.get("locator") is not None else {}),
                     **({"chunk_id": h["chunk_id"]} if h.get("chunk_id") is not None else {}),
                     **({"parent_id": h["parent_id"]} if h.get("parent_id") is not None else {})}
                    for h in valid_es_hits]
        else:
            try:
                offset = int(args.get("offset") or 0)
            except (TypeError, ValueError):
                offset = 0
            offset = max(0, offset)
            # `offset` は LLM が渡す未検証値。絶対上限 `MAX_HITS_ABS_MAX` を「offset＋このページの件数」の合計に適用し、`offset` 自体は巻き戻さない。ページの件数 `used_max_hits` を縮めて天井を守り、`offset` が天井以上なら空を返す
            used_max_hits = min((max_hits or MAX_HITS), max(0, MAX_HITS_ABS_MAX - offset))
            if used_max_hits <= 0:
                # 絶対上限（`MAX_HITS_ABS_MAX`）に達していてこれ以上は取れない＝空は「該当なし」ではなく上限到達
                return ({"hits": [], "truncated": True, "limit_reached": True,
                        "max_hits_abs": MAX_HITS_ABS_MAX}, docs, cites, cards)
            # `truncated_docs`: `_GREP_FILE_CAP_BYTES` で打ち切られた文書の doc_id（ヒット 0 件の打切り文書も載る）
            hits = grep_tool.grep_search(q, world, max_hits=used_max_hits, scope_paths=sp,
                                         deadline=deadline, layer=layer, truncated_docs=truncated_docs,
                                         offset=offset, stats=grep_stats)
            hit_ranks = list(range(offset + 1, offset + len(hits) + 1))  # 続きのページは元の検索の順位で数える
        # 親返し（es_search 限定）: rag チャンク由来のヒット（`chunk_id` あり）は doc_id ごとに束ねて `_resolve_parent_return` へ渡し、legacy ヒット（40 行チャンク由来）は素通しする。
        # ① 引用（cites）・doc 収集・rag_groups の組み立て（1 回だけ）
        # ヒットはスコア降順のまま渡ってくる。`template`（出力順のプレースホルダ列）で各 doc の最初に出現したヒットの位置を予約し、legacy ヒットはその場で確定させる（検索結果全体のスコア降順を保つ）。
        # legacy ヒットの `text` はここではクリップしない（② が per_hit ごとに再構築するため、redaction・引用・rag_groups は 1 回で済ませる）
        parent_return_on = name == "es_search"
        template: list = []  # [("legacy", hit_view_dict) | ("rag", doc_id), ...]（出力順）
        rag_groups: dict = {}
        rag_slot_index: dict[str, int] = {}  # doc_id -> template 内の予約位置（代表ヒットの位置）
        template_scores: list = []  # template と同じ並びの点数（呼び出しの記録だけが使う・結果には出さない）
        template_ranks: list = []  # template と同じ並びの元の順位（同上）
        for h, hit_rank in zip(hits, hit_ranks):
            docs.add(h["doc_id"])
            redacted_text = _redact(h["text"])
            quote = redacted_text[:500]  # citation の quote（出典カードの表示用）だけ固定上限・LLM 向け本文は切らない
            # rag_chunks 由来（locator あり）は位置ヒントを LLM への text にだけ添える（citation の quote は hint 抜き）。hint は本文と結合してから redaction を通す
            hint = citations.locator_hint(h.get("locator"))
            text_for_llm = _redact(f"{h['text']}（位置: {hint}）") if hint else redacted_text
            # 引用（cites）は tier に関わらず子チャンク単位のまま（親返しで粒度を落とさない）
            cites.append(citations.from_grep_hit(h, quote=quote, include_match=False))  # match 無し・整形は citations に集約
            if parent_return_on and h.get("chunk_id"):
                rag_item = {
                    "chunk_id": h["chunk_id"], "parent_id": h.get("parent_id"),
                    "locator": h.get("locator"), "score": h.get("score"), "text": text_for_llm,
                }
                if h.get("keyword_match") is not None:
                    rag_item["keyword_match"] = h["keyword_match"]
                rag_groups.setdefault(h["doc_id"], []).append(rag_item)
                if h["doc_id"] not in rag_slot_index:
                    # 最初に出現した位置＝その doc の最高スコア。2 件目以降は予約済みの枠へ集約するだけ
                    rag_slot_index[h["doc_id"]] = len(template)
                    template.append(("rag", h["doc_id"]))
                    template_scores.append(h.get("score"))
                    template_ranks.append(hit_rank)
                continue
            hit_view = {"doc_id": h["doc_id"], "line": h["line"], "text": text_for_llm}
            if h.get("keyword_match") is not None:
                hit_view["keyword_match"] = h["keyword_match"]
            # grep ヒットが持つ登録者重要度（`grep_search` が条件付きで付ける）を LLM 向け tool result にも転送する（重要文書を優先して精読できるように）。理由が無ければキーを作らない
            if h.get("importance"):
                hit_view["importance"] = h["importance"]
                if h.get("importance_reason"):
                    hit_view["importance_reason"] = h["importance_reason"]
            if h.get("file_truncated"):
                # `file_truncated`（`_GREP_FILE_CAP_BYTES` で打ち切られたファイル由来のヒット）を転送する（読む経路と同じ語彙で、cap より後ろが検索できていない可能性を黙らせない）
                hit_view["file_truncated"] = True
            if h.get("encoding_caution"):
                # 一部が化けているヒットの注意文（`encoding_caution`）も転送する
                hit_view["encoding_caution"] = h["encoding_caution"]
            if h.get("section_truncated"):
                hit_view["section_truncated"] = True  # 節・窓の本文が 64KiB の上限で切れている（全文は read_doc／read_around）
            if name == "es_search":
                hit_view["fragment"] = True  # ES のヒット本文は文書の断片
            template.append(("legacy", hit_view))
            template_scores.append(h.get("score"))
            template_ranks.append(hit_rank)

        # ② per_hit を割り当てて `view` を組み立て、直列化後の実バイト数で収まりを保証する
        def _build_view(per_hit: int) -> dict:
            out = []
            any_clipped = False
            legacy_bytes = 0
            for kind, payload in template:
                if kind == "rag":
                    out.append(None)  # 集約結果が確定するまでの予約枠
                    continue
                hv = dict(payload)  # per_hit ごとに作り直す（元の template は不変に保つ）
                clipped_text = _clip_utf8_bytes(hv["text"], per_hit)
                if clipped_text != hv["text"]:
                    # 末尾だけを切る（MD ヒットの見出し行は section 本文の先頭行のため失われない）
                    hv["text"] = clipped_text
                    hv["text_truncated"] = True
                    any_clipped = True
                legacy_bytes += len(hv["text"].encode("utf-8"))
                out.append(hv)
            if rag_groups:
                # legacy ヒット分（クリップ後の実バイト数）を先に差し引いた残りが rag doc 群の予算
                budget_for_rag = max(0, tr_max_bytes - legacy_bytes)
                resolved = _resolve_parent_return(world, rag_groups, sp, layer, budget_for_rag)
                resolved_by_doc = {r["doc_id"]: r for r in resolved}
                for doc_id, idx in rag_slot_index.items():
                    out[idx] = resolved_by_doc[doc_id]  # 予約した代表位置へ集約結果を差し戻す（順位保持）
                    if resolved_by_doc[doc_id].get("text_truncated"):
                        # 親返し側（1 文書あたりの上限）で切られた doc も最上位の切り詰めフラグへ合流させる
                        any_clipped = True
            view = {"hits": out}
            # ヒット上限（used_max_hits）に達した検索は母集団の一部しか返していないため打ち切りの印を返す。
            # `next_offset` は ripgrep_search だけが持ち、`MAX_HITS_ABS_MAX` 未満のときだけ付ける（ちょうど天井に達するページは「ここで終わり」の合図として省く）
            if cap_reached or len(hits) >= used_max_hits:
                view["truncated"] = True
                if name == "ripgrep_search" and offset + used_max_hits < MAX_HITS_ABS_MAX:
                    view["next_offset"] = offset + used_max_hits
            if name == "es_search":
                view["mode_used"] = mode_used
            if degrade_reason:  # es_search のみ・BM25 継続時の縮退理由
                view["degrade_reason"] = degrade_reason
            if truncated_docs:  # ripgrep_search のみ・打切りで探せていない文書
                view["truncated_docs"] = truncated_docs[:_TRUNCATED_DOCS_MAX]
                if len(truncated_docs) > _TRUNCATED_DOCS_MAX:
                    view["truncated_docs_total"] = len(truncated_docs)  # 一覧は先頭だけ・総数
            if grep_stats.get("unreadable_files"):  # ripgrep_search のみ・開けず／読めずに探せていないファイルの件数（名前は出さない）
                view["unreadable_files"] = grep_stats["unreadable_files"]
            if excluded["not_current"] or excluded["withheld"]:  # es_search のみ・結果から除いたヒット
                view["excluded_hits"] = {k: v for k, v in excluded.items() if v}
            if any(isinstance(v, dict) and v.get("section_truncated") for v in out):
                view["section_truncated"] = True
            if any_clipped:
                # ヒット単位のクリップを最上位にも申告する（`tool_result_clipped` 計測は最上位キーしか見ない）
                view["text_truncated"] = True
            return view

        per_hit = max(_HIT_TEXT_MIN_BYTES, tr_max_bytes // max(1, used_max_hits))
        view = _build_view(per_hit)
        # `per_hit` は text だけの割当てで、付帯情報と JSON 構造分を数えていない。直列化後の実バイト数が `tr_max_bytes` を超えたら、`per_hit` を詰めて再構築する（最大 4 回・下限 `_HIT_TEXT_MIN_BYTES`）。
        # 超過のまま外側の `mcp_server._clip_tool_result` に任せると `next_offset` 等の構造が失われるため。それでも収まらない極端なケースだけ外側のクリップ（fail-open）に委ねる
        attempts = 1
        while (per_hit > _HIT_TEXT_MIN_BYTES
              and len(json.dumps(view, ensure_ascii=False).encode("utf-8")) > tr_max_bytes
              and attempts < 4):
            per_hit = max(_HIT_TEXT_MIN_BYTES, int(per_hit * 0.75))
            view = _build_view(per_hit)
            attempts += 1
        if name == "es_search":
            _set_hit_scores(template_scores)
        _set_hit_ranks(template_ranks)
        return (view, docs, cites, cards)
    if name in ("graph_resolve", "graph_impact"):
        # 起点の候補（graph_resolve）→ 識別子からの影響のたどり（graph_impact）。構造の辺だけをたどり、層 code でも使える（資料は返さない）。層 docs は拒否する
        from ... import graph_tools  # 遅延 import（循環回避）
        from ...ingest.world_neo4j import GraphSchemaEraError  # 遅延 import（循環回避）
        if layer == "docs":
            return ({"error": graph_tools.LAYER_REJECT_MESSAGE}, docs, cites, cards)
        try:
            result = graph_tools.run(name, args, world, sp, layer=layer)
        except GraphSchemaEraError as e:
            return ({"error": GRAPH_REINGEST_ERROR_CODE, "world": e.world, "stored_era": e.stored_era},
                    docs, cites, cards)
        fitted = graph_tools.fit_to_bytes(result, tr_max_bytes)
        if fitted is not None:
            result = fitted[0]
        elif graph_tools.json_bytes(result) > tr_max_bytes:
            result = {"error": "tool_result_budget_too_small"}
        return (result, docs, cites, cards)
    if name == "graph_neighbors":
        if layer not in (None, "both"):
            # 層が限定されている間は `graph_neighbors` 自体を拒否する（グラフ traversal は層フィルタ非適用のため、層外の名前・経路・doc_id が漏れる迂回路になる）
            return ({"error": "指定した探す対象（層）では関係グラフの照会は使えません"}, docs, cites, cards)
        from ... import lens_service  # 遅延 import（循環回避）
        from ...ingest.world_neo4j import GraphSchemaEraError  # 遅延 import（循環回避）
        term = str(args.get("name") or "")
        try:
            if not term:
                raw_cards = []
            elif graph_only:
                raw_cards = lens_service.neighbor_cards_graph_only(world, term, sp)
            else:
                raw_cards = lens_service.neighbor_cards(world, term, sp)
        except GraphSchemaEraError as e:
            # 世代不一致は調査を終端させず、MCP 側と同じ機械可読コードのツール結果へ変換して返す（`_record_tool_result_error_code` が `graph_schema_era_mismatch` を立てる）
            return ({"error": GRAPH_REINGEST_ERROR_CODE, "world": e.world, "stored_era": e.stored_era},
                    docs, cites, cards)
        # `neighbor_cards` が内部で捕捉した障害のコード（`NeighborCardsFailure.error_code`・無ければ None）
        _graph_error_code = getattr(raw_cards, "error_code", None)
        # 4 つ目の戻り値（カードのサイドカー）は件数＋直列化バイト上限でクリップする（超過分は捨てる・fail-closed）。Neo4j 側の取得件数上限（`lens_service`）は触らない
        clipped = _clip_cards(raw_cards, max_bytes=tr_max_bytes)
        # カード単位で裏付け doc の実在（資料フォルダ・範囲内）を検証し、裏付け doc を主張したのに 1 件も実在しないカードは cards と LLM への view の両方から除外する。doc を主張しないカード（グラフ位相情報等）はそのまま通す
        cards = []
        unverified_cards = 0  # 裏付け資料を確認できず一覧から除いたカードの数
        for c in clipped:
            claimed_ids = _card_claimed_doc_ids(c)
            if not claimed_ids:
                cards.append(c)
                continue
            verified_ids = _card_verified_doc_ids(c, world, sp)
            if not verified_ids:
                unverified_cards += 1
                continue
            docs |= verified_ids  # 出典付与は検証済み doc のみ（決定的 troubleshoot と同じく edge doc も含める）
            # 検証済み doc_id をカード自身へ同梱する（呼び出し元が Evidence digest を組むときに再検証しない）
            c = {**c, "_verified_doc_ids": sorted(verified_ids)}
            cards.append(c)
        view = [{"name": c["name"], "role": c.get("role", ""), "category": c.get("category", ""),
                 "path": c.get("path", []), "distance": c.get("distance"),
                 "edges": _card_edges_view(c, world, sp)} for c in cards]
        result = {"neighbors": view}
        unverified_edges = sum(1 for v in view for e in v["edges"] if e.get("unverified"))
        if unverified_cards or unverified_edges:
            # 裏付け資料を確認できなかったカード（一覧から除外）・辺（資料名を伏せて `unverified`）の数。名前は出さない
            result["unverified"] = {k: n for k, n in (("cards", unverified_cards), ("edges", unverified_edges)) if n}
        dataitems_excluded = int(getattr(raw_cards, "dataitems_excluded", 0) or 0)
        if dataitems_excluded:
            result["excluded"] = {"data_items": dataitems_excluded}  # 粒度が細かすぎるため一覧に含めない DataItem（影響の調査で見る）
        # 打ち切りの申告（`coverage`）: 取得時の時間切れ・行数の天井・文書探索の打ち切り（`raw_cards.coverage`）に、カード数の上限を足す。
        # 空・部分結果を「近傍なし」「近傍の全部」と区別させる（`complete:false` のとき `limits[].kind` が理由）。
        coverage = getattr(raw_cards, "coverage", None)
        coverage = coverage.copy() if coverage is not None else graph_coverage.Coverage()
        omitted = None
        if len(clipped) < len(raw_cards):
            # 件数／バイト上限で捨てた分がある＝返した近傍は部分集合。打ち切りの事実と総数を返す（モデルが「すべて」と断定しないように）
            partial = coverage.has(graph_coverage.KIND_TIMEOUT, graph_coverage.KIND_ROW_CAP)
            result["truncated"] = True
            # 取得自体が部分結果（時間切れ・行数の天井）なら総数・省略件数は確定できない
            result["count"] = None if partial else len(raw_cards)
            coverage.add(graph_coverage.KIND_CARD_CAP, graph_coverage.STAGE_CARDS)
            omitted = None if partial else len(raw_cards) - len(clipped)
        elif coverage.has(graph_coverage.KIND_TIMEOUT, graph_coverage.KIND_ROW_CAP):
            result["truncated"] = True  # 時間切れの空・天井の部分結果も「一部しか調べていない」＝続きは取れない（総数は不明）
        if _graph_error_code == graph_coverage.KIND_GRAPH_UNAVAILABLE:
            coverage.add(graph_coverage.KIND_GRAPH_UNAVAILABLE)
        if _graph_error_code in (None, graph_coverage.KIND_GRAPH_UNAVAILABLE):
            # `graph_internal_error` は `error_code` だけ（`kind` に対応する語が無い）。`depth`＝近傍たどりの深さの上限（その先に辺が残るかは判定しない＝truncated null）
            result["coverage"] = coverage.as_dict(omitted=omitted, depth=getattr(raw_cards, "depth", None))
            unresolved = getattr(raw_cards, "unresolved", None)
            if unresolved is not None:
                result["unresolved"] = unresolved             # 起点の名前に一致する未解決の参照（保存が無い旧グラフは available:false）
        if _graph_error_code:
            # `neighbor_cards` が捕捉した障害コード（`"graph_unavailable"`／`"graph_internal_error"`）を返す（`run_tool` 境界が `backend_failures["graph"]` へ反映する）
            result["error_code"] = _graph_error_code
        return (result, docs, cites, cards)
    if name == "read_around":
        doc_id = str(args.get("doc_id") or "")
        try:
            line = int(args.get("line") or 1)
            # LLM が window を省略した既定値にも `window_cap` を使う
            window = int(args.get("window") or (window_cap or READ_WINDOW))
            window_requested = window
        except (TypeError, ValueError):
            return ({"error": "line/window は整数で",
                    "error_code": _READ_INVALID_ARGS_ERROR_CODE}, docs, cites, cards)
        # 上限は 200 を後退させず、`READ_WINDOW`／`window_cap` が 200 を超えたときだけ追随する（既定・LLM 指定のどちらの値にも適用する）
        window = max(1, min(window, max(200, window_cap or READ_WINDOW)))
        # 層外は `_safe_doc_path` に `layer` を渡し、`classify_document` 確定後の判定で拒否する。symlink TOCTOU 対策は `_open_doc_stream` に集約済み
        f, err = _open_doc_stream(world, doc_id, sp, layer)
        if err is not None:
            return (err, docs, cites, cards)
        # ストリーミング窓抽出: 目標の終端 `e_target` に達したら打ち切る（ファイル全体を保持せず、総行数でのクランプも不要）
        s = max(0, line - 1 - window)  # 0-based 窓の開始
        e_target = line - 1 + window + 1  # 0-based 窓の終端（排他）
        collected: list[tuple[int, str]] = []
        reached_eof = True
        total_seen = 0
        key_tracker = redact_keys.KeyBlockRedactor(_redact_plain)  # 窓より前の行を辿って鍵ブロックの内側かを確定する
        try:
            reader, it, encoding_caution = _stream_doc_lines(f)
            for idx, t in enumerate(it):
                if idx >= e_target:
                    reached_eof = False
                    break
                total_seen = idx + 1
                if idx >= s:
                    collected.append((idx + 1, t))
                else:
                    key_tracker.track(t)
        finally:
            f.close()
        red_lines = redact_keys.redact_window_lines([t for _, t in collected], _redact_plain, key_tracker.in_key_block)
        text = "\n".join(f"{i}: {r}" for (i, _), r in zip(collected, red_lines))
        # 返却テキストを UTF-8 バイト上限で切り詰め、実際に短くなった時だけ `text_truncated` を明示する（read_doc／doc_outline と同じ語彙）
        read_around_truncated = len(text.encode("utf-8")) > tr_max_bytes
        text = _clip_utf8_bytes(text, tr_max_bytes)
        docs.add(doc_id)
        # `encoding_caution` は `text` より前に置く（後ろだと大きな本文の陰で切り落とされる）
        result: dict = {"doc_id": doc_id}
        if encoding_caution:  # 対象外にしないが符号化の読み取りが不確実な資料の印
            result["encoding_caution"] = encoding_caution
        result["text"] = text
        if read_around_truncated:
            result["text_truncated"] = True
        if reader.truncated or reader.line_overflowed:
            result["file_truncated"] = True  # 読み取りの上限（ファイル 64MiB・単一行 2MiB）で一部を読めていない
            if reader.line_overflowed:
                result["line_overflowed"] = True
        if window_requested > window:
            result["window_clamped"] = {"requested": window_requested, "used": window}
        if reached_eof and not reader.truncated and line > total_seen:
            result["line_beyond_eof"] = True  # 指定した行はファイルの末尾より後ろ
            result["total_lines"] = total_seen
        return (result, docs, cites, cards)
    if name == "read_doc":
        doc_id = str(args.get("doc_id") or "")
        try:
            start = int(args.get("start_line") or 1)
        except (TypeError, ValueError):
            return ({"error": "start_line は整数で",
                    "error_code": _READ_INVALID_ARGS_ERROR_CODE}, docs, cites, cards)
        if start < 1:
            start = 1
        f, err = _open_doc_stream(world, doc_id, sp, layer)
        if err is not None:
            return (err, docs, cites, cards)
        # 1 回のページ幅は read_around と同じ「200 行フロア」（`max(200, window_cap or READ_WINDOW)`）
        page = max(200, window_cap or READ_WINDOW)
        # `total_lines` の申告には全行数のカウントが要るが、行の内容はページ窓（`[start-1, target_end)`）の外なら保持しない
        target_end = start - 1 + page
        window_lines: list[str] = []
        total = 0
        key_tracker = redact_keys.KeyBlockRedactor(_redact_plain)  # 窓より前の行を辿って鍵ブロックの内側かを確定する
        try:
            reader, it, encoding_caution = _stream_doc_lines(f)
            for idx, t in enumerate(it):
                if start - 1 <= idx < target_end:
                    window_lines.append(t)
                elif idx < start - 1:
                    key_tracker.track(t)
                total += 1
        finally:
            f.close()
        window_lines = redact_keys.redact_window_lines(window_lines, _redact_plain, key_tracker.in_key_block)
        file_truncated = reader.truncated or reader.line_overflowed
        if total and start > total:
            return ({"error": f"range 外です（start_line={start}・全{total}行）",
                    "error_code": _READ_INVALID_ARGS_ERROR_CODE}, docs, cites, cards)
        # 1 行ずつバイト予算を累積し、超える直前の行で止めてそこを実際の `end_line` にする（一括クリップすると `end_line` と `text` が食い違う）。1 行目単独で予算を超える場合だけその行をクリップして `text_truncated` で明示する
        out_lines: list = []
        cum_bytes = 0
        actual_end = start - 1
        text_truncated = False
        for offset, wline in enumerate(window_lines):
            i = start - 1 + offset
            ln = f"{i + 1}: {wline}"
            ln_bytes = len(ln.encode("utf-8"))
            sep_bytes = 1 if out_lines else 0  # 結合する "\n" の分
            if cum_bytes + sep_bytes + ln_bytes > tr_max_bytes:
                if not out_lines:
                    out_lines.append(_clip_utf8_bytes(ln, tr_max_bytes))
                    actual_end = i + 1
                text_truncated = True
                break
            out_lines.append(ln)
            cum_bytes += sep_bytes + ln_bytes
            actual_end = i + 1
        docs.add(doc_id)
        result = {"doc_id": doc_id, "start_line": start, "end_line": actual_end,
                 "total_lines": total, "text": "\n".join(out_lines)}
        if text_truncated:
            result["text_truncated"] = True
        if file_truncated:
            result["file_truncated"] = True
        if encoding_caution:  # 対象外にしないが符号化の読み取りが不確実な資料の印
            result["encoding_caution"] = encoding_caution
        return (result, docs, cites, cards)
    if name == "doc_outline":
        doc_id = str(args.get("doc_id") or "")
        f, err = _open_doc_stream(world, doc_id, sp, layer)
        if err is not None:
            return (err, docs, cites, cards)
        all_headings: list = []
        total = 0
        try:
            reader, it, _enc_caution = _stream_doc_lines(f)  # doc_outline は見出しのみ返す＝注意文は付けない
            for idx, t in enumerate(it):
                m = _HEADING_RE.match(t.lstrip())
                if m:
                    full_title = _redact(m.group(2).strip())
                    heading = {"line": idx + 1, "level": len(m.group(1)), "title": full_title[:_OUTLINE_TITLE_MAX_CHARS]}
                    if len(full_title) > _OUTLINE_TITLE_MAX_CHARS:
                        heading["title_truncated"] = True  # 見出しが長く、途中で切って返している
                    all_headings.append(heading)
                total += 1
        finally:
            f.close()
        file_truncated = reader.truncated or reader.line_overflowed
        # 件数上限（`_OUTLINE_MAX_HEADINGS`）に加え、タイトルの累積 UTF-8 バイト数でも打ち切る。`count` は打ち切り前の総見出し数のまま
        headings: list = []
        cum_bytes = 0
        truncated = len(all_headings) > _OUTLINE_MAX_HEADINGS
        for h in all_headings[:_OUTLINE_MAX_HEADINGS]:
            h_bytes = len(h["title"].encode("utf-8"))
            if cum_bytes + h_bytes > tr_max_bytes:
                truncated = True
                break
            headings.append(h)
            cum_bytes += h_bytes
        docs.add(doc_id)
        result = {"doc_id": doc_id, "total_lines": total, "count": len(all_headings),
                 "headings": headings, "truncated": truncated}
        if any(h.get("title_truncated") for h in headings):
            result["titles_truncated"] = True
        if file_truncated:
            result["file_truncated"] = True
        return (result, docs, cites, cards)
    if name == "compare_documents":
        # 実装本体は独立モジュール（`compare_docs.py`）。ここでは範囲・deadline を渡して呼び、① 出典（docs）への反映、② 予算クリップだけを担う
        from ... import compare_docs
        result = compare_docs.compare(world, args, scope_paths=sp, deadline=deadline)
        status = result.get("status")
        if status == "comparable":
            cc = result.get("compare_conditions") or {}
            for side in ("left", "right"):
                doc_id = (cc.get(side) or {}).get("doc_id")
                if doc_id:
                    docs.add(doc_id)
            diff_text = result.get("diff") or ""
            if len(diff_text.encode("utf-8")) > tr_max_bytes:  # 予算超過のときだけ印の分を先に引いて切る
                clipped = _clip_utf8_bytes(diff_text, max(1, tr_max_bytes - _BYTE_CLIP_MARK_BYTES))
                result = {**result, "diff": clipped, "truncated": True, "byte_clipped": True}
        elif status == "unsupported":
            for key in ("left_doc_id", "right_doc_id"):
                doc_id = result.get(key)
                if doc_id:
                    docs.add(doc_id)
        elif status == "needs_disambiguation":
            src = result.get("source_doc_id")
            if src:
                docs.add(src)
        return (result, docs, cites, cards)
    if name in ("xlsx_sheets", "xlsx_range", "docx_paragraphs", "pptx_slides", "pdf_pages"):
        # 原本読取ツール（`doc_readers.py`）: Office／PDF は常に docs 側扱い。探す対象（層）がソースに限定されたターンでは使わせない（file_head は別分岐で `_safe_original_path` の layer 引数を通す）
        if layer == "code":
            return ({"error": "探す対象がソースに限定されています"}, docs, cites, cards)
        doc_id = str(args.get("doc_id") or "")
        kinds = {"xlsx_sheets": _XLSX_KINDS, "xlsx_range": _XLSX_KINDS,
                 "docx_paragraphs": _DOCX_KINDS, "pptx_slides": _PPTX_KINDS,
                 "pdf_pages": _PDF_KINDS}[name]
        resolved = _safe_original_path(world, doc_id, sp, kinds=kinds)
        if resolved is None:
            return ({"error": "doc_id が無効、または読み取り対象外です"}, docs, cites, cards)
        _root, _resolved_doc_id, rp, st = resolved
        f, open_err = _open_verified_original(_root, _resolved_doc_id, st)  # TOCTOU 再検証
        if open_err is not None:
            return (open_err, docs, cites, cards)
        from ... import doc_readers
        if name == "xlsx_sheets":
            result = doc_readers.xlsx_sheets(f)
        elif name == "xlsx_range":
            sheet = str(args.get("sheet") or "")
            if not sheet:
                _close_quiet_local(f)
                return ({"error": "sheet が必要です"}, docs, cites, cards)
            kwargs = {"clean": _redact}
            if args.get("range"):
                kwargs["range_a1"] = str(args["range"])
            if args.get("max_rows") is not None:
                kwargs["max_rows"] = args["max_rows"]
            if args.get("max_cols") is not None:
                kwargs["max_cols"] = args["max_cols"]
            result = doc_readers.xlsx_range(f, sheet, **kwargs)
        elif name == "docx_paragraphs":
            kwargs = {"clean": _redact}
            if args.get("start") is not None:
                kwargs["start"] = args["start"]
            if args.get("count") is not None:
                kwargs["count"] = args["count"]
            if args.get("table_start") is not None:
                kwargs["table_start"] = args["table_start"]
            if args.get("table_row_start") is not None:
                kwargs["table_row_start"] = args["table_row_start"]
            result = doc_readers.docx_paragraphs(f, **kwargs)
        elif name == "pptx_slides":
            kwargs = {"clean": _redact}
            if args.get("pages"):
                kwargs["pages"] = str(args["pages"])
            result = doc_readers.pptx_slides(f, **kwargs)
        else:  # pdf_pages
            kwargs = {"clean": _redact}
            if args.get("pages"):
                kwargs["pages"] = str(args["pages"])
            result = doc_readers.pdf_pages(f, **kwargs)
        result = _redact_deep(result)
        if not (isinstance(result, dict) and result.get("error")):
            docs.add(doc_id)  # 成功時だけ出典（sources）に載せる
            result = _finish_reader_result(name, result, doc_id, tr_max_bytes)
        return (result, docs, cites, cards)
    if name == "file_head":
        doc_id = str(args.get("doc_id") or "")
        resolved = _safe_original_path(world, doc_id, sp, kinds=_FILE_HEAD_KINDS, layer=layer)
        if resolved is None:
            return ({"error": "doc_id が無効、または読み取り対象外です"}, docs, cites, cards)
        _root, _resolved_doc_id, rp, st = resolved
        f, open_err = _open_verified_original(_root, _resolved_doc_id, st)  # TOCTOU 再検証
        if open_err is not None:
            return (open_err, docs, cites, cards)
        from ... import corpus_docs, doc_readers
        kwargs = {"clean": _redact}
        if args.get("max_bytes") is not None:
            kwargs["max_bytes"] = args["max_bytes"]
        result = _redact_deep(doc_readers.file_head(f, **kwargs))
        if not (isinstance(result, dict) and result.get("error")):
            docs.add(doc_id)
            # file_head も原本読取ツールなので、一部が化けている資料を印なしで黙って返さない
            caution = corpus_docs.encoding_caution_for(rp)
            if caution:
                result["encoding_caution"] = caution
            result = _finish_reader_result(name, result, doc_id, tr_max_bytes)
        return (result, docs, cites, cards)
    return ({"error": f"unknown tool: {name}"}, docs, cites, cards)


def verify_doc_exists(doc_id: str, world: str, scope_paths=None) -> bool:
    """doc_id が資料フォルダ内に文書として実在するかを確認する（`sources`＝出典フッターの DL リンク・graph card の裏付け doc を機械検証で絞る用途）。次の 3 つを全て満たす必要がある。
    (1) 実在: `documents.resolve`（`world_graph.resolve_path`＝root 配下への直接解決）。種別を問わず解決するため、これだけでは秘匿ファイルや鍵・内部設定ファイルも通ってしまう。
    (2) 文書種別: `corpus_docs.status_document_reachable(doc_id, world, allow_content_sniff=True)` が `True`（読めると確定）であること（fail-closed。`None`＝判定不能も「読める」扱いにしない）。画像は派生 MD の有無でなくこの分類で許可する（「本文を読めるか」と「文書として実在するか」は別の問い）。`allow_content_sniff=True` は軽量テキスト枠の第 2 段の文書を「存在しない」と誤判定して出典から除外しないため（本関数は 1 doc_id につき 1 回だけ呼ばれる）。
    (3) 範囲: `scope_paths` を渡した場合は `scope_mod.in_scope(doc_id, scope_paths)` も満たすこと（多層防御）。
    """
    if scope_paths is not None and not scope_mod.in_scope(doc_id, scope_paths):
        return False
    try:
        from ... import corpus_docs, documents
        if corpus_docs.status_document_reachable(doc_id, world, allow_content_sniff=True) is not True:
            return False
        return documents.resolve(doc_id, world) is not None
    except Exception:
        return False

def _card_source_ok(doc, world: str, sp, cache: dict) -> bool:
    """辺の根拠が指す資料が、実在・範囲内・非秘匿か（辺の `doc` と同じ検証。同じ資料は 1 度だけ調べる）。"""
    if not doc or text_kind.is_sensitive_doc_id(str(doc)):
        return False
    if doc not in cache:
        cache[doc] = verify_doc_exists(doc, world, sp)
    return cache[doc]


def _card_sources_view(sources: list, world: str, sp, cache: dict, limit: int = 3) -> tuple:
    """辺の根拠（`sources`）を LLM 向けに写す。資料（`doc_id`・`file`・`from_def.file`）が検証を通らない根拠は除く（伏せた件数は数えない）。"""
    from ... import graph_tools  # 遅延 import（循環回避）
    views = (graph_tools.source_view(s, lambda d: _card_source_ok(d, world, sp, cache)) for s in sources or [])
    out = [v for v in views if v is not None]
    # 検証で除いた後に先頭 `limit` 件へ切る（先に切ると有効な根拠が検証で落ちた根拠に押し出される）。切った分は戻り値の 2 つ目
    return out[:limit], max(0, len(out) - limit)


def _card_edges_view(card: dict, world: str | None = None, sp=None) -> list:
    """1 件の `graph_neighbors` card が持つ代表経路の辺（`evidence.edges`）を、LLM 向け `view` 用に既知キーだけ写して返す。`doc` は検証済み集合にある KB 内 rel_path だけ出す。壊れた・古い形の辺があっても落とさず、あるキーだけ拾う。"""
    ev = card.get("evidence", {}) or {}
    verified = card.get("_verified_doc_ids")
    out = []
    cache: dict = {}
    for e in ev.get("edges", []) or []:
        if not isinstance(e, dict):
            continue
        item = {k: e[k] for k in ("type", "from", "to", "doc", "via", "rule") if e.get(k)}
        if "line" in e and e["line"] is not None:
            item["line"] = e["line"]
        if world is not None and e.get("sources"):
            item["sources"], cut = _card_sources_view(e["sources"], world, sp, cache)
            item["sources_overflow_count"] = int(e.get("sources_overflow_count") or 0) + cut
        # 検証済み集合に無い doc を持つ辺は、doc を落として `unverified` を立てる（辺ごと消すと経路が繋がって見えて確定根拠に化ける。実在しない原本は名指しさせない）
        if item.get("doc") and verified is not None and item["doc"] not in set(verified):
            item.pop("doc", None)
            item["unverified"] = True
        if item:
            out.append(item)
    return out

def _card_claimed_doc_ids(card: dict) -> set:
    """1 件の `graph_neighbors` card（troubleshoot 原因候補）が根拠として主張する（未検証の）doc（`evidence.grep[].doc_id`／`evidence.edges[].doc`）の集合を返す。"""
    ev = card.get("evidence", {}) or {}
    doc_ids = {g.get("doc_id") for g in ev.get("grep", []) if g.get("doc_id")}
    doc_ids |= {e.get("doc") for e in ev.get("edges", []) if e.get("doc")}
    return doc_ids

def _card_verified_doc_ids(card: dict, world: str, scope_paths=None) -> set:
    """1 件の card が主張する doc のうち、資料フォルダ内に実在するものの集合を返す（カード単位の検証）。
    グラフは取り込み時点のスナップショットなので、原本が後から消えても card は残りうる。裏付け doc を主張したのに 1 件も実在しない card は無効（呼び出し元の `run_tool` が除外する）。doc を主張しない card は検証の対象外。
    """
    doc_ids = _card_claimed_doc_ids(card)
    return {d for d in doc_ids if verify_doc_exists(d, world, scope_paths)}
