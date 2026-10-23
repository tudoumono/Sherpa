"""ES の構造の欄と `es_search` の `exact` を、実 ES・実際の Excel 取り込みで確かめる結合テスト。

`fixtures/eval` の Excel と、同じ行に並ぶ独立 2 表の合成ブックを取り込み、次を固定する:
識別子の完全一致で当たる record が 1 件、既存の本文検索の結果が欄を足す前後で同じ、
独立 2 表のあいだで値が混ざらない、索引の版を上げた後の最初の更新で入れ直しが 1 回だけ走り埋め込みを再計算しない。
ES に到達できなければ SKIP。
"""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
import pytest

from sherpa import documents, embeddings, tool_dispatch, worlds
from sherpa import es_index as es
from sherpa.ingest import office_md, worker

ROOT = Path(__file__).resolve().parents[2]
EXCEL = ROOT / "fixtures" / "eval" / "excel_ja" / "inputs"
WORLD = "es-struct-check"
_QUERIES = ("表示名", "現行顧客区分", "A01", "廃止", "有効 状態", "項目ID", "受注入力", "画面ID")


def _two_tables(path: Path) -> None:
    """同じ物理行に並ぶ、互いに無関係な 2 つの表（左: 項目ID／表示名／状態・右: 画面ID／画面名）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "一覧"
    left = [("項目ID", "表示名", "状態"), ("B01", "標準顧客区分", "有効"), ("B02", "旧顧客区分", "廃止"),
            ("B03", "予備区分", "検討中")]
    right = [("画面ID", "画面名"), ("S01", "受注入力"), ("S02", "出荷入力"), ("S03", "請求入力")]
    for r, row in enumerate(left, 1):
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
    for r, row in enumerate(right, 1):
        for c, v in enumerate(row, 5):
            ws.cell(r, c, v)
    wb.save(path)


@pytest.fixture
def es_world(tmp_path, monkeypatch):
    if not es.available():
        pytest.skip("ES 未起動")
    from _world_registry import register_test_world
    src = tmp_path / "src"
    src.mkdir()
    for name in ("JPX-016.xlsx", "JPX-001.xlsx"):
        shutil.copy(EXCEL / name, src / name)
    _two_tables(src / "two.xlsx")
    monkeypatch.setenv("SHERPA_DERIVED_DIR", str(tmp_path / "derived"))
    real_world_dir = worlds.world_dir
    monkeypatch.setattr(worlds, "world_dir", lambda w: src if w == WORLD else real_world_dir(w))
    monkeypatch.setattr(es.embeddings, "cfg", lambda settings=None, **kw: None)
    rep = office_md.build_derived(src, worlds.derived_md_dir(WORLD), world_sig="s1")
    assert rep.get("error") is None, rep
    register_test_world(WORLD)
    yield WORLD
    es.delete_world(WORLD)


def _raw_search(world: str, query: dict, size: int = 200) -> list:
    res = es._req("POST", f"/{es._index(world)}/_search", {"size": size, "query": query})
    return [h["_source"] for h in res["hits"]["hits"]]


def _snapshot(world: str) -> dict:
    out = {}
    for q in _QUERIES:
        hits, reason = es.search(world, q, k=20, vector=False)
        assert reason is None
        out[q] = [(h["doc_id"], h.get("chunk_id"), round(h["score"], 4)) for h in hits]
    return out


def test_exact_finds_the_one_record_that_defines_the_id(es_world):
    assert es.index_world(es_world)["chunks"] > 0
    hits, reason = es.search_exact(es_world, ["A01"], k=20)
    assert reason is None
    assert [h["doc_id"] for h in hits] == ["JPX-016.xlsx"]                      # 該当 record は 1 件
    assert "現行顧客区分" in hits[0]["text"] and "A01" in hits[0]["text"]
    # ツール経由（実在確認・親返しを通る）。exact だけ・exact と query の併用（AND）
    res = tool_dispatch.run_tool("es_search", {"exact": "A01"}, es_world, None)[0]
    assert [h["doc_id"] for h in res["hits"]] == ["JPX-016.xlsx"] and res["mode_used"] == "keyword"
    assert "現行顧客区分" in res["hits"][0]["text"]
    both = tool_dispatch.run_tool("es_search", {"exact": "A01", "query": "表示名", "mode": "keyword"}, es_world, None)[0]
    assert [h["doc_id"] for h in both["hits"]] == ["JPX-016.xlsx"]
    none = tool_dispatch.run_tool("es_search", {"exact": "A01", "query": "受注入力", "mode": "keyword"}, es_world, None)[0]
    assert none["hits"] == []                                                    # exact で絞ってから本文で順位付け
    assert es.search_exact(es_world, ["ZZZ99"])[0] == []
    ids = es.search_exact(es_world, ["A01", "B01"], k=20)[0]
    assert {h["doc_id"] for h in ids} == {"JPX-016.xlsx", "two.xlsx"}


def test_labels_and_region_context_are_searchable_fields(es_world):
    es.index_world(es_world)
    labelled = _raw_search(es_world, {"match": {"labels": "顧客"}})
    assert labelled and all(s["doc_id"] in {"JPX-016.xlsx", "two.xlsx"} for s in labelled)
    region = _raw_search(es_world, {"match_phrase": {"region_context": "項目定義"}})
    assert region and all(s["doc_id"] == "JPX-016.xlsx" for s in region)


def test_existing_text_search_results_do_not_change_with_new_fields(es_world, monkeypatch):
    """欄を足す前（構造の欄を ES へ送らない）と後で、本文検索のヒットの集合・順位・点数が同じ。"""
    with monkeypatch.context() as m:
        m.setattr(es, "_chunk_structure_meta", lambda chunk: {})
        es.index_world(es_world)
        before = _snapshot(es_world)
    assert not _raw_search(es_world, {"exists": {"field": "identifier_values"}})
    es.index_world(es_world)
    assert _raw_search(es_world, {"exists": {"field": "identifier_values"}})
    after = _snapshot(es_world)
    assert any(before.values()) and after == before


def test_independent_tables_on_the_same_rows_do_not_mix_values(es_world):
    es.index_world(es_world)
    records = [s for s in _raw_search(es_world, {"term": {"doc_id": "two.xlsx"}}) if s.get("record_identity_values")]
    left = [s for s in records if s["record_identity_values"][0].startswith("B")]
    right = [s for s in records if s["record_identity_values"][0].startswith("S")]
    assert len(left) == 3 and len(right) == 3
    for s in left:
        assert len(s["record_identity_values"]) == 1 and all(v.startswith("B") for v in s["identifier_values"])
        assert not any(v in s.get("labels", []) for v in ("受注入力", "出荷入力", "請求入力"))
    for s in right:
        assert len(s["record_identity_values"]) == 1 and all(v.startswith("S") for v in s["identifier_values"])
        assert not any(v in s.get("labels", []) for v in ("標準顧客区分", "旧顧客区分", "予備区分", "有効"))
    assert {tuple(s["labels"]) for s in left} >= {("標準顧客区分", "有効"), ("旧顧客区分", "廃止")}
    assert {tuple(s["labels"]) for s in right} == {("受注入力",), ("出荷入力",), ("請求入力",)}


def test_mapping_version_bump_reindexes_once_and_reuses_embeddings(es_world, monkeypatch):
    texts_embedded: list = []

    def stub_embed(texts, c, **kw):
        texts_embedded.extend(texts)
        return [[0.1] * 8 for _ in texts]

    monkeypatch.setattr(es.embeddings, "cfg", lambda settings=None, **kw: {"provider": "stub", "key": "x", "model": "stub", "dim": 8})
    monkeypatch.setattr(es.embeddings, "embed", stub_embed)
    calls: list = []
    real_index_world = es.index_world
    monkeypatch.setattr(es, "index_world", lambda *a, **kw: calls.append(1) or real_index_world(*a, **kw))

    def sync() -> bool:
        """取り込みの更新と同じ経路（`needs_reindex` が真のときだけ、ワーカーの `index_world_with_human_md_holdback`）。"""
        if es.needs_reindex(es_world, "sig-1", settings={}):
            return worker.index_world_with_human_md_holdback(es_world, content_sig="sig-1", settings={}).get("error") is None
        return False

    with monkeypatch.context() as m:
        m.setattr(es, "ES_MAPPING_VERSION", "9")                                 # 入れ替え前の版で索引済みの状態
        assert sync() is True
        assert not es.needs_reindex(es_world, "sig-1", settings={})
    first_embeds = len(texts_embedded)
    assert first_embeds > 0 and calls == [1]

    assert es.needs_reindex(es_world, "sig-1", settings={}) is True              # 入れ替え後の最初の更新
    assert sync() is True and calls == [1, 1]
    assert len(texts_embedded) == first_embeds                                   # 本文が同じ＝キャッシュから取る（再計算なし）
    assert es._index_meta(es_world)["mapping_version"] == es.ES_MAPPING_VERSION
    assert es.needs_reindex(es_world, "sig-1", settings={}) is False             # 以降の更新では入れ直さない
    assert sync() is False and calls == [1, 1]
    assert _raw_search(es_world, {"exists": {"field": "identifier_values"}})


def test_exact_with_vector_finds_the_group_whichever_row_holds_the_id(es_world, monkeypatch):
    """表の行のまとまりの先頭・中間・末尾のどの行の ID でも、exact と意味検索（mode=vector／hybrid）の併用でそのまとまりが候補に入り、当たった行の chunk_id が残る。"""
    import hashlib

    def vec_for(t):
        h = hashlib.sha256(t.strip().encode("utf-8")).digest()
        return [(b - 127.5) / 1000.0 for b in h[:16]]

    monkeypatch.setattr(es.embeddings, "cfg", lambda settings=None, **kw: {"provider": "stub", "key": "x", "model": "stub", "dim": 16})
    monkeypatch.setattr(es.embeddings, "embed", lambda texts, c, **kw: [vec_for(t) for t in texts])
    assert es.index_world(es_world, settings={}).get("vectors") is True
    left = [s for s in _raw_search(es_world, {"term": {"doc_id": "two.xlsx"}}) if (s.get("record_identity_values") or [""])[0].startswith("B")]
    heads = [s for s in left if s.get("vector_text")]
    assert len(heads) == 1 and len(left) == 3                                    # 3 行が 1 つのまとまり（ベクトルは先頭だけ）
    head = heads[0]
    assert head["row_group_identifier_values"] == ["B01", "B02", "B03"]
    assert head["record_identity_values"] == ["B01"]                             # 他の行の ID は混ぜない
    assert {s["row_group_head_chunk_id"] for s in left if s is not head} == {head["chunk_id"]}
    row_chunk = {s["record_identity_values"][0]: s["chunk_id"] for s in left}
    for value in ("B01", "B02", "B03"):                                          # 先頭・中間・末尾
        for hits, reason in (es.search_knn_only(es_world, head["vector_text"], exact=[value], settings={}),
                             es.search(es_world, head["vector_text"], exact=[value], settings={})):
            assert reason is None
            found = [h for h in hits if h["chunk_id"] == head["chunk_id"]]
            assert len(found) == 1 and found[0]["row_group_exact_chunk_ids"] == [row_chunk[value]]
