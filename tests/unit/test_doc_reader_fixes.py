"""原本読取ツール（xlsx_sheets／xlsx_range／pdf_pages）の直接試験で見つかった不具合の回帰テスト。
実害の再現に限る（0 行の返却・応答の倍化・先読みの時間・続きが取れない・非表示や表示形式の取り違え）。DB/ES/Neo4j 不要。
"""
from __future__ import annotations

import json
import os
import tempfile
from pathlib import Path

os.environ.setdefault("SHERPA_USE_FIXTURES", "1")

import openpyxl
import pytest

from sherpa import doc_readers as DR
from sherpa import mcp_server
from sherpa.agentic_search import _redact
from sherpa.parts.read import tools as RT
from sherpa.providers.codex import usage as codex_usage

BEGIN = "-----BEGIN RSA PRIVATE KEY-----"


@pytest.fixture()
def tmp() -> Path:
    return Path(tempfile.mkdtemp())


def _open(p: Path):
    return open(p, "rb")


def _size(obj) -> int:
    return len(json.dumps(obj, ensure_ascii=False).encode("utf-8"))


def _pdf(path: Path, pages: list[dict]) -> None:
    """文字・画像・図形を持つ PDF を直接組む（各ページ `{"lines": [...], "image": bool, "rect": bool}`）。"""
    objs: list[bytes] = []

    def add(body: bytes) -> int:
        objs.append(body)
        return len(objs)

    add(b"")  # 1 Catalog（後で埋める）
    add(b"")  # 2 Pages（後で埋める）
    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    img = add(b"<< /Type /XObject /Subtype /Image /Width 1 /Height 1 /ColorSpace /DeviceGray /BitsPerComponent 8 /Length 1 >>\nstream\n\x80\nendstream")
    kids = []
    for pg in pages:
        ops = []
        lines = pg.get("lines") or []
        if lines:
            size, lead = pg.get("font", 8), pg.get("leading", 10)
            ops.append(f"BT /F1 {size} Tf {lead} TL 5 780 Td " + " ".join(f"({ln}) Tj T*" for ln in lines) + " ET")
        if pg.get("image"):
            ops.append("q 50 0 0 50 100 100 cm /Im1 Do Q")
        if pg.get("rect"):
            ops.append("100 100 200 200 re S")
        data = "\n".join(ops).encode("latin-1")
        content = add(b"<< /Length %d >>\nstream\n" % len(data) + data + b"\nendstream")
        page = add(("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents %d 0 R "
                    "/Resources << /Font << /F1 %d 0 R >> /XObject << /Im1 %d 0 R >> >> >>" % (content, font, img)).encode())
        kids.append(page)
    objs[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objs[1] = ("<< /Type /Pages /Kids [%s] /Count %d >>" % (" ".join(f"{k} 0 R" for k in kids), len(kids))).encode()
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    path.write_bytes(bytes(out))


# ===== xlsx_range: 1 行が大きくても 0 行で返さない =====

def _wide_row_xlsx(path: Path, cell_chars: int, cols: int = 50) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "大"
    for r in (1, 2):
        for c in range(1, cols + 1):
            ws.cell(r, c, "x" * cell_chars)
    wb.save(path)


def test_xlsx_range_row_larger_than_budget_returns_leading_columns_not_empty_rows(tmp):
    p = tmp / "wide.xlsx"
    _wide_row_xlsx(p, cell_chars=1200)  # 1 行 60KB 超
    raw = DR.xlsx_range(_open(p), "大")
    out = RT._finish_reader_result("xlsx_range", raw, "wide.xlsx", 30000)
    assert "error" not in out and _size(out) <= 30000
    assert len(out["rows"]) == 1 and 0 < len(out["rows"][0]) < 50
    assert out["cols_clipped"] is True and out["truncated"] is True and out["byte_clipped"] is True
    n = len(out["rows"][0])
    assert out["range"] == f"A1:{openpyxl.utils.get_column_letter(n)}1" and out["locator"] == f"大!{out['range']}"


def test_xlsx_range_single_cell_larger_than_budget_is_an_error_not_empty_rows(tmp):
    p = tmp / "huge.xlsx"
    _wide_row_xlsx(p, cell_chars=32000, cols=2)
    out = RT._finish_reader_result("xlsx_range", DR.xlsx_range(_open(p), "大"), "huge.xlsx", 30000)
    assert out.get("error") and "rows" not in out


# ===== 応答が本文を 2 回持たない・読取ツールの予算 =====

def test_reader_result_has_no_duplicate_text_and_fits_budget(tmp):
    p = tmp / "t.pdf"
    _pdf(p, [{"lines": ["x" * 200] * 150, "font": 1, "leading": 4} for _ in range(3)])
    raw = DR.pdf_pages(_open(p), pages="1-3")
    out = RT._finish_reader_result("pdf_pages", raw, "t.pdf", 12000)
    assert "text" not in out and out["doc_id"] == "t.pdf" and out["locator"].startswith("pages[")
    assert _size(out) <= 12000 and out["truncated"] is True and out["byte_clipped"] is True


def test_mcp_reader_tools_use_the_smaller_reader_budget(monkeypatch):
    monkeypatch.setenv("SHERPA_MCP_TOOL_BUDGET_BYTES", "65536")
    monkeypatch.setenv("SHERPA_MCP_READER_BUDGET_BYTES", "30000")
    assert mcp_server._tool_budget_bytes("pdf_pages") == 30000
    assert mcp_server._tool_budget_bytes("xlsx_range") == 30000
    assert mcp_server._tool_budget_bytes("ripgrep_search") == 65536
    monkeypatch.delenv("SHERPA_MCP_READER_BUDGET_BYTES")
    assert mcp_server._tool_budget_bytes("pdf_pages") == 65536


def test_codex_reader_budget_stays_under_codex_cut():
    env = codex_usage._resolve_mcp_budget({})
    assert int(env["SHERPA_MCP_READER_BUDGET_BYTES"]) <= 30000 < 40000


# ===== pdf_pages =====

def test_pdf_pages_text_offset_continues_a_long_page(tmp):
    p = tmp / "long.pdf"
    _pdf(p, [{"lines": ["abcdefghij" * 20] * 150, "font": 1, "leading": 4}, {"lines": ["second"]}])
    first = DR.pdf_pages(_open(p), pages="1")["pages"][0]
    assert first["text_truncated"] is True and first["next_text_offset"] == len(first["text"]) == 20000
    rest = DR.pdf_pages(_open(p), pages="1", text_offset=first["next_text_offset"])["pages"][0]
    full = first["text"] + rest["text"]
    assert rest["text_offset"] == 20000 and "text_truncated" not in rest
    assert len(full.replace("\n", "")) == 30000
    assert "error" in DR.pdf_pages(_open(p), pages="1-2", text_offset=5)


def test_pdf_pages_tells_scan_from_blank_and_counts_images(tmp):
    p = tmp / "mix.pdf"
    _pdf(p, [{"lines": ["hello"], "image": True}, {"image": True}, {}, {"rect": True}])
    pages = {pg["no"]: pg for pg in DR.pdf_pages(_open(p), pages="1-4")["pages"]}
    assert pages[1]["image_count"] == 1 and "no_text_layer" not in pages[1]
    assert pages[2]["no_text_layer"] and pages[2]["image_count"] == 1 and "画像あり" in pages[2]["note"]
    assert pages[3]["no_text_layer"] and pages[3]["blank"] is True and "白紙" in pages[3]["note"]
    assert pages[4]["no_text_layer"] and "blank" not in pages[4] and "図形" in pages[4]["note"]


def _count_extract(monkeypatch) -> list[int]:
    import pdfplumber.page as plumber_page
    calls: list[int] = []
    orig = plumber_page.Page.extract_text

    def counting(self, *a, **k):
        calls.append(self.page_number)
        return orig(self, *a, **k)

    monkeypatch.setattr(plumber_page.Page, "extract_text", counting)
    return calls


def test_pdf_pages_lookback_is_skipped_when_earlier_pages_have_no_key_marker(tmp, monkeypatch):
    p = tmp / "many.pdf"
    _pdf(p, [{"lines": [f"page {i}"]} for i in range(1, 61)])
    calls = _count_extract(monkeypatch)
    r = DR.pdf_pages(_open(p), pages="55", clean=_redact)
    assert [pg["text"] for pg in r["pages"]] == ["page 55"]
    assert calls == [55]


def test_pdf_pages_lookback_still_tracks_a_key_block_started_earlier(tmp, monkeypatch):
    p = tmp / "key.pdf"
    pages = [{"lines": [f"page {i}"]} for i in range(1, 61)]
    pages[9] = {"lines": ["-----BEGIN RSA PRIVATE KEY-----", "MIIEabc"]}
    _pdf(p, pages)
    calls = _count_extract(monkeypatch)
    r = DR.pdf_pages(_open(p), pages="55", clean=_redact)
    assert r["pages"][0]["text"] == "[REDACTED]"
    assert calls == [10, 55]  # 記号のあるページと選択ページだけ抽出し、記号の無い先読みページは省く


# ===== ページ指定・range =====

def test_partial_page_range_is_not_reported_as_ignored():
    nos, trunc, err, notes = DR._resolve_pages("11-30", 25, "1-10", 10, 20)
    assert nos == list(range(11, 26)) and err is None and "pages_ignored" not in notes
    _nos, _t, _e, notes = DR._resolve_pages("1,30-40", 25, "1-10", 10, 20)
    assert notes["pages_ignored"] == ["30-40"]


def _simple_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for r in range(1, 6):
        for c in range(1, 5):
            ws.cell(r, c, r * 10 + c)
    wb.save(path)


def test_xlsx_range_rejects_zero_based_range_and_normalizes_reversed(tmp):
    p = tmp / "a.xlsx"
    _simple_xlsx(p)
    assert DR.xlsx_range(_open(p), "S", "A0:B2") == {"error": "range が不正です"}
    r = DR.xlsx_range(_open(p), "S", "C3:A1")
    assert r["range"] == "A1:C3" and len(r["rows"]) == 3 and r["rows"][0] == ["11", "12", "13"]


# ===== 非表示・結合・表示形式 =====

def _layout_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "予算"
    ws["A1"] = "見出し"
    ws.merge_cells("A1:C1")
    ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "率", 0.125, 182400, 5
    ws["B2"].number_format = "0.0%"
    ws["C2"].number_format = "#,##0"  # 桁区切りだけ＝表示値は載せない
    ws["D2"].number_format = '"a)"0'
    ws["A3"], ws["B3"] = "日", __import__("datetime").datetime(2019, 7, 1)
    ws["B3"].number_format = "yyyy/m/d"
    ws["A4"], ws["B4"] = "非表示行", 940
    ws.row_dimensions[4].hidden = True
    ws.column_dimensions["E"].hidden = True
    ws["E1"] = "メモ"
    wb.create_sheet("前年度").sheet_state = "hidden"
    wb.save(path)


def test_xlsx_sheets_marks_hidden_sheet(tmp):
    p = tmp / "l.xlsx"
    _layout_xlsx(p)
    sheets = {s["name"]: s for s in DR.xlsx_sheets(_open(p))["sheets"]}
    assert sheets["前年度"]["hidden"] is True and "hidden" not in sheets["予算"]


def test_xlsx_range_reports_hidden_merged_and_display_without_changing_values(tmp):
    p = tmp / "l.xlsx"
    _layout_xlsx(p)
    r = DR.xlsx_range(_open(p), "予算", "A1:E4")
    assert r["hidden_rows"] == [4] and r["hidden_cols"] == ["E"] and r["merged"] == ["A1:C1"]
    assert r["rows"][1][1] == "0.125" and r["rows"][3] == ["非表示行", "940", "", "", ""]  # 値は従来どおり返す
    assert r["display"]["B2"] == "12.5%" and r["display"]["B3"] == "2019/7/1"
    assert "C2" not in r.get("display", {}) and "C2" not in r.get("display_format", {})
    assert "D2" in r["display"] or "D2" in r["display_format"]
    assert "sheet_hidden" not in r
    # 2 回目（キャッシュ命中）でも同じ・読み位置が狂わない
    assert DR.xlsx_range(_open(p), "予算", "A1:E4") == r
    assert DR.xlsx_range(_open(p), "前年度", "A1:A1")["sheet_hidden"] is True


@pytest.mark.parametrize("chunk", [1 << 20, 97, 7])
def test_xlsx_layout_scan_is_the_same_for_any_chunk_size(tmp, monkeypatch, chunk):
    """シート XML を小さな塊で読んでも（タグや行が塊をまたいでも）同じ結果。"""
    monkeypatch.setattr(DR, "_LAYOUT_CHUNK_BYTES", chunk)
    p = tmp / "l.xlsx"
    _layout_xlsx(p)
    r = DR.xlsx_range(_open(p), "予算", "A1:E4")
    assert r["hidden_rows"] == [4] and r["hidden_cols"] == ["E"] and r["merged"] == ["A1:C1"]
    assert r["display"]["B2"] == "12.5%" and r["display"]["B3"] == "2019/7/1"


def test_xlsx_range_row_clip_drops_extras_of_removed_rows():
    result = {"sheet": "S", "range": "A1:B50", "rows": [[f"v{r}" * 100, "x"] for r in range(50)], "truncated": False,
              "hidden_rows": [2, 40], "display": {"B1": "1%", "B45": "2%"}, "merged": ["A1:B1", "A44:B44"]}
    out = RT._finish_reader_result("xlsx_range", result, "a.xlsx", 8192)
    n = len(out["rows"])
    assert 0 < n < 40 and out["hidden_rows"] == [2] and out["display"] == {"B1": "1%"} and out["merged"] == ["A1:B1"]
