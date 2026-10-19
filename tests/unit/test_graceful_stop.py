"""停止要求での取り込みの中断と、起動時の自動再開（`make stop` / `make restart` で中途半端な状態を作らない）。

- office_md の本体ループ: 停止要求が立つと次のファイルへ進まず、処理中の 1 件は変換キャッシュまで保存し、公開しない
- worker: 中断した run は失敗ではなく「中断」として記録し、署名を確定せず、Webhook を送らない
- 起動時の再開: 中断した run の資料フォルダの同期を 1 回だけ始める（無効化 env で始めない・実行中なら何もしない）
"""
from __future__ import annotations

import contextlib
import threading
from pathlib import Path

import openpyxl
import pytest

from sherpa import corpus_docs, es_index, reconcile, store, webhooks, worlds
from sherpa.ingest import background, office_md, resume, world_neo4j, worker


@pytest.fixture(autouse=True)
def _clean_stop_state(monkeypatch):
    monkeypatch.setattr(background, "_stop_event", threading.Event())
    yield
    if hasattr(background._tls, "run_id"):
        del background._tls.run_id


@contextlib.contextmanager
def _as_background_run():
    """背景実行の run のスレッドとして振る舞う（停止要求は背景実行の中でだけ効く）。"""
    background._tls.run_id = 99
    try:
        yield
    finally:
        del background._tls.run_id


def _make_source(tmp_path, monkeypatch, world_id: str, names: list[str]) -> Path:
    source = tmp_path / "kb" / world_id
    source.mkdir(parents=True)
    monkeypatch.setenv("SHERPA_DERIVED_DIR", str(tmp_path / "derived"))
    for name in names:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"], ws["B1"] = "No", "内容"
        ws["A2"], ws["B2"] = 1, f"NEEDLE_{name}"
        wb.save(source / name)
    return source


def _snapshot(world_id: str) -> dict:
    out: dict[str, bytes] = {}
    for layer, d in (("md", worlds.derived_md_dir(world_id)), ("rag", worlds.derived_rag_dir(world_id)),
                     ("ir", worlds.derived_ir_dir(world_id))):
        if d.is_dir():
            for p in sorted(d.rglob("*")):
                if p.is_file():
                    out[f"{layer}/{p.relative_to(d).as_posix()}"] = p.read_bytes()
    return out


def test_stop_request_finishes_current_file_then_stops_without_publishing(monkeypatch, tmp_path):
    names = ["a.xlsx", "b.xlsx", "c.xlsx"]
    source = _make_source(tmp_path, monkeypatch, "stop-w", names)
    dmd = worlds.derived_md_dir("stop-w")
    calls: list[str] = []
    orig = office_md._convert_with_arms

    def counting(path, enabled):
        name = Path(path).name
        calls.append(name)
        if name == "b.xlsx":
            background.request_stop()        # b の処理中に停止要求が立つ
        return orig(path, enabled)

    monkeypatch.setattr(office_md, "_convert_with_arms", counting)
    with _as_background_run(), pytest.raises(background.StopRequested):
        office_md.build_derived(source, dmd, world_sig="s1")

    assert calls == ["a.xlsx", "b.xlsx"]                       # c へは進まない
    cache = office_md._conv_cache_root_for(dmd)
    assert (cache / "a.xlsx.key.json").is_file() and (cache / "b.xlsx.key.json").is_file()   # 処理中だった b もキャッシュまで済む
    assert not (cache / "c.xlsx.key.json").exists()
    assert not dmd.exists()                                      # 公開していない
    assert not list(dmd.parent.glob("*.staging"))                # staging も残さない

    # 再開: 済んだ a・b はキャッシュから取れ、c だけ変換する。通しでやった場合と派生物が同じ
    background._stop_event.clear()
    calls.clear()
    office_md.build_derived(source, dmd, world_sig="s1")
    assert calls == ["c.xlsx"]
    resumed = _snapshot("stop-w")

    ref = _make_source(tmp_path, monkeypatch, "stop-ref", names)
    monkeypatch.setattr(office_md, "_convert_with_arms", orig)
    office_md.build_derived(ref, worlds.derived_md_dir("stop-ref"), world_sig="s1")
    full = {k.replace("stop-ref", "stop-w"): v for k, v in _snapshot("stop-ref").items()}
    assert resumed == full


def test_stop_request_is_ignored_outside_background_runs(monkeypatch, tmp_path):
    source = _make_source(tmp_path, monkeypatch, "stop-direct", ["a.xlsx"])
    background.request_stop()
    rep = office_md.build_derived(source, worlds.derived_md_dir("stop-direct"), world_sig="s1")
    assert rep.get("error") is None and rep["converted"] == 1


@pytest.fixture
def _stub_pipeline(monkeypatch):
    monkeypatch.setattr(worker, "world_state", lambda world, progress=None: ("sig", {"a": [1, 2, 3]}))
    monkeypatch.setattr(worker, "build_world_graph", lambda world: ([], [], []))
    monkeypatch.setattr(worker, "_ledger_rows", lambda world, *, sig: [])
    monkeypatch.setattr(world_neo4j, "_env", lambda: {"uri": "bolt://x", "user": "u", "pw": "p"})
    loaded = []
    monkeypatch.setattr(world_neo4j, "load_world", lambda *a, **kw: loaded.append(1) or (0, 0))
    monkeypatch.setattr(es_index, "index_world", lambda world, content_sig=None, **kw: {"available": True, "chunks": 0})
    monkeypatch.setattr(reconcile, "reconcile_derivatives", lambda reflect=True: None)
    monkeypatch.setattr(corpus_docs, "scan_report", lambda world, expected_rels=None: {"document_count": 0})
    monkeypatch.setattr(store, "set_scan_report", lambda world, report: None)
    monkeypatch.setattr(store, "replace_documents", lambda world, rows: 0)
    monkeypatch.setattr(store, "update_ingest_run_progress", lambda run_id, progress: None)
    monkeypatch.setattr(store, "downgrade_orphaned_extracting_runs", lambda world=None: [])
    monkeypatch.setattr(store, "start_ingest_run", lambda world, **kw: {"id": 1})

    @contextlib.contextmanager
    def _noop_lock(world_id):
        yield
    monkeypatch.setattr(store, "world_lock", _noop_lock)
    rec = {"sigs": [], "finished": [], "confirmed": [], "webhooks": [], "loaded": loaded}
    monkeypatch.setattr(store, "set_world_sig",
                        lambda world, sig, manifest=None, doc_count=None, scan_report=None: rec["sigs"].append(sig))
    monkeypatch.setattr(store, "finish_ingest_run", lambda run_id, **kw: rec["finished"].append(kw) or {"id": run_id})
    monkeypatch.setattr(store, "finish_ingest_run_and_confirm_world",
                        lambda run_id, world, **kw: rec["confirmed"].append(kw) or {"id": run_id})
    monkeypatch.setattr(webhooks, "notify_run_terminal", lambda *a, **kw: rec["webhooks"].append(a))
    return rec


def test_worker_records_interrupt_without_confirming_signature(monkeypatch, _stub_pipeline):
    def stopped(world, **_kw):
        raise background.StopRequested()
    monkeypatch.setattr(worker, "_build_derived", stopped)

    res = worker.run("w", op="refresh")

    assert res["status"] == "failed" and res["interrupted"] is True
    snap = _stub_pipeline["finished"][0]["extraction_snapshot"]
    assert snap["interrupted"] is True and snap["op"] == "refresh"
    assert _stub_pipeline["sigs"] == [""]                 # 署名は pre-invalidate のまま＝次回の同期が続きを走らせる
    assert _stub_pipeline["confirmed"] == [] and _stub_pipeline["loaded"] == []
    assert _stub_pipeline["webhooks"] == []                # 停止による中断は失敗として通知しない


def test_worker_stops_between_stages_before_touching_graph(monkeypatch, _stub_pipeline):
    monkeypatch.setattr(worker, "_build_derived",
                        lambda world, **_kw: background.request_stop() or {"converted": 0, "failed": 0, "unsupported": 0, "by_ext": {}})
    with _as_background_run():
        res = worker.run("w")
    assert res["status"] == "failed" and res["interrupted"] is True
    assert _stub_pipeline["loaded"] == [] and _stub_pipeline["confirmed"] == []


# ---- 起動時の自動再開 ----

def test_resume_starts_one_sync_per_interrupted_world(monkeypatch):
    monkeypatch.setattr(store, "list_worlds_with_interrupted_latest_run", lambda: [
        {"world_id": "w1", "run_id": 5, "op": "refresh"},
        {"world_id": "w2", "run_id": 6, "op": "delete"},      # 削除の続きを同期で走らせない
        {"world_id": "w3", "run_id": 7, "op": None},           # 操作名の記録が無い古い版（0.18.1 相当）の中断 run
        {"world_id": "w4", "run_id": 8, "op": None}])          # 参照先が解決できない資料フォルダは対象外
    monkeypatch.setattr(worlds, "world_dir", lambda wid: "/x" if wid == "w3" else None)
    started = []
    monkeypatch.setattr(background, "start_or_join",
                        lambda wid, op, fp, create, work, **kw: started.append((wid, op)) or (1, False))

    assert resume.resume_interrupted() == ["w1", "w3"]
    assert started == [("w1", "refresh"), ("w3", "refresh")]


def test_resume_disabled_by_env(monkeypatch):
    monkeypatch.setenv("SHERPA_RESUME_INTERRUPTED", "0")
    monkeypatch.setattr(store, "list_worlds_with_interrupted_latest_run",
                        lambda: pytest.fail("無効のときは検索しない"))
    assert resume.resume_interrupted() == []


def test_resume_skips_world_already_running(monkeypatch):
    monkeypatch.setattr(store, "list_worlds_with_interrupted_latest_run",
                        lambda: [{"world_id": "w1", "run_id": 5, "op": "register"}])
    monkeypatch.setattr(background, "start_or_join", lambda *a, **kw: (9, True))   # 実行中の run に合流＝新しく始めていない
    assert resume.resume_interrupted() == []


def test_interrupted_latest_run_query_against_db():
    """強制終了で格下げされた run（操作名つき）が再開の対象に出て、その後に別の run が走れば出ない。登録解除済みは出ない。"""
    wid = "resume-db-test"
    store.upsert_world(wid, f"/tmp/{wid}", label=wid)
    try:
        run = store.start_ingest_run(wid)
        store.note_ingest_run_op(run["id"], "register")
        store.downgrade_orphaned_extracting_runs(wid)
        got = [r for r in store.list_worlds_with_interrupted_latest_run() if r["world_id"] == wid]
        assert got and got[0]["run_id"] == run["id"] and got[0]["op"] == "register"

        run2 = store.start_ingest_run(wid)
        store.finish_ingest_run(run2["id"], status="auto_published")
        assert not [r for r in store.list_worlds_with_interrupted_latest_run() if r["world_id"] == wid]

        run3 = store.start_ingest_run(wid)
        store.downgrade_orphaned_extracting_runs(wid)
        store.delete_world_row(wid)
        assert not [r for r in store.list_worlds_with_interrupted_latest_run() if r["world_id"] == wid]
    finally:
        store.delete_world_row(wid)


def test_background_writes_and_clears_state_file_for_stop_script(monkeypatch, tmp_path):
    """stop.sh が待ち状況を出せるよう、処理中のファイルを RUN_DIR の状態ファイルへ書き、終われば消す。"""
    import os
    import time
    state = tmp_path / "ingest_active"
    monkeypatch.setenv("SHERPA_INGEST_STATE_FILE", str(state))
    monkeypatch.setattr(store, "fail_close_if_extracting", lambda run_id, reason: False)
    monkeypatch.setattr(store, "note_ingest_run_op", lambda run_id, op: None)
    inside = threading.Event()
    release = threading.Event()
    seen = {}

    def work(run_id):
        background.note_file("sub/a.xls")
        seen["text"] = state.read_text(encoding="utf-8")
        inside.set()
        release.wait(timeout=5)

    background.start_or_join("w-state", "refresh", "", lambda: 7, work)
    assert inside.wait(timeout=5)
    lines = seen["text"].splitlines()
    assert lines[0] == f"pid={os.getpid()}"
    kind, world, since, rel = lines[1].split("\t")
    assert (kind, world, rel) == ("run", "w-state", "sub/a.xls") and int(since) > 0
    release.set()
    deadline = time.monotonic() + 5
    while state.exists() and time.monotonic() < deadline:
        time.sleep(0.02)
    assert not state.exists()
