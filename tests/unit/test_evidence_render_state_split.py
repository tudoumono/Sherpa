"""長文の分割で、要素・セルに結び付いた状態・重なり・前面/背面図形の行を全片へ付けることの契約。

設計: docs/design/rag.md「チャンク化の考え方」
"""
from __future__ import annotations

import hashlib
import itertools
import json
import os
import re
import zipfile
from pathlib import Path

os.environ.setdefault("SHERPA_USE_FIXTURES", "1")

import openpyxl
import pytest
from openpyxl.styles import Font

from sherpa import es_index, json_io, worlds
from sherpa.ingest import context_ir as CIR
from sherpa.ingest import evidence_ir as IR
from sherpa.ingest import evidence_render as R
from sherpa.ingest import evidence_spike, excel_display, office_md
from sherpa.parts.read import tools as read_tools

_HASH = "sha256:" + "0" * 64
_STRIKE = "取り消し線: 「あり」"
_NS_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _prose(n: int, tag: str = "記述") -> str:
    return "".join(f"{tag}{i:04d}です。" for i in range(n // 9 + 1))[:n]


# ---- 組み立て ----

def _element_records(value, *, etype="shape", file_type="pptx", source_name="deck.pptx", relations=(),
                     others=(), visibility="visible", **extension):
    element = IR.EvidenceElement(
        element_id="e1", type=etype, parent_id=None, order=0, value=value,
        locator=IR.Locator(part="ppt/slides/slide1.xml", slide=1), coverage_id="c1",
        visibility=visibility, extension=extension)
    ir = IR.EvidenceIR(
        schema_version=IR.EVIDENCE_IR_SCHEMA_VERSION, parser_profile="test",
        source=IR.EvidenceSource(file_type=file_type, content_hash=_HASH),
        elements=[element, *others], relations=list(relations))
    ctx = CIR.ContextIR(schema_version=CIR.CONTEXT_IR_SCHEMA_VERSION, analyzer_version="test",
                        source_hash=_HASH, source_name=source_name, document_titles={})
    return R._non_table_records(ir, ctx, {}, set())


def _spans(records):
    return [r["citations"][0]["text_span"] for r in records]


def _write_xlsx(path: Path, values: list, *, strike=True, number_format=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "明細"
    ws["A1"], ws["B1"] = "項目", "備考"
    for row, value in enumerate(values, start=2):
        ws.cell(row=row, column=1, value=f"旧プラン{row}")
        cell = ws.cell(row=row, column=2, value=value)
        if strike:
            cell.font = Font(strike=True)
        if number_format:
            cell.number_format = number_format
    wb.save(path)


def _render_xlsx(tmp_path, values, *, mutate=None, enrich=False, **kw):
    path = tmp_path / "s.xlsx"
    _write_xlsx(path, values, **kw)
    ir = evidence_spike.extract(path)
    if enrich:
        excel_display.enrich_evidence(ir, path)
    if mutate is not None:
        mutate(ir)
    return R.render(ir, source_name="s.xlsx")


def _anchored(rendered) -> dict[str, str]:
    bodies, reason = es_index._parse_rag_md_chunks(rendered.markdown)
    assert reason is None
    return bodies


def _cell_chunks(rendered, marker="備考（"):
    texts = _anchored(rendered)
    return [(c, texts[c["chunk_id"]]) for c in rendered.chunks
            if c["content_type"] == "table_record" and marker in texts[c["chunk_id"]]]


def _cell_spans(rendered) -> list[tuple[int, int]]:
    out = []
    for chunk, _text in _cell_chunks(rendered):
        out.extend((c["text_span"]["start"], c["text_span"]["end"]) for c in chunk["citations"]
                   if c.get("text_span") and c["locator"].get("cell_range") == "B2")
    return sorted(out)


def _contiguous(spans, total):
    assert spans and spans[0][0] == 0 and spans[-1][1] == total
    assert all(a[1] == b[0] for a, b in itertools.pairwise(spans))


# ---- (1) セルの長文: 全片に状態行 ----

def test_cell_every_span_carries_state_line_in_semantic_and_body(tmp_path):
    text = _prose(3000)
    rendered = _render_xlsx(tmp_path, [text])
    chunks = _cell_chunks(rendered)
    spans = _cell_spans(rendered)
    assert len(spans) >= 3
    _contiguous(spans, len(text))
    assert sum(t.count(_STRIKE) for _c, t in chunks) == len(spans)
    assert sum(c["body"].count(_STRIKE) for c, _t in chunks) == len(spans)
    body_parts = re.findall(r"備考: (.*?)\n" + _STRIKE, "\n".join(c["body"] for c, _t in chunks), re.DOTALL)
    assert "".join(body_parts) == text
    for _c, shown in chunks:
        assert re.search(r"備考（\d+/\d+）:\n.*?\n" + _STRIKE, shown, re.DOTALL)


def test_cell_tail_span_alone_names_target_position_and_state(tmp_path):
    rendered = _render_xlsx(tmp_path, [_prose(3000)])
    chunk, shown = _cell_chunks(rendered)[-1]
    n = len(_cell_spans(rendered))
    assert f"備考（{n}/{n}）:" in shown
    assert "出所: 原本「s.xlsx」 / シート「明細」" in shown and shown.rstrip().endswith(_STRIKE)
    assert chunk["region_context"]["sheet"] == "明細"
    assert chunk["citations"][0]["locator"]["cell_range"] == "B2"
    assert chunk["body"].endswith(_STRIKE)


def test_cell_excel_value_note_stays_on_first_span_only_while_state_goes_to_all():
    cell = IR.EvidenceElement(
        element_id="e1", type="cell", parent_id=None, order=1, value="45900",
        locator=IR.Locator(part="xl/worksheets/sheet1.xml", sheet="明細", extension={"row": 2, "column": 2}),
        coverage_id="c1",
        extension={"format_state": {"strike": "whole"}, "raw_value": "45900", "number_format": "yyyy/mm/dd",
                   "display_status": "rendered", "display_value": "2025/09/10"})
    pieces = [R._field_piece(cell, "期限", i, 3, i, i + 1, "x", []) for i in range(3)]
    assert "Excel原値" in pieces[0]["semantic"] and "Excel原値" in pieces[0]["exact"]
    for piece in pieces[1:]:
        assert "Excel" not in piece["semantic"] and "Excel" not in piece["exact"]
    for piece in pieces:
        assert piece["semantic"].endswith(_STRIKE) and piece["exact"].endswith(_STRIKE)


def test_cell_excel_value_note_first_span_only_end_to_end(tmp_path):
    rendered = _render_xlsx(tmp_path, [_prose(3000)], number_format="@", enrich=True)
    chunks = _cell_chunks(rendered)
    assert [t.count("Excel原値") for _c, t in chunks] == [1] + [0] * (len(chunks) - 1)
    assert sum(t.count(_STRIKE) for _c, t in chunks) == len(_cell_spans(rendered))


def test_cell_overlap_line_keeps_wording_and_goes_to_every_span(tmp_path):
    overlap = "廃止" * 5

    def mutate(ir):
        cell = next(e for e in ir.elements if e.type == "cell" and e.locator.cell_range == "B2")
        cell.extension["occluded_by"] = {"kind": "shape", "text": overlap}

    rendered = _render_xlsx(tmp_path, [_prose(3000)], mutate=mutate)
    chunks = _cell_chunks(rendered)
    tail = f"\n重なり: 「{overlap}」\n{_STRIKE}"
    assert all(shown.rstrip().endswith(tail) for _c, shown in chunks)
    assert all(c["body"].endswith(tail) for c, _t in chunks)
    for _c, shown in chunks:
        assert "すべて" not in shown and "隠れ" not in shown


def _overlap_line(n: int, ch: str = "重") -> str:
    shown = ch * n if n <= 80 else ch * 80 + f"…（全{n}字）"
    return f"重なり: 「{shown}」"


def _cell_with(extension: dict):
    def mutate(ir):
        cell = next(e for e in ir.elements if e.type == "cell" and e.locator.cell_range == "B2")
        cell.extension.update(extension)
    return mutate


def _assert_cell_within_cap(rendered, text, note_lines):
    spans = _cell_spans(rendered)
    _contiguous(spans, len(text))
    note = "".join("\n" + line for line in note_lines)
    assert all(end - start + len(note) <= R.MAX_TEXT_SPAN_CHARS for start, end in spans)
    chunks = _cell_chunks(rendered)
    assert sum(c["body"].count(note) for c, _t in chunks) == len(spans)
    assert sum(t.count(note) for _c, t in chunks) == len(spans)
    for _c, shown in chunks:
        assert len(shown) <= R.MAX_GROUP_CHARS
    return spans


@pytest.mark.parametrize("overlap_chars", [0, 50, 200, 1000])
def test_cell_state_lines_are_capped_and_subtracted_from_the_split_budget(tmp_path, overlap_chars):
    text = _prose(4000)
    ext = {"occluded_by": {"kind": "shape", "text": "重" * overlap_chars}} if overlap_chars else {}
    rendered = _render_xlsx(tmp_path, [text], mutate=_cell_with(ext))
    lines = ([_overlap_line(overlap_chars)] if overlap_chars else []) + [_STRIKE]
    _assert_cell_within_cap(rendered, text, lines)


def test_cell_with_a_long_overlap_text_and_unbroken_body_stays_within_the_cap(tmp_path):
    text = "あ" * 3000
    rendered = _render_xlsx(tmp_path, [text], mutate=_cell_with({"covered_by_text": {"element_id": "s", "text": "重" * 1000}}))
    spans = _assert_cell_within_cap(rendered, text, [_overlap_line(1000), _STRIKE])
    assert all("…（全1000字）" in t for _c, t in _cell_chunks(rendered))
    assert all(end - start >= 860 for start, end in spans[:-1])


def test_cell_body_that_fits_alone_is_split_when_body_plus_notes_exceed_the_cap(tmp_path):
    text = _prose(1100)
    anchors = [{"behind_doc": False, "name": "確" * 50}, {"behind_doc": True, "name": "背" * 50}]
    rendered = _render_xlsx(
        tmp_path, [text], mutate=_cell_with({"occluded_by": {"kind": "shape", "text": "重" * 200},
                                              "floating_anchors": anchors}))
    lines = [_overlap_line(200), _STRIKE, f"前面図形: 「{'確' * 50}」", f"背面図形: 「{'背' * 50}」"]
    spans = _assert_cell_within_cap(rendered, text, lines)
    assert len(spans) >= 2


def test_cell_many_floating_anchors_are_summarised_and_fixed_lines_stay(tmp_path):
    text = _prose(3000)
    anchors = [{"behind_doc": False, "name": f"{i:02d}" + "図" * 48} for i in range(20)]
    rendered = _render_xlsx(tmp_path, [text], mutate=_cell_with({"floating_anchors": anchors}))
    cell_extension = {"format_state": {"strike": "whole"}, "floating_anchors": anchors}
    lines = R._occlusion_kv_lines(cell_extension)
    assert lines[0] == _STRIKE
    shown = [line for line in lines if line.startswith("前面図形:")]
    assert 0 < len(shown) < 20 and lines[-1] == f"（ほかに重なり・前面/背面図形の行が {20 - len(shown)} 件）"
    _assert_cell_within_cap(rendered, text, lines)


def test_cell_state_goes_to_every_span_even_when_a_bracket_longer_than_a_span_is_cut(tmp_path):
    text = "前置き。" + "【取り消し線：" + "あ" * 2600 + "】後ろ。"
    rendered = _render_xlsx(tmp_path, [text])
    spans = _cell_spans(rendered)
    _contiguous(spans, len(text))
    assert len(spans) >= 3
    assert sum(t.count(_STRIKE) for _c, t in _cell_chunks(rendered)) == len(spans)


def test_cell_partial_bracket_in_text_is_not_promoted_to_a_whole_cell_state(tmp_path):
    text = _prose(1500) + "【取り消し線：旧記述】" + _prose(1500)
    rendered = _render_xlsx(tmp_path, [text], strike=False)
    chunks = _cell_chunks(rendered)
    assert len(chunks) >= 3
    assert all("取り消し線: 「あり」" not in shown for _c, shown in chunks)
    assert sum(shown.count("【取り消し線：旧記述】") for _c, shown in chunks) == 1


# ---- (2) 要素の長文: 全片に状態・重なり・前面/背面図形 ----

_ELEMENT_TAIL = "\n可視性: 「スライド範囲外」\n重なり: 「廃止」\n前面図形: 「確認済スタンプ」\n背面図形: 「透かし」"
_ELEMENT_EXT = {
    "visibility_reason": "off_slide",
    "covered_by_text": {"element_id": "s2", "text": "廃止"},
    "floating_anchors": [{"behind_doc": False, "name": "確認済スタンプ"}, {"behind_doc": True, "name": "透かし"}],
}


def test_element_every_span_carries_state_overlap_and_floating_lines():
    records = _element_records(_prose(3000), visibility="hidden", **_ELEMENT_EXT)
    assert len(records) >= 3
    for record in records:
        assert record["semantic_text"].endswith(_ELEMENT_TAIL)
        assert record["markdown_text"].endswith(_ELEMENT_TAIL)
        assert record["exact_text"] and "重なり" not in record["exact_text"]


def test_element_tail_span_alone_names_kind_position_and_state():
    records = _element_records(_prose(3000), visibility="hidden", **_ELEMENT_EXT)
    tail = records[-1]
    n = len(records)
    assert f"原本「deck.pptx」のスライド1にあるshape内容（{n}/{n}）:" in tail["semantic_text"]
    assert tail["semantic_text"].endswith(_ELEMENT_TAIL)
    assert tail["citations"][0]["locator"]["slide"] == 1


def test_element_relation_sentence_and_shape_fill_stay_on_first_span_only():
    target = IR.EvidenceElement(
        element_id="t1", type="shape", parent_id=None, order=1, value="対象図形",
        locator=IR.Locator(part="ppt/slides/slide1.xml", slide=1), coverage_id="c2")
    relation = IR.EvidenceRelation(relation_id="rel1", type="connects_to", source_id="e1", target_id="t1",
                                   evidence_ids=[], confidence=1.0)
    records = [r for r in _element_records(_prose(3000), relations=[relation], others=[target], **_ELEMENT_EXT)
               if r["evidence_ids"] == ["e1"]]
    assert len(records) >= 3
    assert "対象「対象図形」へ接続している" in records[0]["semantic_text"]
    for record in records[1:]:
        assert "接続している" not in record["semantic_text"]
    for record in records:
        assert "重なり: 「廃止」" in record["semantic_text"]


def test_element_overlap_line_keeps_observation_wording():
    records = _element_records(_prose(3000), covered_by_text={"element_id": "s2", "text": "廃止"})
    for record in records:
        assert record["semantic_text"].endswith("\n重なり: 「廃止」")
        assert "すべて" not in record["semantic_text"] and "隠れ" not in record["semantic_text"]


def _state_tail(value, **kw) -> str:
    element = IR.EvidenceElement(
        element_id="e1", type=kw.pop("etype", "shape"), parent_id=None, order=0, value=value,
        locator=IR.Locator(part="ppt/slides/slide1.xml", slide=1), coverage_id="c1",
        visibility=kw.pop("visibility", "visible"), extension=kw)
    return R._element_piece(element, {"e1": element}, [], "deck.pptx", ["スライド1"],
                            include_native_metadata=True)["state_tail"]


def _assert_element_within_cap(records, text, note):
    spans = _spans(records)
    assert spans[0]["start"] == 0 and spans[-1]["end"] == len(text)
    assert all(a["end"] == b["start"] for a, b in itertools.pairwise(spans))
    assert all(s["end"] - s["start"] + len(note) <= R.MAX_TEXT_SPAN_CHARS for s in spans)
    assert all(r["semantic_text"].endswith(note) for r in records)
    assert "".join(r["exact_text"] for r in records) == text


@pytest.mark.parametrize("overlap_chars", [1, 50, 300, 1000])
def test_element_state_lines_are_capped_and_subtracted_from_the_split_budget(overlap_chars):
    ext = dict(_ELEMENT_EXT)
    ext["covered_by_text"] = {"element_id": "s2", "text": "重" * overlap_chars}
    text = _prose(5000)
    records = _element_records(text, visibility="hidden", **ext)
    note = ("\n可視性: 「スライド範囲外」\n" + _overlap_line(overlap_chars)
            + "\n前面図形: 「確認済スタンプ」\n背面図形: 「透かし」")
    _assert_element_within_cap(records, text, note)


def test_element_unbroken_body_with_a_long_overlap_text_stays_within_the_cap():
    text = "あ" * 3000
    records = _element_records(text, covered_by_text={"element_id": "s2", "text": "重" * 1000})
    note = "\n" + _overlap_line(1000)
    _assert_element_within_cap(records, text, note)
    assert "…（全1000字）" in note and len(records) == 3
    assert all(s["end"] - s["start"] >= 860 for s in _spans(records)[:-1])


def test_element_body_that_fits_alone_is_split_when_body_plus_notes_exceed_the_cap():
    text = _prose(1100)
    ext = {"covered_by_text": {"element_id": "s2", "text": "重" * 200},
           "floating_anchors": [{"behind_doc": False, "name": "確" * 50}, {"behind_doc": True, "name": "背" * 50}]}
    records = _element_records(text, **ext)
    note = "\n" + "\n".join([_overlap_line(200), f"前面図形: 「{'確' * 50}」", f"背面図形: 「{'背' * 50}」"])
    assert len(records) >= 2
    _assert_element_within_cap(records, text, note)


def test_element_many_floating_anchors_are_summarised_and_fixed_lines_stay():
    anchors = [{"behind_doc": False, "name": f"{i:02d}" + "図" * 48} for i in range(20)]
    tail = _state_tail("x", visibility="hidden", format_state={"strike": "whole"}, floating_anchors=anchors)
    lines = tail.split("\n")[1:]
    shown = [line for line in lines if line.startswith("前面図形:")]
    assert lines[0] == "可視性: 「hidden」" and _STRIKE in lines
    assert 0 < len(shown) < 20 and lines[-1] == f"（ほかに重なり・前面/背面図形の行が {20 - len(shown)} 件）"
    assert len(tail) <= R._STATE_NOTE_MAX_CHARS
    text = _prose(4000)
    records = _element_records(text, visibility="hidden", format_state={"strike": "whole"}, floating_anchors=anchors)
    _assert_element_within_cap(records, text, tail)


@pytest.mark.parametrize("body_len", [900, 1000, 1100, 1200, 1201, 3000, 6000])
@pytest.mark.parametrize("anchor_count", [0, 3, 20])
def test_body_plus_notes_never_exceed_the_cap_for_any_note_size(body_len, anchor_count):
    ext = {"covered_by_text": {"element_id": "s2", "text": "重" * 1000}, "format_state": {"strike": "whole"},
           "floating_anchors": [{"behind_doc": False, "name": "図" * 50} for _ in range(anchor_count)]}
    text = "あ" * body_len
    tail = _state_tail(text, visibility="hidden", **ext)
    assert len(tail) <= R._STATE_NOTE_MAX_CHARS
    assert R._split_budget(tail) >= R.MAX_TEXT_SPAN_CHARS - R._STATE_NOTE_MAX_CHARS > R._MIN_SPLIT_BUDGET
    records = _element_records(text, visibility="hidden", **ext)
    if body_len + len(tail) > R.MAX_TEXT_SPAN_CHARS:
        assert len(records) >= 2
    else:
        assert len(records) == 1
    if len(records) > 1:
        _assert_element_within_cap(records, text, tail)


def test_quoted_values_are_cut_only_beyond_the_limit_and_the_same_lines_are_used_unsplit():
    assert R._occlusion_kv_lines({"covered_by_text": {"element_id": "s", "text": "重" * 80}}) == [_overlap_line(80)]
    assert R._occlusion_kv_lines({"covered_by_text": {"element_id": "s", "text": "重" * 81}}) == [_overlap_line(81)]
    assert _overlap_line(81).endswith("…（全81字）」")
    records = _element_records("短い本文", covered_by_text={"element_id": "s", "text": "重" * 1000})
    assert len(records) == 1 and records[0]["semantic_text"].endswith("\n" + _overlap_line(1000))


def test_element_bracket_longer_than_a_span_keeps_state_on_every_cut_piece():
    text = "前置き。" + "【取り消し線：" + "あ" * 2600 + "】後ろ。"
    records = _element_records(text, etype="paragraph", visibility="visible",
                               format_state={"strike": "whole"}, covered_by_text={"element_id": "s2", "text": "廃止"})
    assert len(records) >= 3
    tail = "\n重なり: 「廃止」\n" + _STRIKE
    for record in records:
        assert record["semantic_text"].endswith(tail)
        assert record["semantic_text"].count(_STRIKE) == 1
    # 括りの中の切れ目は閉じ・開き直し、状態行はその外（片の末尾）に付く
    assert records[1]["semantic_text"].endswith("】" + tail)
    assert "\n【取り消し線（続き）：" in records[2]["semantic_text"]


def test_cut_bracket_pieces_with_continued_marks_and_state_stay_within_the_cap():
    """括りの中で切った片も、閉じ・開き直しの記号と全片の状態行を足して上限以内に収まる。"""
    for label in ("取り消し線", "非表示", "非表示・取り消し線"):
        text = f"【{label}：" + "あ" * 2600 + "】"
        spans = R._split_exact(text, R._split_budget("\n" + _STRIKE))
        assert len(spans) >= 3
        assert "".join(part for _s, _e, part in spans) == text
        for start, end, _part in spans:
            shown = R._continued_part(text, start, end)
            assert len(shown) + len("\n" + _STRIKE) <= R.MAX_TEXT_SPAN_CHARS, (label, len(shown))


def test_shortened_cut_does_not_split_a_short_bracket_that_follows():
    """括りの中で切った片を記号の分だけ短くしても、その先の短い括りの途中では切らない。"""
    head = "【取り消し線：" + "あ" * 1600 + "】"
    text = head + "い" * (2380 - len(head)) + "【非表示：短短短短短短短短短短】" + "う" * 600
    short = text.index("【非表示：")
    short_end = text.index("】", short) + 1
    for note in ("", "\n" + _STRIKE):
        spans = R._split_exact(text, R._split_budget(note))
        assert "".join(part for _s, _e, part in spans) == text
        assert not any(short < end < short_end for _s, end, _p in spans), (note, spans)
        for start, end, _part in spans:
            assert len(R._continued_part(text, start, end)) + len(note) <= R.MAX_TEXT_SPAN_CHARS


def test_element_partial_bracket_is_not_promoted_to_a_whole_element_state():
    text = _prose(1500) + "【取り消し線：旧記述】" + _prose(1500)
    records = _element_records(text, etype="paragraph")
    assert len(records) >= 3
    assert all("取り消し線: 「あり」" not in r["semantic_text"] for r in records)
    assert sum(r["semantic_text"].count("【取り消し線：旧記述】") for r in records) == 1


def test_element_whole_state_line_does_not_spread_brackets_to_unrelated_spans():
    text = _prose(1500) + "【取り消し線：旧記述】" + _prose(1500)
    records = _element_records(text, etype="paragraph", format_state={"strike": "whole"})
    holders = [r for r in records if "【取り消し線：旧記述】" in r["semantic_text"]]
    assert len(holders) == 1
    for record in records:
        assert record["semantic_text"].count(_STRIKE) == 1
        if record is not holders[0]:
            assert "【" not in record["semantic_text"]


def test_two_elements_with_the_same_text_and_state_are_not_merged():
    def one(eid, order):
        return IR.EvidenceElement(
            element_id=eid, type="shape", parent_id=None, order=order, value=_prose(2500),
            locator=IR.Locator(part="ppt/slides/slide1.xml", slide=1), coverage_id=f"c{order}",
            extension={"format_state": {"strike": "whole"}})

    first, second = one("e1", 0), one("e2", 1)
    ir = IR.EvidenceIR(
        schema_version=IR.EVIDENCE_IR_SCHEMA_VERSION, parser_profile="test",
        source=IR.EvidenceSource(file_type="pptx", content_hash=_HASH), elements=[first, second])
    ctx = CIR.ContextIR(schema_version=CIR.CONTEXT_IR_SCHEMA_VERSION, analyzer_version="test",
                        source_hash=_HASH, source_name="deck.pptx", document_titles={})
    records = R._non_table_records(ir, ctx, {}, set())
    per_element = {eid: [r for r in records if r["evidence_ids"] == [eid]] for eid in ("e1", "e2")}
    assert len(per_element["e1"]) == len(per_element["e2"]) >= 3
    assert len({r["logical_record_id"] for r in records}) == 2
    assert all(r["semantic_text"].endswith(_STRIKE) for r in records)


def test_cells_with_the_same_text_and_state_are_not_merged(tmp_path):
    text = _prose(2500)
    rendered = _render_xlsx(tmp_path, [text, text])
    chunks = [(c, t) for c, t in _cell_chunks(rendered)]
    by_cell = {}
    for chunk, shown in chunks:
        cell_ranges = {c["locator"].get("cell_range") for c in chunk["citations"] if c.get("text_span")}
        for cell_range in cell_ranges:
            by_cell.setdefault(cell_range, []).append(shown)
    assert set(by_cell) == {"B2", "B3"}
    assert len(by_cell["B2"]) == len(by_cell["B3"]) >= 2
    pieces = sum(t.count("備考（") for t in by_cell["B2"] + by_cell["B3"])
    assert pieces >= 6 and sum(t.count(_STRIKE) for t in by_cell["B2"] + by_cell["B3"]) == pieces


# ---- (3) Excel の取り消し線の要素 ----

def test_xlsx_strike_element_chunk_has_the_strike_line(tmp_path):
    rendered = _render_xlsx(tmp_path, ["旧料金体系"])
    texts = _anchored(rendered)
    strike_chunks = [t for c in rendered.chunks if c["content_type"] == "strike_text" for t in [texts[c["chunk_id"]]]]
    assert len(strike_chunks) == 1
    assert strike_chunks[0].count(_STRIKE) == 1 and "旧料金体系" in strike_chunks[0]
    cell_texts = [t for _c, t in _cell_chunks(rendered, marker="旧料金体系")]
    assert any(_STRIKE in t for t in cell_texts)


def test_xlsx_long_strike_element_carries_the_line_on_every_span(tmp_path):
    rendered = _render_xlsx(tmp_path, [_prose(3000)])
    texts = _anchored(rendered)
    chunks = [texts[c["chunk_id"]] for c in rendered.chunks if c["content_type"] == "strike_text"]
    assert len(chunks) >= 3
    assert all(t.rstrip().endswith(_STRIKE) for t in chunks)


def test_strike_text_without_reason_gets_the_line_but_another_reason_is_left_alone():
    def piece(**ext):
        element = IR.EvidenceElement(
            element_id="e1", type="strike_text", parent_id=None, order=0, value="旧",
            locator=IR.Locator(part="xl/worksheets/sheet1.xml", sheet="明細", cell_range="B2"),
            coverage_id="c1", extension=ext)
        return R._element_piece(element, {"e1": element}, [], "s.xlsx", ["シート「明細」"],
                                include_native_metadata=True)

    assert piece()["semantic"].count(_STRIKE) == 1
    assert piece(format_state={"strike": "whole"})["semantic"].count(_STRIKE) == 1
    assert _STRIKE not in piece(visibility_reason="hidden_run")["semantic"]


def test_docx_struck_run_host_paragraph_does_not_get_the_whole_element_line(tmp_path):
    path = tmp_path / "a.docx"
    body = (f"<w:p><w:r><w:t>{_prose(1500)}</w:t></w:r>"
            "<w:r><w:rPr><w:strike/></w:rPr><w:t>旧記述</w:t></w:r>"
            f"<w:r><w:t>{_prose(1500)}</w:t></w:r></w:p>")
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", f'<?xml version="1.0"?><w:document {_NS_W}><w:body>{body}</w:body></w:document>')
    rendered = R.render(evidence_spike.extract(path), source_name="a.docx")
    texts = _anchored(rendered)
    host = [texts[c["chunk_id"]] for c in rendered.chunks if c["content_type"] == "paragraph"]
    strike = [texts[c["chunk_id"]] for c in rendered.chunks if c["content_type"] == "strike_text"]
    assert len(host) >= 3 and all(_STRIKE not in t for t in host)
    assert sum(t.count("【取り消し線：旧記述】") for t in host) == 1
    assert len(strike) == 1 and strike[0].count(_STRIKE) == 1


# ---- (f) 状態・関係の無い長文は変えない（陰性） ----

def test_split_exact_default_budget_is_unchanged():
    text = _prose(3300)
    assert R._split_exact(text) == R._split_exact(text, R.MAX_TEXT_SPAN_CHARS)
    assert R._split_budget("") == R.MAX_TEXT_SPAN_CHARS


def _digest_records(records) -> str:
    payload = [[r["semantic_text"], r["markdown_text"], r["exact_text"], r["record_id"],
                [c.get("text_span") for c in r["citations"]]] for r in records]
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def _digest_rendered(rendered) -> str:
    md = "\n".join(line for line in rendered.markdown.splitlines()
                   if not line.startswith(("<!-- chunk:", "原本SHA-256", "変換プロファイル")))  # 版の文字列は比べない
    payload = [md, [c["body"] for c in rendered.chunks],
               [[ci.get("text_span") for ci in c["citations"]] for c in rendered.chunks]]
    return hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()


def _plain_long_with_bracket() -> str:
    return _prose(1500) + "【取り消し線：旧記述】" + _prose(1800)


def _golden_element_digest() -> str:
    target = IR.EvidenceElement(
        element_id="t1", type="shape", parent_id=None, order=1, value="対象図形",
        locator=IR.Locator(part="ppt/slides/slide1.xml", slide=1), coverage_id="c2")
    relation = IR.EvidenceRelation(relation_id="rel1", type="connects_to", source_id="e1", target_id="t1",
                                   evidence_ids=[], confidence=1.0)
    plain = _element_records(_plain_long_with_bracket(), etype="paragraph", file_type="docx",
                             source_name="doc.docx")
    related = _element_records(_prose(3000), relations=[relation], others=[target])
    return _digest_records(plain + related)


def _golden_cell_digest(tmp_path) -> str:
    return _digest_rendered(_render_xlsx(tmp_path, [_plain_long_with_bracket()], strike=False))


# 状態・関係の無い入力を変更前の実装へ通して得た出力の SHA-256（本文・分割位置・record_id を含む）
_GOLDEN_ELEMENT = "d90874722bbd33681a7284dfb64cb8f9749d272ac2e3b5735753d75cbf746849"
_GOLDEN_CELL = "fdc2bdbcd52256a553ae8a7195d0252effe0180432034b4d6a1292271f4847b3"


def test_long_text_without_state_or_relation_is_byte_identical_to_before_element():
    assert _golden_element_digest() == _GOLDEN_ELEMENT


def test_long_text_without_state_is_byte_identical_to_before_cell(tmp_path):
    assert _golden_cell_digest(tmp_path) == _GOLDEN_CELL


# ---- (g) 人用 MD は変えない ----

_GOLDEN_HUMAN_XLSX = "dd6065cd7ceff1cd7b94cb4122e63ed1dd45cd9a69b573a19b0d3e29c67e31cd"
_GOLDEN_HUMAN_DOCX = "de50b976412e0db4ff329aeaff11f52158ae2862938a4b27cedc5856ea2f6074"


def _human_md(tmp_path, name: str) -> str:
    src = tmp_path / "src"
    src.mkdir(exist_ok=True)
    if name.endswith(".xlsx"):
        _write_xlsx(src / name, [_prose(3000)])
    else:
        body = (f"<w:p><w:r><w:t>{_prose(1500)}</w:t></w:r>"
                "<w:r><w:rPr><w:strike/></w:rPr><w:t>旧記述</w:t></w:r>"
                f"<w:r><w:t>{_prose(1500)}</w:t></w:r></w:p>")
        with zipfile.ZipFile(src / name, "w") as z:
            z.writestr("word/document.xml",
                       f'<?xml version="1.0"?><w:document {_NS_W}><w:body>{body}</w:body></w:document>')
    der = tmp_path / "derived"
    rep = office_md.build_derived(src, der)
    assert rep["rag_failed"] == 0, rep
    return (der / f"{name}.md").read_text(encoding="utf-8")


def test_human_md_xlsx_keeps_the_body_and_adds_only_the_strike_note(tmp_path):
    md = _human_md(tmp_path, "s.xlsx")
    note = "\n\n※ セル B2の文字には取り消し線が引かれています"
    assert md.endswith(note) and "重なり" not in md
    assert hashlib.sha256(md.removesuffix(note).encode("utf-8")).hexdigest() == _GOLDEN_HUMAN_XLSX


def test_human_md_does_not_carry_the_new_lines_and_is_unchanged_docx(tmp_path):
    md = _human_md(tmp_path, "a.docx")
    assert "取り消し線: " not in md and "重なり" not in md
    assert hashlib.sha256(md.encode("utf-8")).hexdigest() == _GOLDEN_HUMAN_DOCX


# ---- 索引へ渡す本文・束ねたベクトル用表現・道具が返す根拠本文（実 ES・実モデルなし） ----

@pytest.fixture()
def _built_world(monkeypatch, tmp_path):
    root = tmp_path / "kb" / "w1"
    root.mkdir(parents=True)
    monkeypatch.setattr(worlds, "world_dir", lambda w: root if w == "w1" else None)
    monkeypatch.setenv("SHERPA_DERIVED_DIR", str(tmp_path / "derived"))
    text = "".join(f"旧料金体系の記述{i:03d}。" for i in range(260))
    _write_xlsx(root / "s.xlsx", [text])
    rep = office_md.build_derived(root, worlds.derived_md_dir("w1"))
    assert rep["rag_failed"] == 0, rep
    return worlds.derived_rag_dir("w1")


def test_new_lines_reach_index_body_group_vector_text_and_tool_region_text(_built_world):
    rag = _built_world
    md_path, jsonl = rag / "s.xlsx.rag.md", rag / "s.xlsx.rag_chunks.jsonl"
    chunks = [json.loads(line) for line in json_io.read_text_maybe_gzip(jsonl).splitlines()]
    cell_ids = [c["chunk_id"] for c in chunks if c["content_type"] == "table_record" and "備考: " in c["body"]]
    assert len(cell_ids) >= 3

    # ① 索引へ渡す本文（rag.md のアンカー間）
    anchors, reason = es_index._load_rag_md_anchors(md_path)
    assert reason is None
    for chunk_id in cell_ids:
        assert anchors[chunk_id].endswith(_STRIKE)

    # ② 束ねたベクトル用表現（表の行の束の先頭が持つ vector_text）
    _seen, invalid = es_index._rag_chunks_validate(jsonl, anchors, "s.xlsx")
    assert invalid is None
    entries = list(es_index._iter_rag_chunk_entries(jsonl, "s.xlsx", anchors, {"doc_id": "s.xlsx"}))
    heads = [(body, text) for _i, body, text, skip in entries
             if not skip and body["chunk_id"] in cell_ids and "vector_text" in body]
    assert len(heads) == 1
    head, vector_text = heads[0]
    rows = vector_text.split("\n")
    assert "取り消し線" in rows[1]                                  # 列名の行
    assert [row.endswith("「あり」") for row in rows[2:]] == [True] * len(cell_ids)
    assert set(cell_ids) <= set(head["row_group"]["chunk_ids"])

    # ③ 道具が返す根拠本文（親返しの領域本文）
    info: dict = {}
    region = read_tools._rag_md_region_text("w1", "s.xlsx", [], None, set(cell_ids), 10 ** 7, info)
    assert region is not None and info["missing_chunks"] == 0
    assert region.count(_STRIKE) == len(cell_ids)
    for chunk_id in cell_ids:
        assert anchors[chunk_id] in region


# ---- 表のセルの一部だけの状態（セル全体の状態なし）が、分割の幅を超えるとき ----

def _partial_strike_cell_world(monkeypatch, tmp_path, struck_chars: int):
    """表のセルが「前。」＋取り消し線の「旧」×N＋「現行。」の3つの run の架空 DOCX を、実際の抽出→Evidence→描画に通す。"""
    root = tmp_path / "kb" / "w1"
    root.mkdir(parents=True)
    monkeypatch.setattr(worlds, "world_dir", lambda w: root if w == "w1" else None)
    monkeypatch.setenv("SHERPA_DERIVED_DIR", str(tmp_path / "derived"))
    cell = ('<w:p><w:r><w:t>前。</w:t></w:r><w:r><w:rPr><w:strike/></w:rPr><w:t>' + "旧" * struck_chars
            + '</w:t></w:r><w:r><w:t>現行。</w:t></w:r></w:p>')
    table = ('<w:tbl><w:tr><w:tc><w:p><w:r><w:t>項目</w:t></w:r></w:p></w:tc>'
             '<w:tc><w:p><w:r><w:t>内容</w:t></w:r></w:p></w:tc></w:tr>'
             f'<w:tr><w:tc><w:p><w:r><w:t>A01</w:t></w:r></w:p></w:tc><w:tc>{cell}</w:tc></w:tr></w:tbl>')
    with zipfile.ZipFile(root / "表.docx", "w") as z:
        z.writestr("word/document.xml", f'<?xml version="1.0"?><w:document {_NS_W}><w:body>'
                   f'<w:p><w:r><w:t>本文。</w:t></w:r></w:p>{table}</w:body></w:document>')
    rep = office_md.build_derived(root, worlds.derived_md_dir("w1"))
    assert rep["rag_failed"] == 0, rep
    rag = worlds.derived_rag_dir("w1")
    jsonl, md_path = rag / "表.docx.rag_chunks.jsonl", rag / "表.docx.rag.md"
    records = [json.loads(line) for line in json_io.read_text_maybe_gzip(jsonl).splitlines()]
    return jsonl, md_path, [r for r in records if r["content_type"] == "table_record" and "内容" in r["body"]]


def test_partial_strike_in_a_long_cell_reaches_body_index_text_and_region_text_of_every_piece(monkeypatch, tmp_path):
    jsonl, md_path, records = _partial_strike_cell_world(monkeypatch, tmp_path, 2600)
    assert len(records) >= 2
    anchors, reason = es_index._load_rag_md_anchors(md_path)
    assert reason is None
    spans = sorted((c["text_span"]["start"], c["text_span"]["end"])
                   for r in records for c in r["citations"] if c.get("text_span"))
    assert spans[0][0] == 0 and all(a[1] == b[0] for a, b in zip(spans, spans[1:]))   # 文字範囲は欠落・重複なく続く
    for record in records:
        body, indexed = record["body"], anchors[record["chunk_id"]]
        for text in (body, indexed):                                     # JSONL の body と ① 索引へ渡す本文
            assert "【取り消し線" in text, text[:80]
            assert "取り消し線: 「あり」" not in text                       # セル全体の状態にはしない
    tail = records[-1]
    assert "【取り消し線（続き）：" in tail["body"] and tail["body"].rstrip().endswith("】現行。")
    first = records[0]["body"]
    assert "前。" in first and "前。】" not in first and "【取り消し線：前" not in first   # 前の通常文字は括りの外
    # ② 束ねたベクトル用表現
    entries = list(es_index._iter_rag_chunk_entries(jsonl, "表.docx", anchors, {"doc_id": "表.docx"}))
    heads = [text for _i, body, text, skip in entries if not skip and body["chunk_id"] in {r["chunk_id"] for r in records}]
    assert heads and all("取り消し線" in text for text in heads)
    # ③ 道具が返す根拠本文
    info: dict = {}
    region = read_tools._rag_md_region_text("w1", "表.docx", [], None, {r["chunk_id"] for r in records}, 10 ** 7, info)
    assert region is not None and info["missing_chunks"] == 0
    for record in records:
        assert anchors[record["chunk_id"]] in region and "【取り消し線" in anchors[record["chunk_id"]]


def test_partial_strike_in_a_very_long_cell_marks_every_row_group_vector_text(monkeypatch, tmp_path):
    jsonl, md_path, records = _partial_strike_cell_world(monkeypatch, tmp_path, 30000)
    anchors, reason = es_index._load_rag_md_anchors(md_path)
    assert reason is None
    ids = {r["chunk_id"] for r in records}
    entries = list(es_index._iter_rag_chunk_entries(jsonl, "表.docx", anchors, {"doc_id": "表.docx"}))
    groups = [(body, text) for _i, body, text, skip in entries if not skip and body["chunk_id"] in ids]
    assert len(groups) >= 2                                              # 束が複数に分かれる長さ
    for body, text in groups:
        rows = text.split("\n")[2:]
        assert rows and all("取り消し線" in row for row in rows), rows[:3]
    for record in records:
        assert "【取り消し線" in record["body"]
        for citation in record["citations"]:
            span = citation.get("text_span")
            if span:
                assert span["end"] - span["start"] <= R.MAX_TEXT_SPAN_CHARS


def test_long_cell_partial_strike_pieces_keep_the_split_budget(monkeypatch, tmp_path):
    exact = "前。【取り消し線：" + "旧" * 2600 + "】現行。"
    spans = R._split_exact(exact)
    assert len(spans) >= 3
    for start, end, _part in spans:
        assert len(R._continued_part(exact, start, end)) <= R.MAX_TEXT_SPAN_CHARS
