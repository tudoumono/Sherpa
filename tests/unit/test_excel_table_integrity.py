"""Excel の表が、見出しと値の対応を保ったまま rag.md／チャンクへ届くことの回帰テスト。

原本は各テストの中で openpyxl／xlsxwriter により作る（fixtures の原本は使わない）。
経路は `evidence_spike.extract` → `evidence_render.render`（本番の rag.md・rag_chunks と同じ関数）。
非表示の行・列・シートは `test_occ1_kv_visibility_rag_md.py`／`test_visibility_l3.py` が確かめている。
"""
from __future__ import annotations

import datetime
import json

import openpyxl
import pytest
import xlsxwriter
from openpyxl.styles import PatternFill

from sherpa.ingest import context_ir, evidence_render, evidence_spike, excel_display, office_md
from sherpa.ingest.arms import ooxml_arm
from sherpa.ingest.ooxml import excel

_HEADER_FILL = PatternFill(fill_type="solid", fgColor="FFD9EAF7")


def _render(path):
    evidence = evidence_spike.extract(path)
    excel_display.enrich_evidence(evidence, path)
    return evidence_render.render(evidence, source_name=path.name)


def _kinds(chunks, kind):
    return [c for c in chunks if c["content_type"] == kind]


def _chunk_with(chunks, *needles):
    hits = [c for c in chunks if all(n in c["body"] for n in needles)]
    assert hits, f"{needles} を同居させたチャンクが無い: {[c['body'] for c in chunks]}"
    return hits[0]


def _save(wb, tmp_path, name="t.xlsx"):
    path = tmp_path / name
    wb.save(path)
    return path


def test_blank_header_cell_keeps_field_name_and_value_in_one_chunk(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["C1"] = "項目名", "説明"
    ws["B1"].fill = _HEADER_FILL
    ws["A2"], ws["C2"] = "顧客ID", "10桁"
    chunks = _render(_save(wb, tmp_path)).chunks
    record = _chunk_with(chunks, "項目名: 顧客ID", "説明: 10桁")
    assert record["content_type"] == "table_record"
    assert not _kinds(chunks, "coordinate_record")


def test_header_with_blank_cells_inside_body_keeps_row_together(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["項目ID", "名称", "初期値", "備考"])
    ws.append(["A01", "顧客", None, "メモ1"])
    ws.append(["A02", "金額", "100", "メモ2"])
    ws.append(["A03", "税", None, "メモ3"])
    chunks = _render(_save(wb, tmp_path)).chunks
    record = _chunk_with(chunks, "項目ID: A02", "初期値: 100", "備考: メモ2")
    assert record["content_type"] == "table_record"
    first = _chunk_with(chunks, "項目ID: A01")
    assert "備考: メモ1" in first["body"] and "初期値" not in first["body"]
    assert not _kinds(chunks, "coordinate_record")


def test_last_row_with_blank_trailing_cell_keeps_multilevel_header(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["C1"] = "基本", "詳細"
    ws.merge_cells("A1:B1")
    ws.merge_cells("C1:D1")
    ws.append(["項目名", "桁数", "型", "備考"])
    ws.append(["顧客ID", 10, "数字", "必須"])
    ws.append(["金額", 12, "数値", None])
    chunks = _render(_save(wb, tmp_path)).chunks
    last = _chunk_with(chunks, "項目名: 金額", "桁数: 12")
    assert last["content_type"] == "table_record"
    assert not _kinds(chunks, "coordinate_record")


def test_tables_side_by_side_across_a_gap_column_are_not_mixed(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    for col, name in enumerate(["項目ID", "名称", "桁数"], start=1):
        ws.cell(row=1, column=col, value=name)
    for col, name in enumerate(["コード", "区分名", "備考"], start=6):
        ws.cell(row=1, column=col, value=name)
    ws.append(["A01", "顧客", 10, None, None, "01", "法人", "メモ"])
    chunks = _render(_save(wb, tmp_path)).chunks
    left = _chunk_with(chunks, "項目ID: A01")
    right = _chunk_with(chunks, "コード: 01")
    assert left is not right
    assert "区分名" not in left["body"] and "桁数" not in right["body"]


def test_long_description_split_keeps_link_back_to_item_and_note_is_separate(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "項目"
    ws.append(["項目ID", "名称", "説明"])
    ws.append(["A01", "顧客区分", "顧客の区分を表す。" * 400])
    ws.append(["A02", "金額", "税込金額。"])
    ws["F5"] = "注記: 金額は税込で表示する"
    chunks = _render(_save(wb, tmp_path)).chunks
    parts = [c for c in chunks if c["content_type"] == "table_record"
             and any(k["value"] == "A01" for k in c["record_keys"])]
    assert len(parts) >= 2
    assert len({c["logical_record_id"] for c in parts}) == 1
    assert all(c["field_group_count"] == len(parts) for c in parts)
    assert "項目ID: A01" in parts[0]["body"]
    assert all(c["region_context"]["cell_range"] == "A1:C3" for c in parts)
    assert "".join(c["body"] for c in parts).count("顧客の区分を表す。") == 400
    assert _chunk_with(chunks, "注記: 金額は税込で表示する")["content_type"] == "coordinate_record"
    assert _chunk_with(chunks, "項目ID: A02", "説明: 税込金額。")


def test_value_on_second_sheet_keeps_sheet_name_and_header(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "表紙"
    wb.active["A1"], wb.active["A2"] = "改版履歴", "第1版"
    sheet = wb.create_sheet("項目一覧")
    sheet.append(["項目ID", "桁数"])
    sheet.append(["CUST", 10])
    chunks = _render(_save(wb, tmp_path)).chunks
    record = _chunk_with(chunks, "項目ID: CUST", "桁数: 10")
    assert record["section_path"] == ["シート「項目一覧」"]
    assert record["region_context"]["sheet"] == "項目一覧"


def test_leading_zero_date_percent_unit_zero_and_blank_are_kept_verbatim(tmp_path):
    path = tmp_path / "types.xlsx"
    wb = xlsxwriter.Workbook(path)
    ws = wb.add_worksheet("S")
    pct = wb.add_format({"num_format": "0.0%"})
    day = wb.add_format({"num_format": "yyyy/mm/dd"})
    for col, name in enumerate(["項目ID", "コード", "税率", "開始日", "件数", "備考"]):
        ws.write(0, col, name)
    ws.write(1, 0, "A01")
    ws.write_string(1, 1, "00123")
    ws.write_number(1, 2, 0.1, pct)
    ws.write_datetime(1, 3, datetime.datetime(2026, 4, 1), day)
    ws.write_number(1, 4, 0)
    ws.write(1, 5, "10桁")
    ws.write(2, 0, "A02")
    ws.write_string(2, 1, "00456")
    ws.write_number(2, 2, 0.08, pct)
    ws.write_datetime(2, 3, datetime.datetime(2026, 10, 1), day)
    ws.write_blank(2, 4, None)
    ws.write(2, 5, "単位:円")
    wb.close()
    chunks = _render(path).chunks
    zero = _chunk_with(chunks, "項目ID: A01")
    for needle in ("コード: 00123", "税率: 10.0%", "開始日: 2026/04/01", "件数: 0", "備考: 10桁"):
        assert needle in zero["body"]
    blank = _chunk_with(chunks, "項目ID: A02")
    assert "コード: 00456" in blank["body"] and "備考: 単位:円" in blank["body"]
    assert "件数" not in blank["body"]


def test_formula_with_saved_value_keeps_value_and_without_it_is_neither_zero_nor_computed(tmp_path):
    cached = tmp_path / "cached.xlsx"
    wb = xlsxwriter.Workbook(cached)
    ws = wb.add_worksheet("S")
    for col, name in enumerate(["項目", "元値", "換算"]):
        ws.write(0, col, name)
    ws.write(1, 0, "A")
    ws.write_number(1, 1, 5)
    ws.write_formula(1, 2, "=B2*100", None, 500)
    wb.close()
    record = _chunk_with(_render(cached).chunks, "項目: A", "元値: 5")
    assert "換算: 500" in record["body"]

    plain = openpyxl.Workbook()
    plain.active.append(["項目", "元値", "換算"])
    plain.active.append(["B", 7, "=B2*100"])
    chunks = _render(_save(plain, tmp_path, "plain.xlsx")).chunks
    joined = "\n".join(c["body"] for c in chunks)
    assert "=B2*100" in joined
    assert "換算: 0" not in joined and "700" not in joined
    assert "換算: " not in joined


def test_plain_table_stays_one_record_per_row(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["項目ID", "名称", "桁数"])
    ws.append(["A01", "顧客", 10])
    ws.append(["A02", "金額", 12])
    chunks = _render(_save(wb, tmp_path)).chunks
    assert [c["content_type"] for c in chunks] == ["context_summary", "table_record", "table_record"]
    assert "名称: 金額" in chunks[2]["body"]


@pytest.mark.parametrize("rows,pairs", [
    ([["タイムアウト", "30秒"], ["再試行回数", "3回"], ["接続先", "サーバーA"]],
     [("タイムアウト", "30秒"), ("再試行回数", "3回"), ("接続先", "サーバーA")]),
    ([["桁数", "10桁"], ["区分", 2], ["備考", "任意"]],
     [("桁数", "10桁"), ("区分", "2"), ("備考", "任意")]),
], ids=["単位付きの値", "単位付きと数値"])
def test_headerless_vertical_item_value_table_keeps_each_row_pair(tmp_path, rows, pairs):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    result = _render(_save(wb, tmp_path))
    assert [rg.mode for rg in result.context.regions] == ["coordinate_fallback"]
    assert result.context.regions[0].header_row is None
    for label, value in pairs:
        record = _chunk_with(result.chunks, label, value)
        assert f"{label}: {value}" in record["body"]
    bodies = "\n".join(c["body"] for c in result.chunks)
    for label, _ in pairs:
        for other, _ in pairs:
            if other != label:
                assert f"{label}: {other}" not in bodies
    assert not _kinds(result.chunks, "table_record")


@pytest.mark.parametrize("rows", [
    [["項目", "1月", "2月"], ["売上", 100, 200]],
    [["項目", "2025年度"], ["売上", 100], ["原価", 80]],
    [["項目", "2025"], ["売上", 100], ["原価", 80]],
    [["品名", "単価"], ["りんご", 100], ["みかん", 80]],
], ids=["月の見出し", "年度の見出し", "年の見出し", "文字の見出し"])
def test_header_led_tables_with_calendar_or_text_headers_stay_header_records(tmp_path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    result = _render(_save(wb, tmp_path))
    assert [rg.mode for rg in result.context.regions] == ["header_record"]
    assert result.context.regions[0].header_row == 1


def _grid(rows, cols, values):
    grid = [[None] * cols for _ in range(rows)]
    for (r, c), v in values.items():
        grid[r - 1][c - 1] = v
    return grid


@pytest.mark.parametrize("values,filled,expected", [
    pytest.param({(1, 1): "項目名", (1, 3): "説明", (2, 1): "顧客ID", (2, 3): "10桁"}, {(1, 2)},
                 ["A1:C2"], id="見出しの空セルを背景色が埋める"),
    pytest.param({(1, 1): "項目名", (1, 3): "説明", (2, 1): "顧客ID", (2, 3): "10桁"}, set(),
                 ["A1:A2", "C1:C2"], id="背景色の無い空き列は別表のまま"),
    pytest.param({(1, 1): "項目名", (1, 2): "型", (1, 3): "備考",
                  (2, 1): "顧客ID", (2, 2): "数字", (3, 1): "金額", (3, 2): "数値", (3, 3): "税込"}, set(),
                 ["A1:B3", "C1:C3"], id="末尾の列の欠けは列ごとに見出しへ戻す"),
    pytest.param({(1, 1): "項目名", (1, 2): "型", (1, 3): "備考",
                  (2, 1): "a", (2, 2): "b", (2, 3): "c", (3, 1): "d", (3, 2): "e"}, set(),
                 ["A1:C3"], id="最終行の末尾の空欄は同じ表"),
])
def test_regions_rejoin_header_fragments(values, filled, expected):
    out = excel.regions(_grid(3, 3, values), cap_rows=3, cap_cols=3, filled=filled)
    assert sorted(rg.range for rg in out) == sorted(expected)


def test_regions_do_not_rejoin_when_most_header_columns_have_no_body():
    values = {(1, c): f"見出し{c}" for c in range(1, 7)}
    values.update({(r, c): f"v{r}{c}" for r in range(2, 4) for c in range(1, 7)})
    values.update({(r, c): f"w{r}{c}" for r in range(4, 6) for c in range(1, 3)})
    out = excel.regions(_grid(5, 6, values), cap_rows=5, cap_cols=6)
    assert sorted(rg.range for rg in out) == ["A1:F3", "A4:B5"]


def test_xlsx_extractor_version_reaches_rag_signature_so_derived_files_are_rebuilt(tmp_path):
    current = office_md._current_rag_sig()
    assert f"xlsx={ooxml_arm.XLSX_EXTRACTOR_VERSION}" in current
    stale = current.replace(f"xlsx={ooxml_arm.XLSX_EXTRACTOR_VERSION}", "xlsx=xlsx-ooxml-v5")
    (tmp_path / office_md._RAG_SIG_MARKER).write_text(stale, encoding="utf-8")
    assert office_md.rag_sig_drift(tmp_path) is True
    (tmp_path / office_md._RAG_SIG_MARKER).write_text(current, encoding="utf-8")
    assert office_md.rag_sig_drift(tmp_path) is False


def test_rebuilt_derived_files_carry_joined_record_into_rag_md_and_chunks(tmp_path):
    world = tmp_path / "world"
    world.mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["C1"] = "項目名", "説明"
    ws["B1"].fill = _HEADER_FILL
    ws["A2"], ws["C2"] = "顧客ID", "10桁"
    wb.save(world / "a.xlsx")
    derived = tmp_path / "derived"
    office_md.build_derived(world, derived)
    rag_md = (tmp_path / "rag" / "a.xlsx.rag.md").read_text(encoding="utf-8")
    assert "項目名: 「顧客ID」\n\n説明: 「10桁」" in rag_md
    import gzip
    raw = (tmp_path / "rag" / "a.xlsx.rag_chunks.jsonl").read_bytes()
    text = gzip.decompress(raw).decode("utf-8") if raw[:2] == b"\x1f\x8b" else raw.decode("utf-8")
    rows = [json.loads(line) for line in text.splitlines() if line.strip()]
    assert [r["content_type"] for r in rows] == ["context_summary", "table_record"]


_PRICE_ROWS = [
    ("P-30197", "六角ボルト M8×30", 16.2, 17.1, 0.0556, 500, 7, 3, "丸辻工業", None),
    ("P-23456", "六角ボルト M8×35", 17.5, 18.6, 0.0629, 500, 10, 3, "ハヤセ電材", "特急は1回100個まで"),
    ("P-28149", "六角ボルト M10×30", 21, 22.4, 0.0667, 300, 10, 4, "ハヤセ電材", None),
    ("P-10895", "六角ナット M8", 346, 366.8, 0.0601, 300, 12, 4, "協和ベアリング販売", None),
]


def _price_book(tmp_path, *, title=True, year_numbers=False, name="price.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    top = 1
    if title:
        ws["A1"] = "購買部品 単価一覧（単位：円）"
        ws.merge_cells("A1:J1")
        top = 2
    ws.cell(top, 1, "部品コード")
    ws.cell(top, 2, "部品名")
    ws.cell(top, 3, "単価（円）")
    ws.merge_cells(start_row=top, start_column=3, end_row=top, end_column=5)
    ws.cell(top, 7, "リードタイム（営業日）")
    ws.merge_cells(start_row=top, start_column=7, end_row=top, end_column=8)
    ws.cell(top, 9, "仕入先")
    ws.cell(top, 10, "備考")
    for col in (1, 2, 9, 10):
        ws.merge_cells(start_row=top, start_column=col, end_row=top + 1, end_column=col)
    lower = [(3, 2023 if year_numbers else "2023年度"), (4, 2024 if year_numbers else "2024年度"),
             (5, "改定率"), (6, "最小発注数"), (7, "通常"), (8, "特急")]
    for col, text in lower:
        ws.cell(top + 1, col, text)
    for row in _PRICE_ROWS:
        ws.append(list(row))
    path = tmp_path / name
    wb.save(path)
    return path


def _record_labels(chunks):
    labels = set()
    for chunk in chunks:
        if chunk["content_type"] != "table_record":
            continue
        for line in chunk["body"].split("\n"):
            if ": " in line:
                labels.add(line.split(": ", 1)[0])
    return labels


@pytest.mark.parametrize("title", [True, False], ids=["表題行あり", "表題行なし"])
def test_two_level_header_with_vertical_merges_labels_values_not_first_data_row(tmp_path, title):
    chunks = _render(_price_book(tmp_path, title=title)).chunks
    record = _chunk_with(chunks, "部品コード: P-23456", "部品名: 六角ボルト M8×35")
    assert "2023年度: 17.5" in record["body"] and "2024年度: 18.6" in record["body"]
    if title:
        assert "特急: 3" in record["body"] and "仕入先: ハヤセ電材" in record["body"]
        assert "備考: 特急は1回100個まで" in record["body"]
    assert not any(label.split(" > ")[-1].replace(".", "").isdigit() for label in _record_labels(chunks)
                   if not label.endswith("年度"))
    assert not any("P-30197" in label or "16.2" in label for label in _record_labels(chunks))


def test_title_row_is_not_used_as_header_and_group_header_is_inherited(tmp_path):
    chunks = _render(_price_book(tmp_path)).chunks
    record = _chunk_with(chunks, "部品コード: P-23456")
    assert "単価（円） > 2023年度: 17.5" in record["body"]
    assert "リードタイム（営業日） > 通常: 10" in record["body"]
    assert "購買部品" not in "".join(label for label in _record_labels(chunks))


def test_numeric_fiscal_years_in_lower_header_row_stay_header_labels(tmp_path):
    chunks = _render(_price_book(tmp_path, year_numbers=True)).chunks
    record = _chunk_with(chunks, "部品コード: P-23456")
    assert "単価（円） > 2023: 17.5" in record["body"] and "単価（円） > 2024: 18.6" in record["body"]


def test_date_header_cells_stay_labels_over_numeric_rows(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "月次実績"
    ws.merge_cells("A1:D1")
    ws["A2"], ws["B2"], ws["C2"] = "支店", datetime.date(2024, 4, 1), datetime.date(2024, 5, 1)
    ws["D2"] = "備考"
    ws.append(["東京", 120, 130, None])
    ws.append(["大阪", 90, 95, "改装中"])
    path = tmp_path / "dates.xlsx"
    wb.save(path)
    chunks = _render(path).chunks
    record = _chunk_with(chunks, "支店: 大阪", "改装中")
    assert "2024-04-01" in record["body"] and ": 90" in record["body"]
    assert not any(label in {"東京", "120", "大阪", "90"} for label in _record_labels(chunks))


def test_sparse_columns_with_blank_header_cell_stay_in_the_header_led_table(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["項目名", None, "型", "桁数", "必須", "備考"])
    ws["B1"].fill = _HEADER_FILL
    ws.append(["受注番号", "JUCHU_NO", "CHAR", 10, "○", "採番規則は別紙"])
    ws.append(["明細番号", "MEISAI_NO", "NUMBER", 3, None, None])
    ws.append(["品目名", "HINMOKU_NM", "VARCHAR", 60, None, "全角30文字"])
    ws.append(["単価", "TANKA", "NUMBER", 11, "○", None])
    chunks = _render(_save(wb, tmp_path)).chunks
    record = _chunk_with(chunks, "項目名: 品目名", "桁数: 60", "備考: 全角30文字")
    assert record["content_type"] == "table_record"
    _chunk_with(chunks, "項目名: 受注番号", "必須: ○", "備考: 採番規則は別紙")
    assert not any(label.isdigit() or label == "CHAR" for label in _record_labels(chunks))
    assert not _kinds(chunks, "coordinate_record")


def test_regions_unify_banded_table_with_sparse_last_column_and_absorb_islands():
    values = {(1, c): "表題" if c == 1 else None for c in range(1, 5)}
    values.update({(2, 1): "コード", (2, 2): "名称", (2, 3): "数量", (2, 4): "備考"})
    for r in range(3, 8):
        values.update({(r, 1): f"P{r}", (r, 2): f"品{r}", (r, 3): r * 1.5})
    values[(5, 4)] = "注意"
    merged = {(1, c): {"anchor": (1, 1), "row_span": 1, "column_span": 4} for c in range(1, 5)}
    out = excel.regions(_grid(7, 4, values), cap_rows=7, cap_cols=4, merged=merged)
    assert [rg.range for rg in out] == ["A1:D7"]


def test_context_analyzer_version_reaches_rag_signature(tmp_path):
    current = office_md._current_rag_sig()
    assert f"{context_ir.CONTEXT_IR_SCHEMA_VERSION}/{context_ir.CONTEXT_ANALYZER_VERSION}" in current
    stale = current.replace(context_ir.CONTEXT_ANALYZER_VERSION, "xlsx-context-rules-v4")
    (tmp_path / office_md._RAG_SIG_MARKER).write_text(stale, encoding="utf-8")
    assert office_md.rag_sig_drift(tmp_path) is True


def _block_book(path, *, gap=False, blocks=("北陸", "東海", "関西"), rows=6, short_block=None, blank_label_row=None):
    """行見出し（費目コード・費目名）の右に、拠点ごとの月次ブロックが横に並ぶ表。"""
    wb = xlsxwriter.Workbook(path)
    ws = wb.add_worksheet("予算実績")
    month = wb.add_format({"num_format": 'yyyy"年"m"月"'})
    thousand = wb.add_format({"num_format": "#,##0,"})
    pct = wb.add_format({"num_format": "0.0%"})
    ws.merge_range(0, 0, 1, 0, "費目コード")
    ws.merge_range(0, 1, 1, 1, "費目名")
    col = 2
    for name in blocks:
        ws.merge_range(0, col, 0, col + 3, name)
        ws.write_datetime(1, col, datetime.datetime(2025, 4, 1), month)
        ws.write_datetime(1, col + 1, datetime.datetime(2025, 5, 1), month)
        ws.write(1, col + 2, "通期実績")
        ws.write(1, col + 3, "執行率")
        count = rows - 1 if name == short_block else rows
        for r in range(count):
            ws.write_number(2 + r, col, 13157000 + r * 1000000, thousand)
            ws.write_number(2 + r, col + 1, 2000000, thousand)
            ws.write_number(2 + r, col + 2, 15157000 + r * 1000000, thousand)
            ws.write_formula(2 + r, col + 3, f'=IFERROR({xlsxwriter.utility.xl_rowcol_to_cell(2 + r, col + 2)}/20000000,"")',
                             pct, 0.7578)
        col += 4 + (1 if gap else 0)
    for r in range(rows):
        if r != blank_label_row:
            ws.write(2 + r, 0, f"C{r + 1:03d}")
        ws.write(2 + r, 1, f"費目{r + 1}")
    wb.close()


def _block_records(chunks, block):
    return [c for c in chunks if c["content_type"] == "table_record" and f"領域「{block}」" in c["section_path"]]


@pytest.mark.parametrize("gap", [False, True], ids=["空き列なし", "空き列あり"])
def test_side_by_side_blocks_carry_row_labels_date_headers_and_thousand_display(tmp_path, gap):
    path = tmp_path / "blocks.xlsx"
    _block_book(path, gap=gap)
    chunks = _render(path).chunks
    for block in ("北陸", "東海", "関西"):
        records = _block_records(chunks, block)
        assert len(records) == 6
        third = next(c for c in records if "行見出し > 費目名: 費目3" in c["body"] or "費目名: 費目3" in c["body"])
        assert "費目コード: C003" in third["body"]
        assert "2025年4月: 15,157" in third["body"] and "通期実績: 17,157" in third["body"]
        assert "2025年4月のExcel原値: 15157000" in third["body"]
        assert "執行率: 75.8%" in third["body"]
    assert "00:00:00" not in "".join(c["body"] for c in chunks)
    assert not any(": 15,157,000," in c["body"] for c in chunks)


def test_row_labels_are_not_added_to_a_block_with_different_rows_or_blank_labels(tmp_path):
    short = tmp_path / "short.xlsx"
    _block_book(short, blocks=("北陸", "東海"), short_block="東海")
    chunks = _render(short).chunks
    assert all("費目名" not in c["body"] for c in _block_records(chunks, "東海"))
    assert all("費目名: 費目" in c["body"] for c in _block_records(chunks, "北陸"))


def test_independent_tables_side_by_side_get_no_row_labels(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["F1"] = "北陸", "東海"
    ws.merge_cells("A1:C1")
    ws.merge_cells("F1:H1")
    for first, base in (("A", 1), ("F", 6)):
        for offset, name in enumerate(["品目", "数量", "単価"]):
            ws.cell(2, base + offset, name)
        for row in range(3, 6):
            ws.cell(row, base, f"{first}{row}")
            ws.cell(row, base + 1, row)
            ws.cell(row, base + 2, row * 10)
    chunks = _render(_save(wb, tmp_path)).chunks
    assert not any("行見出し" in c["body"] for c in chunks)
    east = [c for c in chunks if c["content_type"] == "table_record" and "領域「東海」" in c["section_path"]]
    assert east and all("品目: F" in c["body"] for c in east)


@pytest.mark.parametrize("value,code,expected", [
    (13157000, "#,##0,", "13,157"),
    (13157000, "#,##0.0,", "13,157.0"),
    (1234567890, '#,##0,,"M"', "1,235M"),
    (500, "#,##0,", "1"),
    (13157000, "#,##0", "13,157,000"),
])
def test_trailing_comma_in_number_format_scales_displayed_value(value, code, expected):
    assert excel_display.format_display(value, code).value == expected


def test_extractor_version_reaches_evidence_signature_so_cached_display_values_are_rebuilt():
    assert f"xlsx={ooxml_arm.XLSX_EXTRACTOR_VERSION}" in office_md._current_evidence_ir_sig()


def test_blocks_under_title_and_note_rows_get_row_labels_even_with_empty_middle_rows(tmp_path):
    path = tmp_path / "titled.xlsx"
    wb = xlsxwriter.Workbook(path)
    ws = wb.add_worksheet("集計")
    month = wb.add_format({"num_format": 'yyyy"年"m"月"'})
    thousand = wb.add_format({"num_format": "#,##0,"})
    ws.write(0, 0, "2025年度 拠点別 月次経費集計・予算比")
    ws.write(1, 0, "（単位：千円）表示は千円未満切り捨て")
    ws.write(2, 0, "※ 執行率は通期実績÷通期予算")
    ws.merge_range(3, 0, 4, 0, "費目コード")
    ws.merge_range(3, 1, 4, 1, "費目名")
    for index, name in enumerate(("北陸", "東海", "関西")):
        col = 2 + index * 3
        ws.merge_range(3, col, 3, col + 2, name)
        ws.write_datetime(4, col, datetime.datetime(2025, 4, 1), month)
        ws.write(4, col + 1, "通期実績")
        ws.write(4, col + 2, "通期予算")
        for row in range(5, 10):
            if name == "関西" and row == 7:
                continue
            ws.write_number(row, col, 13157000 + row, thousand)
            ws.write_formula(row, col + 1, f"=SUM({xlsxwriter.utility.xl_rowcol_to_cell(row, col)})", thousand, 13157000 + row)
            ws.write_number(row, col + 2, 15000000, thousand)
    for row in range(5, 10):
        ws.write(row, 0, f"E-{110 + row}")
        ws.write(row, 1, f"設備保守費{row}")
    wb.close()
    chunks = _render(path).chunks
    for block in ("北陸", "東海", "関西"):
        records = _block_records(chunks, block)
        assert records
        for chunk in records:
            assert "行見出し > 費目コード: E-" in chunk["body"] and "行見出し > 費目名: 設備保守費" in chunk["body"]
            assert chunk["record_keys"][0]["value"].startswith("E-")
    third = next(c for c in _block_records(chunks, "東海") if "E-117" in c["body"])
    assert "2025年4月: 13,157" in third["body"]
    assert len(_block_records(chunks, "関西")) == 4


def test_blocks_with_total_row_whose_name_cell_is_empty_still_get_row_labels(tmp_path):
    path = tmp_path / "total.xlsx"
    wb = xlsxwriter.Workbook(path)
    ws = wb.add_worksheet("集計")
    thousand = wb.add_format({"num_format": "#,##0,"})
    ws.merge_range(0, 0, 1, 0, "費目コード")
    ws.merge_range(0, 1, 1, 1, "費目名")
    for index, name in enumerate(("北陸", "九州")):
        col = 2 + index * 2
        ws.merge_range(0, col, 0, col + 1, name)
        ws.write(1, col, "通期実績")
        ws.write(1, col + 1, "通期予算")
        for row in range(2, 6):
            ws.write_number(row, col, 13157000 + row, thousand)
            ws.write_number(row, col + 1, 14400000, thousand)
    for row in range(2, 5):
        ws.write(row, 0, f"E-{113 + row}")
        ws.write(row, 1, f"設備保守費{row}")
    ws.write(5, 0, "合計")
    wb.close()
    chunks = _render(path).chunks
    kyushu = _block_records(chunks, "九州")
    assert len(kyushu) == 4
    assert any("行見出し > 費目コード: E-115" in c["body"] and "行見出し > 費目名: 設備保守費2" in c["body"] for c in kyushu)
    assert any("行見出し > 費目コード: 合計" in c["body"] for c in kyushu)
