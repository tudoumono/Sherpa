"""Context IR のディスク保存（`ir/{rel}.context.json`）の単体テスト。

保存→読み戻しで同じ Context IR になること、読み戻した Context IR で描画しても出力が変わらないこと、
Evidence・解析規則・原本が変わったら読まずに作り直すこと、再描画（`refresh_rag`）が作り直さず読むことを固定する。
"""
from __future__ import annotations

import dataclasses
from pathlib import Path

import openpyxl
import pytest

from sherpa import worlds
from sherpa.ingest import context_ir, evidence_ir, evidence_render, office_md

ROOT = Path(__file__).resolve().parents[2]
INPUTS = ROOT / "fixtures" / "eval"
SAMPLES = (
    "excel_ja/inputs/JPX-016.xlsx",
    "office_ja/inputs/OJA-DOCX-EASY.docx",
    "office_ja/inputs/OJA-PPTX-EASY.pptx",
    "office_ja/inputs/OJA-PDF-EASY.pdf",
)


def _evidence(rel: str) -> tuple[Path, evidence_ir.EvidenceIR]:
    path = INPUTS / rel
    ir = office_md._extract_canonical_evidence(
        path, extraction_path=path, consume_legacy=False, legacy_conversion=None)
    return path, ir


@pytest.mark.parametrize("rel", SAMPLES)
def test_round_trip_keeps_context_and_render_output(tmp_path, rel):
    path, ir = _evidence(rel)
    built = evidence_render.render(ir, source_name=path.name)
    target = tmp_path / "x.context.json"
    context_ir.write_json_atomic(target, built.context, ir)
    loaded = context_ir.load_current(target, ir, source_name=path.name)
    assert loaded == built.context
    again = evidence_render.render(ir, source_name=path.name, context=loaded)
    assert (again.markdown, again.chunks) == (built.markdown, built.chunks)


def test_saved_bytes_are_deterministic(tmp_path):
    path, ir = _evidence(SAMPLES[0])
    context = context_ir.build(ir, source_name=path.name)
    assert context_ir.to_gzip_bytes(context, ir) == context_ir.to_gzip_bytes(context, ir)


def test_missing_or_broken_file_is_not_reused(tmp_path):
    path, ir = _evidence(SAMPLES[0])
    target = tmp_path / "x.context.json"
    assert context_ir.load_current(target, ir, source_name=path.name) is None
    target.write_bytes(b"not json")
    assert context_ir.load_current(target, ir, source_name=path.name) is None
    target.write_text('{"key": {}, "context": {}}', encoding="utf-8")
    assert context_ir.load_current(target, ir, source_name=path.name) is None


def test_changed_source_rules_or_evidence_are_not_reused(monkeypatch, tmp_path):
    path, ir = _evidence(SAMPLES[0])
    context = context_ir.build(ir, source_name=path.name)
    target = tmp_path / "x.context.json"
    context_ir.write_json_atomic(target, context, ir)
    assert context_ir.load_current(target, ir, source_name=path.name) is not None

    other_source = dataclasses.replace(ir, source=dataclasses.replace(ir.source, content_hash="sha256:other"))
    assert context_ir.load_current(target, other_source, source_name=path.name) is None
    assert context_ir.load_current(target, ir, source_name="renamed.xlsx") is None

    first = ir.elements[0]
    edited = dataclasses.replace(ir, elements=[dataclasses.replace(first, value="別の値"), *ir.elements[1:]])
    assert context_ir.load_current(target, edited, source_name=path.name) is None

    for attr in ("CONTEXT_IR_SCHEMA_VERSION", "CONTEXT_ANALYZER_VERSION", "IDENTIFIER_ROLE_ANALYZER_VERSION",
                 "IDENTIFIER_METADATA_SCHEMA_VERSION", "IDENTIFIER_MAX_MENTIONS_PER_CHUNK"):
        with monkeypatch.context() as m:
            m.setattr(context_ir, attr, "changed" if attr != "IDENTIFIER_MAX_MENTIONS_PER_CHUNK" else 7)
            assert context_ir.load_current(target, ir, source_name=path.name) is None, attr


def _world(tmp_path, monkeypatch, world_id: str, cell: str) -> Path:
    source = tmp_path / "kb" / world_id
    source.mkdir(parents=True)
    monkeypatch.setenv("SHERPA_DERIVED_DIR", str(tmp_path / "derived"))
    _write_book(source / "a.xlsx", cell)
    return source


def _write_book(path: Path, cell: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "項目ID", "表示名"
    ws["A2"], ws["B2"] = "A01", cell
    wb.save(path)


def test_build_derived_writes_context_into_ir_layer_only(monkeypatch, tmp_path):
    source = _world(tmp_path, monkeypatch, "ctx-ir-layer", "現行顧客区分")
    dmd = worlds.derived_md_dir("ctx-ir-layer")
    rep = office_md.build_derived(source, dmd, world_sig="s1")
    assert rep.get("error") is None, rep
    ir_dir = worlds.derived_ir_dir("ctx-ir-layer")
    assert (ir_dir / "a.xlsx.context.json").is_file()
    assert not list(source.rglob("*.context.json"))                      # 登録ディレクトリには書かない
    assert not list(dmd.rglob("*.context.json")) and not list(worlds.derived_rag_dir("ctx-ir-layer").rglob("*.context.json"))
    ir = evidence_ir.read_json_file(ir_dir / "a.xlsx.evidence.json")
    assert context_ir.load_current(ir_dir / "a.xlsx.context.json", ir, source_name="a.xlsx") is not None


def test_refresh_rag_reuses_saved_context_and_rebuilds_when_source_changes(monkeypatch, tmp_path):
    source = _world(tmp_path, monkeypatch, "ctx-refresh", "現行顧客区分")
    dmd = worlds.derived_md_dir("ctx-refresh")
    office_md.build_derived(source, dmd, world_sig="s1")
    ir_dir, rag_dir = worlds.derived_ir_dir("ctx-refresh"), worlds.derived_rag_dir("ctx-refresh")
    before = (rag_dir / "a.xlsx.rag.md").read_bytes(), (rag_dir / "a.xlsx.rag_chunks.jsonl").read_bytes()
    saved = (ir_dir / "a.xlsx.context.json").read_bytes()

    built: list = []
    real_build = context_ir.build
    monkeypatch.setattr(context_ir, "build", lambda *a, **kw: built.append(1) or real_build(*a, **kw))

    rep = office_md.refresh_rag(source, dmd, write_rag_sig_marker=True)
    assert rep["rag_failed"] == 0 and rep["rag_generated"] == 1
    assert built == []                                                    # 保存済みを読んだ（作り直していない）
    assert (rag_dir / "a.xlsx.rag.md").read_bytes() == before[0]
    assert (rag_dir / "a.xlsx.rag_chunks.jsonl").read_bytes() == before[1]
    assert (ir_dir / "a.xlsx.context.json").read_bytes() == saved

    _write_book(source / "a.xlsx", "廃止予定項目")                        # 原本が変わった
    rep = office_md.refresh_rag(source, dmd, write_rag_sig_marker=True)
    assert rep["rag_failed"] == 0
    assert built == [1]                                                   # 作り直した
    assert (ir_dir / "a.xlsx.context.json").read_bytes() != saved
    assert "廃止予定項目" in (rag_dir / "a.xlsx.rag.md").read_text(encoding="utf-8")


def test_refresh_rag_rebuilds_when_saved_context_is_missing(monkeypatch, tmp_path):
    """保存前に取り込まれた資料（`context.json` が無い）は、再描画のときに作って保存する。"""
    source = _world(tmp_path, monkeypatch, "ctx-missing", "現行顧客区分")
    dmd = worlds.derived_md_dir("ctx-missing")
    office_md.build_derived(source, dmd, world_sig="s1")
    target = worlds.derived_ir_dir("ctx-missing") / "a.xlsx.context.json"
    saved = target.read_bytes()
    target.unlink()
    rep = office_md.refresh_rag(source, dmd, write_rag_sig_marker=True)
    assert rep["rag_failed"] == 0 and target.read_bytes() == saved
