"""保守評価セット（eval/maint-v020）の問いと正解を、フィクスチャに照らして検証する。

正解の根拠（パスと行・セル範囲・段落）が実在すること、件数と種類の要件、フィクスチャが生成スクリプトと
一致することを固定する。モデルや API は呼ばない。説明: eval/maint-v020/README.md。
"""
from __future__ import annotations

import copy
import importlib.util
import sys
from pathlib import Path

import docx
import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[2]
EVAL_DIR = ROOT / "eval" / "maint-v020"


def _load(name: str):
    spec = importlib.util.spec_from_file_location(f"maint_v020_{name}", EVAL_DIR / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)
    return mod


verify = _load("verify")
build_fixtures = _load("build_fixtures")


@pytest.fixture(scope="module")
def dataset():
    return verify.load_systems(EVAL_DIR), verify.load_questions(EVAL_DIR)


def _content(path: Path):
    """ファイルの中身。xlsx・docx は ZIP のバイト列ではなく、セルと段落の内容で比べる（ライブラリの版に依らない）。"""
    if path.suffix == ".xlsx":
        wb = openpyxl.load_workbook(path)
        return [(ws.title, c.coordinate, c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None]
    if path.suffix == ".docx":
        d = docx.Document(path)
        return ([(p.style.name, p.text) for p in d.paragraphs if p.text.strip()]
                + [[c.text for c in r.cells] for t in d.tables for r in t.rows])
    return path.read_bytes()


def _contents(root: Path) -> dict[str, object]:
    return {p.relative_to(root).as_posix(): _content(p) for p in sorted(root.rglob("*")) if p.is_file()}


def test_every_question_and_its_evidence_is_verified(dataset):
    systems, questions = dataset
    problems = verify.check_all(ROOT, systems, questions) + verify.check_coverage(questions, systems)
    assert not problems, "\n".join(problems)


def test_requirements_of_the_set(dataset):
    systems, questions = dataset
    kinds = {k: sum(q["kind"] == k for q in questions) for k in verify.KINDS.values()}
    assert len(questions) >= 50 and all(n >= 10 for n in kinds.values()), kinds
    assert any(q["answerability"] == "unanswerable" for q in questions)
    assert {systems[q["system"]]["split"] for q in questions} == {"dev", "eval"}


def test_fixtures_match_the_generator(tmp_path):
    build_fixtures.build_all(tmp_path / "fixtures")
    assert _contents(tmp_path / "fixtures") == _contents(ROOT / "fixtures" / "eval_v020")


def _mutate_line(q):
    q["gold"]["evidence"][0]["at"] = {"type": "line", "start": 9999, "end": 9999}


def _mutate_contains(q):
    q["gold"]["evidence"][0]["contains"] = ["この語は原本のその箇所に無い"]


def _mutate_path(q):
    q["gold"]["evidence"][0]["path"] += ".missing"


def _mutate_scope(q):
    q["scope"] = ["fixtures/eval_v020/経費精算"]


def _mutate_claim_ref(q):
    q["gold"]["claims"][0]["supported_by"] = [["e99"]]


@pytest.mark.parametrize("mutate", [_mutate_line, _mutate_contains, _mutate_path, _mutate_scope, _mutate_claim_ref])
def test_a_broken_evidence_location_is_reported(dataset, mutate):
    systems, questions = dataset
    q = copy.deepcopy(next(x for x in questions if x["id"] == "SF-01"))
    mutate(q)
    assert verify.check_question(ROOT, q, systems)


def test_xlsx_docx_and_absence_locators_are_reported_when_wrong(dataset):
    systems, questions = dataset
    by_id = {q["id"]: q for q in questions}

    q = copy.deepcopy(by_id["SF-07"])
    for ev in q["gold"]["evidence"]:
        if ev["at"]["type"] == "cell":
            ev["at"]["sheet"] = "存在しないシート"
            break
    assert any("シートが無い" in p for p in verify.check_question(ROOT, q, systems))

    q = copy.deepcopy(by_id["SF-07"])
    for ev in q["gold"]["evidence"]:
        if ev["at"]["type"] == "paragraph":
            ev["at"]["index"] = 99
            break
    assert any("段落が範囲外" in p for p in verify.check_question(ROOT, q, systems))

    q = copy.deepcopy(by_id["SF-04"])
    q["gold"]["absence"][0]["pattern"] = "手数料率"
    assert any("「資料に無い」はずの" in p for p in verify.check_question(ROOT, q, systems))


def test_pair_conclusions_must_agree_with_the_expectation(dataset):
    systems, questions = dataset
    broken = copy.deepcopy(questions)
    variant = next(q for q in broken if q["id"] == "SF-16")
    variant["pair"]["expect"] = "same"
    assert any("SF-16" in p for p in verify.check_pairs(broken, systems))
