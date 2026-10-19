"""Excel の MD 化で、作り方を変えた中間表現・JSON の直列化が、従来の作り方と同じ結果になることの固定。

従来の作り方（`dataclasses.asdict` を経由する直列化・空セルまで作る走査・全文を読んでから型付けする読み戻し）を
テストの中に基準として持ち、同じ入力で結果が一致することを確かめる。派生物の byte が変わると、
再変換が走る・ID が食い違う・古い派生物が読めなくなる、のいずれかになる。
"""
from __future__ import annotations

import dataclasses
import datetime
import gzip
import json

import openpyxl
import pytest
from openpyxl.cell.cell import MergedCell
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink

from sherpa.ingest import context_ir, document_ir, evidence_ir, evidence_render, evidence_spike, excel_display
from sherpa.ingest.arms import ooxml_arm
from sherpa.ingest.ooxml import excel


def _compact(value) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _book(tmp_path, rows: int = 40):
    """見出し付きの表（識別子らしい値・日付・数式・結合セル・コメント・取り消し線・リンク・背景色を含む）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "一覧"
    ws.append(["ID", "名称", "呼び出し先", "日付", "金額"])
    for i in range(1, rows):
        ws.append([f"ID-{i:04d}", f"品名\"{i}\"\\n😀", f"JOB-{i} PRG_A.{i}", datetime.date(2024, 1 + i % 12, 1 + i % 27), i * 1.5])
        ws.cell(row=i + 1, column=5).number_format = '"¥"#,##0.00'
    ws["F2"] = "=SUM(E2:E5)"
    ws["B3"].comment = Comment("確認済み", "山田")
    ws["B4"].comment = Comment("   ", "空")
    ws["B5"].font = Font(strike=True)
    ws["B6"].hyperlink = "https://example.com/x"
    ws["B7"].hyperlink = Hyperlink(ref="B7", location="一覧!A1")
    ws["D3"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["D4"].fill = PatternFill("solid", fgColor="FFFFFF")
    ws.merge_cells("H1:J2")
    ws["H1"] = "結合"
    ws["H1"].fill = PatternFill("solid", fgColor="00FF00")
    ws["Z90"] = "遠い"
    wb.create_sheet("空").sheet_state = "hidden"
    path = tmp_path / "book.xlsx"
    wb.save(path)
    return path


def _evidence(path):
    evidence = evidence_spike.extract(path)
    excel_display.enrich_evidence(evidence, path)
    return evidence


def test_evidence_json_parts_match_asdict_serialization(tmp_path):
    ir = _evidence(_book(tmp_path))
    ir.elements[3].extension["odd"] = (1, {"b": 2, "a": (3, None)})          # asdict は tuple を tuple のまま複製する
    expected = (
        '{"coverage":[' + ",".join(_compact(dataclasses.asdict(v)) for v in ir.coverage)
        + '],"elements":[' + ",".join(_compact(dataclasses.asdict(v)) for v in ir.elements)
        + '],"parser_profile":' + _compact(ir.parser_profile)
        + ',"relations":[' + ",".join(_compact(dataclasses.asdict(v)) for v in ir.relations)
        + '],"schema_version":' + _compact(ir.schema_version)
        + ',"source":' + _compact(dataclasses.asdict(ir.source)) + "}\n"
    )
    assert "".join(evidence_ir._json_parts(ir)) == expected
    assert evidence_ir.to_json_str(ir) == expected


def test_evidence_json_with_nested_dataclass_falls_back_to_asdict(tmp_path):
    ir = _evidence(_book(tmp_path, rows=5))
    ir.elements[2].extension["nested"] = evidence_ir.Locator(part="p", extension={"k": 1})
    expected_element = _compact(dataclasses.asdict(ir.elements[2]))
    assert expected_element in evidence_ir.to_json_str(ir)


def test_evidence_read_matches_original_across_read_chunk_boundaries(tmp_path, monkeypatch):
    ir = _evidence(_book(tmp_path))
    target = tmp_path / "x.evidence.json"
    evidence_ir.write_json_atomic(target, ir)
    monkeypatch.setattr(evidence_ir, "_READ_CHUNK", 97)                  # 要素・文字列・数値の途中で必ず読み足しが起きる大きさ
    loaded = evidence_ir.read_json_file(target)
    assert loaded == ir
    assert evidence_ir.from_json_str(evidence_ir.to_json_str(ir)) == ir


@pytest.mark.parametrize("batch_items,read_chunk", [(3, 97), (1, 64), (256, 1 << 20), (7, 1 << 20)])
def test_evidence_read_batches_elements_without_changing_the_result(tmp_path, monkeypatch, batch_items, read_chunk):
    ir = _evidence(_book(tmp_path))
    lookalike = next(i for i, e in enumerate(ir.elements) if e.type == "cell")
    ir.elements[lookalike] = dataclasses.replace(ir.elements[lookalike], value='x,{"coverage_id":"y"},z')
    target = tmp_path / "x.evidence.json"
    evidence_ir.write_json_atomic(target, ir)
    monkeypatch.setattr(evidence_ir, "_BATCH_ITEMS", batch_items)
    monkeypatch.setattr(evidence_ir, "_READ_CHUNK", read_chunk)
    assert evidence_ir.read_json_file(target) == ir


def test_evidence_read_accepts_json_written_with_other_formatting(tmp_path):
    ir = _evidence(_book(tmp_path, rows=8))
    pretty = json.dumps(json.loads(evidence_ir.to_json_str(ir)), ensure_ascii=False, indent=2)
    assert evidence_ir.from_json_str(pretty) == ir


def test_evidence_read_rejects_truncated_trailing_and_shape_errors(tmp_path):
    ir = _evidence(_book(tmp_path, rows=5))
    text = evidence_ir.to_json_str(ir)
    with pytest.raises(ValueError):
        evidence_ir.from_json_str(text[:-40])
    with pytest.raises(ValueError):
        evidence_ir.from_json_str(text + "x")
    payload = json.loads(text)
    del payload["elements"][0]["locator"]
    with pytest.raises(ValueError):
        evidence_ir.from_json_str(json.dumps(payload))
    broken = tmp_path / "broken.evidence.json"
    broken.write_bytes(gzip.compress(text.encode("utf-8"))[:-20])
    with pytest.raises(OSError):
        evidence_ir.read_json_file(broken)


@pytest.mark.parametrize("locator", [
    evidence_ir.Locator(part="xl/workbook.xml", sheet="シート\"1\"", cell_range="A1",
                        extension={"row": 1, "column": 2, "row_span": 1, "column_span": 1}),
    evidence_ir.Locator(part="xl/workbook.xml", sheet=None, cell_range=None, object_id="sheet:1",
                        extension={"document_ir_source_map": {"sheet": "明細"}}),
    evidence_ir.Locator(part="ppt/slides/slide1.xml", slide=2, bbox=[1, 2.5, 3, 4], extension={}),
    evidence_ir.Locator(part="word/document.xml", object_id=7, extension={"flag": True, "none": None, "ふ": "é\n"}),
    evidence_ir.Locator(part="x", extension={"ratio": 0.5}),
])
def test_element_id_fast_path_equals_generic_id(locator):
    expected = evidence_spike._stable_id("evidence", "cell", dataclasses.asdict(locator), 3)
    assert evidence_spike._element_id("cell", locator, 3) == expected


def test_document_json_matches_asdict_serialization(tmp_path):
    document = ooxml_arm._build_xlsx_ir(_book(tmp_path))
    assert document is not None
    assert document_ir.to_json_str(document) == json.dumps(
        dataclasses.asdict(document), ensure_ascii=False, sort_keys=True, indent=2) + "\n"


@pytest.mark.parametrize("rows", [20, 700])                                # 512 件ごとの分割をまたぐ件数を含む
def test_context_json_matches_asdict_serialization(tmp_path, rows):
    path = _book(tmp_path, rows=rows)
    ir = _evidence(path)
    context = context_ir.build(ir, source_name=path.name)
    expected = json.dumps(
        {"key": context_ir._current_key(ir, context.source_name), "context": dataclasses.asdict(context)},
        ensure_ascii=False, separators=(",", ":")) + "\n"
    assert context_ir.to_json_str(context, ir) == expected


def _dense(ws, max_row, max_col):
    return [cell for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col) for cell in row]


def test_sparse_scans_match_dense_scans(tmp_path):
    wb = openpyxl.load_workbook(_book(tmp_path), data_only=True)
    ws = wb["一覧"]
    max_row, max_col = min(ws.max_row, 200), min(ws.max_column, 30)
    resolver = excel.ColorResolver(wb)
    anchor_of = {(r, c): (m.min_row, m.min_col) for m in ws.merged_cells.ranges
                 for r in range(m.min_row, m.max_row + 1) for c in range(m.min_col, m.max_col + 1)}
    expected_filled = set()
    for c in _dense(ws, max_row, max_col):                                # 空セルまで作って数える従来の走査
        fill = c.fill
        if isinstance(c, MergedCell) and (c.row, c.column) in anchor_of:
            fill = ws.cell(*anchor_of[(c.row, c.column)]).fill
        if excel._fill_is_colored(fill, resolver):
            expected_filled.add((c.row, c.column))
    assert excel.filled_cells(ws, max_row, max_col, resolver=resolver) >= {(3, 4), (1, 8)}
    expected_strike = sorted(
        (c.row, c.column) for c in _dense(ws, max_row, max_col)
        if c.font is not None and c.font.strike and c.value is not None)
    assert [(e["row"], e["column"]) for e in excel.strike_cells(ws)] == expected_strike == [(5, 2)]
    assert [(e["row"], e["column"]) for e in excel.cell_comments(ws)] == [(3, 2)]
    assert [(e["row"], e["column"]) for e in excel.cell_hyperlinks(ws)] == [(6, 2), (7, 2)]
    grid = excel.value_grid(ws, max_row, max_col)
    assert grid == [list(row) for row in ws.iter_rows(
        min_row=1, max_row=max_row, min_col=1, max_col=max_col, values_only=True)]
    assert excel.filled_cells(ws, max_row, max_col, resolver=resolver) == expected_filled


@pytest.mark.parametrize("text", [
    "", "abc", "12-34", "2024-01-05", "ABC-123", "ＡＢＣ－１２３", "x1", "A/B", "SYS:ID01", "T-1,T-2", "呼出 = JOB-01、JOB-02",
    "１２３４", "ａ", "a_b", "品名12-1", "a" * 300 + "1", "1.5", "No.1", "ＡＢ", "id", "①②",
])
def test_identifier_prefilter_never_hides_a_mention(text):
    if not context_ir.has_identifier_candidate(text):
        batch = context_ir.identifier_mentions_with_metadata(
            text, field_label="ID", evidence_id="e", locator={"part": "p"}, allow_record_identity=True)
        assert batch.mention_count == 0 and batch.mentions == ()


def test_context_is_the_same_without_the_candidate_prefilter(tmp_path, monkeypatch):
    path = _book(tmp_path)
    ir = _evidence(path)
    with_filter = context_ir.to_json_str(context_ir.build(ir, source_name=path.name), ir)
    monkeypatch.setattr(context_ir, "has_identifier_candidate", lambda text: True)
    without_filter = context_ir.to_json_str(context_ir.build(ir, source_name=path.name), ir)
    assert with_filter == without_filter
    rendered = evidence_render.render(ir, source_name=path.name)
    assert any(chunk["identifier_mentions"] for chunk in rendered.chunks)


def test_enrich_marks_cells_of_a_missing_sheet_as_not_found(tmp_path):
    path = _book(tmp_path, rows=5)
    ir = evidence_spike.extract(path)
    moved = next(e for e in ir.elements if e.type == "cell")
    ir.elements[ir.elements.index(moved)] = dataclasses.replace(
        moved, locator=dataclasses.replace(moved.locator, sheet="無いシート"))
    excel_display.enrich_evidence(ir, path)
    ghost = next(e for e in ir.elements if e.locator.sheet == "無いシート")
    assert ghost.extension["display_reason"] == "cell_not_found"
    assert ghost.extension["display_status"] == "unsupported"


def _formula_book(tmp_path, name, mutate=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    for row in range(1, 6):
        ws.cell(row=row, column=1, value=row * 10)
    path = tmp_path / name
    wb.save(path)
    if mutate is not None:
        import zipfile
        with zipfile.ZipFile(path) as source:
            parts = {entry: source.read(entry) for entry in source.namelist()}
        parts["xl/worksheets/sheet1.xml"] = mutate(parts["xl/worksheets/sheet1.xml"].decode("utf-8")).encode("utf-8")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
            for entry, body in parts.items():
                target.writestr(entry, body)
    return path


def _formulas_of(path):
    wb_values, wb_formula = excel.load_two(path)
    try:
        return excel.formulas(wb_formula["S"], wb_values["S"])
    finally:
        wb_values.close()
        wb_formula.close()


@pytest.mark.parametrize("name,mutate,expect_count", [
    ("plain.xlsx", None, 0),
    ("formula.xlsx", lambda x: x.replace('<c r="A1" t="n"><v>10</v></c>', '<c r="A1"><f>A2*2</f><v>20</v></c>'), 1),
    ("prefixed.xlsx", lambda x: x.replace('<c r="A1" t="n"><v>10</v></c>', '<c r="A1"><x:f>A2*2</x:f><v>20</v></c>'), None),
    ("strvalue.xlsx", lambda x: x.replace('<c r="A1" t="n"><v>10</v></c>', '<c r="A1" t="str"><v>=A2</v></c>'), 1),
    ("inline.xlsx", lambda x: x.replace('<c r="A1" t="n"><v>10</v></c>', '<c r="A1" t="inlineStr"><is><t>=A2</t></is></c>'), 1),
    ("entity.xlsx", lambda x: x.replace('<c r="A1" t="n"><v>10</v></c>', '<c r="A1" t="str"><v>&#61;A2</v></c>'), 1),
])
def test_formula_scan_skip_never_hides_a_formula(tmp_path, monkeypatch, name, mutate, expect_count):
    path = _formula_book(tmp_path, name, mutate)
    skipped = None
    try:
        skipped = _formulas_of(path)
    except Exception:                                   # 名前空間つきの接頭辞は openpyxl が読めないことがある（比較は両方とも同じ結果になること）
        skipped = "error"
    monkeypatch.setattr(excel, "_may_have_formulas", lambda ws: True)
    try:
        full = _formulas_of(path)
    except Exception:
        full = "error"
    assert skipped == full
    if expect_count is not None:
        assert len(full) == expect_count


def test_formula_scan_skip_with_shared_string_starting_with_equals(tmp_path, monkeypatch):
    wb = openpyxl.Workbook()
    wb.active.title = "S"
    wb.active["A1"] = "=abc"
    wb.active["A2"] = "plain"
    path = tmp_path / "sst.xlsx"
    wb.save(path)
    with_skip = _formulas_of(path)
    monkeypatch.setattr(excel, "_may_have_formulas", lambda ws: True)
    assert with_skip == _formulas_of(path)
    assert [item["cell"] for item in with_skip] == ["A1"]
