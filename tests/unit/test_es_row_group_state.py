"""束ねたベクトルの入力（表の行のまとまりの `vector_text`）が、値を短縮しても状態の種類と対象を残す契約。

実際の抽出→描画→JSONL→`_iter_rag_chunk_entries` を通す（実 ES・実モデルなし）。
設計: docs/design/rag.md「表の行・本文のまとまり」
"""
from __future__ import annotations

import hashlib
import os
import re
import shutil
import zipfile
from pathlib import Path

os.environ.setdefault("SHERPA_USE_FIXTURES", "1")

import pytest

from sherpa import es_index, worlds
from sherpa.ingest import office_md

_NS_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_ROOT = Path(__file__).resolve().parents[2]
_MARKERS_XLSX = _ROOT / "fixtures" / "eval" / "deprecation_markers" / "inputs" / "DEP-XLSX-MARKERS.xlsx"


def _w(text: str, strike: bool = False) -> str:
    pr = "<w:rPr><w:strike/></w:rPr>" if strike else ""
    return f"<w:r>{pr}<w:t>{text}</w:t></w:r>"


def _wr(text: str, *props: str) -> str:
    return f"<w:r><w:rPr>{''.join(props)}</w:rPr><w:t>{text}</w:t></w:r>"


def _tr(*cells: str) -> str:
    return "<w:tr>" + "".join(f"<w:tc><w:p>{c}</w:p></w:tc>" for c in cells) + "</w:tr>"


def _vector_texts(monkeypatch, tmp_path, name: str, build) -> list[str]:
    """`build(root)` が置いた原本を取り込み、表の行の束（列名の行を持つもの）の先頭の `vector_text` を出現順で返す。"""
    root = tmp_path / "kb" / "w1"
    root.mkdir(parents=True)
    monkeypatch.setattr(worlds, "world_dir", lambda w: root if w == "w1" else None)
    monkeypatch.setenv("SHERPA_DERIVED_DIR", str(tmp_path / "derived"))
    build(root)
    rep = office_md.build_derived(root, worlds.derived_md_dir("w1"))
    assert rep["rag_failed"] == 0, rep
    rag = worlds.derived_rag_dir("w1")
    jsonl, md_path = rag / f"{name}.rag_chunks.jsonl", rag / f"{name}.rag.md"
    anchors, reason = es_index._load_rag_md_anchors(md_path)
    assert reason is None
    return [body["vector_text"]
            for _i, body, _t, skip in es_index._iter_rag_chunk_entries(jsonl, name, anchors, {"doc_id": name})
            if not skip and "vector_text" in body and any(ln.startswith("列: ") for ln in body["vector_text"].split("\n"))]


def _docx_table(root: Path, rows: list[str]) -> None:
    table = "<w:tbl>" + _tr(_w("項目"), _w("内容")) + "".join(rows) + "</w:tbl>"
    with zipfile.ZipFile(root / "表.docx", "w") as z:
        z.writestr("word/document.xml", f'<?xml version="1.0"?><w:document {_NS_W}><w:body>'
                   f"<w:p>{_w('本文。')}</w:p>{table}</w:body></w:document>")


def test_partial_strike_in_a_long_cell_keeps_kind_and_target_in_the_bundle_vector_text(monkeypatch, tmp_path):
    cell = _w("通常の文字" * 10 + "あい") + _w("旧い承認条件", strike=True) + _w("現在は個別承認。")

    def build(root):
        _docx_table(root, [_tr(_w("A01"), cell), _tr(_w("A02"), _w("通常のデータ"))])

    (vector,) = _vector_texts(monkeypatch, tmp_path, "表.docx", build)
    rows = vector.split("\n")
    row = next(r for r in rows if r.startswith("A01"))
    assert "旧い承認条件" in row and "取り消し線" in row                       # 状態の種類と対象
    assert row.startswith("A01 | 通常の文字通常の文字")                          # 通常の値の短縮は残る
    assert "セル全体" not in row and "「あり」" not in row                      # セル全体の状態にはしない
    assert any(r.startswith("A02 | 通常のデータ") for r in rows)
    assert all(len(r.encode("utf-8")) <= es_index._ROW_GROUP_ROW_MAX_BYTES for r in rows[2:])


def test_unresolved_overlap_and_whole_cell_strike_lines_survive_the_cut(monkeypatch, tmp_path):
    def build(root):
        shutil.copy(_MARKERS_XLSX, root / _MARKERS_XLSX.name)

    (vector,) = _vector_texts(monkeypatch, tmp_path, _MARKERS_XLSX.name, build)
    rows = vector.split("\n")
    item01 = next(r for r in rows if r.startswith("ITEM01"))
    item02 = next(r for r in rows if r.startswith("ITEM02"))
    assert "「あり」" in item01                                              # セル全体の取り消し線の行
    assert "前面図形" in item02 and "未確定" in item02                       # 覆われているかは未確定、が落ちない
    assert all(len(r.encode("utf-8")) <= es_index._ROW_GROUP_ROW_MAX_BYTES for r in rows[2:])
    assert len(vector.encode("utf-8")) <= es_index._ROW_GROUP_MAX_BYTES


def test_state_lines_stay_inside_the_row_budget_when_ordinary_text_is_long(monkeypatch, tmp_path):
    cells = [_w(f"{i}列目の長い通常の文字" * 6) + _w(f"旧条件{i}", strike=True) + _w("現行") for i in range(1, 4)]

    def build(root):
        table = ("<w:tbl>" + _tr(_w("項目"), _w("甲"), _w("乙"), _w("丙"))
                 + _tr(_w("A01"), *cells) + _tr(_w("A02"), _w("a"), _w("b"), _w("c")) + "</w:tbl>")
        with zipfile.ZipFile(root / "表.docx", "w") as z:
            z.writestr("word/document.xml", f'<?xml version="1.0"?><w:document {_NS_W}><w:body>'
                       f"<w:p>{_w('本文。')}</w:p>{table}</w:body></w:document>")

    (vector,) = _vector_texts(monkeypatch, tmp_path, "表.docx", build)
    row = next(r for r in vector.split("\n") if r.startswith("A01"))
    assert len(row.encode("utf-8")) <= es_index._ROW_GROUP_ROW_MAX_BYTES
    assert row.count("取り消し線") >= 3 and all(f"旧条件{i}" in row for i in (1, 2, 3))


def test_plain_text_in_a_column_named_like_a_state_line_is_cut_as_before():
    long_value = "確認済みの承認手順に従って、担当者が内容を確認し直す必要があるため注意する。"
    assert len(long_value) > es_index._ROW_GROUP_CELL_MAX_CHARS
    fields = [("項目", "A01"), ("状態", long_value)]
    columns = es_index._row_group_columns(fields)
    assert es_index._row_group_line({}, fields, columns) == "A01 | " + long_value[:31] + "…"


def test_overlap_line_keeps_its_unresolved_tail_even_when_a_wide_row_overflows_the_budget():
    ref = "ERROR_NODE（図形 ID 31・シート「統合設計」のセル範囲 AL8:AS9）"
    columns = [(f"列{i}", 1) for i in range(10)]
    fields = [(f"列{i}", f"通常の長い値{i}" * 8) for i in range(10)]
    fields += [("前面図形", ref + "（覆われているかは未確定）", True), ("前面図形", ref + "（覆われているかは未確定）", True)]
    line = es_index._row_group_line({}, fields, columns)
    assert len(line.encode("utf-8")) <= es_index._ROW_GROUP_ROW_MAX_BYTES
    assert line.count("（覆われているかは未確定）") == 2


def test_short_state_line_missing_from_the_head_row_still_gets_its_heading(monkeypatch, tmp_path):
    import openpyxl
    from openpyxl.styles import Font

    def build(root):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "明細"
        ws["A1"], ws["B1"] = "項目", "内容"
        ws["A2"], ws["B2"] = "A01", "通常"
        ws["A3"], ws["B3"] = "A02", "旧記述"
        ws["B3"].font = Font(strike=True)
        wb.save(root / "s.xlsx")

    (vector,) = _vector_texts(monkeypatch, tmp_path, "s.xlsx", build)
    row = next(r for r in vector.split("\n") if r.startswith("A02"))
    assert row == "A02 | 旧記述 | 取り消し線: 「あり」"
    assert next(r for r in vector.split("\n") if r.startswith("A01")) == "A01 | 通常"


def test_every_partial_state_note_is_kept_while_the_row_has_room(monkeypatch, tmp_path):
    cell = (_w("通常の文字" * 10 + "あい") + _w("甲", strike=True) + _w("、") + _wr("乙", "<w:vanish/>") + _w("、")
            + _wr("丙", "<w:strike/>", "<w:vanish/>") + _w("。"))

    def build(root):
        _docx_table(root, [_tr(_w("A01"), cell), _tr(_w("A02"), _w("通常のデータ"))])

    (vector,) = _vector_texts(monkeypatch, tmp_path, "表.docx", build)
    row = next(r for r in vector.split("\n") if r.startswith("A01"))
    assert "（一部「甲」に取り消し線）" in row and "（一部「乙」に非表示）" in row and "（一部「丙」に非表示・取り消し線）" in row
    assert len(row.encode("utf-8")) <= es_index._ROW_GROUP_ROW_MAX_BYTES


def test_notes_that_do_not_fit_are_counted_instead_of_silently_dropped():
    brackets = "".join(f"【取り消し線：{'対象' * 6}{i}】" for i in range(12))
    value = "通常の文字" * 10 + brackets
    fields = [("項目", "A01"), ("内容", value)]
    line = es_index._row_group_line({}, fields, es_index._row_group_columns(fields))
    assert len(line.encode("utf-8")) <= es_index._ROW_GROUP_ROW_MAX_BYTES
    assert "（一部「対象対象対象対象対象対象0」に取り消し線）" in line
    kept = line.count("（一部「")
    assert 1 <= kept < 12 and f"（ほか{12 - kept}件）" in line


def test_plain_column_named_like_a_connection_line_gets_no_state_heading():
    columns = [("項目", 1)]
    fields = [("項目", "A01"), ("接続先", "通常値")]
    assert es_index._row_group_line({}, fields, columns) == "A01 | 通常値"
    fields = [("項目", "A01"), ("接続先", "「通常値」")]
    assert es_index._row_group_line({}, fields, columns) == "A01 | 「通常値」"
    rendered = [("項目", "A01", False), ("接続先", "図形A（始点→終点の指定。矢印の向きは未確認）", True)]
    assert es_index._row_group_line({}, rendered, columns) == "A01 | 接続先: 図形A（始点→終点の指定。矢印の向きは未確認）"


def test_summary_of_dropped_notes_survives_when_every_note_is_dropped():
    columns = [(f"列{i}", 1) for i in range(32)]
    fields = [(f"列{i}", "abcdef") for i in range(31)] + [("列31", "x" * 33 + "【取り消し線：対象】")]
    line = es_index._row_group_line({}, fields, columns)
    assert len(line.encode("utf-8")) <= es_index._ROW_GROUP_ROW_MAX_BYTES
    assert "（ほか1件）" in line


# ---- 最終の予算配分の後でも、状態と対象・観測と「未確定」が組で残り、省略は件数で示される ----

def _balanced(row: str) -> bool:
    """行の中の状態の括りが、開いたら必ず閉じている（途中で切れていない）。"""
    return row.count("【") == row.count("】")


def test_short_partial_state_in_an_early_column_survives_the_row_budget(monkeypatch, tmp_path):
    first = _w("通常") + _w("旧A", strike=True) + _w("現在")
    middle = [_w("別の通常説明です。" * 4) for _ in range(8)]
    last = _w("別の通常説明です。" * 4) + _w("旧条件B", strike=True)

    def build(root):
        table = ("<w:tbl>" + _tr(*[_w(f"見出し{i}") for i in range(10)])
                 + _tr(first, *middle, last) + _tr(*[_w(f"値{i}") for i in range(10)]) + "</w:tbl>")
        with zipfile.ZipFile(root / "表.docx", "w") as z:
            z.writestr("word/document.xml", f'<?xml version="1.0"?><w:document {_NS_W}><w:body>'
                       f"<w:p>{_w('本文。')}</w:p>{table}</w:body></w:document>")

    (vector,) = _vector_texts(monkeypatch, tmp_path, "表.docx", build)
    row = next(r for r in vector.split("\n") if r.startswith("通常"))
    assert "通常【取り消し線：旧A】現在" in row                          # 短い括りは括りごと残る
    assert "旧条件B" in row and "取り消し線" in row                      # 後ろの括りも対象とともに残る
    assert _balanced(row) and len(row.encode("utf-8")) <= es_index._ROW_GROUP_ROW_MAX_BYTES


@pytest.mark.parametrize("name_suffix", ["", "&#10;補足"])
def test_every_overlap_line_keeps_its_unresolved_tail_or_is_counted(monkeypatch, tmp_path, name_suffix):
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    from test_occlusion_fill_kinds import _A_NS, _GRADIENT, _PNG, _R_NS, _XDR_NS, _write_zip

    def build(root):
        seed = tmp_path / "seed.png"
        seed.write_bytes(_PNG)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "明細"
        ws["A1"], ws["B1"] = "項目", "内容"
        ws["A2"], ws["B2"] = "A01", "通常"
        ws["A3"], ws["B3"] = "A02", "後ろの行の値"
        ws.add_image(XLImage(str(seed)), "B3")
        path = tmp_path / "front.xlsx"
        wb.save(path)
        with zipfile.ZipFile(path) as z:
            entries = {name: z.read(name) for name in z.namelist()}
        drawing = next(n for n in entries if n.startswith("xl/drawings/drawing") and n.endswith(".xml"))
        anchors = "".join(
            '<xdr:twoCellAnchor><xdr:from><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>2</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
            '<xdr:to><xdr:col>2</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>3</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
            f'<xdr:sp><xdr:nvSpPr><xdr:cNvPr id="{i}" name="説明図形{i}{name_suffix}"/><xdr:cNvSpPr/></xdr:nvSpPr>'
            f"<xdr:spPr>{_GRADIENT}</xdr:spPr></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>"
            for i in range(1, 5))
        entries[drawing] = (f'<?xml version="1.0"?><xdr:wsDr xmlns:xdr="{_XDR_NS}" xmlns:a="{_A_NS}" xmlns:r="{_R_NS}">'
                            f"{anchors}</xdr:wsDr>").encode()
        _write_zip(root / "s.xlsx", entries)

    (vector,) = _vector_texts(monkeypatch, tmp_path, "s.xlsx", build)
    row = next(r for r in vector.split("\n") if r.startswith("A02"))
    assert len(row.encode("utf-8")) <= es_index._ROW_GROUP_ROW_MAX_BYTES
    tail = "（覆われているかは未確定）"
    lines = re.findall(r"前面図形: [^|]*", row)
    assert lines and all(line.rstrip().endswith(tail) for line in lines)      # 残った観測の行は「未確定」まで組で残る
    omitted = re.search(r"（ほか(\d+)件）", row)
    assert len(lines) + (int(omitted.group(1)) if omitted else 0) == 4        # 外した分は件数で示される


def test_wide_row_keeps_state_pairs_and_counts_what_it_omits():
    columns = [(f"列{i}", 1) for i in range(40)]
    fields = [(f"列{i}", "通常の値" * 3, False) for i in range(40)]
    fields[3] = ("列3", "通常" + "【取り消し線：旧A】" + "現在", False)
    fields += [("前面図形", f"図形{i}（図形 ID {i}・シート「明細」のセル範囲 B3）（覆われているかは未確定）", True) for i in range(6)]
    line = es_index._row_group_line({}, fields, columns)
    assert len(line.encode("utf-8")) <= es_index._ROW_GROUP_ROW_MAX_BYTES
    assert _balanced(line)
    omitted = re.search(r"（ほか(\d+)件）", line)
    assert omitted and int(omitted.group(1)) == 7                              # 外した括り 1 件と観測の行 6 件が件数で示される
    assert line.count("前面図形: ") == line.count("（覆われているかは未確定）") == 0


def test_moderately_wide_row_keeps_every_pair_without_a_summary():
    columns = [(f"列{i}", 1) for i in range(12)]
    fields = [(f"列{i}", "通常の値" * 3, False) for i in range(12)]
    fields[3] = ("列3", "通常" + "【取り消し線：旧A】" + "現在", False)
    fields += [("前面図形", f"図形{i}（図形 ID {i}・シート「明細」のセル範囲 B3）（覆われているかは未確定）", True) for i in range(2)]
    line = es_index._row_group_line({}, fields, columns)
    assert len(line.encode("utf-8")) <= es_index._ROW_GROUP_ROW_MAX_BYTES and "（ほか" not in line
    assert "通常【取り消し線：旧A】現在" in line
    assert line.count("前面図形: ") == line.count("（覆われているかは未確定）") == 2


def test_original_value_written_like_the_summary_is_kept_and_the_summary_is_single():
    columns = [(f"列{i}", 1) for i in range(40)]
    ref = "図形A（図形 ID 1・シート「明細」のセル範囲 B3）（覆われているかは未確定）"
    for literal, summaries in (("（ほか1件）", 2), ("（ほか7件）", 1)):
        fields = [(f"列{i}", literal if i == 5 else "通常の値" * 3, False) for i in range(40)]
        fields += [("前面図形", ref, True)]
        line = es_index._row_group_line({}, fields, columns)
        assert len(line.encode("utf-8")) <= es_index._ROW_GROUP_ROW_MAX_BYTES
        assert line.count(literal) == 1 + (summaries - 1 if literal == "（ほか1件）" else 0)   # 原本の値は消えない
        assert line.count("（ほか1件）") == summaries                                              # 要約は 1 件だけ


def test_names_and_quotes_in_observation_lines_never_break_the_line_in_rag_and_human_wording():
    from sherpa.ingest import observation_vocab as vocab

    ref = "説明図形1\n補足（図形 ID 1・シート「架空表」のセル範囲 B3:B3）"
    for human in (False, True):
        line = vocab.say("cover_back_undetermined", human=human, ref=ref, subject="B3")
        assert "\n" not in line and "説明図形1 補足（図形 ID 1" in line and line.endswith("（覆われているかは未確定）")
        assert "\n" not in vocab.say("notice_picture", human=human, name="図\n1")
    assert vocab.say("anchor_front", value="甲\r\n乙") == "前面図形: 「甲 乙」"
    assert vocab.say("overlap_text", value="本文の\n引用") == "重なり: 「本文の 引用」"


# ---- 原本の通常の列を観測の行と取り違えない ----

def test_plain_state_column_whose_value_looks_quoted_is_cut_as_before(monkeypatch, tmp_path):
    value = "「" + "通常の値で特別な書式の状態はありません。" * 4 + "」"

    def build(root):
        table = "<w:tbl>" + _tr(_w("項目"), _w("状態")) + _tr(_w("A01"), _w(value)) + _tr(_w("A02"), _w("短い")) + "</w:tbl>"
        with zipfile.ZipFile(root / "表.docx", "w") as z:
            z.writestr("word/document.xml", f'<?xml version="1.0"?><w:document {_NS_W}><w:body>'
                       f"<w:p>{_w('本文。')}</w:p>{table}</w:body></w:document>")

    (vector,) = _vector_texts(monkeypatch, tmp_path, "表.docx", build)
    assert "A01 | " + value[:31] + "…" in vector.split("\n")
    assert "状態: " not in vector


def test_chunks_without_any_observation_line_do_not_get_the_provenance_field(monkeypatch, tmp_path):
    from sherpa import json_io

    def build(root):
        _docx_table(root, [_tr(_w("A01"), _w("通常")), _tr(_w("A02"), _w("通常"))])

    _vector_texts(monkeypatch, tmp_path, "表.docx", build)
    jsonl = worlds.derived_rag_dir("w1") / "表.docx.rag_chunks.jsonl"
    assert all("observation_lines" not in line for line in json_io.read_text_maybe_gzip(jsonl).splitlines())


# ---- 状態の無い表の入力は、バイト単位で変えない ----

_GOLDEN_STATE_FREE = "c7bada9bac7ef03ffd1932d6c8e58bfc43883a70faf3a07c622dd2cb4ab57c72"


def test_table_without_state_keeps_the_bundle_vector_text_byte_identical(monkeypatch, tmp_path):
    def build(root):
        _docx_table(root, [
            _tr(_w("A01"), _w("通常の文字" * 12 + "。【補足：旧い運用の注記】" + "続き" * 20)),
            _tr(_w("A02"), _w("短い値")),
            _tr(_w("A03"), _w("確認済みの承認手順に従って、担当者が内容を確認し直す必要がある。")),
        ])

    (vector,) = _vector_texts(monkeypatch, tmp_path, "表.docx", build)
    assert hashlib.sha256(vector.encode("utf-8")).hexdigest() == _GOLDEN_STATE_FREE
