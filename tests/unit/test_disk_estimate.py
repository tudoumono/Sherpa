"""取り込み前のディスク容量の見積もり（`sherpa.disk_estimate`）。

外部境界（`shutil.disk_usage`・docker・ES）だけを差し替え、原本の木・ハードリンク・st_dev は実ファイルで確かめる。
実害の再現: 同じディスクに乗る置き場の合算漏れ、更新で今の使用分を二重に数える、確かめられない置き場を「足りる」と出す、
見積もりの失敗で取り込みを妨げる、圧縮された OOXML の派生物とチャンク数を原本の大きさから外す、の 5 点。
"""
from __future__ import annotations

import logging
import os
import random
import subprocess
import time
import zipfile

import pytest

from sherpa import disk_estimate as de
from sherpa import es_index, json_io
from sherpa.ingest import office_md

MB = de.MB
GIB = de.GIB


def _counts(**fam):
    c = de._new_counts()
    for name, (n, size) in fam.items():
        c["families"][name] = [n, size]
        if name == "legacy":
            c["legacy_max"] = size // max(n, 1)
    return c


@pytest.fixture(autouse=True)
def _no_embed_no_ocr(monkeypatch):
    monkeypatch.setenv("SHERPA_DISABLE_EMBED", "1")
    monkeypatch.setenv("SHERPA_OCR_ENABLED", "0")


# ---- 見積もりの計算 ----

def test_register_estimate_follows_family_coefficients():
    # PDF 100MB・50 件（ベクトル無し）: 派生物 = 2.4×100MB + 40KB×50 + 512B×50 を 2% 増し、ES = チャンク×1KB を取り込み中の最大（1.5 倍）。
    need = de.compute_needs(_counts(pdf=(50, 100 * MB)), None, mode="register", dim=0)
    derived = (2.4 * 100 * MB + 50 * de._DERIVED_FILE_BYTES + 50 * de._DERIVED_ALL_FILE_BYTES) * 1.02
    assert need["derived"] == pytest.approx(derived, rel=1e-6)
    chunks = 1900 * 100 + 4 * 50
    assert need["es"] == pytest.approx(chunks * 1000 * 1.5, rel=1e-6)
    assert need["neo4j"] >= de._NEO4J_FIXED_INITIAL


def test_vectors_are_bundled_and_scale_with_dimension():
    # ベクトルを持つのは束ねた後の本数だけ（docx は 1/2）。ES は 1 次元 3.5B、埋め込みキャッシュは 1 次元 1B ＋ 1 行 700B。
    kw = dict(mode="register")
    base = de.compute_needs(_counts(docx=(10, 10 * MB)), None, dim=0, **kw)
    with_vec = de.compute_needs(_counts(docx=(10, 10 * MB)), None, dim=1536, **kw)
    vecs = (650 * 10 + 4 * 10) * 0.5
    assert with_vec["es"] - base["es"] == pytest.approx(vecs * 3.5 * 1536 * 1.5, rel=1e-6)
    assert with_vec["embed"] == pytest.approx(vecs * (1 * 1536 + 700), rel=1e-6)
    assert base["embed"] == 0
    half = de.compute_needs(_counts(docx=(10, 10 * MB)), None, dim=768, **kw)
    assert half["embed"] < with_vec["embed"] and half["es"] < with_vec["es"]


def test_peak_is_judged_but_settled_size_is_smaller():
    c = _counts(xlsx=(10, 10 * MB))
    peak = de.compute_needs(c, None, mode="register", dim=1536)
    settled = de.compute_needs(c, None, mode="register", dim=1536, settled=True)
    assert peak["es"] == pytest.approx(settled["es"] * de._ES_PEAK_FACTOR, rel=1e-6)
    fss = {p: de.Fs(("dev", 1), 100 * GIB, 200 * GIB) for p in de.PLACES}
    needs = {p: 0 for p in de.PLACES}
    needs["es"] = 90 * GIB
    res = de.assess(needs, fss, es_settled=60 * GIB)
    assert res["status"] == "tight" and "60GB" in res["lines"][0]["text"] and "取り込み中" in res["lines"][0]["text"]


def test_text_is_indexed_without_vectors_and_sources_are_not_indexed():
    need = de.compute_needs(_counts(text=(10, 5 * MB), source=(100, 3 * MB)), None, mode="register", dim=1536)
    assert need["embed"] == 0                         # テキスト資料はベクトルを通さない
    assert need["es"] == pytest.approx((200 * 5 + 10) * 6000 * 1.5, rel=1e-6)   # ソースは ES に入れない
    assert need["neo4j"] == de._NEO4J_FIXED_INITIAL + 100 * de._NEO4J_SOURCE_FILE_BYTES + 10 * de._NEO4J_OTHER_FILE_BYTES


def test_update_does_not_count_what_is_already_in_use():
    total = _counts(xlsx=(1000, 500 * MB))
    delta = _counts(xlsx=(10, 5 * MB))
    es_now = 3 * GIB
    upd = de.compute_needs(total, delta, mode="update", dim=0, linked=True, es_actual=es_now)
    reg = de.compute_needs(total, None, mode="register", dim=0)
    # 派生物は変わった分だけ（キャッシュとリンクしているので変わらない分は新しい実体を作らない）
    assert upd["derived"] < reg["derived"] / 50
    # ES は作り直し中に新旧 2 世代が並ぶ: 新しい 1 世代ぶんは全ファイルのぶん、今の実測（詰めた後）がそれより大きければ最大に直したそちら
    assert upd["es"] == max(reg["es"], es_now * 1.5)
    bigger = de.compute_needs(total, delta, mode="update", dim=0, linked=True, es_actual=reg["es"] * 2)
    assert bigger["es"] == reg["es"] * 3


def test_update_without_link_needs_a_full_new_set():
    total = _counts(docx=(100, 100 * MB))
    delta = _counts(docx=(1, 1 * MB))
    linked = de.compute_needs(total, delta, mode="update", dim=0, linked=True)
    unlinked = de.compute_needs(total, delta, mode="update", dim=0, linked=False)
    full = de.compute_needs(total, None, mode="register", dim=0)
    assert linked["derived"] < full["derived"] / 10
    assert unlinked["derived"] == full["derived"]


def test_update_with_no_change_asks_for_nothing():
    total = _counts(pdf=(100, 100 * MB))
    need = de.compute_needs(total, _counts(), mode="update", dim=1536, es_actual=10 * GIB)
    assert sum(need.values()) == 0


def test_legacy_files_reserve_libreoffice_temp_by_largest_file():
    small = de.compute_needs(_counts(legacy=(1, 10 * MB)), None, mode="register", dim=0)
    none = de.compute_needs(_counts(pdf=(1, 10 * MB)), None, mode="register", dim=0)
    assert small["tmp"] == de._TMP_BASE_BYTES + de._TMP_LEGACY_BYTES + 2 * 10 * MB
    assert none["tmp"] == de._TMP_BASE_BYTES


def test_logs_are_limited_by_remaining_room_under_the_total_cap(monkeypatch):
    from sherpa import log_setup
    nearly_full = log_setup._TOTAL_MAX_BYTES - 5 * MB
    assert de._logs_need(100_000, nearly_full) == 5 * MB
    assert de._logs_need(10, 0) == de._LOG_BASE_BYTES + 10 * de._LOG_PER_FILE_BYTES


# ---- 原本を数える ----

def test_count_sources_counts_by_family_and_reports_added_and_changed(tmp_path):
    root = tmp_path / "docs"
    (root / "sub").mkdir(parents=True)
    (root / ".git").mkdir()
    (root / ".git" / "obj.pdf").write_bytes(b"x" * 100)          # 版管理の記録は数えない
    (root / "a.pdf").write_bytes(b"p" * 1000)
    (root / "sub" / "b.docx").write_bytes(b"d" * 2000)
    (root / "sub" / "c.doc").write_bytes(b"l" * 3000)
    (root / "note.txt").write_text("hello")
    (root / "link.pdf").symlink_to(root / "a.pdf")                 # symlink は辿らない
    st_a = (root / "a.pdf").stat()
    st_c = (root / "sub" / "c.doc").stat()
    manifest = {
        "a.pdf": [st_a.st_mtime_ns, st_a.st_ctime_ns, st_a.st_size],         # 変更なし
        "sub/c.doc": [1, 1, st_c.st_size],                                    # 変更あり
        "gone.txt": [1, 1, 5],                                                # 削除（数えない）
    }                                                                         # b.docx と note.txt は追加
    total, delta = de.count_sources(root, manifest=manifest)
    fam = total["families"]
    assert fam["pdf"] == [1, 1000] and fam["docx"] == [1, 2000] and fam["legacy"] == [1, 3000] and fam["text"][0] == 1
    assert total["legacy_max"] == 3000
    assert delta["families"]["pdf"] == [0, 0]
    assert delta["families"]["docx"] == [1, 2000] and delta["families"]["legacy"] == [1, 3000]
    assert de._file_total(delta) == 3


def test_count_sources_does_not_write_into_the_source_tree(tmp_path):
    root = tmp_path / "docs"
    root.mkdir()
    (root / "a.xlsx").write_bytes(b"x" * 10)
    before = sorted(p.name for p in root.rglob("*"))
    de.count_sources(root)
    assert sorted(p.name for p in root.rglob("*")) == before


# ---- 置き場と空き ----

def test_places_on_the_same_filesystem_are_summed_before_comparing(tmp_path, monkeypatch):
    # 派生物 6GB と一時領域 6GB は、それぞれ空き 10GB に収まるが、同じディスクなので合わせて 12GB 要る。
    monkeypatch.setattr(de.shutil, "disk_usage", lambda p: type("U", (), {"free": 10 * GIB, "total": 100 * GIB})())
    fs_a = de.fs_of(tmp_path / "derived" / "w")           # まだ無い場所は最も近い親で判定する
    fs_b = de.fs_of(tmp_path / "tmp")
    assert fs_a.key == fs_b.key and fs_a.key is not None
    other = de.Fs(("dev", -1), 500 * GIB, 600 * GIB)
    fss = {p: other for p in de.PLACES}
    fss["derived"], fss["tmp"] = fs_a, fs_b
    needs = {p: 0 for p in de.PLACES}
    needs["derived"], needs["tmp"] = 6 * GIB, 6 * GIB
    res = de.assess(needs, fss)
    assert res["status"] == "short"
    assert len(res["lines"]) == 1 and set(res["lines"][0]["places"]) == {"derived", "tmp"}
    assert res["lines"][0]["need_bytes"] == 12 * GIB


def test_different_filesystems_are_judged_separately(tmp_path):
    a, b = de.Fs(("dev", 1), 100 * GIB, 200 * GIB), de.Fs(("dev", 2), 1 * GIB, 200 * GIB)
    fss = {p: a for p in de.PLACES}
    fss["es"] = b
    needs = {p: 0 for p in de.PLACES}
    needs["derived"], needs["es"] = 10 * GIB, 5 * GIB
    res = de.assess(needs, fss)
    by = {tuple(ln["places"]): ln["status"] for ln in res["lines"]}
    assert by == {("es",): "short", ("derived",): "ok"}
    assert res["status"] == "short"


@pytest.mark.parametrize("free,expected", [
    (9 * GIB, "short"),        # 空き < 必要量
    (10 * GIB, "tight"),       # 空き == 必要量 は余裕なし
    (11 * GIB, "tight"),       # 必要量×1.2 未満
    (12 * GIB, "ok"),          # 必要量×1.2 ちょうどは十分
    (50 * GIB, "ok"),
])
def test_judge_boundaries(free, expected):
    assert de.judge(10 * GIB, free) == expected


def test_es_is_tight_when_it_would_pass_the_disk_watermark():
    fs = de.Fs(("es",), 40 * GIB, 100 * GIB)             # 取り込み後の使用率 = (60+20)/100 = 80%
    fss = {p: fs for p in de.PLACES}
    needs = {p: 0 for p in de.PLACES}
    needs["es"] = 20 * GIB
    assert de.assess(needs, fss)["status"] == "ok"
    fs2 = de.Fs(("es",), 40 * GIB, 100 * GIB)
    fss = {p: fs2 for p in de.PLACES}
    needs["es"] = 32 * GIB                                 # (60+32)/100 = 92%、空きは必要量の 1.25 倍で十分に見える
    res = de.assess(needs, fss)
    assert res["status"] == "tight" and "9割" in res["lines"][0]["text"]


def test_unreadable_docker_volumes_are_reported_not_assumed_ok(monkeypatch):
    def no_docker(*a, **k):
        raise FileNotFoundError("docker")
    monkeypatch.setattr(de.subprocess, "run", no_docker)
    monkeypatch.setattr(es_index, "_req", lambda *a, **k: (_ for _ in ()).throw(OSError("es down")))
    es_fs, neo_fs = de.docker_filesystems()
    assert es_fs.key is None and neo_fs.key is None
    fss = {p: de.Fs(("dev", 1), 500 * GIB, 600 * GIB) for p in de.PLACES}
    fss["es"], fss["neo4j"] = es_fs, neo_fs
    needs = {p: 0 for p in de.PLACES}
    needs["es"], needs["neo4j"] = 3 * GIB, 2 * GIB
    res = de.assess(needs, fss)
    assert res["status"] == "unknown"
    assert all(ln["status"] == "unknown" and ln["free_bytes"] is None for ln in res["lines"])
    assert any("確かめられません" in ln["text"] for ln in res["lines"])


def test_docker_volumes_under_the_same_docker_root_share_one_filesystem(monkeypatch):
    # ホストから中身を読めない（root 専用）ボリュームでも、ES が自分の空きを報告し、Neo4j は同じ保存先なら同じ場所とみなす。
    def fake_run(cmd, **k):
        name = cmd[-1]
        return subprocess.CompletedProcess(cmd, 0, stdout=f"/nonexistent-docker-root/volumes/{name}/_data\n", stderr="")
    monkeypatch.setattr(de.subprocess, "run", fake_run)
    monkeypatch.setattr(es_index, "fs_stats", lambda: {"total": 200 * GIB, "free": 80 * GIB})
    es_fs, neo_fs = de.docker_filesystems()
    assert es_fs.key == ("es",) and neo_fs.key == ("es",)
    assert neo_fs.free == 80 * GIB


def test_docker_volumes_in_different_roots_are_not_assumed_the_same(monkeypatch):
    mounts = {"sherpa-mvp_es": "/nonexistent-a/volumes/sherpa-mvp_es/_data",
              "sherpa-mvp_neo4j": "/nonexistent-b/volumes/sherpa-mvp_neo4j/_data"}
    monkeypatch.delenv("COMPOSE_PROJECT_NAME", raising=False)
    monkeypatch.setattr(de.subprocess, "run", lambda cmd, **k: subprocess.CompletedProcess(
        cmd, 0, stdout=mounts[cmd[-1]] + "\n", stderr=""))
    monkeypatch.setattr(es_index, "fs_stats", lambda: {"total": 200 * GIB, "free": 80 * GIB})
    es_fs, neo_fs = de.docker_filesystems()
    assert es_fs.key == ("es",) and neo_fs.key is None


# ---- キャッシュのリンク ----

def test_conv_cache_link_detection_uses_real_hardlinks(tmp_path, monkeypatch):
    from sherpa import worlds
    monkeypatch.setenv("SHERPA_DERIVED_DIR", str(tmp_path))
    rag = worlds.derived_rag_dir("w")
    cache = worlds.derived_dir("w") / "_conv_cache"
    rag.mkdir(parents=True)
    cache.mkdir(parents=True)
    assert de.conv_cache_linked("w") is False             # 派生物が無ければ「共有していない」側へ倒す
    for i in range(4):
        (rag / f"f{i}.rag.md").write_text("x")
    assert de.conv_cache_linked("w") is False             # リンクされていない（コピー運用）
    for i in range(4):
        os.link(rag / f"f{i}.rag.md", cache / f"f{i}.rag.md")
    assert de.conv_cache_linked("w") is True


# ---- 実行・失敗・ログ ----

def _wait(root, wid=None, timeout=10.0):
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        snap = de.request_estimate(root, wid)
        if snap["state"] != "counting":
            return snap
        time.sleep(0.05)
    raise AssertionError("見積もりが終わらない")


@pytest.fixture
def _clear_jobs():
    de._JOBS.clear()
    yield
    de._JOBS.clear()


def test_estimate_failure_is_reported_and_never_raises(tmp_path, monkeypatch, _clear_jobs):
    def boom(*a, **k):
        raise OSError("share gone")
    monkeypatch.setattr(de, "count_sources", boom)
    snap = _wait(str(tmp_path))
    assert snap["state"] == "failed"
    assert "取り込みはそのまま始められます" in snap["message"]


def test_slow_walk_returns_counting_without_blocking_the_caller(tmp_path, monkeypatch, _clear_jobs):
    import threading
    gate = threading.Event()

    def slow(root, **k):
        gate.wait(5)
        return de._new_counts(), de._new_counts()
    monkeypatch.setattr(de, "count_sources", slow)
    monkeypatch.setattr(de, "place_filesystems", lambda wid: {p: de.Fs(("dev", 1), 100 * GIB, 200 * GIB) for p in de.PLACES})
    t0 = time.monotonic()
    first = de.request_estimate(str(tmp_path), None)
    again = de.request_estimate(str(tmp_path), None)      # 同じ場所は合流（二重に数えない）
    assert time.monotonic() - t0 < 1.0
    assert first["state"] == again["state"] == "counting"
    assert len(de._JOBS) == 1
    gate.set()
    assert _wait(str(tmp_path))["state"] == "ready"


def test_register_estimate_end_to_end_on_a_real_tree(tmp_path, monkeypatch, _clear_jobs):
    root = tmp_path / "docs"
    root.mkdir()
    (root / "a.xlsx").write_bytes(b"x" * 3 * MB)
    (root / "b.pdf").write_bytes(b"y" * 1 * MB)
    monkeypatch.setattr(de, "place_filesystems", lambda wid: {p: de.Fs(("dev", 1), 500 * GIB, 600 * GIB) for p in de.PLACES})
    snap = _wait(str(root))
    assert snap["state"] == "ready" and snap["mode"] == "register"
    assert snap["files"] == 2 and snap["bytes"] == 4 * MB
    assert snap["status"] == "ok" and len(snap["lines"]) == 1
    assert snap["need_bytes"] == sum(snap["needs"].values()) > 0


def test_start_log_is_warning_only_when_started_with_a_warning(tmp_path, caplog, _clear_jobs):
    now = time.monotonic()
    ok = {"status": "ok", "lines": [{"places": ["derived"], "need_bytes": GIB, "free_bytes": 50 * GIB, "status": "ok"}],
          "need_bytes": GIB}
    short = {"status": "short", "lines": [{"places": ["es"], "need_bytes": 9 * GIB, "free_bytes": GIB, "status": "short"},
                                          {"places": ["neo4j"], "need_bytes": GIB, "free_bytes": None, "status": "unknown"}],
             "need_bytes": 10 * GIB}
    for name, res in (("ok", ok), ("short", short)):
        job = de._Job(root=f"/r/{name}", world_id=None, mode="register", state="ready", result=res, finished=now)
        de._JOBS[f"-|/r/{name}"] = job
    with caplog.at_level(logging.INFO, logger="sherpa"):
        de.log_start("register", "/r/ok", None)
        de.log_start("register", "/r/short", None)
        de.log_start("refresh", "/r/none", "w")
    recs = {r.getMessage(): r.levelno for r in caplog.records if "disk_estimate" in r.getMessage()}
    levels = {("判定=ok" in m, "判定=short" in m, "見積もりなし" in m): lv for m, lv in recs.items()}
    assert levels[(True, False, False)] == logging.INFO
    assert levels[(False, True, False)] == logging.WARNING
    assert levels[(False, False, True)] == logging.INFO
    short_msg = next(m for m in recs if "判定=short" in m)
    assert "空き不明" in short_msg and "空き1GB" in short_msg


# ---- OOXML の中身を数える ----

_WORDS = ("処理", "確認", "受注", "請求", "在庫", "顧客", "更新", "登録", "削除", "出力", "帳票", "画面")


def _sentence(rng: random.Random, n: int) -> str:
    out = ""
    while len(out) < n:
        out += rng.choice(_WORDS) + rng.choice("をのにが")
    return out[:n] + "。"


def _make_xlsx(path, rows: int, cols: int = 6) -> None:
    import openpyxl
    rng = random.Random(1)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "一覧"
    ws.append([f"項目{c}" for c in range(cols)])
    for i in range(rows):
        ws.append([f"ID-{i:05d}" if c == 0 else (_sentence(rng, 12) if c % 2 else i * c + 0.5) for c in range(cols)])
    wb.save(path)


def _make_docx(path, paras: int, table_rows: int = 0) -> None:
    import docx
    rng = random.Random(2)
    d = docx.Document()
    for _ in range(paras):
        d.add_paragraph(_sentence(rng, 60))
    if table_rows:
        t = d.add_table(rows=table_rows, cols=4)
        for i in range(table_rows):
            for j in range(4):
                t.cell(i, j).text = _sentence(rng, 20)
    d.save(path)


def _make_pptx(path, slides: int) -> None:
    import pptx
    rng = random.Random(3)
    pr = pptx.Presentation()
    for _ in range(slides):
        s = pr.slides.add_slide(pr.slide_layouts[1])
        s.shapes.title.text = _sentence(rng, 12)
        s.placeholders[1].text = "\n".join(_sentence(rng, 30) for _ in range(5))
    pr.save(path)


_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_NOTE_XML = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:{root} xmlns:w="' + _W_NS + '">{items}</w:{root}>')


def _inject_docx_parts(path, parts: dict) -> None:
    """既存の docx に、脚注・文末脚注の部品を足す（python-docx には書く API が無い）。"""
    with zipfile.ZipFile(path) as zin:
        members = {n: zin.read(n) for n in zin.namelist()}
    ctypes = members["[Content_Types].xml"].decode("utf-8")
    rels = members["word/_rels/document.xml.rels"].decode("utf-8")
    for name, (xml, ctype, rtype) in parts.items():
        members[name] = xml.encode("utf-8")
        ctypes = ctypes.replace("</Types>", f'<Override PartName="/{name}" ContentType="{ctype}"/></Types>')
        rels = rels.replace("</Relationships>", f'<Relationship Id="rId{900 + len(members)}" Type="{rtype}" '
                                                f'Target="{name.removeprefix("word/")}"/></Relationships>')
    members["[Content_Types].xml"] = ctypes.encode("utf-8")
    members["word/_rels/document.xml.rels"] = rels.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, b in members.items():
            zout.writestr(n, b)


def _make_docx_with_extras(path, paras: int = 20, notes: int = 15, comments: int = 10) -> None:
    """本文のほかにヘッダー・フッター・脚注・文末脚注・コメントを持つ docx。"""
    import docx
    rng = random.Random(4)
    d = docx.Document()
    runs = [d.add_paragraph(_sentence(rng, 60)).runs[0] for _ in range(paras)]
    d.sections[0].header.paragraphs[0].text = _sentence(rng, 80)
    d.sections[0].footer.paragraphs[0].text = _sentence(rng, 80)
    for i in range(comments):
        d.add_comment(runs[i], text=_sentence(rng, 150), author="作成者", initials="作")
    d.save(path)

    def items(tag: str, n: int) -> str:
        sep = (f'<w:{tag} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{tag}>'
               f'<w:{tag} w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:{tag}>')
        return sep + "".join(
            f'<w:{tag} w:id="{i}"><w:p><w:r><w:t>{_sentence(rng, 300)}</w:t></w:r></w:p>'
            f'<w:p><w:r><w:t>{_sentence(rng, 100)}</w:t></w:r></w:p></w:{tag}>' for i in range(1, n + 1))
    ns = "application/vnd.openxmlformats-officedocument.wordprocessingml"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    _inject_docx_parts(path, {
        "word/footnotes.xml": (_NOTE_XML.format(root="footnotes", items=items("footnote", notes)),
                               f"{ns}.footnotes+xml", f"{rel}/footnotes"),
        "word/endnotes.xml": (_NOTE_XML.format(root="endnotes", items=items("endnote", notes // 3)),
                              f"{ns}.endnotes+xml", f"{rel}/endnotes"),
    })


def _make_pptx_with_notes(path, slides: int = 10, notes: bool = True) -> None:
    """スライドのほかに発表者ノートを持つ pptx。"""
    import pptx
    rng = random.Random(5)
    pr = pptx.Presentation()
    for _ in range(slides):
        s = pr.slides.add_slide(pr.slide_layouts[1])
        s.shapes.title.text = _sentence(rng, 12)
        s.placeholders[1].text = _sentence(rng, 30)
        if notes:
            s.notes_slide.notes_text_frame.text = "\n".join(_sentence(rng, 1000) for _ in range(3))
    pr.save(path)


_MAKERS = {"xlsx": _make_xlsx, "docx": _make_docx, "pptx": _make_pptx}


def _convert(tmp_path, monkeypatch, name: str, make) -> tuple[int, int]:
    """実際に変換して `(チャンク数, 派生物のバイト数)` を返す（DB には触れない）。"""
    monkeypatch.setenv("SHERPA_PG_DSN", "postgresql://invalid:1/none")
    wd = tmp_path / "world"
    wd.mkdir()
    make(wd / name)
    derived = tmp_path / "derived"
    rep = office_md.build_derived(wd, derived)
    assert rep["evidence_ir_failed"] == 0 and rep["rag_failed"] == 0
    chunks = size = 0
    seen = set()
    for p in derived.parent.rglob("*"):
        if not p.is_file() or p.is_symlink() or wd in p.parents:
            continue
        st = p.stat()
        if (st.st_dev, st.st_ino) in seen:       # 変換結果のキャッシュとはハードリンクで同じ実体
            continue
        seen.add((st.st_dev, st.st_ino))
        size += st.st_size
        if p.name.endswith(".rag_chunks.jsonl"):
            with json_io.open_text_maybe_gzip(p) as f:
                chunks += sum(1 for _ in f)
    return chunks, size


@pytest.mark.parametrize("family,name,make", [
    ("xlsx", "a100.xlsx", lambda p: _make_xlsx(p, 100)),
    ("xlsx", "a1000.xlsx", lambda p: _make_xlsx(p, 1000)),
    ("xlsx", "a5000.xlsx", lambda p: _make_xlsx(p, 5000)),
    ("docx", "a100.docx", lambda p: _make_docx(p, 100)),
    ("docx", "a1000.docx", lambda p: _make_docx(p, 1000)),
    ("docx", "table.docx", lambda p: _make_docx(p, 20, table_rows=200)),
    ("pptx", "a20.pptx", lambda p: _make_pptx(p, 20)),
    ("pptx", "a100.pptx", lambda p: _make_pptx(p, 100)),
])
def test_estimated_chunks_and_derived_bytes_match_a_real_conversion(tmp_path, monkeypatch, family, name, make):
    # 実害: 原本の大きさ（ZIP で圧縮済み）から係数で出すと、文字の多い表で派生物もチャンク数も桁で外れる。
    chunks, size = _convert(tmp_path, monkeypatch, name, make)
    est_derived, est_chunks = de.measure_ooxml(tmp_path / "world" / name, family, deadline=time.monotonic() + 60)
    assert chunks * 0.75 <= est_chunks <= chunks * 1.25
    assert size * 0.75 <= est_derived <= size * 1.5


@pytest.mark.parametrize("family,name,make", [
    ("docx", "extras.docx", _make_docx_with_extras),
    ("pptx", "notes.pptx", _make_pptx_with_notes),
])
def test_estimated_chunks_include_the_parts_ingest_reads_besides_the_body(tmp_path, monkeypatch, family, name, make):
    # 実害: 本文（document.xml／スライド）だけを数えると、ヘッダー・脚注・コメント・発表者ノートの分のチャンクが見積もりから抜ける。
    chunks, _size = _convert(tmp_path, monkeypatch, name, make)
    est_derived, est_chunks = de.measure_ooxml(tmp_path / "world" / name, family, deadline=time.monotonic() + 60)
    assert chunks * 0.75 <= est_chunks <= chunks * 1.25
    plain = {"docx": lambda p: _make_docx(p, 20), "pptx": lambda p: _make_pptx_with_notes(p, notes=False)}[family]
    (tmp_path / "plain").mkdir()
    body_only, _ = _convert(tmp_path / "plain", monkeypatch, name, plain)
    assert body_only < chunks * 0.75                                   # 本文だけでは足りない文書（抜けると 25% の許容を超える）


def test_parts_that_cannot_be_counted_keep_the_coefficient_estimate(tmp_path, monkeypatch):
    _make_docx_with_extras(tmp_path / "e.docx")
    path = tmp_path / "e.docx"
    full = de.measure_ooxml(path, "docx", deadline=time.monotonic() + 60)
    with zipfile.ZipFile(path) as zf:
        body = zf.getinfo("word/document.xml").file_size
        extras = sum(i.compress_size for i in zf.infolist()
                     if i.filename.startswith(("word/header", "word/footer", "word/footnotes", "word/endnotes", "word/comments")))
        feat_full = de._scan_docx(zf, time.monotonic() + 10)
        monkeypatch.setattr(de, "_MEASURE_FILE_BYTES", body)                 # 本文で 1 ファイルの上限を使い切る
        feat = de._scan_docx(zf, time.monotonic() + 10)
    assert feat_full["uncounted_bytes"] == 0 and feat["uncounted_bytes"] == extras
    c = de.COEFS["docx"]
    limited = de.measure_ooxml(path, "docx", deadline=time.monotonic() + 60)
    assert limited[1] == pytest.approx(feat["chunks"] + feat["media_n"] + extras / MB * c.chunks_per_mb)
    assert limited[0] > 0 and limited[1] < full[1]                           # 数えた分より少ない（係数の式は 1MB あたりなので小さな部品は小さく出る）


def test_sampling_extrapolates_when_the_xml_is_larger_than_the_read_limit(tmp_path, monkeypatch):
    path = tmp_path / "big.xlsx"
    _make_xlsx(path, 5000)
    full = de.measure_ooxml(path, "xlsx", deadline=time.monotonic() + 60)
    monkeypatch.setattr(de, "_MEASURE_PART_BYTES", 64 * 1024)
    monkeypatch.setattr(de, "_MEASURE_READ_BYTES", 16 * 1024)
    sampled = de.measure_ooxml(path, "xlsx", deadline=time.monotonic() + 60)
    with zipfile.ZipFile(path) as zf:
        assert zf.getinfo("xl/worksheets/sheet1.xml").file_size > 64 * 1024     # 本当に途中で切って外挿している
    assert abs(sampled[1] / full[1] - 1) < 0.1
    assert abs(sampled[0] / full[0] - 1) < 0.2


@pytest.mark.parametrize("family", ["xlsx", "docx", "pptx"])
def test_unreadable_ooxml_falls_back_to_the_coefficient_formula(tmp_path, family):
    root = tmp_path / "docs"
    root.mkdir()
    (root / f"garbage.{family}").write_bytes(b"x" * 5000)                       # ZIP ではない
    (root / f"empty_file.{family}").write_bytes(b"")
    with zipfile.ZipFile(root / f"no_parts.{family}", "w") as zf:               # ZIP だが本体の XML が無い
        zf.writestr("readme.txt", "x")
    good = root / f"truncated.{family}"
    _MAKERS[family](good, 50)
    good.write_bytes(good.read_bytes()[:-200])                                  # 中央ディレクトリが欠けた壊れた ZIP
    for p in root.iterdir():
        assert de.measure_ooxml(p, family, deadline=time.monotonic() + 10) is None
    total, _ = de.count_sources(root)
    n, size = total["families"][family]
    assert n == 4 and total["measured"][family][0] == 0
    assert de._family_needs(total, dim=1536) == de._family_needs(_counts(**{family: (n, size)}), dim=1536)


def test_measured_files_replace_the_coefficients_and_the_rest_keep_them(tmp_path):
    root = tmp_path / "docs"
    root.mkdir()
    _make_xlsx(root / "good.xlsx", 500)
    (root / "broken.xlsx").write_bytes(b"PK\x03\x04" + b"x" * 3000)
    total, _ = de.count_sources(root)
    n, size = total["families"]["xlsx"]
    m_n, m_size, m_derived, m_chunks = total["measured"]["xlsx"]
    assert (n, m_n) == (2, 1)
    need = de._family_needs(total, dim=0)
    c = de.COEFS["xlsx"]
    rest = size - m_size
    derived = c.derived_ratio * rest + c.derived_file * 2 + m_derived + 2 * de._DERIVED_ALL_FILE_BYTES
    assert need["derived"] == pytest.approx(derived * (1 + de._CONV_CACHE_OVERHEAD), rel=1e-6)
    chunks = m_chunks + c.chunks_per_mb * rest / MB + c.chunks_file
    assert need["es"] == pytest.approx(chunks * c.es_chunk_bytes, rel=1e-6)


def test_update_measures_only_added_and_changed_files(tmp_path):
    root = tmp_path / "docs"
    root.mkdir()
    _make_xlsx(root / "same.xlsx", 50)
    _make_xlsx(root / "new.xlsx", 50)
    st = (root / "same.xlsx").stat()
    manifest = {"same.xlsx": [st.st_mtime_ns, st.st_ctime_ns, st.st_size]}
    total, delta = de.count_sources(root, manifest=manifest)
    assert delta["families"]["xlsx"][0] == 1 and delta["measured"]["xlsx"][0] == 1
    assert total["measured"]["xlsx"][0] == 1


def test_measuring_stops_when_the_time_budget_is_used_up(tmp_path, monkeypatch):
    root = tmp_path / "docs"
    root.mkdir()
    for i in range(3):
        _make_xlsx(root / f"b{i}.xlsx", 20)
    monkeypatch.setattr(de, "_MEASURE_BUDGET_SEC", 0)
    total, _ = de.count_sources(root)
    assert total["families"]["xlsx"][0] == 3 and total["measured"]["xlsx"][0] == 0     # 残りは係数の式
    assert de.measure_ooxml(root / "b0.xlsx", "xlsx", deadline=time.monotonic() - 1) is None


def test_measuring_leaves_half_of_a_short_walk_deadline_for_walking(tmp_path, monkeypatch):
    root = tmp_path / "docs"
    root.mkdir()
    _make_xlsx(root / "a.xlsx", 20)
    seen = []
    real = de.measure_ooxml
    monkeypatch.setattr(de, "measure_ooxml", lambda *a, deadline, **k: seen.append(deadline) or real(*a, deadline=deadline, **k))
    now = time.monotonic()
    de.count_sources(root, deadline=now + 20)
    assert len(seen) == 1 and seen[0] <= now + 10.5


def test_measuring_does_not_touch_the_source_tree(tmp_path):
    root = tmp_path / "docs"
    root.mkdir()
    _make_xlsx(root / "a.xlsx", 30)
    _make_docx(root / "b.docx", 30)
    _make_pptx(root / "c.pptx", 3)
    before = {p.name: (p.stat().st_mtime_ns, p.read_bytes()) for p in root.iterdir()}
    total, _ = de.count_sources(root)
    assert [total["measured"][f][0] for f in ("docx", "xlsx", "pptx")] == [1, 1, 1]
    assert {p.name: (p.stat().st_mtime_ns, p.read_bytes()) for p in root.iterdir()} == before


def test_embedded_images_are_counted_once_by_content(tmp_path):
    import io

    import docx
    from PIL import Image
    buf = io.BytesIO()
    Image.frombytes("RGB", (64, 64), os.urandom(64 * 64 * 3)).save(buf, "PNG")
    d = docx.Document()
    d.add_paragraph("本文")
    for _ in range(3):
        d.add_picture(io.BytesIO(buf.getvalue()))                                # 同じ画像は ZIP の中に 1 つ
    d.save(tmp_path / "img.docx")
    with zipfile.ZipFile(tmp_path / "img.docx") as zf:
        feat = de._scan_docx(zf, time.monotonic() + 10)
    assert feat["media_n"] == 1 and feat["media_bytes"] == len(buf.getvalue())


def test_estimate_output_shape_is_unchanged_with_measured_files(tmp_path, monkeypatch):
    root = tmp_path / "docs"
    root.mkdir()
    _make_xlsx(root / "a.xlsx", 200)
    monkeypatch.setattr(de, "place_filesystems", lambda wid: {p: de.Fs(("dev", 1), 500 * GIB, 600 * GIB) for p in de.PLACES})
    res = de.estimate_now(str(root), None)
    assert set(res) == {"status", "lines", "need_bytes", "es_settled_bytes", "mode", "files", "bytes", "changed_files",
                        "needs", "unchanged"}
    assert set(res["needs"]) == {"derived", "embed", "es", "neo4j", "ocr", "logs", "tmp"}
    assert res["files"] == 1 and res["mode"] == "register"


def test_file_read_limit_extrapolates_from_the_sheets_that_were_read(tmp_path, monkeypatch):
    import openpyxl
    rng = random.Random(4)
    wb = openpyxl.Workbook()
    for k in range(4):
        ws = wb.active if k == 0 else wb.create_sheet()
        ws.append(["No", "内容", "備考"])
        for i in range(400):
            ws.append([i, _sentence(rng, 20), _sentence(rng, 20)])
    wb.save(tmp_path / "sheets.xlsx")
    full = de.measure_ooxml(tmp_path / "sheets.xlsx", "xlsx", deadline=time.monotonic() + 60)
    with zipfile.ZipFile(tmp_path / "sheets.xlsx") as zf:
        one = zf.getinfo("xl/worksheets/sheet1.xml").file_size
    monkeypatch.setattr(de, "_MEASURE_FILE_BYTES", one + one // 4)           # 3 枚目以降は読まない
    limited = de.measure_ooxml(tmp_path / "sheets.xlsx", "xlsx", deadline=time.monotonic() + 60)
    assert full[1] == 4 * 401
    assert abs(limited[1] / full[1] - 1) < 0.1 and abs(limited[0] / full[0] - 1) < 0.1
