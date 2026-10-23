"""Office 観測の共通の扱い（語彙の置き場所・図形 ID での対応・取り消し線の format_state・covers の正本・状態の由来）の単体テスト。

正典: docs/proposals/2026-10-22-Office観測の共通の扱い.md §1。判定の条件（遮蔽の閾値・塗りの区分）は変えない土台なので、
ここでは「今の判定の結果が正しい要素に付くこと」と「表し方（extension のキー）」だけを固定する。
"""
from __future__ import annotations

import itertools
import pathlib
import zipfile

import openpyxl
import pytest
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Font

import _image_builders
from sherpa.ingest import evidence_ir, evidence_render, evidence_spike, observation_vocab
from sherpa.ingest.arms import ooxml_arm

_ROOT = pathlib.Path(__file__).resolve().parents[2]
_STAMP_ASSET = _ROOT / "fixtures/eval/excel_ja/assets/廃止スタンプ.png"
_EXCEL_JA_INPUTS = _ROOT / "fixtures/eval/excel_ja/inputs"

_PPTX_NS = (
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
)
_DOCX_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_STRIKE_LINE = "取り消し線: 「あり」"


def _write_zip(path: pathlib.Path, entries: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path, "w") as z:
        for name, data in entries.items():
            z.writestr(name, data)


# ---- pptx の手組み ----

_SHAPE_IDS = itertools.count(2)


def _xfrm(x, y, cx, cy) -> str:
    return f'<a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'


def _sp_text(text, *, shape_id, name, x=0, y=0, cx=4000000, cy=2000000) -> str:
    return (
        f'<p:sp><p:nvSpPr><p:cNvPr id="{shape_id}" name="{name}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        f"<p:spPr>{_xfrm(x, y, cx, cy)}</p:spPr>"
        f"<p:txBody><a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp>"
    )


def _sp_fill(*, shape_id, name, x=0, y=0, cx=4000000, cy=2000000) -> str:
    return (
        f'<p:sp><p:nvSpPr><p:cNvPr id="{shape_id}" name="{name}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr>{_xfrm(x, y, cx, cy)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        '<a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill></p:spPr></p:sp>'
    )


def _pic(*, shape_id, name, x, y, cx, cy) -> str:
    return (
        f'<p:pic><p:nvPicPr><p:cNvPr id="{shape_id}" name="{name}"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>'
        f"<p:spPr>{_xfrm(x, y, cx, cy)}</p:spPr></p:pic>"
    )


def _connector(*, shape_id, name, start, end) -> str:
    return (
        f'<p:cxnSp><p:nvCxnSpPr><p:cNvPr id="{shape_id}" name="{name}"/>'
        f'<p:cNvCxnSpPr><a:stCxn id="{start}" idx="0"/><a:endCxn id="{end}" idx="0"/></p:cNvCxnSpPr><p:nvPr/></p:nvCxnSpPr>'
        f"<p:spPr>{_xfrm(0, 0, 100, 100)}</p:spPr></p:cxnSp>"
    )


def _pptx(tmp_path: pathlib.Path, shapes: str, name: str = "base.pptx") -> pathlib.Path:
    slide = f"""<?xml version="1.0"?>
<p:sld {_PPTX_NS}><p:cSld><p:spTree>{shapes}</p:spTree></p:cSld></p:sld>"""
    path = tmp_path / name
    _write_zip(path, {"ppt/slides/slide1.xml": slide.encode("utf-8")})
    return path


def _by_object_id(ir: evidence_ir.EvidenceIR, object_id: int) -> evidence_ir.EvidenceElement:
    return next(
        el for el in ir.elements
        if el.locator.object_id == object_id and el.extension.get("origin") != "document-ir-v2-adapter")


def _covers(ir: evidence_ir.EvidenceIR) -> list[evidence_ir.EvidenceRelation]:
    return [r for r in ir.relations if r.type == observation_vocab.RELATION_COVERS]


def _image_pptx_with_cover_after_leading_objects(tmp_path: pathlib.Path) -> pathlib.Path:
    """最初の文字付き図形より前に画像とコネクタがあり、文字付き図形が後ろの塗り図形に覆われるスライド。"""
    return _pptx(
        tmp_path,
        _pic(shape_id=10, name="先頭の画像", x=9000000, y=9000000, cx=100000, cy=100000)
        + _connector(shape_id=11, name="先頭の線", start=12, end=13)
        + _sp_text("覆われる文字", shape_id=12, name="覆われる図形")
        + _sp_fill(shape_id=13, name="覆う図形"),
    )


# ---- ①: 図形 ID での対応 ----

def test_cover_state_lands_on_the_covered_shape_when_images_and_connectors_precede_it(tmp_path):
    p = _image_pptx_with_cover_after_leading_objects(tmp_path)
    ir = evidence_spike.extract(p)
    assert evidence_ir.validation_errors(ir) == []

    covered = _by_object_id(ir, 12)
    leading_picture = _by_object_id(ir, 10)
    assert covered.visibility == "hidden"
    assert covered.extension["visibility_reason"] == "occluded"
    assert leading_picture.visibility == "visible"
    assert "visibility_reason" not in leading_picture.extension       # 先頭の画像に覆いの状態が付かない
    assert _by_object_id(ir, 11).visibility == "visible"

    covers = _covers(ir)
    assert [(r.source_id, r.target_id) for r in covers] == [(_by_object_id(ir, 13).element_id, covered.element_id)]


def test_alias_binds_the_adapted_shape_to_its_own_native_object_by_shape_id(tmp_path):
    p = _image_pptx_with_cover_after_leading_objects(tmp_path)
    ir = evidence_spike.extract(p)
    elements = {el.element_id: el for el in ir.elements}
    aliases, suppressed = evidence_render._element_aliases(ir, elements)
    covered = _by_object_id(ir, 12)
    assert list(aliases) == [covered.element_id]
    assert len(suppressed) == 1


def test_document_ir_shape_carries_the_original_shape_id_and_name(tmp_path):
    p = _image_pptx_with_cover_after_leading_objects(tmp_path)
    document = ooxml_arm._build_pptx_ir(p)
    shape = next(el for el in document.elements if el.type == "shape")
    assert shape.source_map["shape_id"] == 12 and shape.source_map["shape_name"] == "覆われる図形"
    assert shape.source_map["occluded_by"]["shape_id"] == 13


def test_covered_by_text_lands_on_the_right_shape_and_keeps_visibility(tmp_path):
    p = _pptx(
        tmp_path,
        _pic(shape_id=10, name="先頭の画像", x=9000000, y=9000000, cx=100000, cy=100000)
        + _sp_text("下の文字", shape_id=12, name="下の図形")
        + _sp_text("廃止", shape_id=13, name="上の図形"),
        name="text_cover.pptx",
    )
    ir = evidence_spike.extract(p)
    under = _by_object_id(ir, 12)
    assert under.visibility == "visible"
    assert under.extension["covered_by_text"]["text"] == "廃止"
    assert "covered_by_text" not in _by_object_id(ir, 10).extension
    (cover,) = _covers(ir)
    assert cover.extension["fill"] == observation_vocab.FILL_NONE_WITH_TEXT
    assert cover.extension["occlusion"] == observation_vocab.OCCLUSION_NOT_OCCLUDING


# ---- ②: 要素全体の取り消し線は format_state ----

def _strike_docx(tmp_path: pathlib.Path) -> pathlib.Path:
    body = (
        '<w:p><w:r><w:t>通常本文</w:t></w:r>'
        '<w:r><w:rPr><w:strike/></w:rPr><w:t>取り消し線本文</w:t></w:r></w:p>'
    )
    doc = f'<?xml version="1.0"?><w:document {_DOCX_NS}><w:body>{body}</w:body></w:document>'
    path = tmp_path / "strike.docx"
    _write_zip(path, {"word/document.xml": doc.encode("utf-8")})
    return path


def _strike_xlsx(tmp_path: pathlib.Path) -> pathlib.Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "一覧"
    ws["A1"] = "通常値"
    ws["A2"] = "廃止された項目"
    ws["A2"].font = Font(strike=True)
    path = tmp_path / "strike.xlsx"
    wb.save(path)
    return path


@pytest.mark.parametrize("make", [_strike_docx, _strike_xlsx])
def test_whole_element_strike_is_a_format_state_not_a_visibility_reason(tmp_path, make):
    p = make(tmp_path)
    ir = evidence_spike.extract(p)
    assert evidence_ir.validation_errors(ir) == []
    struck = [el for el in ir.elements if el.extension.get("format_state")]
    assert struck and all(el.extension["format_state"] == {"strike": "whole"} for el in struck)
    assert all(el.extension.get("visibility_reason") != "strike" for el in ir.elements)
    assert all(el.visibility == "visible" for el in struck)

    rendered = evidence_render.render(ir, source_name=str(p))
    assert _STRIKE_LINE in rendered.markdown


# ---- ③: covers の正本 ----

def _one_cell_stamp_xlsx(tmp_path: pathlib.Path, image: pathlib.Path | None = None) -> pathlib.Path:
    """A2 だけを覆う画像つきの xlsx。`image` を省くと不透明の画像を使う。"""
    if image is None:
        image = tmp_path / "opaque.png"
        image.write_bytes(_image_builders.opaque_png())
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "一覧"
    ws["A1"] = "見出し"
    ws["A2"] = "廃止された値"
    img = XLImage(str(image))
    img.anchor = TwoCellAnchor(
        editAs="oneCell",
        _from=AnchorMarker(col=0, colOff=0, row=1, rowOff=0),
        to=AnchorMarker(col=1, colOff=0, row=2, rowOff=0),
    )
    ws.add_image(img)
    path = tmp_path / "stamp.xlsx"
    wb.save(path)
    return path


def test_xlsx_covers_is_the_single_source_of_the_covered_cell_state(tmp_path):
    ir = evidence_spike.extract(_one_cell_stamp_xlsx(tmp_path))
    cell = next(el for el in ir.elements if el.type == "cell" and el.locator.cell_range == "A2")
    (cover,) = _covers(ir)
    assert cover.target_id == cell.element_id
    assert cover.extension["basis"] == observation_vocab.BASIS_GEOMETRY
    assert cover.extension["fill"] == observation_vocab.FILL_OPAQUE
    assert cover.extension["occlusion"] == observation_vocab.OCCLUSION_CONFIRMED
    assert cover.extension["cover_ratio"] == pytest.approx(1.0)
    assert cell.visibility == "hidden"
    assert cell.extension["visibility_reason"] == "occluded_by_picture"
    assert cell.extension["occluded_by"]["element_id"] == cover.source_id


def test_transparent_picture_cover_keeps_the_cell_visible(tmp_path):
    ir = evidence_spike.extract(_one_cell_stamp_xlsx(tmp_path, _STAMP_ASSET))
    cell = next(el for el in ir.elements if el.type == "cell" and el.locator.cell_range == "A2")
    (cover,) = _covers(ir)
    assert cover.target_id == cell.element_id
    assert cover.extension["fill"] == observation_vocab.FILL_TRANSPARENT_PNG
    assert cover.extension["occlusion"] == observation_vocab.OCCLUSION_UNDETERMINED
    assert cell.visibility == "visible"
    assert "visibility_reason" not in cell.extension and "occluded_by" not in cell.extension


def test_apply_cover_state_derives_the_back_state_from_covers_only():
    source = evidence_ir.EvidenceSource(file_type="xlsx", content_hash="sha256:" + "0" * 64)
    builder = evidence_spike._Builder(source)
    locator = evidence_ir.Locator(part="xl/worksheets/sheet1.xml", sheet="S", cell_range="A1")
    back = builder.add_element("cell", locator, parent_id=None, order=1, value="背面", extension={})
    front = builder.add_element(
        "shape", evidence_ir.Locator(part="xl/drawings/drawing1.xml", sheet="S", object_id=1),
        parent_id=None, order=2, value="前面", extension={"z_order": 3})
    evidence_spike._apply_cover_state(builder, back, file_type="xlsx")
    assert builder.element(back).visibility == "visible"                 # covers が無ければ何も書かない

    builder.add_covers(front, back, fill=observation_vocab.FILL_OPAQUE, cover_ratio=1.0, coordinate_system="x")
    evidence_spike._apply_cover_state(builder, back, file_type="xlsx")
    derived = builder.element(back)
    assert derived.visibility == "hidden"
    assert derived.extension["visibility_reason"] == "occluded_by_shape"
    assert derived.extension["occluded_by"] == {"kind": "shape", "element_id": front, "z_order": 3, "text": "前面"}


def test_not_occluding_cover_does_not_hide_the_back():
    source = evidence_ir.EvidenceSource(file_type="pptx", content_hash="sha256:" + "0" * 64)
    builder = evidence_spike._Builder(source)
    back = builder.add_element(
        "shape", evidence_ir.Locator(part="ppt/slides/slide1.xml", slide=1, object_id=1),
        parent_id=None, order=1, value="背面", extension={})
    front = builder.add_element(
        "shape", evidence_ir.Locator(part="ppt/slides/slide1.xml", slide=1, object_id=2),
        parent_id=None, order=2, value="前面", extension={})
    builder.add_covers(front, back, fill=observation_vocab.FILL_NONE_WITH_TEXT, cover_ratio=None, coordinate_system="x")
    evidence_spike._apply_cover_state(builder, back, file_type="pptx")
    derived = builder.element(back)
    assert derived.visibility == "visible" and "visibility_reason" not in derived.extension
    assert derived.extension["covered_by_text"] == {"element_id": front, "text": "前面"}


# ---- ④: 状態の由来 ----

def test_state_basis_separates_direct_specification_from_geometry(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "一覧"
    ws["A1"] = "通常行"
    ws["A2"] = "非表示行の値"
    ws["A3"] = "取り消し線の行"
    ws["A3"].font = Font(strike=True)
    ws.row_dimensions[2].hidden = True
    path = tmp_path / "basis.xlsx"
    wb.save(path)
    ir = evidence_spike.extract(path)
    cells = {el.locator.cell_range: el for el in ir.elements if el.type == "cell"}
    assert cells["A2"].extension["state_basis"] == {"visibility": "ooxml_direct"}
    assert cells["A3"].extension["state_basis"] == {"format": "ooxml_direct"}
    assert "state_basis" not in cells["A1"].extension

    covered = evidence_spike.extract(_one_cell_stamp_xlsx(tmp_path))
    stamped = next(el for el in covered.elements if el.type == "cell" and el.locator.cell_range == "A2")
    assert stamped.extension["state_basis"] == {"visibility": "geometry"}

    occluded_pptx = evidence_spike.extract(_image_pptx_with_cover_after_leading_objects(tmp_path))
    assert _by_object_id(occluded_pptx, 12).extension["state_basis"] == {"visibility": "geometry"}


def test_state_basis_for_each_known_reason():
    for reason in ("occluded", "occluded_by_picture", "occluded_by_shape", "off_slide"):
        assert observation_vocab.state_basis_for(reason) == observation_vocab.STATE_BASIS_GEOMETRY
    for reason in ("hidden_sheet", "very_hidden", "hidden_row", "hidden_column", "hidden_run",
                   "hidden_slide", "hidden_slide_inherited"):
        assert observation_vocab.state_basis_for(reason) == observation_vocab.STATE_BASIS_OOXML_DIRECT
    assert observation_vocab.state_basis_for("strike") is None
    assert observation_vocab.state_basis_for(None) is None


# ---- ⑤: 関係の extension ----

def test_existing_relations_carry_basis_direction_and_ends(tmp_path):
    p = _pptx(
        tmp_path,
        _sp_text("始点", shape_id=2, name="始点", x=0, y=0, cx=3000000, cy=1000000)
        + _sp_text("終点", shape_id=3, name="終点", x=2000000, y=500000, cx=3000000, cy=1000000)
        + _connector(shape_id=4, name="線", start=2, end=3),
        name="relations.pptx",
    )
    ir = evidence_spike.extract(p)
    connects = [r for r in ir.relations if r.type == "connects_to"]
    overlaps = [r for r in ir.relations if r.type == "overlaps"]
    assert connects and overlaps
    for relation in connects:
        assert relation.extension["basis"] == observation_vocab.BASIS_OOXML_CONNECTOR
        assert relation.extension["direction"] == observation_vocab.DIRECTION_ENDPOINT_ORDER_ONLY
        assert relation.extension["ends"] == observation_vocab.ENDS_BOTH
    for relation in overlaps:
        assert relation.extension["basis"] == observation_vocab.BASIS_GEOMETRY
        assert relation.extension["direction"] == observation_vocab.DIRECTION_UNDIRECTED
        assert relation.extension["ends"] == observation_vocab.ENDS_BOTH


def test_xlsx_connector_and_overlap_relations_carry_the_common_keys():
    p = _EXCEL_JA_INPUTS / "JPX-021.xlsx"
    if not p.is_file():
        pytest.skip(f"fixture が無い環境: {p}")
    ir = evidence_spike.extract(p)
    connects = [r for r in ir.relations if r.type == "connects_to"]
    assert connects
    assert all(r.extension["direction"] == observation_vocab.DIRECTION_ENDPOINT_ORDER_ONLY for r in connects)
    for relation in ir.relations:
        if relation.type in {"connects_to", "overlaps"}:
            assert {"basis", "direction", "ends"} <= set(relation.extension)


# ---- 語彙 ----

def test_occlusion_follows_the_fill_class():
    assert observation_vocab.occlusion_of(observation_vocab.FILL_OPAQUE) == observation_vocab.OCCLUSION_CONFIRMED
    for fill in (observation_vocab.FILL_UNKNOWN, observation_vocab.FILL_TRANSPARENT_PNG):
        assert observation_vocab.occlusion_of(fill) == observation_vocab.OCCLUSION_UNDETERMINED
    for fill in (observation_vocab.FILL_TRANSLUCENT, observation_vocab.FILL_TRANSPARENT,
                 observation_vocab.FILL_NONE_WITH_TEXT):
        assert observation_vocab.occlusion_of(fill) == observation_vocab.OCCLUSION_NOT_OCCLUDING


def test_wording_table_gives_the_existing_lines():
    say = observation_vocab.say
    assert say("visibility", label="図形に覆われている") == "可視性: 「図形に覆われている」"
    assert say("strike_whole") == _STRIKE_LINE
    assert say("overlap_text", value="廃止") == "重なり: 「廃止」"
    assert say("anchor_behind", value="透かし") == "背面図形: 「透かし」"
    assert say("anchor_front", value="図形") == "前面図形: 「図形」"
    assert say("state_lines_omitted", count=3) == "（ほかに重なり・前面/背面図形の行が 3 件）"
