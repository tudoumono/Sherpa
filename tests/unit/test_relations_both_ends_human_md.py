"""関係を両端から参照できること（RAG の描画）と、人用 Markdown の状態・関係・未解析の通知。

正典: docs/proposals/2026-10-22-Office観測の共通の扱い.md §1.3・§1.5・§1.6、§3.2。
判定の条件（遮蔽・塗り）は変えない。ここでは「判定の結果を、両側のチャンクと人用 Markdown にどう書くか」だけを固定する。
"""
from __future__ import annotations

import os
import pathlib
import re
import zipfile

os.environ.setdefault("SHERPA_USE_FIXTURES", "1")

import openpyxl
import pytest
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.styles import Font
from _image_builders import opaque_png
from test_observation_base import _connector, _pptx, _sp_fill, _sp_text

from sherpa.ingest import context_ir as CIR
from sherpa.ingest import evidence_ir as IR
from sherpa.ingest import evidence_render as R
from sherpa.ingest import evidence_spike, observation_vocab, office_md

_ROOT = pathlib.Path(__file__).resolve().parents[2]
_V021 = _ROOT / "fixtures/eval_v021/備品貸出/02_設計"
_HASH = "sha256:" + "0" * 64
# 人用 Markdown に出してはいけない内部の語・ID・RAG 向けの「キー: 「値」」の書き方
_INTERNAL = re.compile(
    r"図形 ID|前面図形|背面図形|可視性|covers|covered|overlaps|visibility|occluded|ooxml_|geometry|_x0000|docPr|element_id"
    r"|[^\n]*[:：] 「|取り消し線: |: 「あり」")


def _chunks(markdown: str) -> list[str]:
    return markdown.split("<!-- chunk:")[1:]


def _chunk_with(markdown: str, *needles: str) -> str:
    found = [c for c in _chunks(markdown) if all(n in c for n in needles)]
    assert len(found) == 1, (needles, len(found))
    return found[0]


def _rag(path: pathlib.Path) -> str:
    return R.render(evidence_spike.extract(path), source_name=path.name).markdown


# ---- pptx: 同名の対象 2 つを 1 つの塗り図形が覆う ----

def _two_same_name_covered(tmp_path) -> pathlib.Path:
    shapes = (
        _sp_text("左の説明", shape_id=5, name="同名", x=0, y=3000000, cx=1000000, cy=500000)
        + _sp_text("右の説明", shape_id=6, name="同名", x=5000000, y=3000000, cx=1000000, cy=500000)
        + _sp_fill(shape_id=7, name="覆いの図形", x=0, y=2900000, cx=7000000, cy=700000)
    )
    path = _pptx(tmp_path, shapes, name="same_name.pptx")
    with zipfile.ZipFile(path, "a") as z:                    # スライドの大きさ（位置の言葉に使う）
        z.writestr("ppt/presentation.xml",
                   '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
                   '<p:sldSz cx="9144000" cy="6858000"/></p:presentation>')
    return path


def test_front_chunk_alone_identifies_each_covered_target_by_name_id_and_position(tmp_path):
    md = _rag(_two_same_name_covered(tmp_path))
    front = _chunk_with(md, "覆っている対象")
    assert "同名（図形 ID 5・スライド1）" in front and "同名（図形 ID 6・スライド1）" in front
    assert "（覆われていると判定）" in front


def test_back_chunk_carries_the_same_relation_and_the_evidence_relation_stays_single(tmp_path):
    p = _two_same_name_covered(tmp_path)
    md = _rag(p)
    back = _chunk_with(md, "図形: 「左の説明」")
    assert "前面図形: 覆いの図形（図形 ID 7・スライド1）（覆われていると判定）" in back
    covers = [r for r in evidence_spike.extract(p).relations if r.type == observation_vocab.RELATION_COVERS]
    assert len(covers) == 2 and len({(r.source_id, r.target_id) for r in covers}) == 2
    assert all(r.source_id != covers[0].target_id for r in covers)          # 逆向きの複製が無い


# ---- 覆いの確からしさで文を書き分ける ----

def _two_shapes(relation_ext, *, front_value="前面の文字"):
    def shape(element_id, object_id, name, value):
        return IR.EvidenceElement(
            element_id=element_id, type="shape", parent_id=None, order=object_id, value=value,
            locator=IR.Locator(part="ppt/slides/slide2.xml", slide=2, object_id=object_id),
            coverage_id="c1", extension={"name": name})

    front, back = shape("front", 11, "前面の図形", front_value), shape("back", 12, "背面の図形", "背面の文字")
    relation = IR.EvidenceRelation(
        relation_id="r1", type=observation_vocab.RELATION_COVERS, source_id="front", target_id="back",
        evidence_ids=["front", "back"], confidence=1.0, extension=relation_ext)
    ir = IR.EvidenceIR(
        schema_version=IR.EVIDENCE_IR_SCHEMA_VERSION, parser_profile="test",
        source=IR.EvidenceSource(file_type="pptx", content_hash=_HASH), elements=[front, back], relations=[relation])
    ctx = CIR.ContextIR(schema_version=CIR.CONTEXT_IR_SCHEMA_VERSION, analyzer_version="test",
                        source_hash=_HASH, source_name="deck.pptx", document_titles={})
    records = {r["evidence_ids"][0]: r["semantic_text"] for r in R._non_table_records(ir, ctx, {}, set())}
    return records["front"], records["back"]


@pytest.mark.parametrize("fill,occlusion,front_text,back_text,must_not", [
    (observation_vocab.FILL_OPAQUE, observation_vocab.OCCLUSION_CONFIRMED,
     "覆っている対象: 背面の図形（図形 ID 12・スライド2）（覆われていると判定）",
     "前面図形: 前面の図形（図形 ID 11・スライド2）（覆われていると判定）", None),
    (observation_vocab.FILL_UNKNOWN, observation_vocab.OCCLUSION_UNDETERMINED,
     "重なっている対象: 背面の図形（図形 ID 12・スライド2）（覆っているかは未確定）",
     "前面図形: 前面の図形（図形 ID 11・スライド2）（覆われているかは未確定）", "覆われていると判定"),
    (observation_vocab.FILL_TRANSLUCENT, observation_vocab.OCCLUSION_NOT_OCCLUDING,
     "重なっている対象: 背面の図形（図形 ID 12・スライド2）（覆ってはいない：半透明）",
     "前面図形: 前面の図形（図形 ID 11・スライド2）（覆われてはいない：半透明）", "覆われていると判定"),
])
def test_sentences_differ_by_occlusion(fill, occlusion, front_text, back_text, must_not):
    ext = observation_vocab.relation_extension(
        observation_vocab.BASIS_GEOMETRY, observation_vocab.DIRECTION_UNDIRECTED, fill=fill, occlusion=occlusion)
    front, back = _two_shapes(ext)
    assert front_text in front and back_text in back
    if must_not:
        assert must_not not in front and must_not not in back


# ---- 接続: 始点・終点・コネクタの 3 つから参照できる ----

def test_connects_to_is_referenced_from_start_end_and_connector_without_claiming_an_arrow(tmp_path):
    p = _pptx(tmp_path, (
        _sp_text("始点の文字", shape_id=2, name="始点", x=0, y=0, cx=1000000, cy=500000)
        + _sp_text("終点の文字", shape_id=3, name="終点", x=2000000, y=500000, cx=1000000, cy=500000)
        + _connector(shape_id=4, name="線", start=2, end=3)), name="connect.pptx")
    md = _rag(p)
    tail = "（始点→終点の指定。矢印の向きは未確認）"
    assert f"接続先: 終点（図形 ID 3・スライド1）{tail}" in _chunk_with(md, "図形: 「始点の文字」")
    assert f"接続元: 始点（図形 ID 2・スライド1）{tail}" in _chunk_with(md, "図形: 「終点の文字」")
    connector = _chunk_with(md, "接続: ")
    assert f"接続: 始点（図形 ID 2・スライド1） → 終点（図形 ID 3・スライド1）{tail}" in connector
    assert "へ接続している" not in md                                    # 向きを断定する旧い文は、指定だけを読んだ関係には使わない


# ---- 長文分割の後半の片にも参照が残り、行は上限内 ----

def test_every_span_of_a_long_text_keeps_the_cover_lines_within_the_limit():
    ext = observation_vocab.relation_extension(
        observation_vocab.BASIS_GEOMETRY, observation_vocab.DIRECTION_UNDIRECTED,
        fill=observation_vocab.FILL_OPAQUE, occlusion=observation_vocab.OCCLUSION_CONFIRMED)
    long_text = "".join(f"長文{i:04d}です。" for i in range(400))
    element = IR.EvidenceElement(
        element_id="front", type="shape", parent_id=None, order=1, value=long_text,
        locator=IR.Locator(part="ppt/slides/slide1.xml", slide=1, object_id=11), coverage_id="c1",
        extension={"name": "長文の前面"})
    back = IR.EvidenceElement(
        element_id="back", type="shape", parent_id=None, order=2, value="背面",
        locator=IR.Locator(part="ppt/slides/slide1.xml", slide=1, object_id=12), coverage_id="c1",
        extension={"name": "背面"})
    relation = IR.EvidenceRelation(
        relation_id="r1", type=observation_vocab.RELATION_COVERS, source_id="front", target_id="back",
        evidence_ids=["front", "back"], confidence=1.0, extension=ext)
    ir = IR.EvidenceIR(
        schema_version=IR.EVIDENCE_IR_SCHEMA_VERSION, parser_profile="test",
        source=IR.EvidenceSource(file_type="pptx", content_hash=_HASH), elements=[element, back], relations=[relation])
    ctx = CIR.ContextIR(schema_version=CIR.CONTEXT_IR_SCHEMA_VERSION, analyzer_version="test",
                        source_hash=_HASH, source_name="deck.pptx", document_titles={})
    records = [r for r in R._non_table_records(ir, ctx, {}, set()) if r["evidence_ids"] == ["front"]]
    assert len(records) >= 3
    for record in records:
        text = record["semantic_text"]
        assert "覆っている対象: 背面（図形 ID 12・スライド1）（覆われていると判定）" in text
        tail = text[text.index("\n覆っている対象"):]
        assert len(tail) <= R._STATE_NOTE_MAX_CHARS


# ---- 1 つの図形が 100 個のセルに重なる ----

def _opaque_image(tmp_path) -> XLImage:
    """覆っていると判定される（画素がすべて不透明な）画像。"""
    path = tmp_path / "opaque.png"
    path.write_bytes(opaque_png())
    return XLImage(str(path))


def _hundred_cells_xlsx(tmp_path) -> pathlib.Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "一覧"
    for column in range(1, 11):
        ws.cell(1, column, f"見出し{column}")
    for row in range(2, 12):
        for column in range(1, 11):
            ws.cell(row, column, f"値{row}-{column}")
    image = _opaque_image(tmp_path)
    image.anchor = TwoCellAnchor(
        editAs="oneCell", _from=AnchorMarker(col=0, colOff=0, row=1, rowOff=0),
        to=AnchorMarker(col=10, colOff=0, row=11, rowOff=0))
    ws.add_image(image)
    path = tmp_path / "hundred.xlsx"
    wb.save(path)
    return path


def test_one_figure_over_a_hundred_cells_lists_ten_and_counts_the_rest(tmp_path):
    p = _hundred_cells_xlsx(tmp_path)
    ir = evidence_spike.extract(p)
    assert len([r for r in ir.relations if r.type == observation_vocab.RELATION_COVERS]) == 100
    md = R.render(ir, source_name=p.name).markdown
    chunks = _chunks(md)
    front = _chunk_with(md, "覆っている対象")
    assert front.count("セル") == 10 and "ほか 90 件" in front
    relation_lines = [line for line in front.split("\n") if line.startswith(("覆っている対象", "ほか"))]
    assert sum(1 + len(line) for line in relation_lines) <= R._STATE_NOTE_MAX_CHARS
    cell_chunks = [c for c in chunks if "前面図形:" in c]
    assert sum(c.count("前面図形:") for c in cell_chunks) == 100          # 各セルに増えるのは前面を指す 1 行だけ
    for chunk in cell_chunks:
        assert chunk.count("前面図形:") == chunk.count("可視性: 「画像に覆われている」")
        assert all(len(line) <= R._STATE_NOTE_MAX_CHARS for line in chunk.split("\n"))


def test_relation_block_trims_to_the_budget_and_counts_the_omitted_lines():
    entries = [observation_vocab.FrontEntry(2, "cover_front_not_occluding", "半透明", f"対象{i}（図形 ID {i}）")
               for i in range(30)]
    entries += [f"前面図形: 相手{i}（図形 ID {i}）" for i in range(30)]
    lines = observation_vocab.relation_block(entries, 250)
    assert sum(1 + len(line) for line in lines) <= 250
    assert lines[-1].startswith("ほか ")
    shown_fronts = lines[0].count("（図形 ID")
    assert shown_fronts <= observation_vocab.MAX_FRONT_ITEMS
    assert int(re.search(r"ほか (\d+) 件", lines[-1]).group(1)) + shown_fronts + sum(
        1 for line in lines if line.startswith("前面図形:")) == 60


# ---- 人用 Markdown ----

def _docx_with_ole() -> pathlib.Path:
    p = _V021 / "04_添付/機器一覧と添付.docx"
    if not p.is_file():
        pytest.skip(f"fixture が無い環境: {p}")
    return p


def test_docx_human_md_notes_the_unparsed_ole_after_its_paragraph_without_internal_ids():
    md = office_md.to_markdown(_docx_with_ole())
    anchor, tail = md.index("貸出対象の機器は"), md.index("機器台帳の更新は")
    note = md.index("埋め込み（OLE）")
    assert anchor < note < tail
    assert "の中身は読み取っていません" in md and "は未対応で、内容を読めていません" in md
    assert not _INTERNAL.search(md)


def test_pptx_human_md_puts_front_and_back_sentences_after_each_shape_and_hides_internal_ids(tmp_path):
    md = office_md.to_markdown(_two_same_name_covered(tmp_path))
    left, right, front_note = md.index("左の説明"), md.index("右の説明"), md.index("を覆っています")
    assert left < right < front_note
    back_note = md[left:right]
    assert "※ この図形は「覆いの図形」（スライド 1 の 3 番目の図形・" in back_note
    assert "に覆われています（覆われていると判定）" in back_note
    assert "※ この図形は「同名」（スライド 1 の 1 番目の図形・" in md and "「同名」（スライド 1 の 2 番目の図形・" in md
    assert not _INTERNAL.search(md)


def test_pptx_human_md_notes_the_unparsed_ole_in_its_place():
    p = _V021 / "04_添付/貸出台帳の添付.pptx"
    if not p.is_file():
        pytest.skip(f"fixture が無い環境: {p}")
    md = office_md.to_markdown(p)
    assert md.index("2025年度の貸出台帳を添付する。") < md.index("※ この埋め込み（OLE）「台帳ファイル」の中身は読み取っていません")
    assert not _INTERNAL.search(md)


def _stamp_xlsx(tmp_path, *, strike=True) -> pathlib.Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "一覧"
    ws["A1"] = "見出し"
    ws["A2"] = "廃止された値"
    ws["A3"] = "取り消された値"
    if strike:
        ws["A3"].font = Font(strike=True)
    image = _opaque_image(tmp_path)
    image.anchor = TwoCellAnchor(
        editAs="oneCell", _from=AnchorMarker(col=0, colOff=0, row=1, rowOff=0),
        to=AnchorMarker(col=1, colOff=0, row=2, rowOff=0))
    ws.add_image(image)
    path = tmp_path / "stamp.xlsx"
    wb.save(path)
    return path


def test_xlsx_human_md_notes_state_relation_and_unparsed_image_after_the_table(tmp_path):
    md = office_md.to_markdown(_stamp_xlsx(tmp_path))
    notes = md[md.index("取り消された値 |"):]
    assert "※ セル A2は画像「Image 1」（シート「一覧」のセル範囲 A2:A2）に覆われています（覆われていると判定）" not in notes
    assert "※ セル A2は「Image 1」（シート「一覧」のセル範囲 A2:A2）に覆われています（覆われていると判定）" in notes
    assert "※ セル A3の文字には取り消し線が引かれています" in notes
    assert "は「セル A2」" not in notes and "セル A2を覆っています（覆われていると判定）" in notes
    assert "のセル範囲 A2:A2にある画像「Image 1」の内容は読み取っていません" in notes
    assert not _INTERNAL.search(md.replace("図形 ID", ""))


def test_hidden_slide_and_off_slide_use_plain_sentences(tmp_path):
    p = _two_same_name_covered(tmp_path)
    with zipfile.ZipFile(p) as z:
        entries = {n: z.read(n) for n in z.namelist()}
    entries["ppt/slides/slide1.xml"] = entries["ppt/slides/slide1.xml"].replace(b"<p:sld ", b'<p:sld show="0" ', 1)
    hidden = tmp_path / "hidden.pptx"
    with zipfile.ZipFile(hidden, "w") as z:
        for name, data in entries.items():
            z.writestr(name, data)
    md = office_md.to_markdown(hidden)
    assert "※ このスライドは非表示です" in md and not _INTERNAL.search(md)


# ---- RAG と人用 Markdown が同じ観測を同じ鍵・同じ語で書く ----

def test_rag_and_human_md_take_each_sentence_from_the_wording_tables(tmp_path, monkeypatch):
    rag_words = dict(observation_vocab.WORDING)
    rag_words["cover_back_confirmed"] = "RAG前面 {ref}"
    rag_words["cover_front_confirmed"] = "RAG背面 {refs}"
    rag_words["notice_ole"] = "RAG未解析 {name}"
    human_words = dict(observation_vocab.HUMAN_WORDING)
    human_words["cover_back_confirmed"] = "人前面 {subject} {ref}"
    human_words["cover_front_confirmed"] = "人背面 {subject} {refs}"
    human_words["notice_ole"] = "人未解析 {name}"
    monkeypatch.setattr(observation_vocab, "WORDING", rag_words)
    monkeypatch.setattr(observation_vocab, "HUMAN_WORDING", human_words)
    p = _two_same_name_covered(tmp_path)
    rag, human = _rag(p), office_md.to_markdown(p)
    assert "RAG前面 覆いの図形（図形 ID 7・スライド1）" in rag and "RAG背面" in rag and "人前面" not in rag
    assert "人前面 この図形 「覆いの図形」（スライド 1 の 3 番目の図形・" in human and "人背面" in human and "RAG前面" not in human
    ole = _V021 / "04_添付/貸出台帳の添付.pptx"
    if ole.is_file():
        assert "RAG未解析 台帳ファイル" in _rag(ole) and "人未解析 台帳ファイル" in office_md.to_markdown(ole)


@pytest.mark.parametrize("key,values,shared", [
    ("cover_back_confirmed", {"ref": "X", "fill": "不透明", "subject": "Y"}, ["覆われ", "判定"]),
    ("cover_back_undetermined", {"ref": "X", "fill": "", "subject": "Y"}, ["覆われ", "未確定"]),
    ("cover_back_not_occluding", {"ref": "X", "fill": "半透明", "subject": "Y"}, ["覆われてはいない", "半透明"]),
    ("cover_front_confirmed", {"refs": "X", "fill": "", "subject": "Y"}, ["覆", "判定"]),
    ("cover_front_undetermined", {"refs": "X", "fill": "", "subject": "Y"}, ["重なって", "未確定"]),
    ("cover_front_not_occluding", {"refs": "X", "fill": "半透明", "subject": "Y"}, ["重なって", "覆ってはいない", "半透明"]),
    ("connect_start", {"refs": "X", "ref": "X", "subject": "Y"}, ["接続", "矢印の向きは未確認"]),
    ("connect_end", {"ref": "X", "subject": "Y"}, ["接続", "矢印の向きは未確認"]),
    ("connect_connector", {"start": "X", "end": "Z", "subject": "Y"}, ["接続", "矢印の向きは未確認"]),
    ("connect_start_unknown", {"subject": "Y"}, ["接続先", "不明"]),
    ("connect_end_unknown", {"subject": "Y"}, ["接続元", "不明"]),
    ("strike_whole", {"subject": "Y"}, ["取り消し線"]),
    ("relation_more", {"count": 5}, ["ほか 5 件"]),
])
def test_observation_words_agree_between_rag_and_human_sentences(key, values, shared):
    rag, human = observation_vocab.say(key, **values), observation_vocab.say(key, human=True, **values)
    for word in shared:
        assert word in rag and word in human, (key, word)


def test_visibility_reasons_share_their_words_with_the_rag_labels():
    for reason, label in observation_vocab.VISIBILITY_REASON_LABELS.items():
        word = "覆われ" if reason.startswith("occluded") else "範囲外" if reason == "off_slide" else "非表示"
        sentence = observation_vocab.visibility_line(reason, label, human=True, subject="Y")
        assert word in label and word in sentence, reason
        assert "可視性" not in sentence and "「" not in sentence


@pytest.mark.parametrize("status,reason,kind", [
    ("metadata_only", "image_content_uninterpreted", ""),
    ("metadata_only", "ole_binary_metadata_only", ""),
    ("unsupported", "unsupported_chart_attributes", "chart_unparsed_attributes"),
    ("failed", "parse_failed", "smartart_data"),
])
def test_notices_say_the_same_observation_in_plain_words(status, reason, kind):
    rag = observation_vocab.notice_line(status, reason, kind, "N")
    human = observation_vocab.notice_line(status, reason, kind, "N", human=True)
    assert rag and human and human.startswith("この") and "N" in human
    assert not _INTERNAL.search(human)


# ---- 状態・関係の無い資料は変わらない ----

def test_plain_documents_are_unchanged_in_human_md_and_rag(tmp_path):
    pptx = _pptx(tmp_path, _sp_text("本文A", shape_id=2, name="a", x=0, y=0, cx=100000, cy=100000)
                 + _sp_text("本文B", shape_id=3, name="b", x=2000000, y=2000000, cx=100000, cy=100000), name="plain.pptx")
    assert office_md.to_markdown(pptx) == "## スライド 1\n\n本文A\n\n本文B"
    docx = tmp_path / "plain.docx"
    with zipfile.ZipFile(docx, "w") as z:
        z.writestr("word/document.xml",
                   '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                   "<w:body><w:p><w:r><w:t>段落A</w:t></w:r></w:p><w:p><w:r><w:t>段落B</w:t></w:r></w:p></w:body></w:document>")
    assert office_md.to_markdown(docx) == "段落A\n\n段落B"
    xlsx = _stamp_xlsx(tmp_path, strike=False)
    wb = openpyxl.load_workbook(xlsx)
    ws = wb.active
    ws._images = []
    plain_xlsx = tmp_path / "plain.xlsx"
    wb.save(plain_xlsx)
    assert "※" not in office_md.to_markdown(plain_xlsx)
    for path in (pptx, docx, plain_xlsx):
        rag = _rag(path)
        assert not re.search(r"覆っている対象|重なっている対象|前面図形|接続先|接続元|ほか \d+ 件|未解析", rag)
