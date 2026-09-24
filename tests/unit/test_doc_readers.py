"""原本読取ツールの中核（`sherpa/doc_readers.py`）の受け入れテスト。

正典: `docs/proposals/2026-09-10-Codex原本直読と調査スキル.md` §2-9・§4。DB/ES/Neo4j 不要
（`doc_readers` は**開いたバイナリファイルオブジェクト**だけを受ける純関数——world/scope/秘匿判定・
path→fd の TOCTOU 対策済み open は呼び出し元 `agentic_search.run_tool`/`_safe_original_path` の
責務であり、ここでは対象外）。

RV1巡目是正（`docs/20-開発ハーネス.md` の敵対 RV・S3b 台帳）: 公開関数の入力が「パス」から
「開いた file object」へ変わった（RV#1・TOCTOU 対策）ため、既存テストは全て `open(p, "rb")` を
経由するよう契約変更している。
"""
from __future__ import annotations

import os
import struct
import tempfile
import zipfile
from pathlib import Path

os.environ.setdefault("SHERPA_USE_FIXTURES", "1")

import docx
import openpyxl
import pptx
import pytest
from pypdf import PdfWriter

from sherpa import doc_readers as DR


@pytest.fixture()
def tmp() -> Path:
    return Path(tempfile.mkdtemp())


def _open(p: Path):
    return open(p, "rb")


# ===== xlsx_sheets / xlsx_range =====

def _make_xlsx(path: Path, rows: int = 11, cols: int = 5) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            ws.cell(row=r, column=c, value=f"r{r}c{c}")
    wb.save(path)


def test_xlsx_sheets_returns_name_and_size(tmp):
    p = tmp / "a.xlsx"
    _make_xlsx(p)
    result = DR.xlsx_sheets(_open(p))
    assert result == {"sheets": [{"name": "Sheet1", "max_row": 11, "max_col": 5}]}


def test_xlsx_range_default_and_explicit_a1(tmp):
    p = tmp / "a.xlsx"
    _make_xlsx(p)
    default = DR.xlsx_range(_open(p), "Sheet1")
    assert default["range"] == "A1:E11"
    assert default["rows"][0] == ["r1c1", "r1c2", "r1c3", "r1c4", "r1c5"]
    assert default["truncated"] is False

    ranged = DR.xlsx_range(_open(p), "Sheet1", "B3:D5")
    assert ranged == {"sheet": "Sheet1", "range": "B3:D5",
                      "rows": [["r3c2", "r3c3", "r3c4"], ["r4c2", "r4c3", "r4c4"],
                              ["r5c2", "r5c3", "r5c4"]],
                      "truncated": False}


def test_xlsx_range_truncates_when_exceeding_max_rows_cols(tmp):
    p = tmp / "a.xlsx"
    _make_xlsx(p, rows=10, cols=10)
    r = DR.xlsx_range(_open(p), "Sheet1", max_rows=3, max_cols=2)
    assert r["truncated"] is True
    assert r["range"] == "A1:B3"
    assert len(r["rows"]) == 3 and all(len(row) == 2 for row in r["rows"])


def test_xlsx_range_none_cell_becomes_empty_string(tmp):
    p = tmp / "a.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "x"
    ws["C1"] = "y"
    # B1 は未設定＝None（シート次元は A1:C1 まで広がるので明示 range で拾う）。
    wb.save(p)
    r = DR.xlsx_range(_open(p), "Sheet1", "A1:C1")
    assert r["rows"] == [["x", "", "y"]]


def test_xlsx_range_unknown_sheet_and_bad_range(tmp):
    p = tmp / "a.xlsx"
    _make_xlsx(p)
    assert DR.xlsx_range(_open(p), "NoSuchSheet") == {"error": "シートが見つかりません"}
    assert DR.xlsx_range(_open(p), "Sheet1", "not-a-range") == {"error": "range が不正です"}


def test_xlsx_sheets_broken_file_returns_error(tmp):
    p = tmp / "broken.xlsx"
    p.write_text("this is not a zip/xlsx file")
    assert DR.xlsx_sheets(_open(p)) == {"error": "Excel を開けませんでした"}


# ===== RV#10: max_rows/max_cols は引数で無制限にできない（200/50 の上限にクランプ）=============

def test_xlsx_range_max_rows_cols_are_clamped_even_if_caller_asks_for_more(tmp):
    p = tmp / "a.xlsx"
    _make_xlsx(p, rows=300, cols=60)
    r = DR.xlsx_range(_open(p), "Sheet1", max_rows=10_000, max_cols=10_000)
    assert len(r["rows"]) == 200                 # 上限 200 行にクランプ（引数で超えられない）
    assert all(len(row) == 50 for row in r["rows"])   # 上限 50 列にクランプ
    assert r["truncated"] is True


# ===== RV#2: セルの伏せ字は 200 文字への切り詰めより先に掛かる（境界をまたぐ秘密も検出）========

def test_xlsx_range_clean_is_applied_before_cell_truncation(tmp):
    p = tmp / "a.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    # PEM 形式の秘密鍵らしき長いセル（300字）——先に 200 字へ切ってから伏せ字を掛けると
    # 正規表現が完全な形の秘密パターンを見られず検出をすり抜ける（RV#2 の実害）。
    pem_body = "\n".join("A" * 60 for _ in range(5))
    secret = f"-----BEGIN PRIVATE KEY-----\n{pem_body}\n-----END PRIVATE KEY-----"
    assert len(secret) > 200
    ws["A1"] = secret
    wb.save(p)

    from sherpa.agentic_search import _redact
    r = DR.xlsx_range(_open(p), "Sheet1", "A1:A1", clean=_redact)
    cell = r["rows"][0][0]
    assert "[REDACTED]" in cell
    assert "BEGIN PRIVATE KEY" not in cell
    assert len(cell) <= 200


def test_xlsx_range_key_block_state_survives_cell_truncation(tmp):
    """状態付き伏せ字は「原本の順序で・切断より前」に掛ける——A1（190字＋BEGIN、200字に切ると
    BEGIN 自体が途中で欠ける長さ）・A2（鍵本文のみ）・A3（END）という並びで、A2 の鍵本文が
    そのまま漏れない（clean が呼ぶたびに前セルの状態を引き継がない bare な関数だと、A1 の BEGIN
    が切断で欠けて検出できず、A2 が伏せられないまま 200 字切り詰めで漏れていた）。"""
    p = tmp / "key.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "x" * 190 + "-----BEGIN RSA PRIVATE KEY-----"
    ws["A2"] = "A" * 300
    ws["A3"] = "-----END RSA PRIVATE KEY----- suffix"
    wb.save(p)

    from sherpa.agentic_search import _redact
    r = DR.xlsx_range(_open(p), "Sheet1", "A1:A3", clean=_redact)
    a1, a2, a3 = (row[0] for row in r["rows"])
    assert "BEGIN" not in a1 and "[REDACTED]" in a1
    assert a2 == "[REDACTED]"           # 鍵本文セルは丸ごと伏せる（漏れない）
    assert "AAAA" not in a2
    assert "END" not in a3 and "suffix" in a3


def test_xlsx_range_selected_subrange_still_tracks_key_block_from_sheet_head(tmp):
    """選択範囲（`A2:A3`）より**前**の行（A1）で始まった鍵ブロックは、選択範囲だけを読むと
    BEGIN を見落として状態が立たないまま鍵本文（A2）が漏れる——状態はシート先頭行から選択範囲
    の終わりまでを辿って確定し、出力だけを選択範囲に絞る。"""
    p = tmp / "key_subrange.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "-----BEGIN RSA PRIVATE KEY-----"
    ws["A2"] = "A" * 100
    ws["A3"] = "-----END RSA PRIVATE KEY-----"
    wb.save(p)

    from sherpa.agentic_search import _redact
    r = DR.xlsx_range(_open(p), "Sheet1", "A2:A3", clean=_redact)
    assert r["range"] == "A2:A3"
    a2, a3 = (row[0] for row in r["rows"])
    assert a2 == "[REDACTED]"           # 選択範囲外（A1）の BEGIN を辿らないと漏れていた鍵本文
    assert "AAAA" not in a2


# ===== redact_keys.KeyBlockRedactor: BEGIN 検出は base_clean より前（生文字列）に行う ============

def test_key_block_redactor_detects_begin_before_base_clean_mangles_it():
    """`base_clean`（kv 秘密パターン）を先に掛けると、鍵ラベルの直後に改行を挟んで BEGIN が
    続く場合（例: `"secret:\\n-----BEGIN..."`）、kv パターンの `\\s*` が改行をまたいで
    `\\S+` が BEGIN の文字列そのものを1トークンとして飲み込み、`[REDACTED]` に置き換えてしまう
    ——結果 BEGIN マーカーの文字列自体が消え、以後どの要素にも状態が立たず鍵本文が漏れる。
    BEGIN/END の検出を生文字列に対して先に行うことで防ぐ。"""
    from sherpa.agentic_search import _redact
    from sherpa.redact_keys import KeyBlockRedactor

    redactor = KeyBlockRedactor(_redact)
    out1 = redactor("secret: x\n-----BEGIN RSA PRIVATE KEY-----")
    assert "BEGIN" not in out1
    assert "[REDACTED]" in out1

    out2 = redactor("A" * 100)          # 鍵本文（次の要素）
    assert out2 == "[REDACTED]"         # BEGIN を見失っていなければここが伏せられる
    assert "AAAA" not in out2

    out3 = redactor("-----END RSA PRIVATE KEY----- tail")
    assert "END" not in out3 and "tail" in out3


def test_xlsx_range_without_clean_keeps_raw_cell_truncated_at_200(tmp):
    p = tmp / "a.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "x" * 300
    wb.save(p)
    r = DR.xlsx_range(_open(p), "Sheet1", "A1:A1")
    assert r["rows"][0][0] == "x" * 200


# ===== RV#3: zip 展開サイズ上限（openpyxl 等に渡す前に central directory だけで見積もる）========

def test_xlsx_sheets_rejects_zip_with_huge_declared_uncompressed_entry(tmp, monkeypatch):
    # 圧縮後のファイル自体は小さいが、中の1エントリの展開後サイズ（zip の central directory に
    # 記録される実際の file_size）が展開上限を超える zip（0 埋めデータは deflate で激しく圧縮できる
    # ため、ディスク上は数KBのまま「展開したら巨大」を再現できる＝zip bomb 対策の実害シナリオ）。
    monkeypatch.setenv("SHERPA_DOC_READ_MAX_UNZIP_BYTES", str(10 * 1024 * 1024))   # 10 MiB
    p = tmp / "bomb.xlsx"
    huge = b"\x00" * (64 * 1024 * 1024)                     # 64 MiB
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/worksheets/sheet1.xml", huge)
    assert os.path.getsize(p) < 1024 * 1024                 # 圧縮後は 1MiB 未満（_too_big_fd は通る）
    assert DR.xlsx_sheets(_open(p)) == {"error": "大きすぎて開けません（展開サイズ）"}


def test_xlsx_sheets_rejects_zip_with_too_many_entries(tmp, monkeypatch):
    monkeypatch.setattr(DR, "_MAX_ZIP_ENTRIES", 5)
    p = tmp / "manyentries.xlsx"
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zf:
        for i in range(10):
            zf.writestr(f"part{i}.xml", b"x")
    assert DR.xlsx_sheets(_open(p)) == {"error": "大きすぎて開けません（展開サイズ）"}


def test_xlsx_sheets_missing_file_returns_error(tmp):
    with pytest.raises(FileNotFoundError):
        _open(tmp / "missing.xlsx")


def _spoof_central_dir_uncompressed_size(path: Path, member: str, fake_size: int) -> None:
    """`path` 内の `member` の central directory レコードの「展開後サイズ」フィールド（sig 直後
    24 バイト目）だけを直接書き換える。compress_size/CRC は元のまま（実データも無改変）——
    宣言 `file_size` だけが実際より小さい zip を作る（RV2巡目#2 是正の再現用）。"""
    raw = bytearray(path.read_bytes())
    name_b = member.encode()
    sig = b"PK\x01\x02"
    idx = 0
    while True:
        idx = raw.find(sig, idx)
        if idx == -1:
            raise AssertionError(f"central directory entry not found: {member}")
        name_len = struct.unpack_from("<H", raw, idx + 28)[0]
        cand = bytes(raw[idx + 46:idx + 46 + name_len])
        if cand == name_b:
            struct.pack_into("<I", raw, idx + 24, fake_size)
            path.write_bytes(bytes(raw))
            return
        idx += 4


def test_xlsx_sheets_rejects_zip_with_spoofed_small_declared_size(tmp):
    # central directory の file_size を実際より小さく偽装した zip（compress_size/CRC は元のまま
    # ＝実データは無改変の 2MB）。宣言値だけを見る `_zip_bounds_error` は素通りするが、実際に
    # 読み流す `_zip_actual_size_error` は CRC 不一致（zipfile 自身が申告サイズで打ち切って検出）
    # として拒否する（RV2巡目#2）。
    p = tmp / "spoofed.xlsx"
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/worksheets/sheet1.xml", b"x" * (2 * 1024 * 1024))
    _spoof_central_dir_uncompressed_size(p, "xl/worksheets/sheet1.xml", 500)
    assert DR.xlsx_sheets(_open(p)) == {"error": "大きすぎて開けません（展開サイズ）"}


def test_docx_paragraphs_rejects_zip_with_spoofed_small_declared_size(tmp):
    # 同じ細工を docx にも適用（読取ツール別の実装差で穴が残らないことの確認）。
    p = tmp / "spoofed.docx"
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", b"y" * (2 * 1024 * 1024))
    _spoof_central_dir_uncompressed_size(p, "word/document.xml", 500)
    assert DR.docx_paragraphs(_open(p)) == {"error": "大きすぎて開けません（展開サイズ）"}


def _spoof_central_dir_size_and_crc(path: Path, member: str, fake_size: int, fake_crc: int) -> None:
    """central directory の「展開後サイズ」（sig+24 バイト目）と CRC（sig+16 バイト目）の**両方**を
    書き換える（宣言サイズと CRC の偽装の再現用）。`fake_crc` を「宣言サイズぶんの実データ」の CRC に
    合わせると、宣言値に達した時点で打ち切って CRC だけ突き合わせる検査（旧 `zf.open()` 依存）は
    矛盾を検知できない。"""
    raw = bytearray(path.read_bytes())
    name_b = member.encode()
    sig = b"PK\x01\x02"
    idx = 0
    while True:
        idx = raw.find(sig, idx)
        if idx == -1:
            raise AssertionError(f"central directory entry not found: {member}")
        name_len = struct.unpack_from("<H", raw, idx + 28)[0]
        cand = bytes(raw[idx + 46:idx + 46 + name_len])
        if cand == name_b:
            struct.pack_into("<I", raw, idx + 16, fake_crc)
            struct.pack_into("<I", raw, idx + 24, fake_size)
            path.write_bytes(bytes(raw))
            return
        idx += 4


def _duplicate_central_dir_entry_same_header_offset(path: Path, member: str, new_name: str) -> None:
    """`member` の central directory レコードを複製し、`new_name`（`member` と同じ長さ）を
    名前に持つ2件目の entry として central directory 内（複製元の直後・EOCD より前）へ
    挿入する——複製元と複製先は同じ `header_offset`（＝同じ1つの local file header／同じ実データ）
    を指す。EOCD のエントリ数・central directory サイズも整合するよう書き換える
    （100 件の entry が同じ実データを指す zip bomb の再現を、2件で最小化したもの）。"""
    assert len(new_name) == len(member), "この再現には同じ長さの名前が必要"
    raw = bytearray(path.read_bytes())
    sig_cd = b"PK\x01\x02"
    name_b = member.encode()
    idx = 0
    rec_start = None
    while True:
        idx = raw.find(sig_cd, idx)
        if idx == -1:
            raise AssertionError(f"central directory entry not found: {member}")
        name_len, extra_len, comment_len = struct.unpack_from("<HHH", raw, idx + 28)
        cand = bytes(raw[idx + 46:idx + 46 + name_len])
        rec_len = 46 + name_len + extra_len + comment_len
        if cand == name_b:
            rec_start = idx
            break
        idx += rec_len
    name_len, extra_len, comment_len = struct.unpack_from("<HHH", raw, rec_start + 28)
    rec_len = 46 + name_len + extra_len + comment_len
    record = bytearray(raw[rec_start:rec_start + rec_len])
    record[46:46 + name_len] = new_name.encode()   # header_offset フィールドはコピーのまま＝同一を指す

    eocd_idx = raw.rfind(b"PK\x05\x06")
    assert eocd_idx != -1, "EOCD が見つからない"
    cd_size, _cd_offset = struct.unpack_from("<II", raw, eocd_idx + 12)
    entries_this_disk, entries_total = struct.unpack_from("<HH", raw, eocd_idx + 8)

    insert_at = rec_start + rec_len
    new_raw = raw[:insert_at] + record + raw[insert_at:]
    new_eocd_idx = eocd_idx + len(record)
    struct.pack_into("<H", new_raw, new_eocd_idx + 8, entries_this_disk + 1)
    struct.pack_into("<H", new_raw, new_eocd_idx + 10, entries_total + 1)
    struct.pack_into("<I", new_raw, new_eocd_idx + 12, cd_size + len(record))
    path.write_bytes(bytes(new_raw))


def test_xlsx_sheets_rejects_zip_with_duplicate_entries_pointing_to_same_local_header(tmp):
    """central directory に同じ `header_offset`（＝同じ1つの local file header・同じ実データ）を
    指す entry を複数持つ zip は、宣言サイズどうしの整合性検査を entry 単位で見ているだけでは
    通り得る（実データは1件ぶんのままなので合計も小さい）が、実際に読み流す段になると同じ実
    データを entry の数だけ繰り返し伸長させられる（100 entry なら 100 倍の増幅）。伸長を始める
    前に圧縮データ区間の重複を見て拒否する。"""
    p = tmp / "dup.xlsx"
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("a.xml", b"hello world" * 100)
    _duplicate_central_dir_entry_same_header_offset(p, "a.xml", "b.xml")
    assert DR.xlsx_sheets(_open(p)) == {"error": "大きすぎて開けません（展開サイズ）"}


def test_xlsx_sheets_accepts_normal_multi_entry_zip_without_overlap(tmp):
    """複数 entry が別々の実データを指す通常の zip は、重複検査に引っかからず正常に開ける
    （新しい重複検査が誤検知しないことの確認）。"""
    p = tmp / "normal.xlsx"
    _make_xlsx(p)
    assert "error" not in DR.xlsx_sheets(_open(p))


def test_xlsx_sheets_rejects_zip_with_declared_size_and_crc_spoofed_to_match_truncated_prefix(tmp, monkeypatch):
    """旧 `_zip_actual_size_error` は `zf.open()`（central directory の宣言
    `file_size` に達した時点で decompress を打ち切り、そこまでの実データの CRC を central
    directory の CRC と突き合わせるだけ）に依存していた——宣言 `file_size` を実際より小さい値へ
    偽装し、CRC も「その宣言サイズぶんの実データ」の CRC に合わせて偽装すれば、`zf.open()` は
    何の矛盾も検知しないまま宣言サイズぶんだけ読んで正常終了していた（実際の展開後サイズは
    無関係に巨大なまま＝zip bomb がすり抜ける）。宣言値に一切依存しない実測（local header から
    圧縮データの開始位置を求め、`compress_size` バイトを実際に伸長してバイト数を数える）へ
    切り替えたことで、この偽装は検出される。
    """
    monkeypatch.setenv("SHERPA_DOC_READ_MAX_UNZIP_BYTES", str(1 * 1024 * 1024))   # 1 MiB
    p = tmp / "bomb2.xlsx"
    real_xml = b"<worksheet>" + b"x" * 526 + b"</worksheet>"
    assert len(real_xml) == 549
    content = real_xml + b" " * (16 * 1024 * 1024)   # 実データは 549 バイト + 16 MiB
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/worksheets/sheet1.xml", content)
    fake_crc = zipfile.crc32(content[:549]) & 0xFFFFFFFF   # 「宣言サイズぶんの実データ」の CRC
    _spoof_central_dir_size_and_crc(p, "xl/worksheets/sheet1.xml", 549, fake_crc)
    assert DR.xlsx_sheets(_open(p)) == {"error": "大きすぎて開けません（展開サイズ）"}


def _set_encrypted_flag(path: Path, member: str) -> None:
    """central directory と local header の両方の general purpose flag に暗号化ビット
    （bit 0）を立てる（暗号化フラグ偽装の再現用：`zf.open()` は暗号化 entry を `RuntimeError(
    "File %r is encrypted...")` で拒否し、この `%r` に zip 内のファイル名（攻撃者が自由に
    埋め込める）がそのまま入る）。"""
    raw = bytearray(path.read_bytes())
    name_b = member.encode()
    sig_cd = b"PK\x01\x02"
    idx = 0
    cd_found = False
    while True:
        idx = raw.find(sig_cd, idx)
        if idx == -1:
            break
        name_len = struct.unpack_from("<H", raw, idx + 28)[0]
        cand = bytes(raw[idx + 46:idx + 46 + name_len])
        if cand == name_b:
            flags = struct.unpack_from("<H", raw, idx + 8)[0]
            struct.pack_into("<H", raw, idx + 8, flags | 0x1)
            cd_found = True
            break
        idx += 4
    assert cd_found, f"central directory entry not found: {member}"
    sig_lh = b"PK\x03\x04"
    idx = 0
    lh_found = False
    while True:
        idx = raw.find(sig_lh, idx)
        if idx == -1:
            break
        name_len = struct.unpack_from("<H", raw, idx + 26)[0]
        cand = bytes(raw[idx + 30:idx + 30 + name_len])
        if cand == name_b:
            flags = struct.unpack_from("<H", raw, idx + 6)[0]
            struct.pack_into("<H", raw, idx + 6, flags | 0x1)
            lh_found = True
            break
        idx += 4
    assert lh_found, f"local header not found: {member}"
    path.write_bytes(bytes(raw))


def test_xlsx_sheets_encrypted_entry_with_secret_name_does_not_leak_or_raise(tmp):
    """暗号化フラグ付き entry を含む zip の実地再現: 旧 `_zip_actual_size_error`
    が事前検査の段階で（すべての entry に対して無条件に）`zf.open()` を呼んでいたため、
    `RuntimeError("File %r is encrypted...")`（member 名＝ここでは `api_key=SECRET` を含む）が
    `_precheck_office`/`xlsx_sheets` のどちらにも捕まらず素通りしていた。検査自体を `zf.open()`
    非依存の実測へ書き換えたため事前検査は例外を出さず通過し、実パーサ（openpyxl）が暗号化を
    検知して失敗しても `xlsx_sheets` の既存の固定文言（`except Exception` で名前を含めない）に
    丸まる——いずれの経路でも例外が伝播せず、秘密っぽい名前が結果に出ない。"""
    p = tmp / "encrypted.xlsx"
    member = "api_key=SECRET"
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(member, b"dummy content")
    _set_encrypted_flag(p, member)
    result = DR.xlsx_sheets(_open(p))
    assert "SECRET" not in repr(result)
    assert "error" in result


def test_precheck_office_wraps_unexpected_inspection_exception_into_generic_error(tmp, monkeypatch):
    """`_precheck_office` は `_zip_bounds_error`/`_zip_actual_size_error` からの
    **想定外の例外**（`BadZipFile`/`OSError` 以外・ここでは暗号化 zip の実際の例外文言を模す）を
    `except Exception` で捕まえ、member 名を含まない固定文言（`{"error": "Office ファイルを
    開けませんでした"}`）に丸める——多層防御（`_zip_actual_size_error` 自身が想定していない
    経路で例外を出しても、ここで最後に受け止める）。"""
    p = tmp / "a.xlsx"
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("x.xml", b"data")
    secret = "api_key=SECRET"

    def _boom(f):
        raise RuntimeError(f"File {secret!r} is encrypted, password required for extraction")

    monkeypatch.setattr(DR, "_zip_actual_size_error", _boom)
    result = DR.xlsx_sheets(_open(p))
    assert result == {"error": "Office ファイルを開けませんでした"}
    assert "SECRET" not in repr(result)


# ===== docx_paragraphs =====

def test_docx_paragraphs_returns_indexed_paragraphs_and_tables(tmp):
    p = tmp / "a.docx"
    d = docx.Document()
    for i in range(5):
        d.add_paragraph(f"para {i}")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "h1"
    t.rows[0].cells[1].text = "h2"
    d.save(p)

    r = DR.docx_paragraphs(_open(p), start=0, count=3)
    assert r["total"] == 5
    assert r["paragraphs"] == [{"i": 0, "style": "Normal", "text": "para 0"},
                               {"i": 1, "style": "Normal", "text": "para 1"},
                               {"i": 2, "style": "Normal", "text": "para 2"}]
    assert r["tables"] == [{"i": 0, "row_start": 0, "total_rows": 2, "rows": [["h1", "h2"], ["", ""]]}]
    assert r["truncated"] is True                # start+count(3) < total(5)

    r2 = DR.docx_paragraphs(_open(p), start=3, count=200)
    assert r2["paragraphs"][0]["i"] == 3
    assert r2["truncated"] is False


def test_docx_paragraphs_table_and_paragraph_limits(tmp):
    p = tmp / "big.docx"
    d = docx.Document()
    d.add_paragraph("only one")
    for _ in range(25):                          # > _DOCX_TABLE_MAX (20)
        d.add_table(rows=1, cols=1)
    d.save(p)
    r = DR.docx_paragraphs(_open(p))
    assert len(r["tables"]) == 20
    assert r["truncated"] is True


def test_docx_paragraphs_broken_file_returns_error(tmp):
    p = tmp / "broken.docx"
    p.write_bytes(b"not a docx")
    assert DR.docx_paragraphs(_open(p)) == {"error": "Word を開けませんでした"}


def test_docx_paragraphs_applies_clean_to_paragraph_and_table_text(tmp):
    p = tmp / "a.docx"
    d = docx.Document()
    d.add_paragraph("token=SECRET123")
    d.save(p)
    from sherpa.agentic_search import _redact
    r = DR.docx_paragraphs(_open(p), clean=_redact)
    assert "[REDACTED]" in r["paragraphs"][0]["text"]


def test_docx_paragraphs_key_block_spans_paragraph_table_paragraph(tmp):
    """本文の並びが「BEGIN を含む段落→鍵本文の表セル→END を含む段落」の docx で、表のセルの
    鍵本文が漏れない。`doc.paragraphs`/`doc.tables` は別々のフラットなリストで本文の出現順を
    保たないため、本文どおりの順序（`doc.iter_inner_content()`）を辿らないと、表が段落の間に
    挟まっている事実を無視して「段落を全部処理してから表を処理する」順になり、鍵ブロックの
    状態（BEGIN の後・END の前）を表のセルへ引き継げない。"""
    p = tmp / "keyblock.docx"
    d = docx.Document()
    d.add_paragraph("prefix -----BEGIN RSA PRIVATE KEY-----")
    t = d.add_table(rows=1, cols=1)
    t.rows[0].cells[0].text = "A" * 200
    d.add_paragraph("-----END RSA PRIVATE KEY----- suffix")
    d.save(p)

    from sherpa.agentic_search import _doc_reader_text_locator, _redact
    r = DR.docx_paragraphs(_open(p), clean=_redact)
    assert "BEGIN" not in r["paragraphs"][0]["text"]
    assert "[REDACTED]" in r["paragraphs"][0]["text"] and "prefix" in r["paragraphs"][0]["text"]
    assert r["tables"][0]["rows"][0][0] == "[REDACTED]"          # 鍵本文セルは丸ごと伏せる
    assert "END" not in r["paragraphs"][1]["text"]
    assert "suffix" in r["paragraphs"][1]["text"]

    # 合成 text（read_evidence 用・`agentic_search._doc_reader_text_locator`）からも漏れない。
    text, _locator = _doc_reader_text_locator("docx_paragraphs", r)
    assert "AAAA" not in text
    assert "BEGIN" not in text and "END" not in text


def test_docx_paragraphs_vertically_merged_cell_end_is_not_double_supplied(tmp):
    """python-docx の `row.cells` は縦結合セルについて同一の `_tc`（XML 要素）を結合範囲の
    行数ぶん重複して返す。重複を無視してそのまま状態付き伏せ字へ通すと、結合セル（ここでは
    END マーカー）が2回供給され、2回目の出現で状態が誤って閉じてしまい、その直後にある
    本当の鍵ブロック（B1 の BEGIN〜B2 の鍵本文）が伏せられずに漏れる。並びは「段落＝未閉の
    BEGIN→2×2 表（A列＝縦結合セルに END・B1＝BEGIN・B2＝鍵本文）」。"""
    p = tmp / "merged.docx"
    d = docx.Document()
    d.add_paragraph("prefix -----BEGIN RSA PRIVATE KEY-----")
    t = d.add_table(rows=2, cols=2)
    merged = t.cell(0, 0).merge(t.cell(1, 0))
    merged.text = "-----END RSA PRIVATE KEY-----"
    t.cell(0, 1).text = "-----BEGIN RSA PRIVATE KEY-----"
    t.cell(1, 1).text = "A" * 100
    d.save(p)

    from sherpa.agentic_search import _redact
    r = DR.docx_paragraphs(_open(p), clean=_redact)
    rows = r["tables"][0]["rows"]
    assert rows[1][1] == "[REDACTED]"          # B2: B1 の BEGIN 以降の鍵本文（漏れてはいけない）
    assert "AAAA" not in rows[1][1]


# ===== pptx_slides =====

def _make_pptx(path: Path, n: int = 3) -> None:
    prs = pptx.Presentation()
    for i in range(n):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"slide {i}"
    prs.save(path)


def test_pptx_slides_default_and_explicit_pages(tmp):
    p = tmp / "a.pptx"
    _make_pptx(p, n=3)
    r = DR.pptx_slides(_open(p), pages="1-2")
    assert r["total"] == 3
    assert [s["no"] for s in r["slides"]] == [1, 2]
    assert r["slides"][0]["texts"] == ["slide 0"]
    assert r["truncated"] is False


def test_pptx_slides_page_spec_variants(tmp):
    p = tmp / "a.pptx"
    _make_pptx(p, n=5)
    assert [s["no"] for s in DR.pptx_slides(_open(p), pages="3")["slides"]] == [3]
    assert [s["no"] for s in DR.pptx_slides(_open(p), pages="1,3,5")["slides"]] == [1, 3, 5]
    # 範囲外のページ番号は無視される。
    assert [s["no"] for s in DR.pptx_slides(_open(p), pages="4,9")["slides"]] == [4]


def test_pptx_slides_selected_page_still_tracks_key_block_from_first_slide(tmp):
    """選択（`pages="2"`）より**前**の1枚目で始まった鍵ブロックは、1枚目を読み飛ばすと BEGIN を
    見落として状態が立たないまま鍵本文（2枚目）が漏れる——状態は1枚目から選択の最終スライド
    まで辿って確定し、出力だけを選択したスライドに絞る。"""
    p = tmp / "key.pptx"
    prs = pptx.Presentation()
    for text in ("-----BEGIN RSA PRIVATE KEY-----", "A" * 100, "-----END RSA PRIVATE KEY-----"):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = text
    prs.save(p)

    from sherpa.agentic_search import _redact
    r = DR.pptx_slides(_open(p), pages="2", clean=_redact)
    assert [s["no"] for s in r["slides"]] == [2]
    assert r["slides"][0]["texts"] == ["[REDACTED]"]     # 1枚目の BEGIN を辿らないと漏れていた本文


def test_pptx_slides_broken_file_returns_error(tmp):
    p = tmp / "broken.pptx"
    p.write_bytes(b"not a pptx")
    assert DR.pptx_slides(_open(p)) == {"error": "PowerPoint を開けませんでした"}


# ===== pdf_pages =====

def _make_pdf(path: Path, n: int = 7) -> None:
    w = PdfWriter()
    for _ in range(n):
        w.add_blank_page(width=200, height=200)
    with path.open("wb") as f:
        w.write(f)


def test_pdf_pages_default_and_range(tmp):
    p = tmp / "a.pdf"
    _make_pdf(p, n=7)
    r = DR.pdf_pages(_open(p))                              # 既定 "1-5"
    assert r["total"] == 7
    assert [pg["no"] for pg in r["pages"]] == [1, 2, 3, 4, 5]
    assert r["truncated"] is False
    assert all(pg["text"] == "" for pg in r["pages"])  # 空白ページ＝抽出結果は空文字

    ranged = DR.pdf_pages(_open(p), pages="2-4")
    assert [pg["no"] for pg in ranged["pages"]] == [2, 3, 4]


def test_pdf_pages_missing_file_returns_error(tmp):
    with pytest.raises(FileNotFoundError):
        _open(tmp / "missing.pdf")


# ===== RV#4: ページ指定で長時間占有しない（範囲外の巨大 hi は列挙前に total へ切る）=============

def test_parse_page_spec_huge_range_with_small_total_resolves_instantly():
    pages, truncated = DR._parse_page_spec("1-1000000000000", total=1, max_count=10)
    assert pages == [1]
    assert truncated is False


def test_parse_page_spec_huge_range_stops_enumeration_at_max_count():
    pages, truncated = DR._parse_page_spec("1-1000000000000", total=1_000_000, max_count=5)
    assert pages == [1, 2, 3, 4, 5]
    assert truncated is True


# ===== file_head =====

def test_file_head_reads_within_cap_and_redacts_nothing_itself(tmp):
    p = tmp / "note.txt"
    p.write_text("hello\nworld\n", encoding="utf-8")
    r = DR.file_head(_open(p), max_bytes=1024)
    assert r == {"size": 12, "text": "hello\nworld\n", "truncated": False}


def test_file_head_truncates_at_max_bytes(tmp):
    p = tmp / "note.txt"
    p.write_text("x" * 100, encoding="utf-8")
    r = DR.file_head(_open(p), max_bytes=10)
    assert r["truncated"] is True
    assert r["text"] == "x" * 10
    assert r["size"] == 100


def test_file_head_missing_file_returns_error(tmp):
    with pytest.raises(FileNotFoundError):
        _open(tmp / "missing.txt")


def test_file_head_applies_clean(tmp):
    p = tmp / "note.txt"
    p.write_text("password=hunter2", encoding="utf-8")
    from sherpa.agentic_search import _redact
    r = DR.file_head(_open(p), clean=_redact)
    assert "[REDACTED]" in r["text"]


def test_file_head_read_os_error_carries_read_io_error_code(tmp):
    """open 成功後（TOCTOU 検証済み）の `read()` 自体が `OSError` で失敗した場合、結果に固定理由
    コード `error_code: "read_io_failed"` が付く——`agentic_search._record_tool_result_error_code`
    が名前非依存で拾い `InvestigationState.backend_failures["read_io"]` へ反映する経路。"""
    p = tmp / "note.txt"
    p.write_text("hello\n", encoding="utf-8")
    f = _open(p)

    def boom_read(n):
        raise OSError("boom")

    f.read = boom_read           # open 自体は成功済み・read() 段だけ壊す
    r = DR.file_head(f)
    assert r == {"error": "ファイルを開けませんでした", "error_code": "read_io_failed"}


# ===== サイズ上限（`SHERPA_DOC_READ_MAX_BYTES`）=====

def test_size_limit_blocks_before_opening(tmp, monkeypatch):
    p = tmp / "note.txt"
    p.write_text("x" * 100, encoding="utf-8")
    monkeypatch.setenv("SHERPA_DOC_READ_MAX_BYTES", "10")
    assert DR.file_head(_open(p)) == {"error": "大きすぎて開けません（サイズ）"}
    assert DR.xlsx_sheets(_open(p)) == {"error": "大きすぎて開けません（サイズ）"}


def test_size_limit_invalid_env_falls_back_to_default(tmp, monkeypatch):
    p = tmp / "note.txt"
    p.write_text("hello", encoding="utf-8")
    monkeypatch.setenv("SHERPA_DOC_READ_MAX_BYTES", "not-a-number")
    assert DR.file_head(_open(p)) == {"size": 5, "text": "hello", "truncated": False}


# ===== ページ指定パース（pptx/pdf 共通の `_parse_page_spec`）=====

def test_parse_page_spec_default_empty_and_truncation():
    pages, truncated = DR._parse_page_spec(None, total=25, max_count=10)
    assert pages == list(range(1, 11))
    assert truncated is True

    pages2, truncated2 = DR._parse_page_spec("", total=3, max_count=10)
    assert pages2 == [1, 2, 3]
    assert truncated2 is False


def test_parse_page_spec_dedupes_and_sorts():
    pages, truncated = DR._parse_page_spec("5,1,3-4,3", total=10, max_count=10)
    assert pages == [1, 3, 4, 5]
    assert truncated is False


# ===== キャッシュ（fstat (dev, ino, mtime, size) キー・LRU）=====

def test_cached_load_reuses_object_for_same_file(tmp):
    p = tmp / "a.xlsx"
    _make_xlsx(p)
    calls = []

    def loader(f):
        calls.append(f)
        return object()

    obj1 = DR._cached_load(_open(p), loader)
    obj2 = DR._cached_load(_open(p), loader)
    assert obj1 is obj2
    assert len(calls) == 1


def test_cached_load_reloads_after_file_modified(tmp):
    p = tmp / "a.xlsx"
    _make_xlsx(p)
    calls = []

    def loader(f):
        calls.append(f)
        return object()

    obj1 = DR._cached_load(_open(p), loader)
    # mtime を確実に進める（同一秒内の書き込みで stat が変わらない環境への対策）。
    os.utime(p, (os.stat(p).st_atime, os.stat(p).st_mtime + 5))
    obj2 = DR._cached_load(_open(p), loader)
    assert obj1 is not obj2
    assert len(calls) == 2


# ===== RV#8: キャッシュ追い出しが並走中の読取を壊さない（close せず参照を落とすだけ）===========

class _FakeLoaded:
    def __init__(self, f):
        self._f = f            # loader が f の所有権を引き継いだ体（openpyxl 等と同じ形）
        self.closed = False

    def close(self):
        self.closed = True


def test_cache_eviction_does_not_close_object_still_held_by_another_caller(tmp):
    DR._cache.clear()
    p_a = tmp / "a.xlsx"
    _make_xlsx(p_a)

    def loader(f):
        return _FakeLoaded(f)

    # スレッド A 相当: 参照を保持し続ける。
    held = DR._cached_load(_open(p_a), loader)
    assert isinstance(held, _FakeLoaded)

    # スレッド B 相当: 別ファイルを 9 件ロードして LRU（8件）から A を追い出す。
    for i in range(9):
        p_b = tmp / f"b{i}.xlsx"
        _make_xlsx(p_b)
        DR._cached_load(_open(p_b), loader)

    assert held.closed is False   # 追い出されても close されていない＝A の読取は引き続き成功する


def test_cache_eviction_does_not_break_concurrent_workbook_read(tmp):
    """実際の openpyxl Workbook で再現: 追い出し後も保持側の実読み取り（iter_rows）が成功する
    （以前の実装は追い出し時に `close()` していたため、共有中の Workbook の以後の読み取りが
    `ValueError: I/O operation on closed file` 等で壊れた）。"""
    DR._cache.clear()
    p_a = tmp / "a.xlsx"
    _make_xlsx(p_a, rows=3, cols=2)
    held = DR._cached_load(_open(p_a), DR._load_workbook)

    for i in range(9):
        p_b = tmp / f"b{i}.xlsx"
        _make_xlsx(p_b)
        DR._cached_load(_open(p_b), DR._load_workbook)

    ws = held["Sheet1"]
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    assert rows == [["r1c1", "r1c2"], ["r2c1", "r2c2"], ["r3c1", "r3c2"]]


def test_docx_paragraphs_tables_can_be_paged_with_table_start_and_row_start(tmp):
    """21 表目以降・51 行目以降も table_start／table_row_start で続きを取れる（読めない表を作らない）。"""
    p = tmp / "paged.docx"
    d = docx.Document()
    for ti in range(25):
        t = d.add_table(rows=60, cols=1)
        for ri in range(60):
            t.rows[ri].cells[0].text = f"T{ti}R{ri}"
    d.save(p)
    r = DR.docx_paragraphs(_open(p))
    assert r["total_tables"] == 25 and [t["i"] for t in r["tables"]] == list(range(20)) and r["truncated"]
    r2 = DR.docx_paragraphs(_open(p), table_start=20)
    assert [t["i"] for t in r2["tables"]] == [20, 21, 22, 23, 24]
    r3 = DR.docx_paragraphs(_open(p), table_start=3, table_row_start=50)
    assert r3["tables"][0]["i"] == 3 and r3["tables"][0]["rows"][0] == ["T3R50"] and r3["tables"][0]["total_rows"] == 60


def test_cache_skips_large_inputs_and_bounds_total_bytes(tmp, monkeypatch):
    """大きな入力はキャッシュしない・保持中の入力バイト合計に上限（展開後のメモリを積み上げない）。"""
    monkeypatch.setattr(DR, "_CACHE_ENTRY_MAX_BYTES", 10)
    DR._cache.clear()
    p = tmp / "x.docx"; docx.Document().save(p)
    DR.docx_paragraphs(_open(p))
    assert not DR._cache, "上限超えの入力がキャッシュされている"


def test_docx_paragraphs_stops_scanning_after_output_window(tmp):
    """巨大な表を持つ docx でも、出力する窓を過ぎたら走査を打ち切る（1 回の呼び出しが秒単位にならない）。"""
    import time
    p = tmp / "huge.docx"
    d = docx.Document()
    d.add_paragraph("先頭")
    t = d.add_table(rows=1200, cols=4)
    for ri in range(1200):
        t.rows[ri].cells[0].text = f"r{ri}"
    d.save(p)
    t0 = time.time()
    r = DR.docx_paragraphs(_open(p), start=0, count=5, clean=lambda x: x)
    assert r["tables"][0]["rows"][0][0] == "r0" and len(r["tables"][0]["rows"]) == 50
    assert time.time() - t0 < 8.0


def test_xlsx_range_tail_window_does_not_scan_whole_sheet(tmp):
    """末尾の範囲を読むときも時間予算内に返り、出力は要求範囲だけ。"""
    import time
    from openpyxl import Workbook
    p = tmp / "big.xlsx"
    wb = Workbook(); ws = wb.active
    for i in range(60000):
        ws.append([f"r{i}", "x"])
    wb.save(p)
    t0 = time.time()
    r = DR.xlsx_range(_open(p), "Sheet", "A59990:B60000", clean=lambda x: x)
    assert len(r["rows"]) == 11 and r["rows"][0][0] == "r59989"
    assert time.time() - t0 < 8.0


def test_xlsx_range_state_scan_includes_columns_left_of_window(tmp):
    """列窓より左のセル（A1）の BEGIN も状態に効く＝B1 だけを読んでも鍵本文が伏せられる。"""
    from openpyxl import Workbook
    from sherpa.agentic_search import _redact
    p = tmp / "cols.xlsx"
    wb = Workbook(); ws = wb.active
    ws.append(["-----BEGIN RSA PRIVATE KEY-----", "MIIEowIBAAKCAQEAsecretbody", "-----END RSA PRIVATE KEY-----"])
    wb.save(p)
    r = DR.xlsx_range(_open(p), "Sheet", "B1:B1", clean=_redact)
    assert "secretbody" not in r["rows"][0][0] and r["range"] == "B1:B1"


def test_xlsx_range_rejects_row_numbers_beyond_excel_limit(tmp):
    from openpyxl import Workbook
    p = tmp / "b.xlsx"; wb = Workbook(); wb.active["A1"] = "x"; wb.save(p)
    r = DR.xlsx_range(_open(p), "Sheet", "A1000000000000:A1000000000000", clean=lambda x: x)
    assert "error" in r


def test_docx_paragraphs_table_tail_row_state_reaches_following_paragraph(tmp):
    """表の最終行（出力行の窓の外）にある BEGIN が、その後に出力する段落の鍵本文に効く。"""
    from sherpa.agentic_search import _redact
    p = tmp / "tail.docx"
    d = docx.Document()
    d.add_paragraph("前置き")
    t = d.add_table(rows=51, cols=1)
    for ri in range(50):
        t.rows[ri].cells[0].text = f"r{ri}"
    t.rows[50].cells[0].text = "-----BEGIN RSA PRIVATE KEY-----"
    d.add_paragraph("MIIEowIBAAKCAQEAsecretbody\n-----END RSA PRIVATE KEY-----")
    d.save(p)
    r = DR.docx_paragraphs(_open(p), clean=_redact)
    assert all("secretbody" not in q["text"] for q in r["paragraphs"])


def test_scan_pages_with_lookback_is_per_interval():
    """`1,300` のような離れた指定は各区間の直前 50 ページだけ先読みする（全ページを辿らない）。"""
    got = DR._scan_pages_with_lookback([1, 300])
    assert got[0] == 1 and 300 in got and 150 not in got and len(got) == 1 + 51


def test_time_budget_returns_error_instead_of_hanging(tmp, monkeypatch):
    """時間予算を超えたら明示エラー（巨大なシートを延々と解析しない）。"""
    from openpyxl import Workbook
    p = tmp / "slow.xlsx"; wb = Workbook(); ws = wb.active
    for i in range(3000):
        ws.append([i, "x"])
    wb.save(p)
    monkeypatch.setenv("SHERPA_DOC_READ_MAX_SECONDS", "0.0001")
    r = DR.xlsx_range(_open(p), "Sheet", "A2900:B3000", clean=lambda x: x)
    assert r == DR._TIME_ERROR


def test_xlsx_range_state_scan_ignores_wrong_dimension_record(tmp):
    """<dimension> が実データより狭い／無いブックでも、列窓の右の BEGIN が状態に効き、既定範囲も実寸になる。"""
    import zipfile, re
    from openpyxl import Workbook
    from sherpa.agentic_search import _redact
    p = tmp / "dim.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "S"
    ws["AD3"] = "-----BEGIN RSA PRIVATE KEY-----"; ws["B5"] = "MIIEpAIBAAKCAQEAsecretbody"; ws["AD9"] = "-----END RSA PRIVATE KEY-----"
    ws["A20"] = "tail"
    wb.save(p)
    q = tmp / "nodim.xlsx"
    with zipfile.ZipFile(p) as zin, zipfile.ZipFile(q, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb"<dimension[^>]*/>", b"", data)
            zout.writestr(item.filename, data)
    r = DR.xlsx_range(_open(q), "S", "A1:B20", clean=_redact)
    assert all("secretbody" not in c for row in r["rows"] for c in row)
    assert DR.xlsx_sheets(_open(q))["sheets"][0]["max_row"] >= 20
    r2 = DR.xlsx_range(_open(q), "S", None, clean=_redact)
    assert r2["range"].endswith("20") and len(r2["rows"]) == 20


def test_xlsx_range_scans_beyond_recorded_dimension_for_key_state(tmp):
    """<dimension> が A1:B20 と過小でも AD 列の BEGIN が状態に効く（reset_dimensions 後に全列を辿る）。"""
    import zipfile, re
    from openpyxl import Workbook
    from sherpa.agentic_search import _redact
    p = tmp / "small_dim.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "S"
    ws["AD3"] = "-----BEGIN RSA PRIVATE KEY-----"; ws["B5"] = "MIIEpAIBAAKCAQEAsecretbody"; ws["AD9"] = "-----END RSA PRIVATE KEY-----"
    wb.save(p)
    q = tmp / "small_dim2.xlsx"
    with zipfile.ZipFile(p) as zin, zipfile.ZipFile(q, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]*"/>', b'<dimension ref="A1:B20"/>', data)
            zout.writestr(item.filename, data)
    r = DR.xlsx_range(_open(q), "S", "A1:B20", clean=_redact)
    assert all("secretbody" not in c for row in r["rows"] for c in row)


def test_xlsx_sheets_dimension_calculation_respects_time_budget(tmp, monkeypatch):
    import zipfile, re
    from openpyxl import Workbook
    p = tmp / "d.xlsx"; wb = Workbook(); ws = wb.active
    for i in range(2000):
        ws.append([i])
    wb.save(p)
    q = tmp / "d2.xlsx"
    with zipfile.ZipFile(p) as zin, zipfile.ZipFile(q, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb"<dimension[^>]*/>", b"", data)
            zout.writestr(item.filename, data)
    monkeypatch.setenv("SHERPA_DOC_READ_MAX_SECONDS", "0.0001")
    r = DR.xlsx_sheets(_open(q))
    assert r["sheets"][0]["name"] == "Sheet" and r["sheets"][0]["dims_estimated"]
    assert "max_row" not in r["sheets"][0]            # 記録が無い＝不明（0 行と偽らない）


def test_xlsx_sheets_recorded_dims_survive_cached_workbook_reset(tmp, monkeypatch):
    """記録寸法は reset_dimensions の前に控える＝キャッシュ済み Workbook の 2 回目でも記録値で縮退できる。"""
    from openpyxl import Workbook
    p = tmp / "rec.xlsx"; wb = Workbook(); ws = wb.active
    for _ in range(29):
        ws.append(["a", "b", "c"])
    wb.save(p)
    assert DR.xlsx_sheets(_open(p))["sheets"][0]["max_row"] == 29     # 1 回目（走査成功・reset 済み）
    monkeypatch.setenv("SHERPA_DOC_READ_MAX_SECONDS", "0.00001")
    r = DR.xlsx_sheets(_open(p))                                        # 2 回目（キャッシュ命中・予算超過）
    assert r["sheets"][0]["dims_estimated"] and r["sheets"][0]["max_row"] == 29


def test_pptx_slides_non_contiguous_pages_do_not_carry_key_state_across_gap(tmp):
    """1 枚目 BEGIN・2 枚目 END の正常終端した鍵の後、離れた 60 枚目の通常本文は伏せない。"""
    from pptx import Presentation
    from pptx.util import Inches
    from sherpa.agentic_search import _redact
    prs = Presentation()
    texts = ["-----BEGIN RSA PRIVATE KEY-----", "MIIEkey\n-----END RSA PRIVATE KEY-----"] + ["normal"] * 58
    for t in texts:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)); tb.text_frame.text = t
    p = tmp / "gap.pptx"; prs.save(p)
    r = DR.pptx_slides(_open(p), pages="1,60", clean=_redact)
    by_no = {sl["no"]: sl for sl in r["slides"]}
    assert "normal" in " ".join(by_no[60]["texts"])


def test_xlsx_sheets_and_default_range_ignore_undersized_dimension_record(tmp):
    """<dimension> が過小（A1:B20）でも、シート一覧と既定範囲は実データの寸法になる。"""
    import zipfile, re
    from openpyxl import Workbook
    p = tmp / "u.xlsx"; wb = Workbook(); ws = wb.active; ws.title = "S"
    ws["AD3"] = "x"; ws["A20"] = "y"; wb.save(p)
    q = tmp / "u2.xlsx"
    with zipfile.ZipFile(p) as zin, zipfile.ZipFile(q, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]*"/>', b'<dimension ref="A1:B20"/>', data)
            zout.writestr(item.filename, data)
    assert DR.xlsx_sheets(_open(q))["sheets"][0]["max_col"] == 30
    assert DR.xlsx_range(_open(q), "S", None, clean=lambda x: x)["range"] == "A1:AD20"


def test_xlsx_sheets_time_budget_is_shared_across_sheets_and_checked_every_row(tmp, monkeypatch):
    """行数が少なく幅の広いシートが複数あっても、呼び出し全体で 1 つの予算・毎行判定で止まる。"""
    from openpyxl import Workbook
    p = tmp / "wide.xlsx"; wb = Workbook()
    for k in range(3):
        ws = wb.active if k == 0 else wb.create_sheet(f"S{k}")
        for _ in range(5):
            ws.append(["x"] * 300)
    wb.save(p)
    monkeypatch.setenv("SHERPA_DOC_READ_MAX_SECONDS", "0.00001")
    r = DR.xlsx_sheets(_open(p))
    assert [sh["name"] for sh in r["sheets"]] == ["Sheet", "S1", "S2"]      # シート名は失わない
    assert all(sh.get("dims_estimated") for sh in r["sheets"])            # 寸法は記録値（推定）


def test_xlsx_sheets_recorded_dims_survive_prior_explicit_range_read(tmp, monkeypatch):
    """明示範囲の xlsx_range（別経路の reset）を先に行っても、記録寸法の控えは失われない。"""
    from openpyxl import Workbook
    p = tmp / "rec2.xlsx"; wb = Workbook(); ws = wb.active
    for _ in range(29):
        ws.append(["a", "b", "c"])
    wb.save(p)
    assert len(DR.xlsx_range(_open(p), "Sheet", "A1:C1", clean=lambda x: x)["rows"]) == 1
    monkeypatch.setenv("SHERPA_DOC_READ_MAX_SECONDS", "0.00001")
    r = DR.xlsx_sheets(_open(p))
    assert r["sheets"][0]["dims_estimated"] and r["sheets"][0]["max_row"] == 29
