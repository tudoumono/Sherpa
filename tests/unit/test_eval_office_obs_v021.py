"""Office 観測の評価（eval/office-obs-v021）の問いと正解を、フィクスチャに照らして検証する。

正解の根拠（ファイル・位置・原本の状態）が実在すること、場面ごとの件数、フィクスチャが生成スクリプトと一致すること、
runner の照合・採点・記録の純関数、dry 実行が外部に接続しないことを固定する。ES・埋め込み・回答の生成は呼ばない。
説明: eval/office-obs-v021/README.md。
"""
from __future__ import annotations

import copy
import hashlib
import importlib.util
import json
import sys
import zipfile
from pathlib import Path

import docx
import pytest
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
EVAL_DIR = ROOT / "eval" / "office-obs-v021"
FIXTURES = ROOT / "fixtures" / "eval_v021"


def _load(name: str):
    spec = importlib.util.spec_from_file_location(f"office_obs_v021_{name}", EVAL_DIR / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = mod
    spec.loader.exec_module(mod)
    return mod


runner = _load("runner")
verify = _load("verify")
build_fixtures = _load("build_fixtures")


@pytest.fixture(scope="module")
def questions():
    return runner.load_questions(EVAL_DIR)


def _by_id(questions, qid):
    return copy.deepcopy(next(q for q in questions if q["id"] == qid))


# ---------------------------------------------------------------- 問いと正解の検証

def test_every_question_and_its_original_is_verified(questions):
    problems = verify.check_all(ROOT, questions, runner) + verify.check_coverage(questions)
    assert not problems, "\n".join(problems)


def test_nine_scenes_with_the_expected_number_of_questions(questions):
    counts = {s: sum(q["scene"] == s for q in questions) for s in "ABCDEFGHI"}
    assert counts == {"A": 3, "B": 3, "C": 3, "D": 3, "E": 3, "F": 2, "G": 2, "H": 1, "I": 2}, counts
    assert len(questions) == 22
    assert {q["format"] for q in questions if q["scene"] == "A"} >= {"docx", "pptx"}
    for scene in "CDI":
        assert {q["direction"] for q in questions if q["scene"] == scene} == {"front", "back"}
    assert {q["id"] for q in questions if q["scene"] in "FGHI"} == {"OO-F1", "OO-F2", "OO-G1", "OO-G2", "OO-H1", "OO-I1", "OO-I2"}


def test_scene_bounds_are_enforced(questions):
    fewer = [q for q in questions if q["id"] != "OO-H1"]
    assert any("場面 H" in p for p in verify.check_coverage(fewer))
    extra = questions + [dict(_by_id(questions, "OO-G2"), id="OO-G3")]
    assert any("場面 G" in p for p in verify.check_coverage(extra))
    no_control = [q for q in questions if q["id"] != "OO-F2"]
    assert any("場面 F" in p for p in verify.check_coverage(no_control))


def _broken(qid, mutate):
    def run(questions):
        q = _by_id(questions, qid)
        mutate(q)
        return verify.check_question(ROOT, q, runner)
    return run


def _run_strike_reversed(q):
    q["gold"]["evidence"][0]["original"][0]["strike"] = False


def _missing_paragraph(q):
    q["gold"]["evidence"][0]["at"] = {"type": "paragraph", "index": 99}


def _word_not_in_place(q):
    q["gold"]["evidence"][0]["contains"] = ["この語は原本のその箇所に無い"]


def _unknown_claim_evidence(q):
    q["gold"]["claims"][0]["evidence"] = ["e99"]


def _bad_regex(q):
    q["observations"][0]["all"] = ["(未閉じ"]


def _observation_outside_evidence(q):
    q["observations"][0]["doc"] = "備品貸出/02_設計/02_運用/運用手引き.docx"


def _gold_answer_unscorable(q):
    q["gold"]["answer"] = "分かりません。"


def _unknown_probe_tool(q):
    q["probes"][0]["tool"] = "web_search"


@pytest.mark.parametrize("qid, mutate, expected", [
    ("OO-A1", _run_strike_reversed, "取り消し線が正解"),
    ("OO-A1", _missing_paragraph, "段落が範囲外"),
    ("OO-A1", _word_not_in_place, "指定の箇所に"),
    ("OO-A1", _unknown_claim_evidence, "存在しない根拠"),
    ("OO-A1", _bad_regex, "正規表現が不正"),
    ("OO-A1", _observation_outside_evidence, "正解の根拠のどれにも無い"),
    ("OO-A1", _gold_answer_unscorable, "模範の回答が自分の採点を通らない"),
    ("OO-A1", _unknown_probe_tool, "tool が不正"),
])
def test_a_broken_question_is_reported(questions, qid, mutate, expected):
    assert any(expected in p for p in _broken(qid, mutate)(questions))


def test_a_wrong_original_state_is_reported(questions):
    q = _by_id(questions, "OO-D1")
    q["gold"]["evidence"][1]["original"][0]["percent"] = 80
    assert any("不透明度" in p for p in verify.check_question(ROOT, q, runner))

    q = _by_id(questions, "OO-C1")
    q["gold"]["evidence"][0]["original"][1].update(front="旧承認ボタンの説明", back="廃止スタンプ")
    assert any("前面にない" in p for p in verify.check_question(ROOT, q, runner))

    q = _by_id(questions, "OO-E1")
    q["gold"]["evidence"][0]["original"][0]["embedded_rows"] = 38
    assert any("行数" in p for p in verify.check_question(ROOT, q, runner))

    q = _by_id(questions, "OO-E1")
    q["gold"]["evidence"][0]["original"][1]["pattern"] = "台帳"
    assert any("本文に出ている" in p for p in verify.check_question(ROOT, q, runner))

    q = _by_id(questions, "OO-C3")
    q["gold"]["evidence"][0]["original"][0]["cell"] = "B2"
    assert any("アンカー" in p for p in verify.check_question(ROOT, q, runner))


def test_the_new_original_states_are_checked(questions):
    q = _by_id(questions, "OO-F1")
    q["gold"]["evidence"][1]["original"][0]["layout_prst"] = "ellipse"
    assert any("レイアウト側の形" in p for p in verify.check_question(ROOT, q, runner))

    q = _by_id(questions, "OO-F1")
    q["gold"]["evidence"][1]["original"][2].update(front="穴の中の文字", back="輪の枠")
    assert any("前面にない" in p for p in verify.check_question(ROOT, q, runner))

    q = _by_id(questions, "OO-F1")
    q["gold"]["evidence"][1]["path"] = q["gold"]["evidence"][1]["path"].replace("避難経路の案内", "承認画面の構成")
    assert verify.check_question(ROOT, q, runner)               # 別の資料の「輪の枠」は無い

    q = _by_id(questions, "OO-G1")
    q["gold"]["evidence"][0]["original"][0].update(struck="条件乙", plain="条件甲")
    assert any("取り消し線が正解" in p for p in verify.check_question(ROOT, q, runner))

    q = _by_id(questions, "OO-G2")
    q["gold"]["evidence"][0]["original"][0]["group"] = "groupY"
    assert any("グループ" in p for p in verify.check_question(ROOT, q, runner))

    q = _by_id(questions, "OO-H1")
    q["gold"]["evidence"][0]["original"][0]["paragraphs"] = ["1020"]
    assert any("段落が" in p for p in verify.check_question(ROOT, q, runner))

    q = _by_id(questions, "OO-I1")
    q["gold"]["evidence"][0]["original"][0]["count"] = 3
    assert any("グラデーションの図形が 4 個" in p for p in verify.check_question(ROOT, q, runner))
    q = _by_id(questions, "OO-I1")
    q["gold"]["evidence"][0]["original"][0]["name_pattern"] = "説明図形[1-4]補足"   # 名前に改行を含む
    assert any("名前が" in p for p in verify.check_question(ROOT, q, runner))


def test_transparency_is_read_from_the_png():
    opaque = build_fixtures._png_rgba(8, 8, lambda x, y: (1, 2, 3, 255))
    holed = build_fixtures._png_rgba(8, 8, lambda x, y: (1, 2, 3, 0 if x == 0 else 255))
    assert verify.png_has_transparent_pixels(holed) and not verify.png_has_transparent_pixels(opaque)


# ---------------------------------------------------------------- フィクスチャ

def _content(path: Path):
    """ファイルの意味の中身（書式の直列化に依らない）。docx は段落・run の取り消し線・図形の文字と埋め込み、xlsx はセルと図形、
    pptx は図形の名前・位置・文字・塗り・画像・OLE とレイアウト側の枠。
    """
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main", "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
    if path.suffix == ".docx":
        d = docx.Document(path)
        body = [(p.style.name, [(r.text, bool(r.font.strike)) for r in p.runs]) for p in d.paragraphs]
        boxes = [etree.tostring(b, method="c14n") for b in d.element.body.iter(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}txbxContent")]
        with zipfile.ZipFile(path) as z:
            embedded = sorted((n, hashlib.sha256(z.read(n)).hexdigest()) for n in z.namelist() if "embeddings/" in n)
        return body, boxes, embedded
    if path.suffix == ".xlsx":
        import openpyxl
        cells = [(ws.title, c.coordinate, c.value) for ws in openpyxl.load_workbook(path) for row in ws.iter_rows() for c in row
                 if c.value is not None]
        with zipfile.ZipFile(path) as z:
            drawings = sorted((n, etree.tostring(etree.fromstring(z.read(n)), method="c14n")) for n in z.namelist() if "drawings/drawing" in n)
        return cells, drawings
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    out = []
    prs = Presentation(path)
    for slide in prs.slides:
        for s in slide.shapes:
            runs = []
            if s.has_text_frame:
                for p in s.text_frame.paragraphs:
                    for r in p.runs:
                        rpr = r._r.find("a:rPr", ns)
                        runs.append((r.text, None if rpr is None else rpr.get("strike")))
            fills = [etree.tostring(f) for f in s._element.xpath("./p:spPr/a:solidFill")]
            geometry = [g.get("prst") for g in s._element.xpath("./p:spPr/a:prstGeom")]
            image = hashlib.sha256(s.image.blob).hexdigest() if s.shape_type == MSO_SHAPE_TYPE.PICTURE else None
            ole = None
            if s.shape_type == MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT:
                ole = (s.ole_format.prog_id, hashlib.sha256(s.ole_format.blob).hexdigest())
            out.append((s.name, s.left, s.top, s.width, s.height, runs, fills, geometry, image, ole, s.is_placeholder))
    layout_frames = [(p.name, p.placeholder_format.idx, [g.get("prst") for g in p._element.xpath("./p:spPr/a:prstGeom")])
                     for layout in prs.slide_layouts for p in layout.placeholders if p.placeholder_format.idx == 13]
    return out, layout_frames


def _contents(root: Path) -> dict[str, object]:
    return {p.relative_to(root).as_posix(): _content(p) for p in sorted(root.rglob("*")) if p.is_file()}


def test_fixtures_match_the_generator(tmp_path):
    build_fixtures.build_all(tmp_path / "fixtures")
    assert _contents(tmp_path / "fixtures") == _contents(FIXTURES)


def test_the_generator_is_deterministic(tmp_path):
    build_fixtures.build_all(tmp_path / "a")
    build_fixtures.build_all(tmp_path / "b")
    def digest(root: Path) -> dict[str, str]:
        return {p.relative_to(root).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
                for p in sorted(root.rglob("*")) if p.is_file()}

    assert digest(tmp_path / "a") == digest(tmp_path / "b")


# ---------------------------------------------------------------- 観測の照合

STRUCK = "【取り消し線：保証金は5万円とする。】保証金は3万円とする。"


def test_an_observation_needs_the_struck_sentence_and_the_plain_one(questions):
    q = _by_id(questions, "OO-A1")
    struck, plain = q["observations"]
    assert runner.match_text(STRUCK, struck) and runner.match_text(STRUCK, plain)
    unmarked = "保証金は5万円とする。保証金は3万円とする。"
    assert not runner.match_text(unmarked, struck)                      # 状態が落ちた本文では、取り消された文が分からない
    assert runner.match_text("【取り消し線（続き）：保証金は5万円とする。】", struck)   # 分割の後半の片の括り
    assert not runner.match_text("【取り消し線：保証金は3万円とする。】", plain)      # 3万円まで括りに入っていたら、現行の文が消える


def test_highlight_marks_do_not_hide_an_observation(questions):
    obs = _by_id(questions, "OO-A1")["observations"][0]
    assert runner.match_text("【取り消し線：<em>保証金</em>は5万円とする。】", obs)


def test_a_cover_judgement_on_the_back_element_breaks_the_not_occluding_observation(questions):
    obs = _by_id(questions, "OO-D1")["observations"][0]
    clean = "出所: 原本「受付案内図.pptx」 / スライド1\n図形: 「受付窓口は3階の総務課です」"
    assert runner.match_text(clean, obs)
    assert not runner.match_text(clean + "\n可視性: 「図形に覆われている」", obs)


def test_the_front_observation_must_name_the_target_in_the_front_chunk(questions):
    obs = _by_id(questions, "OO-C1")["observations"][0]
    front_without_target = "出所: 原本「承認画面の構成.pptx」 / スライド1\n図形: 「廃止」\nこの要素は別要素の原本領域に重なっている。"
    back = "出所: 原本「承認画面の構成.pptx」 / スライド1\n図形: 「旧承認ボタン: 申請者本人が承認できる」\n重なり: 「廃止」"
    front_with_target = front_without_target + "\nこの要素は対象「旧承認ボタン」の原本領域に重なっている"
    assert not runner.match_text(front_without_target, obs)
    assert not runner.match_text(back, obs)                             # 背面のチャンクの『重なり: 「廃止」』は前面の観測にならない
    assert runner.match_text(front_with_target, obs)


def test_the_ole_notice_is_recognised(questions):
    obs = _by_id(questions, "OO-E1")["observations"][0]
    notice = "未対応: kind=ole_embedding, reason=unsupported_ole\n内容を抽出済みとは扱わない。"
    assert runner.match_text(notice, obs)
    assert not runner.match_text("原本のスライド1に画像が存在する。", obs)


def test_rag_md_is_split_by_anchors_even_with_line_numbers():
    md = "# 見出し\n\n<!-- chunk:a -->\n出所: x\n本文: 「甲」\n\n<!-- chunk:b -->\n出所: x\n本文: 「乙」\n"
    numbered = "\n".join(f"{i}: {line}" for i, line in enumerate(md.splitlines(), start=1))
    assert runner.split_rag_md_chunks(md) == runner.split_rag_md_chunks(numbered) == ["出所: x\n本文: 「甲」", "出所: x\n本文: 「乙」"]


def test_search_hits_are_split_into_chunks_by_the_source_line():
    hit = "## 節\n\n出所: 甲\n本文: 「1」\n\n出所: 乙\n本文: 「2」<em>語</em>"
    assert runner.hit_blocks(hit) == ["## 節\n\n", "出所: 甲\n本文: 「1」\n\n", "出所: 乙\n本文: 「2」語"]


def test_hits_without_a_source_line_are_split_at_the_chunk_boundary():
    hit = "出所: 甲\n本文: 「1」\n\nこの要素は別要素の原本領域に重なっている。\n重なっている対象: セルB3\n\nこの要素は別要素の原本領域に重なっている。\n重なっている対象: セルB4"
    assert runner.hit_blocks(hit) == ["出所: 甲\n本文: 「1」\n\n", "この要素は別要素の原本領域に重なっている。\n重なっている対象: セルB3\n\n",
                                      "この要素は別要素の原本領域に重なっている。\n重なっている対象: セルB4"]
    inside = "出所: 甲\n図形: 「廃止」\nこの要素は別要素の原本領域に重なっている。\n覆っている対象: 乙"
    assert runner.hit_blocks(inside) == [inside]                       # チャンクの途中の『この要素は』では切らない


def test_hits_that_carry_the_rag_md_are_split_at_the_anchors():
    hit = "前のチャンクの末尾\n<!-- chunk:a -->\n出所: 甲\n本文: 「1」\n\n<!-- chunk:b -->\n出所: 乙\n本文: 「2」"
    assert runner.hit_blocks(hit) == ["前のチャンクの末尾", "出所: 甲\n本文: 「1」", "出所: 乙\n本文: 「2」"]


# ---------------------------------------------------------------- 段 1・2・3

def _chunk(text, *, embedded=True, vector_text=None, head=None):
    return {"chunk_id": text[:6], "text": text, "vector_text": vector_text, "row_group_head_chunk_id": head, "embedded": embedded}


OBS = {"id": "o", "doc": "d", "all": ["甲.*乙"], "required": True}


def test_stage1_separates_the_es_text_from_the_bundled_vector_input():
    # 束の先頭は vector_text（束全体）で埋め込まれ、2 件目以降は自分のベクトルを持たない
    head = _chunk("甲", vector_text="甲\n乙")
    member = _chunk("乙", embedded=False, head=head["chunk_id"])
    res = runner.stage1_observation(OBS, [head, member])
    assert res == {"doc_indexed": True, "es_text": False, "vector_input": True}      # 本文は別々のチャンクだが、束には両方ある
    joined = runner.stage1_observation(OBS, [_chunk("甲と乙", vector_text="甲と乙")])
    assert joined["es_text"] and joined["vector_input"]
    cut = runner.stage1_observation(OBS, [_chunk("甲と乙", vector_text="甲だけ")])
    assert cut["es_text"] and cut["vector_input"] is False and not runner.stage1_ok(cut)


def test_stage1_has_no_vector_verdict_when_nothing_was_embedded():
    res = runner.stage1_observation(OBS, [_chunk("甲と乙", embedded=False)])
    assert res["es_text"] and res["vector_input"] is None and runner.stage1_ok(res)
    assert runner.stage1_observation(OBS, []) == {"doc_indexed": False, "es_text": False, "vector_input": None}


def test_stage2_reports_which_call_surfaced_the_observation():
    probes = [
        {"index": 1, "tool": "es_search", "units": [{"doc_id": "d", "text": "乙のみ", "rank": 1}, {"doc_id": "x", "text": "甲と乙", "rank": 2}]},
        {"index": 2, "tool": "ripgrep_search", "units": [{"doc_id": "d", "text": "甲と乙", "rank": 1}]},
    ]
    assert runner.surfaced(OBS, probes) == [{"probe": 2, "tool": "ripgrep_search", "rank": 1}]    # 別の文書のヒットは数えない


def test_dropped_stage_is_the_first_stage_that_fails():
    assert runner.drop_stage(False, False, False) == "index"
    assert runner.drop_stage(True, False, None) == "retrieval"
    assert runner.drop_stage(True, True, False) == "answer"
    assert runner.drop_stage(True, True, True) == "none"
    assert runner.drop_stage(True, True, None) == "answer_not_run"


# ---------------------------------------------------------------- 段 3 の採点

def test_the_model_answer_passes_and_a_wrong_assertion_does_not(questions):
    q = _by_id(questions, "OO-A1")
    assert runner.score_answer(q, q["gold"]["answer"], ["貸出条件書.docx"])["passed"]
    wrong = runner.score_answer(q, "保証金は5万円です。出典: 貸出条件書.docx", ["貸出条件書.docx"])
    assert not wrong["passed"] and wrong["wrong"] and not all(c["ok"] for c in wrong["claims"])
    no_source = runner.score_answer(q, "保証金は3万円。5万円の文には取り消し線が付いている。")
    assert no_source["cited"] is False and not no_source["passed"]
    assert runner.score_answer(q, "保証金は3万円。5万円の文には取り消し線が付いている。", ["貸出条件書.docx"])["passed"]


def test_a_quoted_sentence_stays_with_the_explanation_that_negates_it():
    sentences = runner._sentences("同じ段落の「保証金は5万円とする。」には取り消し線が付いている。現行は3万円。")
    assert sentences == ["同じ段落の「保証金は5万円とする。」には取り消し線が付いている。", "現行は3万円。"]


def test_an_unparsed_ole_answer_must_not_state_the_count(questions):
    q = _by_id(questions, "OO-E1")
    ok = runner.score_answer(q, "分からない。添付の台帳ファイルの中身は解析していないため、件数は読み取れない。", ["貸出台帳の添付.pptx"])
    assert ok["passed"], ok
    guess = runner.score_answer(q, "貸出は37件載っています。出典: 貸出台帳の添付.pptx", ["貸出台帳の添付.pptx"])
    assert not guess["passed"] and {w["id"] for w in guess["wrong"]} >= {"w1"}


def test_translucent_overlap_answers_are_told_apart(questions):
    q = _by_id(questions, "OO-D1")
    ok = runner.score_answer(q, "3階の総務課です。半透明の枠が重なっているだけで、文字は隠れておらず読めます。出典: 受付案内図.pptx")
    assert ok["passed"], ok
    assert runner.score_answer(q, "3階の総務課です。枠は重なっていますが、文字は隠れていません。出典: 受付案内図.pptx")["passed"]
    hidden = runner.score_answer(q, "3階の総務課ですが、文字は枠に覆われていて読めません。出典: 受付案内図.pptx")
    assert hidden["wrong"] and not hidden["passed"]

    front = _by_id(questions, "OO-D2")
    assert runner.score_answer(front, "隠していません。強調枠は半透明で重なっているだけです。出典: 受付案内図.pptx")["passed"]
    assert runner.score_answer(front, "はい、強調枠が文字を隠しています。出典: 受付案内図.pptx")["wrong"]

    image = _by_id(questions, "OO-D3")
    assert runner.score_answer(image, "隠れていません。社章は透明な部分を含む画像で、重なっているだけです。出典: 受付案内図.pptx")["passed"]
    assert runner.score_answer(image, "社章の画像に隠れていて読めません。出典: 受付案内図.pptx")["wrong"]


def test_a_stamp_answer_keeps_the_business_deprecation_open(questions):
    q = _by_id(questions, "OO-C2")
    assert runner.score_answer(q, q["gold"]["answer"], ["承認画面の構成.pptx"])["passed"]
    flat = runner.score_answer(q, "はい、現在も有効です。出典: 承認画面の構成.pptx")
    assert flat["wrong"] and not flat["passed"]


def test_scoring_a_file_merges_with_the_earlier_stages(questions):
    q = _by_id(questions, "OO-A3")
    results = {"OO-A3": {"id": "OO-A3", "scene": "A", "stage1": {"ok": False}, "stage2": {"ok": False}}}
    answers = {"OO-A3": {"id": "OO-A3", "answer": q["gold"]["answer"], "sources": ["貸出条件の改定案.pptx"]}}
    merged = runner.score_file([q], answers, results)
    assert merged[0]["stage3"]["passed"] and merged[0]["dropped_at"] == "index"             # 前の段で落ちていれば、その段を記録する
    alone = runner.score_file([q], answers, None)
    assert alone[0]["dropped_at"] == "none"


def test_an_answer_template_lists_every_question(questions):
    lines = [json.loads(x) for x in runner.answer_template(questions).splitlines()]
    assert [x["id"] for x in lines] == [q["id"] for q in questions] and all(x["answer"] == "" for x in lines)


# ---------------------------------------------------------------- dry 実行

def test_dry_run_shows_how_each_stage_is_checked_without_connecting(capsys, monkeypatch):
    def boom(*a, **k):
        raise AssertionError("dry 実行は ES・埋め込みに接続しない")
    monkeypatch.setattr(runner, "run_all", boom)
    monkeypatch.setattr(runner.LiveWorld, "__enter__", boom)
    assert runner.main([]) == 0
    out = capsys.readouterr().out
    for q in runner.load_questions(EVAL_DIR):
        assert q["id"] in out
    assert out.count("段 1 索引") == out.count("段 2 返却する根拠") == out.count("段 3 回答") == 22
    assert "es_search(" in out and "ripgrep_search(" in out and "read_doc(" in out and "score_answer(" in out


def test_dry_run_can_be_narrowed(capsys):
    assert runner.main(["--scene", "E"]) == 0
    out = capsys.readouterr().out
    assert "OO-E1" in out and "OO-E3" in out and "OO-A1" not in out
    assert runner.main(["--id", "OO-A1"]) == 0
    assert "OO-A2" not in capsys.readouterr().out
    assert runner.main(["--id", "OO-A1", "--scene", "E"]) == 2           # 絞り込みは積。当たる問いが無ければ失敗


# ---------------------------------------------------------------- 実際の取り込みの出力に対する観測（Word）

WORD_QUESTIONS = ("OO-A1", "OO-A2", "OO-B1", "OO-B2", "OO-B3")


def test_word_observations_hold_on_the_real_rag_md(tmp_path, questions):
    """Word の取り消し線は取り込みが括りで書く。実際の rag.md のチャンクに、Word の場面の観測が出ていること。"""
    from sherpa.ingest import office_md
    world = tmp_path / "world"
    for doc in ("貸出条件書.docx", "運用手引き.docx"):
        src = next(FIXTURES.rglob(doc))
        dst = world / src.relative_to(FIXTURES)
        dst.parent.mkdir(parents=True, exist_ok=True)
        dst.write_bytes(src.read_bytes())
    derived = tmp_path / "derived" / "md"
    report = office_md.build_derived(world, derived)
    assert report["rag_failed"] == 0 and report["evidence_ir_failed"] == 0
    checked = 0
    for qid in WORD_QUESTIONS:
        q = _by_id(questions, qid)
        for obs in q["observations"]:
            rag = tmp_path / "derived" / "rag" / (obs["doc"] + ".rag.md")
            chunks = runner.split_rag_md_chunks(rag.read_text(encoding="utf-8"))
            assert any(runner.match_text(c, obs) for c in chunks), (qid, obs["id"])
            checked += 1
    assert checked == sum(len(_by_id(questions, qid)["observations"]) for qid in WORD_QUESTIONS)


REPRODUCTION_QUESTIONS = ("OO-F1", "OO-F2", "OO-G1", "OO-G2", "OO-H1", "OO-I1", "OO-I2")


def _index_chunks(rag_dir: Path, doc: str) -> list[dict]:
    """取り込んだ rag.md・rag_chunks から、ES に入るチャンク（本文・ベクトルの入力・ベクトルの有無）を実際の束ね方で作る。"""
    from sherpa import es_index
    jsonl = rag_dir / f"{doc}.rag_chunks.jsonl"
    anchors, reason = es_index._load_rag_md_anchors(rag_dir / f"{doc}.rag.md")
    assert reason is None, reason
    out = []
    for _es_id, body, _text, skip_embed in es_index._iter_rag_chunk_entries(jsonl, doc, anchors, {"doc_id": doc}):
        out.append({"chunk_id": body["chunk_id"], "text": body["text"], "vector_text": body.get("vector_text"),
                    "row_group_head_chunk_id": body.get("row_group_head_chunk_id"), "embedded": not skip_embed})
    return out


def test_the_reproduction_observations_hold_on_the_real_index_input(tmp_path, questions):
    """F〜I の観測が、実際の取り込みの出力（rag.md の本文と、束ねたベクトルの入力）に出る。ES・埋め込みには接続しない。"""
    from sherpa.ingest import office_md
    world = tmp_path / "world"
    for doc in ("避難経路の案内.pptx", "条件の比較図.docx", "基準値の枠.docx", "品目別の値.xlsx"):
        src = next(FIXTURES.rglob(doc))
        dst = world / src.relative_to(FIXTURES)
        dst.parent.mkdir(parents=True, exist_ok=True)
        dst.write_bytes(src.read_bytes())
    report = office_md.build_derived(world, tmp_path / "derived" / "md")
    assert report["rag_failed"] == 0 and report["evidence_ir_failed"] == 0
    checked = 0
    for qid in REPRODUCTION_QUESTIONS:
        for obs in _by_id(questions, qid)["observations"]:
            result = runner.stage1_observation(obs, _index_chunks(tmp_path / "derived" / "rag", obs["doc"]))
            assert result == {"doc_indexed": True, "es_text": True, "vector_input": True}, (qid, obs["id"], result)
            checked += 1
    assert checked == 8


def test_a_front_shape_without_text_is_named_in_its_own_chunk_and_vector_input(tmp_path):
    """文字の無い前面の図形 4 つは、名前・図形 ID・出所を自分のチャンクに持ち、束ねたベクトルの入力でも同じ文が並ばない。"""
    from sherpa.ingest import office_md
    src = next(FIXTURES.rglob("品目別の値.xlsx"))
    dst = tmp_path / "world" / src.relative_to(FIXTURES)
    dst.parent.mkdir(parents=True)
    dst.write_bytes(src.read_bytes())
    office_md.build_derived(tmp_path / "world", tmp_path / "derived" / "md")
    chunks = _index_chunks(tmp_path / "derived" / "rag", "備品貸出/02_設計/05_検証/品目別の値.xlsx")
    for n in range(1, 5):
        own = [c for c in chunks if f"この要素: 説明図形{n} 補足（図形 ID {n + 1}・シート「一覧」のセル範囲 B3:B3）" in c["text"]]
        assert len(own) == 1, n
        assert "出所: 原本「品目別の値.xlsx」 / シート「一覧」" in own[0]["text"]
    vectors = [c["vector_text"] or c["text"] for c in chunks if c["embedded"] and "この要素は別要素" in c["text"]]
    joined = "\n".join(vectors)
    assert joined.count("重なっている対象: セルB3（シート「一覧」）（覆っているかは未確定）") == 4
    assert all(f"この要素: 説明図形{n} 補足" in joined for n in range(1, 5))


# ---------------------------------------------------------------- 外部レビューの再現（F〜I）

def _obs(questions, qid, oid):
    return next(o for o in _by_id(questions, qid)["observations"] if o["id"] == oid)


def test_the_inherited_donut_stays_overlapping_and_undetermined(questions):
    back = _obs(questions, "OO-F1", "o1")
    text = "出所: 原本「避難経路の案内.pptx」 / スライド1\n図形: 「非常口は北側」\n前面図形: 輪の枠（図形 ID 2・スライド1）（覆われているかは未確定）"
    assert runner.match_text(text, back)
    assert not runner.match_text(text.replace("覆われているかは未確定", "覆われていると判定"), back)
    assert not runner.match_text(text + "\n可視性: 「図形に覆われている」", back)
    assert not runner.match_text(text.replace("前面図形: 輪の枠（図形 ID 2・スライド1）（覆われているかは未確定）", ""), back)   # 重なりが落ちた
    front = _obs(questions, "OO-F1", "o2")
    assert runner.match_text("この要素は別要素の原本領域に重なっている。\n重なっている対象: 穴の中の文字（図形 ID 4・スライド1）（覆っているかは未確定）", front)
    assert not runner.match_text("この要素は別要素の原本領域に重なっている。", front)


def test_the_opaque_control_is_judged_covered(questions):
    obs = _obs(questions, "OO-F2", "o1")
    covered = "図形: 「休憩室は2階」\n可視性: 「図形に覆われている」\n前面図形: 目隠しの四角（図形 ID 5・スライド2）（覆われていると判定）"
    assert runner.match_text(covered, obs)
    assert not runner.match_text("図形: 「休憩室は2階」\n前面図形: 目隠しの四角（覆われているかは未確定）", obs)


def test_each_vml_shape_keeps_its_own_strike(questions):
    a, b = _obs(questions, "OO-G1", "o1"), _obs(questions, "OO-G2", "o1")
    rect = "テキストボックス: 「【取り消し線：条件甲】条件乙」"
    shape = "テキストボックス: 「条件甲【取り消し線：条件乙】」"
    assert runner.match_text(rect, a) and runner.match_text(shape, b)
    assert not runner.match_text(rect + "\n" + shape, a) and not runner.match_text(rect + "\n" + shape, b)   # 2 つの図形が混ざる
    for qid in ("OO-G1", "OO-G2"):                                     # 作成者の付けた ID は要求せず、問いにも出さない
        q = _by_id(questions, qid)
        assert [o["id"] for o in q["observations"]] == ["o1"]
        assert "rectA" not in q["question"] and "shapeB" not in q["question"]
        assert not any("rectA" in p or "shapeB" in p for o in q["observations"] for p in o["all"])


def test_two_paragraphs_are_not_joined(questions):
    obs = _obs(questions, "OO-H1", "o1")
    assert runner.match_text("テキストボックス: 「10\n20」", obs)
    assert not runner.match_text("テキストボックス: 「10\n20」\n前面図形: 「1020」", obs)
    assert not runner.match_text("テキストボックス: 「1020」", obs)


def test_four_stacked_shapes_stay_named_one_per_line_and_undetermined(questions):
    back, front = _obs(questions, "OO-I1", "o1"), _obs(questions, "OO-I2", "o1")
    row = "品目: 「品目乙」\n値: 「200」\n" + "\n".join(
        f"前面図形: 説明図形{n} 補足（図形 ID {n + 1}・シート「一覧」のセル範囲 B3:B3）（覆われているかは未確定）" for n in range(1, 5))
    assert runner.match_text(row, back)
    three = row.replace("前面図形: 説明図形4 補足（図形 ID 5・シート「一覧」のセル範囲 B3:B3）（覆われているかは未確定）", "")
    assert not runner.match_text(three, back)
    assert not runner.match_text(row.replace("説明図形1 補足", "説明図形1\n補足"), back)               # 名前の改行が行を割る
    assert not runner.match_text(row.replace("覆われているかは未確定", "覆われていると判定", 1), back)
    chunk = "この要素は別要素の原本領域に重なっている。\n重なっている対象: セルB3（シート「一覧」）（覆っているかは未確定）"
    assert not runner.match_text(chunk, front)                                                         # 名前が前面側のチャンクに無い
    assert runner.match_text("この要素: 説明図形1 補足（図形 ID 2・シート「一覧」のセル範囲 B3:B3）\n" + chunk, front)
    assert not runner.match_text(row, front)                                                           # 背面側の行は前面の観測にならない


def test_the_bundled_vector_input_may_summarise_the_names_but_keeps_the_state(questions):
    """ベクトルの入力は予算の都合で名前の列挙を『（ほかK件）』に要約する。4 つの名前は本文でだけ要求し、ベクトルの入力では値と未確定を見る。"""
    back = _obs(questions, "OO-I1", "o1")
    summarised = ("列: 品目 | 値\n品目乙 | 200 | 前面図形: 説明図形1 補足…（覆われているかは未確定） | "
                  "前面図形: 説明図形2 補足…（覆われているかは未確定） | （ほか2件）")
    assert not runner.match_text(summarised, back)                                         # 本文の観測としては足りない
    assert runner.match_text(summarised, runner.vector_view(back))
    judged = summarised.replace("覆われているかは未確定", "覆われていると判定", 1)
    assert not runner.match_text(judged, runner.vector_view(back))                         # 覆われていると判定は要約しても通らない
    assert not runner.match_text(summarised.replace("未確定", "確定"), runner.vector_view(back))
    body = {"text": "品目: 「品目乙」\n値: 「200」\n" + "\n".join(
        f"前面図形: 説明図形{n} 補足（図形 ID {n + 1}・シート「一覧」のセル範囲 B3:B3）（覆われているかは未確定）" for n in range(1, 5)),
        "vector_text": summarised}
    assert runner.stage1_observation(back, [body]) == {"doc_indexed": True, "es_text": True, "vector_input": True}
    plain = {"id": "o9", "all": ["甲", "乙"]}
    assert runner.vector_view(plain) is plain                                              # vector が無い観測は本文と同じ条件


def test_new_scene_answers_are_scored(questions):
    f1 = _by_id(questions, "OO-F1")
    assert runner.score_answer(f1, f1["gold"]["answer"], ["避難経路の案内.pptx"])["passed"]
    covered = runner.score_answer(f1, "「非常口は北側」の文字は前面の枠に覆われていて読めません。出典: 避難経路の案内.pptx")
    assert covered["wrong"] and not covered["passed"]

    f2 = _by_id(questions, "OO-F2")
    assert runner.score_answer(f2, f2["gold"]["answer"], ["避難経路の案内.pptx"])["passed"]
    assert runner.score_answer(f2, "休憩室は2階の表示は覆われておらず読めます。出典: 避難経路の案内.pptx")["wrong"]

    g1, g2 = _by_id(questions, "OO-G1"), _by_id(questions, "OO-G2")
    assert runner.score_answer(g1, g1["gold"]["answer"], ["条件の比較図.docx"])["passed"]
    assert runner.score_answer(g2, g2["gold"]["answer"], ["条件の比較図.docx"])["passed"]
    swapped = runner.score_answer(g1, "条件甲に取り消し線が付いている図形では、条件乙にも取り消し線が付いています。出典: 条件の比較図.docx")
    assert swapped["wrong"] and not swapped["passed"]
    assert runner.score_answer(g2, "条件乙に取り消し線が付いている図形では、条件甲にも取り消し線が付いています。出典: 条件の比較図.docx")["wrong"]

    h1 = _by_id(questions, "OO-H1")
    assert runner.score_answer(h1, h1["gold"]["answer"], ["基準値の枠.docx"])["passed"]
    joined = runner.score_answer(h1, "基準値は 1020 です。出典: 基準値の枠.docx")
    assert joined["wrong"] and not joined["passed"]

    i1, i2 = _by_id(questions, "OO-I1"), _by_id(questions, "OO-I2")
    assert runner.score_answer(i1, i1["gold"]["answer"], ["品目別の値.xlsx"])["passed"]
    assert runner.score_answer(i2, i2["gold"]["answer"], ["品目別の値.xlsx"])["passed"]
    assert runner.score_answer(i1, "品目乙の値は 200 ですが、図形に覆われていて確認できません。出典: 品目別の値.xlsx")["wrong"]
    assert runner.score_answer(i2, "説明図形1 は B3 を覆っています。出典: 品目別の値.xlsx")["wrong"]


def test_fixture_facts_behind_the_new_scenes():
    """資料そのものに、再現の前提（レイアウト側のドーナツ・不透明な塗り・VML のグループ・2 段落・名前の改行）がある。"""
    from pptx import Presentation
    prs = Presentation(FIXTURES / build_fixtures.LAYOUT_PPTX)
    ring = next(s for s in prs.slides[0].shapes if s.name == "輪の枠")
    assert ring.is_placeholder and not ring._element.spPr.xpath("./a:prstGeom")
    assert [g.get("prst") for g in next(p for p in prs.slides[0].slide_layout.placeholders if p.placeholder_format.idx == 13)
            ._element.spPr.xpath("./a:prstGeom")] == ["donut"]
    with zipfile.ZipFile(FIXTURES / build_fixtures.STACK_XLSX) as z:
        names = [n for n in z.namelist() if n.endswith("drawing1.xml")]
        assert names and z.read(names[0]).decode("utf-8").count("&#10;補足") + z.read(names[0]).decode("utf-8").count("\n補足") == 4
