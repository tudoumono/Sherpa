"""「今すぐ更新」の自己修復（R1 グラフ・R2 ES）を実 Neo4j・実 ES・実 Postgres で検証する。

docs/proposals/2026-09-23-今すぐ更新で索引とグラフを直す.md の受入条件 1〜4 を、
`worker.sync()` の不変分岐（"今すぐ更新" と同じ入口）を通して固定する。world は本テスト専用に
`worlds.register()` で登録し（フィクスチャは使わない・COBOL の CALL で実際のエッジを1本作る）、
Neo4j/ES を直接 corrupt してから sync 1回で戻ることを確かめる。world_id はレーン注入の
`TEST_WORLD_ID`（`_world_setup.py`）から派生させ、並行レーンと衝突しないようにする。

要 Neo4j・要 ES（ES 不可は SKIP・既存の結合テストの流儀）。埋め込みは `sherpa.embeddings`
の外部境界だけスタブ化する（`tests/integration/test_es_index.py::test_vector_hybrid_stubbed`
と同じ手段）。
"""
from __future__ import annotations

import pathlib
import shutil
import tempfile
import urllib.parse

import pytest
from _world_setup import TEST_WORLD_ID, driver

from sherpa import es_index, store, worlds
from sherpa.ingest import worker, world_neo4j


def _mk_world(root: pathlib.Path) -> None:
    """MAINPROG が FOOPROG を CALL する2本＋無関係な BAZPROG の3本（ES 削除後も0件にならない余白）
    ＋埋め込み対象になる xlsx を1本（rag_chunks を持つのは Office/PDF のみ・`_rag_chunk_source_exts`
    参照——COBOL ソースだけでは embed 対象チャンクが無く reused/embedded の検証にならない）。
    """
    d = root / "4期" / "03_開発" / "01_ソース"
    d.mkdir(parents=True, exist_ok=True)
    (d / "MAINPROG.cbl").write_text(
        "       IDENTIFICATION DIVISION.\n"
        "       PROGRAM-ID. MAINPROG.\n"
        "       PROCEDURE DIVISION.\n"
        "           CALL 'FOOPROG'.\n"
        "           STOP RUN.\n", encoding="utf-8")
    (d / "FOOPROG.cbl").write_text(
        "       IDENTIFICATION DIVISION.\n"
        "       PROGRAM-ID. FOOPROG.\n"
        "       PROCEDURE DIVISION.\n"
        "           STOP RUN.\n", encoding="utf-8")
    (d / "BAZPROG.cbl").write_text(
        "       IDENTIFICATION DIVISION.\n"
        "       PROGRAM-ID. BAZPROG.\n"
        "       PROCEDURE DIVISION.\n"
        "           STOP RUN.\n", encoding="utf-8")
    import openpyxl
    dd = root / "4期" / "02_設計" / "01_基本設計"
    dd.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "明細"
    ws["A1"], ws["B1"] = "No", "内容"
    ws["A2"], ws["B2"] = 1, "サンプル内容"
    wb.save(dd / "税計算仕様書.xlsx")


def _graph_counts(drv, world: str):
    with drv.session() as s:
        nc = s.run("MATCH (n:Entity {world_id:$w}) RETURN count(n) AS c", w=world).single()["c"]
        ec = s.run("MATCH (a:Entity {world_id:$w})-[r]->() WHERE r.world_id=$w RETURN count(r) AS c",
                   w=world).single()["c"]
    return nc, ec


def _teardown_world(world: str, root, drv) -> None:
    try:
        worlds.delete(world)
    except Exception:
        pass
    store.delete_world_row(world)
    with drv.session() as s:
        s.run("MATCH (n) WHERE n.world_id=$w AND (n:Entity OR n:SherpaMeta) DETACH DELETE n", w=world)
    shutil.rmtree(root, ignore_errors=True)


def test_sync_unchanged_self_heals_graph_and_es_corruption_then_stabilizes(monkeypatch):
    """受入 1・2・4: 資料不変のまま (a)グラフ全消去 (b)世代スタンプ改ざん (c)ノード一部消去
    (ES)文書一部消去の各状態から、不変分岐の sync を1回呼ぶと元に戻る。埋め込みは実呼び出し0回
    （キャッシュ全ヒット）。壊れていない状態で2回 sync してもグラフ/ES いずれも作り直さない。
    """
    if not es_index.available():
        pytest.skip("ES 未起動")
    from sherpa import embeddings
    embed_calls: list[int] = []
    monkeypatch.setattr(embeddings, "cfg",
                        lambda settings=None, **kw: {"provider": "stub", "key": "x", "model": "stub-m", "dim": 8})

    def _stub_embed(texts, c, **kw):
        embed_calls.append(len(texts))
        return [[0.1] * 8 for _ in texts]
    monkeypatch.setattr(embeddings, "embed", _stub_embed)

    W = f"{TEST_WORLD_ID}-rr-heal"
    root = pathlib.Path(tempfile.mkdtemp())
    _mk_world(root)
    drv = driver()
    try:
        reg = worlds.register(W, str(root.resolve()))
        assert reg["status"] in ("auto_published", "auto_published_with_flags")
        from _world_registry import register_test_world
        register_test_world(W)

        nc0, ec0 = _graph_counts(drv, W)
        assert nc0 >= 2 and ec0 >= 1   # 前提: 3プログラム・CALL 1本以上

        doc_count0 = es_index.count(W)
        assert doc_count0 and doc_count0 >= 2   # 前提: 削除後も0件にならない複数チャンク
        embed_calls.clear()   # 登録時の初回埋め込みは対象外（以降だけ「実呼び出し0回」を見る）

        # --- (a) グラフを丸ごと消した状態から ---
        with drv.session() as s:
            s.run("MATCH (n) WHERE n.world_id=$w AND (n:Entity OR n:SherpaMeta) DETACH DELETE n", w=W)
        rid = store.start_ingest_run(W)["id"]
        res = worker.sync(W, run_id=rid)
        assert res["changed"] is False and res["status"] == "unchanged"
        rec = store.get_latest_run_summary(W)
        assert rec["id"] == rid and rec["status"] == "auto_published"
        assert _graph_counts(drv, W) == (nc0, ec0)

        # --- (b) 世代スタンプを古い値へ書き換えた状態から ---
        with drv.session() as s:
            s.run("MATCH (m:SherpaMeta {world_id:$w}) SET m.schema_era=$old", w=W, old="stale-era-xyz")
        rid = store.start_ingest_run(W)["id"]
        res = worker.sync(W, run_id=rid)
        assert res["changed"] is False
        rec = store.get_latest_run_summary(W)
        assert rec["id"] == rid and rec["status"] == "auto_published"
        with drv.session() as s:
            era = s.run("MATCH (m:SherpaMeta {world_id:$w}) RETURN m.schema_era AS e", w=W).single()["e"]
        assert era == world_neo4j.GRAPH_SCHEMA_ERA
        assert _graph_counts(drv, W) == (nc0, ec0)

        # --- (c) ノードを1件だけ消した状態から ---
        with drv.session() as s:
            s.run("MATCH (n:Entity {world_id:$w}) WITH n LIMIT 1 DETACH DELETE n", w=W)
        rid = store.start_ingest_run(W)["id"]
        res = worker.sync(W, run_id=rid)
        assert res["changed"] is False
        rec = store.get_latest_run_summary(W)
        assert rec["id"] == rid and rec["status"] == "auto_published"
        assert _graph_counts(drv, W) == (nc0, ec0)

        assert embed_calls == []   # グラフ修復は静的解析＋投入のみ＝埋め込みを一切呼ばない

        # --- ES: 文書を1件だけ消した状態から ---
        hits = es_index._req("GET", f"/{es_index._index(W)}/_search?size=1")
        one_id = hits["hits"]["hits"][0]["_id"]
        es_index._req("DELETE", f"/{es_index._index(W)}/_doc/{urllib.parse.quote(one_id, safe='')}?refresh=true")
        assert es_index.count(W) == doc_count0 - 1

        rid = store.start_ingest_run(W)["id"]
        res = worker.sync(W, run_id=rid)
        assert res["changed"] is False
        assert es_index.count(W) == doc_count0
        assert embed_calls == []   # 資料不変＝キャッシュ全ヒットで実埋め込み呼び出しは0回

        rec = store.get_latest_run_summary(W)
        assert rec["id"] == rid and rec["status"] == "auto_published"
        counts = (rec["extraction_snapshot"] or {}).get("counts") or {}
        assert counts.get("reused_chunks", 0) > 0
        assert counts.get("embedded_chunks", 0) == 0
        # 不変分岐で ES を実際に張り直した run は、全段の経路と同じ形で
        # extraction_snapshot.es にも reused/embedded が残る（routers/worlds.py の
        # stage_summary.es はこれをそのまま渡すだけ・web/ingest.js の表示元）。
        es_stage = (rec["extraction_snapshot"] or {}).get("es") or {}
        assert es_stage.get("reused", 0) > 0
        assert es_stage.get("embedded", 0) == 0

        # --- ES: content_sig/doc_count のスタンプ書き込み失敗（外部境界）は次回 sync が必ず張り直す ---
        # 1件消してから、スタンプの PUT（`_confirm_content_sig` の `_mapping` 書き込み）だけを
        # 失敗させる——bulk 本体（delete/create/_bulk）は成功するので件数は戻るが、スタンプが
        # 確定しないため needs_reindex は「要張り直し」のまま残るはず（content_sig と doc_count を
        # 同じ書き込みにまとめた契約の直接確認）。
        current_sig, _ = worker.world_state(W)
        hits2 = es_index._req("GET", f"/{es_index._index(W)}/_search?size=1")
        two_id = hits2["hits"]["hits"][0]["_id"]
        es_index._req("DELETE",
                      f"/{es_index._index(W)}/_doc/{urllib.parse.quote(two_id, safe='')}?refresh=true")
        assert es_index.count(W) == doc_count0 - 1

        orig_req = es_index._req

        def _boom_on_mapping_put(method, path, *a, **kw):
            if method == "PUT" and path.endswith("/_mapping"):
                raise RuntimeError("simulated ES mapping PUT failure")
            return orig_req(method, path, *a, **kw)
        monkeypatch.setattr(es_index, "_req", _boom_on_mapping_put)
        rid = store.start_ingest_run(W)["id"]
        res = worker.sync(W, run_id=rid)
        monkeypatch.setattr(es_index, "_req", orig_req)   # 以降は正しい接続に戻す
        assert res["changed"] is False
        assert es_index.count(W) == doc_count0            # bulk 本体は成功して件数は戻る
        assert es_index.needs_reindex(W, current_sig) is True   # スタンプ確定は失敗＝要張り直しのまま

        rid = store.start_ingest_run(W)["id"]
        res = worker.sync(W, run_id=rid)
        assert res["changed"] is False
        assert es_index.count(W) == doc_count0
        assert es_index.needs_reindex(W, current_sig) is False   # 今度はスタンプも確定し収束する

        # --- 受入4: 壊れていない状態で sync を2回呼んでも作り直さない（照合が安定・ループしない）---
        orig_load_world = world_neo4j.load_world
        orig_index_world = es_index.index_world
        rebuild_calls = {"graph": 0, "es": 0}

        def _counting_load_world(*a, **kw):
            rebuild_calls["graph"] += 1
            return orig_load_world(*a, **kw)

        def _counting_index_world(*a, **kw):
            rebuild_calls["es"] += 1
            return orig_index_world(*a, **kw)
        monkeypatch.setattr(world_neo4j, "load_world", _counting_load_world)
        monkeypatch.setattr(es_index, "index_world", _counting_index_world)
        for _ in range(2):
            rid = store.start_ingest_run(W)["id"]
            res = worker.sync(W, run_id=rid)
            assert res["changed"] is False
            rec = store.get_latest_run_summary(W)
            assert rec["id"] == rid and rec["status"] == "auto_published"
        assert rebuild_calls == {"graph": 0, "es": 0}
    finally:
        _teardown_world(W, root, drv)
        drv.close()


def test_sync_unchanged_reports_graph_repair_failed_when_neo4j_unreachable(monkeypatch):
    """受入 3: Neo4j に繋がらないとき、run は failed になり理由に graph_repair_failed が入る。
    例外は worker.sync() の外へ漏れない。"""
    W = f"{TEST_WORLD_ID}-rr-fail"
    root = pathlib.Path(tempfile.mkdtemp())
    _mk_world(root)
    drv = driver()
    try:
        reg = worlds.register(W, str(root.resolve()))
        assert reg["status"] in ("auto_published", "auto_published_with_flags")
        from _world_registry import register_test_world
        register_test_world(W)

        # 以降だけ Neo4j を不到達にする（`check_graph_counts` が最初に触れる接続先）。cleanup は
        # 正しい接続先で行うため、monkeypatch fixture の自動 undo に任せず自前で復元する。
        orig_env = world_neo4j._env
        world_neo4j._env = lambda: {"uri": "bolt://127.0.0.1:1", "user": "neo4j", "pw": "x"}
        try:
            rid = store.start_ingest_run(W)["id"]
            res = worker.sync(W, run_id=rid)   # 例外は外へ漏れない
        finally:
            world_neo4j._env = orig_env
        assert res["changed"] is False

        rec = store.get_latest_run_summary(W)
        assert rec["id"] == rid and rec["status"] == "failed"
        flags = (rec["extraction_snapshot"] or {}).get("flags") or []
        reasons = [f.get("reason", "") for f in flags]
        assert any(r.startswith("graph_repair_failed:") for r in reasons)
    finally:
        _teardown_world(W, root, drv)
        drv.close()
