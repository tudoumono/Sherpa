"""WMF/EMF/Office パッケージのテスト用合成データ（実環境の資料は使わない）。"""
from __future__ import annotations

import io
import struct
import zipfile


def _wmf_record(function: int, params: bytes) -> bytes:
    if len(params) % 2:
        params += b"\x00"
    return struct.pack("<IH", (6 + len(params)) // 2, function) + params


def wmf(records: list[bytes], *, placeable: bool = True, objects: int = 4) -> bytes:
    body = b"".join(records) + _wmf_record(0x0000, b"")
    header = struct.pack("<HHHIHIH", 1, 9, 0x0300, (18 + len(body)) // 2, objects, 20, 0)
    prefix = struct.pack("<IHhhhhHIH", 0x9AC6CDD7, 0, 0, 0, 1000, 1000, 1440, 0, 0) if placeable else b""
    return prefix + header + body


def wmf_font(charset: int) -> bytes:
    return _wmf_record(0x02FB, struct.pack("<hhhhhBBBB", -12, 0, 0, 0, 400, 0, 0, 0, charset) + b"\x00" * 24)


def wmf_select(index: int) -> bytes:
    return _wmf_record(0x012D, struct.pack("<H", index))


def wmf_textout(text: bytes, x: int, y: int) -> bytes:
    pad = b"\x00" if len(text) % 2 else b""
    return _wmf_record(0x0521, struct.pack("<H", len(text)) + text + pad + struct.pack("<hh", y, x))


def wmf_exttextout(text: bytes, x: int, y: int) -> bytes:
    pad = b"\x00" if len(text) % 2 else b""
    return _wmf_record(0x0A32, struct.pack("<hhHH", y, x, len(text), 0) + text + pad)


def dib(width: int, height: int, *, color: tuple[int, int, int] = (200, 30, 30)) -> bytes:
    row = (width * 3 + 3) // 4 * 4
    pixels = bytes(color[::-1]) * width + b"\x00" * (row - width * 3)
    info = struct.pack("<IiiHHIIiiII", 40, width, height, 1, 24, 0, row * height, 0, 0, 0, 0)
    return info + pixels * height


def wmf_stretchdib(bitmap: bytes) -> bytes:
    params = struct.pack("<I", 0x00CC0020) + struct.pack("<9h", 0, 0, 0, 0, 32, 32, 0, 0, 0)
    return _wmf_record(0x0F43, params[:22] + bitmap)


def _emf_record(kind: int, payload: bytes) -> bytes:
    payload += b"\x00" * (-len(payload) % 4)
    return struct.pack("<II", kind, 8 + len(payload)) + payload


def emf(records: list[bytes]) -> bytes:
    header_payload = bytearray(80)
    struct.pack_into("<8i", header_payload, 0, 0, 0, 100, 100, 0, 0, 1000, 1000)
    header_payload[32:36] = b" EMF"
    struct.pack_into("<I", header_payload, 36, 0x10000)
    return (_emf_record(1, bytes(header_payload)) + b"".join(records)
            + _emf_record(14, struct.pack("<III", 0, 16, 20)))


def emf_font(handle: int, charset: int) -> bytes:
    logfont = struct.pack("<5i4B", -12, 0, 0, 0, 400, 0, 0, 0, charset) + b"\x00" * 64
    return _emf_record(82, struct.pack("<I", handle) + logfont + b"\x00" * 100)


def emf_select(handle: int) -> bytes:
    return _emf_record(37, struct.pack("<I", handle))


def _exttext(kind: int, text: bytes, count: int, x: int, y: int) -> bytes:
    fixed = struct.pack("<4iIff", 0, 0, 0, 0, 1, 1.0, 1.0)            # bounds + graphics mode + scales（記録先頭 8 から 28 バイト）
    off_string = 8 + len(fixed) + 40
    emrtext = struct.pack("<iiII", x, y, count, off_string) + struct.pack("<I4iI", 0, 0, 0, 0, 0, 0)
    return _emf_record(kind, fixed + emrtext + text)


def emf_exttextout_w(text: str, x: int, y: int) -> bytes:
    return _exttext(84, text.encode("utf-16-le"), len(text), x, y)


def emf_exttextout_a(text: bytes, x: int, y: int) -> bytes:
    return _exttext(83, text, len(text), x, y)


def emf_plus_drawstring(text: str, x: float, y: float) -> bytes:
    raw = text.encode("utf-16-le")
    body = struct.pack("<IIIffff", 0, 0, len(text), x, y, 100.0, 20.0) + raw
    body += b"\x00" * (-len(body) % 4)
    plus = struct.pack("<HHII", 0x401C, 0, 12 + len(body), len(body)) + body
    return _emf_record(70, struct.pack("<I", 4 + len(plus)) + b"EMF+" + plus)


def emf_stretchdibits(bitmap: bytes) -> bytes:
    header_len = 40
    off_bmi = 8 + 72
    fixed = struct.pack("<4i", 0, 0, 100, 100) + struct.pack(
        "<6i4I2I2i", 0, 0, 0, 0, 32, 32, off_bmi, header_len, off_bmi + header_len, len(bitmap) - header_len,
        0, 0x00CC0020, 32, 32)
    return _emf_record(81, fixed + bitmap)


def docx_with_media(media_name: str, media: bytes) -> bytes:
    ns = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
          'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
          'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
          'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"')
    document = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {ns}><w:body>'
        '<w:p><w:r><w:t>前の段落</w:t></w:r></w:p>'
        '<w:p><w:r><w:drawing><wp:inline><wp:extent cx="990000" cy="720000"/>'
        '<wp:docPr id="1" name="図 1"/><a:graphic><a:graphicData '
        'uri="http://schemas.openxmlformats.org/drawingml/2006/picture"><pic:pic>'
        '<pic:nvPicPr><pic:cNvPr id="1" name="図 1"/><pic:cNvPicPr/></pic:nvPicPr>'
        '<pic:blipFill><a:blip r:embed="rId5"/></pic:blipFill>'
        '<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="990000" cy="720000"/></a:xfrm></pic:spPr>'
        '</pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>'
        '<w:p><w:r><w:t>後の段落</w:t></w:r></w:p></w:body></w:document>')
    rels = ('<?xml version="1.0" encoding="UTF-8"?><Relationships '
            'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId5" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
            f'Target="media/{media_name}"/></Relationships>')
    content_types = ('<?xml version="1.0" encoding="UTF-8"?><Types '
                     'xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                     '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                     '<Default Extension="xml" ContentType="application/xml"/>'
                     '<Default Extension="emf" ContentType="image/x-emf"/><Default Extension="wmf" ContentType="image/x-wmf"/>'
                     '<Override PartName="/word/document.xml" '
                     'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                     '</Types>')
    root_rels = ('<?xml version="1.0" encoding="UTF-8"?><Relationships '
                 'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                 '<Relationship Id="rId1" '
                 'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
                 'Target="word/document.xml"/></Relationships>')
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as package:
        package.writestr("[Content_Types].xml", content_types)
        package.writestr("_rels/.rels", root_rels)
        package.writestr("word/document.xml", document)
        package.writestr("word/_rels/document.xml.rels", rels)
        package.writestr(f"word/media/{media_name}", media)
    return buffer.getvalue()


def xlsx_with_media(media_name: str, media: bytes, *, anchor_row0: int = 11, anchor_col0: int = 1) -> bytes:
    """A1:B3 と A20:B22 に値のあるブックへ、指定アンカー（0 始まり）に図を 1 つ置く。"""
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    for row in range(1, 4):
        sheet.cell(row=row, column=1, value=f"上{row}")
        sheet.cell(row=row, column=2, value=row)
    for row in range(20, 23):
        sheet.cell(row=row, column=1, value=f"下{row}")
        sheet.cell(row=row, column=2, value=row)
    raw = io.BytesIO()
    workbook.save(raw)
    source = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    drawing = (
        '<?xml version="1.0" encoding="UTF-8"?><xdr:wsDr '
        'xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<xdr:twoCellAnchor><xdr:from><xdr:col>{anchor_col0}</xdr:col><xdr:colOff>0</xdr:colOff>'
        f'<xdr:row>{anchor_row0}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
        f'<xdr:to><xdr:col>{anchor_col0 + 2}</xdr:col><xdr:colOff>0</xdr:colOff>'
        f'<xdr:row>{anchor_row0 + 3}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
        '<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="2" name="図 1"/><xdr:cNvPicPr/></xdr:nvPicPr>'
        '<xdr:blipFill><a:blip r:embed="rId1"/></xdr:blipFill>'
        '<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1" cy="1"/></a:xfrm></xdr:spPr></xdr:pic>'
        '<xdr:clientData/></xdr:twoCellAnchor></xdr:wsDr>')
    drawing_rels = (
        '<?xml version="1.0" encoding="UTF-8"?><Relationships '
        'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
        f'Target="../media/{media_name}"/></Relationships>')
    sheet_rels = (
        '<?xml version="1.0" encoding="UTF-8"?><Relationships '
        'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" '
        'Target="../drawings/drawing1.xml"/></Relationships>')
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                text = text.replace(
                    "</worksheet>",
                    '<drawing xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
                    'r:id="rId1"/></worksheet>')
                data = text.encode("utf-8")
            if info.filename == "[Content_Types].xml":
                text = data.decode("utf-8").replace(
                    "</Types>",
                    '<Default Extension="emf" ContentType="image/x-emf"/>'
                    '<Override PartName="/xl/drawings/drawing1.xml" '
                    'ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>')
                data = text.encode("utf-8")
            target.writestr(info.filename, data)
        target.writestr("xl/worksheets/_rels/sheet1.xml.rels", sheet_rels)
        target.writestr("xl/drawings/drawing1.xml", drawing)
        target.writestr("xl/drawings/_rels/drawing1.xml.rels", drawing_rels)
        target.writestr(f"xl/media/{media_name}", media)
    return out.getvalue()


def wmf_create_pen() -> bytes:
    return _wmf_record(0x02FA, struct.pack("<HhhI", 0, 1, 0, 0))


def wmf_delete(index: int) -> bytes:
    return _wmf_record(0x01F0, struct.pack("<H", index))


def pptx_with_media(media_name: str, media: bytes) -> bytes:
    ns = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
          'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')

    def shape(name: str, text: str) -> str:
        return (f'<p:sp><p:nvSpPr><p:cNvPr id="2" name="{name}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr/>'
                f'<p:txBody><a:bodyPr/><a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp>')

    picture = ('<p:pic><p:nvPicPr><p:cNvPr id="3" name="図"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>'
               '<p:blipFill><a:blip r:embed="rId2"/></p:blipFill><p:spPr/></p:pic>')
    slide = (f'<?xml version="1.0" encoding="UTF-8"?><p:sld {ns}><p:cSld><p:spTree>'
             '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/>'
             f'{shape("t1", "前の文")}{picture}{shape("t2", "後の文")}</p:spTree></p:cSld></p:sld>')
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    office = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as package:
        package.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        package.writestr("ppt/presentation.xml", f'<?xml version="1.0"?><p:presentation {ns}><p:sldIdLst><p:sldId id="256" r:id="rId1"/></p:sldIdLst></p:presentation>')
        package.writestr("ppt/_rels/presentation.xml.rels", f'<?xml version="1.0"?><Relationships xmlns="{rel_ns}"><Relationship Id="rId1" Type="{office}/slide" Target="slides/slide1.xml"/></Relationships>')
        package.writestr("ppt/slides/slide1.xml", slide)
        package.writestr("ppt/slides/_rels/slide1.xml.rels", f'<?xml version="1.0"?><Relationships xmlns="{rel_ns}"><Relationship Id="rId2" Type="{office}/image" Target="../media/{media_name}"/></Relationships>')
        package.writestr(f"ppt/media/{media_name}", media)
    return buffer.getvalue()


def emf_smalltextout(text: bytes | str, x: int, y: int, *, small: bool, no_rect: bool) -> bytes:
    options = (0x200 if small else 0) | (0x100 if no_rect else 0)
    raw = text if small else text.encode("utf-16-le")
    count = len(raw) if small else len(text)
    fixed = struct.pack("<iiIIIff", x, y, count, options, 1, 1.0, 1.0)
    bounds = b"" if no_rect else struct.pack("<4i", 0, 0, 0, 0)
    return _emf_record(108, fixed + bounds + raw)


def emf_blit(kind: int, bitmap: bytes) -> bytes:
    """AlphaBlend(114)/TransparentBlt(116)/MaskBlt(78)/PlgBlt(79) の元ビットマップだけを持つ記録。"""
    base = 96 if kind == 79 else 84
    header_len = 40
    off_bmi = base + 16
    record = bytearray(base - 8)                                  # 記録先頭 8 バイトは _emf_record が付ける
    record += struct.pack("<4I", off_bmi, header_len, off_bmi + header_len, len(bitmap) - header_len)
    return _emf_record(kind, bytes(record) + bitmap)
