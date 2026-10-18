"""使い方: python3 scripts/bench/preserve_check.py <dir>
<dir>/world/設計書 に原本（xlsx/docx/pptx）、<dir>/derived に `office_md.build_derived` の出力を置く。
原本の値（セル・段落・図形の文字）が rag.md のチャンク／人向け MD にどれだけ残るか（値の保持率）と、
同じ行の値が 1 つのチャンク／MD の 1 行に同居するか（行のまとまりの保持率）を数える。
"""
import gzip
import json
import re
import sys
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


def norm(s) -> str:
    return re.sub(r"\s+", "", str(s))


def chunks(derived: Path, name: str) -> list:
    raw = (derived / "rag" / "設計書" / f"{name}.rag_chunks.jsonl").read_bytes()
    try:
        raw = gzip.decompress(raw)
    except OSError:
        pass
    return [json.loads(line)["body"] for line in raw.decode("utf-8").splitlines() if line.strip()]


def values_and_rows(f: Path) -> tuple[list, list]:
    values: list = []
    rowsets: list = []
    ext = f.suffix.lower()
    if ext == ".xlsx":
        for ws in load_workbook(f, data_only=True).worksheets:
            for row in ws.iter_rows():
                vals = [norm(c.value) for c in row if c.value is not None and norm(c.value)]
                values += vals
                if len(vals) >= 2:
                    rowsets.append(vals)
    elif ext == ".docx":
        d = Document(f)
        values += [norm(p.text) for p in d.paragraphs if norm(p.text)]
        for t in d.tables:
            for r in t.rows:
                vals = list(dict.fromkeys(norm(c.text) for c in r.cells if norm(c.text)))
                values += vals
                if len(vals) >= 2:
                    rowsets.append(vals)
    elif ext == ".pptx":
        for sl in Presentation(f).slides:
            for sh in sl.shapes:
                if sh.has_text_frame and norm(sh.text_frame.text):
                    values.append(norm(sh.text_frame.text))
                if getattr(sh, "has_table", False) and sh.has_table:
                    for r in sh.table.rows:
                        vals = [norm(c.text) for c in r.cells if norm(c.text)]
                        values += vals
                        if len(vals) >= 2:
                            rowsets.append(vals)
    return [v for v in values if v], rowsets


def main() -> int:
    base = Path(sys.argv[1])
    world, derived = base / "world" / "設計書", base / "derived"
    res = []
    for f in sorted(world.iterdir()):
        try:
            ch = chunks(derived, f.name)
            hm = (derived / "md" / "設計書" / f"{f.name}.md").read_text(encoding="utf-8")
        except FileNotFoundError:
            continue
        values, rowsets = values_and_rows(f)
        if not values and f.suffix.lower() not in (".xlsx", ".docx", ".pptx"):
            continue
        alltext, hmn = norm("\n".join(ch)), norm(hm)
        kept_rag = sum(1 for v in values if v in alltext)
        kept_hm = sum(1 for v in values if v in hmn)
        row_rag = sum(1 for vals in rowsets if any(all(v in norm(c) for v in vals) for c in ch))
        row_hm = sum(1 for vals in rowsets if any(all(v in norm(line) for v in vals) for line in hm.splitlines()))
        res.append((len(values), kept_rag, kept_hm, len(rowsets), row_rag, row_hm))
        nv, nr = max(1, len(values)), max(1, len(rowsets))
        print(f"{f.name:<28} values={len(values):>5} in_rag={kept_rag / nv:.2f} in_md={kept_hm / nv:.2f} "
              f"rows(>=2cells)={len(rowsets):>4} row_in_one_chunk={row_rag / nr:.2f} row_on_one_md_line={row_hm / nr:.2f} chunks={len(ch)}")
    tv, tr = max(1, sum(r[0] for r in res)), max(1, sum(r[3] for r in res))
    print(f"TOTAL values={tv} in_rag={sum(r[1] for r in res) / tv:.3f} in_md={sum(r[2] for r in res) / tv:.3f} "
          f"rows={tr} row_in_one_chunk={sum(r[4] for r in res) / tr:.3f} row_on_one_md_line={sum(r[5] for r in res) / tr:.3f}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
